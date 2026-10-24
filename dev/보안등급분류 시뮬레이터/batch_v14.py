"""v14 Excel batch processing with isolated clients and validated resumption.

No credentials or mutable global configuration enter this module. The caller
supplies a zero-argument client factory; every task obtains its own client.
Only the coordinator writes checkpoint/audit/workbook files. JSONL writes are
serialized and flushed per result; workbook replacement is atomic and periodic.
"""
from __future__ import annotations

from collections import Counter
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from copy import deepcopy
from datetime import datetime, timezone
import json
import math
import os
from pathlib import Path
import re
import tempfile
import time
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from classification_v14 import PERSONAS, classify_task
from hwpx_writer import export_documents
from v14_domain import (
    CONFIDENCES, CRITERIA, DEFAULT_MODELS, GRADES, GUIDELINE_ID, ITEM_NAMES,
    LEVELS, OPTIONAL_TASK_COLUMNS, SCHEMA_VERSION, STATUSES, TASK_COLUMNS,
    calculate_committee_average, calculate_reviewer_score, canonical_items,
    eligible_items, format_score, make_result, normalize_options, options_fingerprint,
    parse_score, rule_grade, task_fingerprint, task_snapshot, validate_weights,
)


RESULT_COLUMNS = ["연번", "과제명", "상태", "최종등급(참고)", "신뢰도", "종합평균",
                  "채점항목(선택)", "항목선택방식", "핵심고려지표", "AI 종합 검토의견",
                  "실제 심의위원회 논의 필요사항", "분류방식", "문서상태", "문서예정수",
                  "문서성공수", "문서실패수", "문서누락수", "문서경로"]
_RESULT_KEYS = {
    "schema_version", "guideline_id", "task_snapshot", "task_fingerprint", "options_fingerprint",
    "task_name", "method", "provider", "model", "status", "candidate", "criteria", "selected_items",
    "weights", "reviewers", "overall_avg", "rule_grade", "final_grade", "confidence", "rationale",
    "discussion_points", "unresolved_reasons", "documents",
}
_REVIEWER_KEYS = {"reviewer_id", "persona", "status", "item_reviews", "scores", "reviewer_score", "grade",
                  "confidence", "rationale", "discussion_points", "unresolved_reasons"}


class BatchError(RuntimeError):
    """Safe, actionable batch-level input or persistence failure."""


def _require(condition):
    if not condition:
        raise ValueError("저장 결과의 구조 또는 일치 검증에 실패했습니다.")


def _strings(value):
    _require(type(value) is list and all(type(item) is str for item in value))


def _number(value, nullable=False):
    _require((nullable and value is None) or (type(value) in (int, float) and math.isfinite(value)))


def _items(value, detailed=False):
    _require(type(value) is list and len(value) == len(ITEM_NAMES))
    for index, entry in enumerate(value):
        keys = {"item", "level", "note", "law_refs"}
        if detailed:
            keys |= {"matched_reference", "needs_expert_review"}
        _require(type(entry) is dict and set(entry) == keys)
        _require(entry["item"] == ITEM_NAMES[index] and entry["level"] in LEVELS)
        _require(type(entry["note"]) is str)
        _strings(entry["law_refs"])
        allowed = set(CRITERIA[index]["security_refs"] + CRITERIA[index]["sensitive_refs"])
        _require(len(set(entry["law_refs"])) == len(entry["law_refs"]) and set(entry["law_refs"]) <= allowed)
        if detailed:
            _require(type(entry["matched_reference"]) is str and type(entry["needs_expert_review"]) is bool)


def _validate_documents(documents, reviewers):
    _require(type(documents) is dict and set(documents) == {"status", "expected_files", "files", "errors"})
    _require(documents["status"] in ("미생성", "완료", "일부실패", "실패"))
    expected = documents["expected_files"]
    _require(type(expected) is int and expected >= 1)
    expected_identities = {("reviewer", reviewer["reviewer_id"]) for reviewer in reviewers} | {("summary", None)}
    if reviewers or documents["status"] != "미생성":
        _require(expected == len(expected_identities))
    else:
        # Before any reviewer is evaluated, the engine's ungenerated result can
        # retain the intended panel+summary count. Actual exports must contain
        # only the summary; invented individual opinions are never permitted.
        _require(expected in (1, 2, 4))
    _strings(documents["errors"])
    _require(type(documents["files"]) is list and len(documents["files"]) <= expected)
    identities = set()
    for entry in documents["files"]:
        _require(type(entry) is dict and set(entry) == {"kind", "reviewer_id", "path", "status", "error"})
        _require(entry["kind"] in ("summary", "reviewer") and entry["status"] in ("완료", "오류"))
        _require(type(entry["path"]) is str and bool(entry["path"]) and type(entry["error"]) is str)
        _require((entry["kind"] == "summary" and entry["reviewer_id"] is None)
                 or (entry["kind"] == "reviewer" and type(entry["reviewer_id"]) is str and bool(entry["reviewer_id"])))
        identity = (entry["kind"], entry["reviewer_id"])
        _require(identity in expected_identities and identity not in identities)
        identities.add(identity)
    completed = sum(entry["status"] == "완료" for entry in documents["files"])
    if documents["status"] == "미생성":
        _require(not documents["files"] and not documents["errors"])
    elif documents["status"] == "완료":
        _require(completed == expected and identities == expected_identities and not documents["errors"])
    elif documents["status"] == "일부실패":
        _require(0 < completed < expected and bool(documents["errors"]))
    else:
        _require(completed == 0 and bool(documents["errors"]))


def _validate_result(result, opts=None, task=None):
    """Validate the entire saved contract, including complete reviewer evidence.

    This is structural/integrity validation, not cryptographic authenticity.
    Fingerprints identify inputs/settings; they do not authenticate local files.
    Public run options are deliberately absent from persisted records.
    """
    _require(type(result) is dict and set(result) == _RESULT_KEYS)
    _require(type(result["schema_version"]) is int and result["schema_version"] == SCHEMA_VERSION)
    _require(result["guideline_id"] == GUIDELINE_ID)
    snapshot = result["task_snapshot"]
    _require(type(snapshot) is dict and set(snapshot) == set(TASK_COLUMNS + OPTIONAL_TASK_COLUMNS))
    _require(all(type(value) is str for value in snapshot.values()) and snapshot == task_snapshot(snapshot))
    _require(result["task_name"] == snapshot["과제명"])
    for key in ("task_fingerprint", "options_fingerprint"):
        _require(type(result[key]) is str and re.fullmatch(r"[0-9a-f]{64}", result[key]) is not None)
    _require(result["task_fingerprint"] == task_fingerprint(snapshot))
    _require(result["method"] in ("impact", "checklist") and result["provider"] in DEFAULT_MODELS)
    _require(type(result["model"]) is str)
    if result["provider"] in ("claude", "openai"):
        _require(result["model"] == DEFAULT_MODELS[result["provider"]])
    _require(result["status"] in STATUSES)
    for key in ("rationale", "discussion_points"):
        _require(type(result[key]) is str)
    _strings(result["unresolved_reasons"])
    _require(result["final_grade"] in GRADES + (None,) and result["rule_grade"] in GRADES + (None,))
    _require(result["confidence"] in CONFIDENCES + (None,))
    _number(result["overall_avg"], nullable=True)
    candidate = result["candidate"]
    _require(type(candidate) is dict and set(candidate) == {"is_candidate", "reason", "uncertainty"})
    _require(candidate["is_candidate"] is None or type(candidate["is_candidate"]) is bool)
    _require(type(candidate["reason"]) is str and type(candidate["uncertainty"]) is str)
    _items(result["criteria"], detailed=True)
    selected = result["selected_items"]
    _strings(selected)
    _require(list(canonical_items(selected)) == selected and len(set(selected)) == len(selected))
    _require(type(result["weights"]) is dict)
    for name, value in result["weights"].items():
        _require(name in selected)
        _number(value)
        _require(.8 <= value <= 1.2)
    reviewers = result["reviewers"]
    _require(type(reviewers) is list and len(reviewers) in (0, 1, 3))
    for index, reviewer in enumerate(reviewers):
        _require(type(reviewer) is dict and set(reviewer) == _REVIEWER_KEYS)
        _require(reviewer["reviewer_id"] == f"reviewer-{index + 1}")
        _require(reviewer["persona"] == PERSONAS[index] and reviewer["status"] in ("완료", "검토필요", "미실시"))
        for key in ("rationale", "discussion_points"):
            _require(type(reviewer[key]) is str)
        _strings(reviewer["unresolved_reasons"])
        _require(reviewer["grade"] in GRADES + (None,) and reviewer["confidence"] in CONFIDENCES + (None,))
        _number(reviewer["reviewer_score"], nullable=True)
        _require(type(reviewer["scores"]) is dict)
        for name, score in reviewer["scores"].items():
            _require(name in selected and type(score) is int and parse_score(score) == score)
        if reviewer["item_reviews"]:
            _items(reviewer["item_reviews"])
        else:
            _require(type(reviewer["item_reviews"]) is list and reviewer["status"] != "완료")
        if result["method"] == "checklist":
            _require(not reviewer["scores"] and reviewer["reviewer_score"] is None)
    _validate_documents(result["documents"], reviewers)
    if result["method"] == "checklist":
        _require(not result["weights"] and result["overall_avg"] is None and result["rule_grade"] is None)
    if result["status"] == "완료":
        _require(candidate["is_candidate"] is True and bool(candidate["reason"].strip()) and not candidate["uncertainty"])
        _require(not result["unresolved_reasons"] and bool(result["rationale"].strip()))
        _require(result["final_grade"] in GRADES and result["confidence"] in CONFIDENCES)
        _require(len(reviewers) in (1, 3))
        for reviewer in reviewers:
            _require(reviewer["status"] == "완료" and not reviewer["unresolved_reasons"])
            _require(reviewer["grade"] in GRADES and reviewer["confidence"] in CONFIDENCES and bool(reviewer["rationale"].strip()))
            _require(all(entry["note"].strip() for entry in reviewer["item_reviews"]))
            if result["method"] == "impact":
                _require(bool(selected) and set(reviewer["scores"]) == set(selected))
                score = calculate_reviewer_score(reviewer["scores"], selected, result["weights"])
                _require(reviewer["reviewer_score"] is not None and abs(score - reviewer["reviewer_score"]) < 1e-12)
        if result["method"] == "impact":
            _require(validate_weights(selected, result["weights"]) == result["weights"])
            average = calculate_committee_average(reviewers, len(reviewers))
            _require(result["overall_avg"] is not None and abs(average - result["overall_avg"]) < 1e-12)
            _require(result["rule_grade"] == rule_grade(average))
    elif result["status"] == "후보제외":
        _require(candidate["is_candidate"] is False and not candidate["uncertainty"] and not reviewers)
        _require(result["final_grade"] is None)
        _require(result["documents"]["expected_files"] == 1)
    else:
        _require(result["final_grade"] is None)
    if opts is not None:
        _require(result["options_fingerprint"] == options_fingerprint(opts))
        _require((result["method"], result["provider"], result["model"]) == (opts.method, opts.provider, opts.model))
        if result["status"] == "완료":
            _require(len(reviewers) == opts.num_reviewers)
            policy = snapshot["부처 세부기준"] or opts.ministry_policy
            eligible = eligible_items(policy)
            for item in result["criteria"]:
                if item["item"] in eligible:
                    _require(item["level"] not in ("판단곤란", "미검토") and not item["needs_expert_review"] and bool(item["note"].strip()))
                    if item["level"] in ("해당", "일부"):
                        _require(bool(item["law_refs"]))
            for reviewer in reviewers:
                for item in reviewer["item_reviews"]:
                    if item["item"] in eligible:
                        _require(item["level"] not in ("판단곤란", "미검토"))
                        if item["level"] in ("해당", "일부"):
                            _require(bool(item["law_refs"]))
            expected = (list(opts.selected_items) if opts.selection_mode == "manual" else
                        [entry["item"] for entry in result["criteria"] if entry["item"] in eligible
                         and entry["level"] in ("해당", "일부", "판단곤란")])
            _require(selected == expected)
            if opts.method == "impact":
                _require(result["weights"] == validate_weights(selected, dict(opts.weights), opts.weight_tolerance))
    if task is not None:
        _require(snapshot == task_snapshot(task))
    return result


def _atomic_write(path, writer):
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix=".v14-", suffix=".tmp", dir=path.parent, delete=False) as stream:
            temporary = Path(stream.name)
        writer(temporary)
        for attempt in range(3):
            try:
                os.replace(temporary, path)
                break
            except PermissionError:
                if attempt == 2:
                    raise
                time.sleep(.05 * (attempt + 1))
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def make_template(path):
    """Create the existing input flow using only clearly synthetic example rows."""
    target = Path(path).resolve()
    book = Workbook()
    try:
        guide = book.active
        guide.title = "작성안내"
        for row in [
            ["연구보안 등급분류 입력 양식"],
            ["분류대상 시트에 과제요약을 입력하세요. 과제명 열은 필수이며 나머지는 선택 입력입니다."],
            ["빈 과제명과 (예시)로 시작하는 과제명은 제외하고, 중복 과제명은 첫 행만 처리합니다."],
            ["추가 검토정보에는 실제 제공한 사실만, 부처 세부기준에는 적용할 기준이 있을 때만 입력하세요."],
            ["AI 에이전트 검토의견은 참고 의견이며 실제 심의위원회 심의 및 최종 등급 확정을 대신하지 않습니다."],
        ]:
            guide.append(row)
        sheet = book.create_sheet("분류대상")
        sheet.append([name + ("(필수)" if name == "과제명" else "(선택)")
                      for name in TASK_COLUMNS + OPTIONAL_TASK_COLUMNS])
        sheet.append(["합성 세부사업 A", "(예시) 합성 과제 A", "시험용 목표", "실제 과제정보가 아닌 합성 내용", "시험용 효과", "", ""])
        sheet.append(["합성 세부사업 B", "(예시) 합성 과제 B", "시험용 목표", "실제 과제정보가 아닌 합성 내용", "시험용 효과", "합성 검토정보", "합성 부처 기준"])
        for worksheet in book:
            worksheet.freeze_panes = "A2"
            for row in worksheet:
                for cell in row:
                    cell.font = Font(name="맑은 고딕", size=10)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="243B53")
                cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        guide.column_dimensions["A"].width = 110
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 32
        sheet.auto_filter.ref = sheet.dimensions
        target.parent.mkdir(parents=True, exist_ok=True)
        _atomic_write(target, lambda temporary: book.save(temporary))
    except Exception:
        raise BatchError("입력 양식을 저장하지 못했습니다. 열린 파일과 저장 위치의 쓰기 권한을 확인하세요.") from None
    finally:
        book.close()
    return str(target)


def _append_checkpoint(path, encoded):
    # Unbuffered binary output lets a failed write/fsync be truncated exactly
    # to the previous valid UTF-8 record boundary, without a later close flush.
    with path.open("ab", buffering=0) as stream:
        previous_size = stream.tell()
        try:
            data = encoded.encode("utf-8")
            if stream.write(data) != len(data):
                raise OSError("Incomplete checkpoint write")
            os.fsync(stream.fileno())
        except Exception:
            try:
                stream.truncate(previous_size)
            except OSError:
                # The caller reports persistence failure; a subsequent resume
                # must validate every line rather than trust a damaged suffix.
                pass
            raise


def _read_tasks(path):
    book = None
    try:
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["분류대상"] if "분류대상" in book.sheetnames else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        columns = [str(value or "").replace("(필수)", "").replace("(선택)", "").strip() for value in header]
        if "과제명" not in columns:
            raise BatchError("입력 엑셀에 '과제명' 열이 필요합니다. 입력 양식을 확인하세요.")
        relevant = TASK_COLUMNS + OPTIONAL_TASK_COLUMNS
        if any(columns.count(name) > 1 for name in relevant):
            raise BatchError("입력 엑셀의 같은 이름을 가진 열이 중복됩니다. 제목 행을 확인하세요.")
        tasks, seen = [], set()
        excluded = {"blank": 0, "example": 0, "duplicate": 0}
        for values in rows:
            row = {name: "" if value is None else str(value).strip() for name, value in zip(columns, values)}
            name = row.get("과제명", "")
            reason = "blank" if not name else "example" if name.startswith("(예시)") else "duplicate" if name in seen else None
            if reason:
                excluded[reason] += 1
                continue
            seen.add(name)
            tasks.append(task_snapshot(row))
        return tasks, excluded
    except BatchError:
        raise
    except Exception:
        raise BatchError("입력 엑셀을 읽지 못했습니다. 파일 형식과 읽기 권한을 확인하세요.") from None
    finally:
        if book is not None:
            book.close()


def _load_checkpoint(path, opts, tasks, warn):
    if not path.exists():
        return {}
    records = {}
    by_id = {task_fingerprint(task): task for task in tasks}

    def unique_object(pairs):
        value = {}
        for key, item in pairs:
            _require(key not in value)
            value[key] = item
        return value

    try:
        with path.open("r", encoding="utf-8") as stream:
            for line in stream:
                # Blank lines and non-finite constants are corruption, not valid records.
                record = json.loads(line, object_pairs_hook=unique_object,
                                    parse_constant=lambda value: (_ for _ in ()).throw(ValueError()))
                _validate_result(record)
                identity = record["task_fingerprint"]
                if record["options_fingerprint"] == options_fingerprint(opts) and identity in by_id:
                    _validate_result(record, opts, by_id[identity])
                    records[identity] = record
        return {identity: record for identity, record in records.items() if record["status"] == "완료"}
    except Exception:
        warn("중간 저장파일이 구버전이거나 손상·불일치 구조를 포함합니다. 이전 완료 평가를 재사용하지 않고 새로 처리합니다.")
        return {}


def _expected_files(result, opts=None):
    if result["reviewers"]:
        return len(result["reviewers"]) + 1
    if result["status"] == "후보제외":
        return 1
    return opts.num_reviewers + 1 if opts is not None else result["documents"]["expected_files"]


def regenerate_documents(saved_result, output_dir):
    """Return a copy with new document status; never call a model or mutate input."""
    _validate_result(saved_result)
    result = deepcopy(saved_result)
    expected = len(result["reviewers"]) + 1
    try:
        documents = export_documents(deepcopy(result), output_dir)
        _validate_documents(documents, result["reviewers"])
        _require(documents["expected_files"] == expected)
        result["documents"] = documents
    except Exception:
        result["documents"] = {"status": "실패", "expected_files": expected, "files": [],
                               "errors": ["HWPX 문서 저장에 실패했습니다. 저장 위치와 템플릿을 확인하세요."]}
    return result


def _document_counts(result):
    documents = result["documents"]
    completed = sum(entry["status"] == "완료" for entry in documents["files"])
    missing = documents["expected_files"] - completed
    failed = missing if documents["status"] in ("일부실패", "실패") else 0
    return documents["expected_files"], completed, failed, missing


def _counts(results):
    counts = {status: 0 for status in STATUSES}
    counts.update(documents_expected=0, documents_completed=0, documents_failed=0, documents_missing=0)
    for result in results:
        counts[result["status"]] += 1
        for key, value in zip(("documents_expected", "documents_completed", "documents_failed", "documents_missing"), _document_counts(result)):
            counts[key] += value
    return counts


def persist_regenerated_documents(report, saved_result, regenerated_result):
    """Persist one document-only update without evaluating any task again.

    Checkpoint, workbook and audit are separate stores. A successful checkpoint
    remains recoverable when Excel is locked. Returned document results describe
    actual generation; per-store status describes what was durably recorded.
    Neither this function nor its report contains credentials or client options.
    """
    def classification(result):
        return {key: value for key, value in result.items() if key != "documents"}

    try:
        _validate_result(saved_result)
        _validate_result(regenerated_result)
        _require(classification(saved_result) == classification(regenerated_result))
        previous = deepcopy(report)
        matches = [index for index, result in enumerate(previous["results"])
                   if result["task_fingerprint"] == saved_result["task_fingerprint"]]
        _require(len(matches) == 1 and classification(previous["results"][matches[0]]) == classification(saved_result))
    except Exception:
        raise BatchError("문서 재생성 전후의 분류 결과가 일치하지 않습니다. 기존 분류 결과와 기록은 보존됩니다.") from None
    result = deepcopy(regenerated_result)
    updated = deepcopy(previous)
    updated["results"][matches[0]] = result
    updated["counts"] = _counts(updated["results"])
    persistence = {"status": "실패", "checkpoint": "미시도", "workbook": "미시도", "audit": "미시도", "warnings": []}

    def warn(message):
        persistence["warnings"].append(message)
        updated.setdefault("warnings", []).append(message)

    def finish():
        states = [persistence[key] for key in ("checkpoint", "workbook", "audit")]
        persistence["status"] = "완료" if all(state == "완료" for state in states) else ("일부실패" if "완료" in states else "실패")
        updated["document_persistence"] = deepcopy(persistence)
        return {"result": result, "report": updated, "persistence": persistence}

    try:
        source, output = Path(previous["input_path"]).resolve(), Path(previous["output_path"]).resolve()
        checkpoint, audit = Path(previous["checkpoint_path"]).resolve(), Path(previous["audit_path"]).resolve()
        _require(source != output and output.is_file() and checkpoint.is_file())
        _require(not os.path.samefile(source, output))
        _require(checkpoint == Path(str(output) + ".checkpoint.jsonl"))
        _require(audit == Path(str(output) + ".audit.jsonl"))
        _require(previous["selection_mode"] in ("auto", "manual"))
        tasks, _ = _read_tasks(source)
        ordered_ids = [task_fingerprint(task) for task in tasks]
        _require(ordered_ids == previous["input_fingerprints"])
        by_id = {task_fingerprint(task): task for task in tasks}
        report_ids = [entry["task_fingerprint"] for entry in previous["results"]]
        _require(len(set(report_ids)) == len(report_ids))
        _require(report_ids == [identity for identity in ordered_ids if identity in set(report_ids)])
        for entry in previous["results"]:
            _validate_result(entry)
            _require(entry["task_snapshot"] == by_id[entry["task_fingerprint"]])
            _require(entry["options_fingerprint"] == previous["options_fingerprint"])
        latest = {}
        def unique_object(pairs):
            value = {}
            for key, item in pairs:
                _require(key not in value)
                value[key] = item
            return value
        with checkpoint.open("r", encoding="utf-8") as stream:
            for line in stream:
                entry = json.loads(line, object_pairs_hook=unique_object,
                                   parse_constant=lambda value: (_ for _ in ()).throw(ValueError()))
                _validate_result(entry)
                _require(entry["options_fingerprint"] == previous["options_fingerprint"])
                latest[entry["task_fingerprint"]] = entry
        _require(set(latest) == set(report_ids))
        for entry in previous["results"]:
            _require(classification(latest[entry["task_fingerprint"]]) == classification(entry))
    except Exception:
        warn("문서는 생성했지만 결과 기록 갱신을 중단했습니다. 원래 입력·실행 설정·중간 저장의 일치 여부를 확인하세요. 기존 기록은 변경하지 않았습니다.")
        return finish()

    try:
        encoded = json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n"
        _append_checkpoint(checkpoint, encoded)
        persistence["checkpoint"] = "완료"
    except Exception:
        persistence["checkpoint"] = "실패"
        warn("문서는 생성했지만 중간 저장 갱신에 실패했습니다. 이전 중간 저장과 결과 엑셀을 보존했습니다.")
    if persistence["checkpoint"] == "완료":
        try:
            _save_workbook(output, updated["results"], SimpleNamespace(selection_mode=previous["selection_mode"]))
            persistence["workbook"] = "완료"
        except Exception:
            persistence["workbook"] = "실패"
            warn("문서 상태는 중간 저장했지만 결과 엑셀 갱신에 실패했습니다. 열린 엑셀을 닫고 다시 저장하세요. 이전 엑셀은 보존되며 중간 저장으로 복구할 수 있습니다.")
    try:
        event = {"schema_version": SCHEMA_VERSION, "run_id": uuid4().hex,
                 "timestamp": datetime.now(timezone.utc).isoformat(), "event": "batch_documents_regenerated",
                 "task_fingerprint": result["task_fingerprint"], "options_fingerprint": result["options_fingerprint"],
                 "status": result["status"], "document_status": result["documents"]["status"],
                 "documents_expected": result["documents"]["expected_files"],
                 "documents_completed": _document_counts(result)[1],
                 "checkpoint_status": persistence["checkpoint"], "workbook_status": persistence["workbook"]}
        _append_checkpoint(audit, json.dumps(event, ensure_ascii=False, allow_nan=False) + "\n")
        persistence["audit"] = "완료"
    except Exception:
        persistence["audit"] = "실패"
        updated["audit_status"] = "실패"
        warn("문서 재생성 감사기록 저장에 실패했습니다. 문서 생성과 중간 저장·엑셀 갱신 상태를 각각 확인하세요.")
    return finish()


def _save_workbook(path, results, opts):
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "분류결과"
        sheet.append(RESULT_COLUMNS)
        for index, result in enumerate(results, 1):
            expected, completed, failed, missing = _document_counts(result)
            rationale = result["rationale"]
            if result["unresolved_reasons"]:
                rationale += ("\n" if rationale else "") + "\n".join(result["unresolved_reasons"])
            sheet.append([index, result["task_name"], result["status"], result["final_grade"], result["confidence"],
                          result["overall_avg"], ", ".join(result["selected_items"]),
                          "직접 선택" if opts.selection_mode == "manual" else "자동 선택", "",
                          rationale, result["discussion_points"], "체크리스트" if result["method"] == "checklist" else "국가영향 지표",
                          result["documents"]["status"], expected, completed, failed, missing,
                          "\n".join(entry["path"] for entry in result["documents"]["files"])])
            if result["overall_avg"] is not None:
                # Keep the underlying numeric value; only adjust its display.
                # Below-band values must never look equal to 3.50 or 4.50.
                display = format_score(result["overall_avg"])
                places = max(2, len(display.split(".", 1)[1]) if "." in display else 0)
                sheet.cell(index + 1, 6).number_format = "0." + "0" * places
        summary = book.create_sheet("요약")
        summary.append(["구분", "건수", "비율"])
        counts = _counts(results)
        summary.append(["전체 과제 수", len(results), 1 if results else 0])
        grade_counts = Counter(result["final_grade"] for result in results if result["status"] == "완료")
        for grade in GRADES:
            count = grade_counts[grade]
            summary.append([grade + "(참고)", count, count / len(results) if results else 0])
        for status in STATUSES:
            summary.append([status, counts[status], counts[status] / len(results) if results else 0])
        for label, key in (("문서 예정 수", "documents_expected"), ("문서 성공 수", "documents_completed"),
                           ("문서 실패 수", "documents_failed"), ("문서 누락 수", "documents_missing")):
            summary.append([label, counts[key]])
        for worksheet in book:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for row in worksheet:
                for cell in row:
                    # Formula-looking user/model text must remain literal text.
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.font = Font(name="맑은 고딕", size=10)
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="243B53")
                cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            for column in worksheet.columns:
                worksheet.column_dimensions[column[0].column_letter].width = 22
        sheet.column_dimensions["J"].width = 65
        sheet.column_dimensions["K"].width = 45
        _atomic_write(path, lambda temporary: book.save(temporary))
    except Exception:
        raise BatchError("결과 엑셀을 저장하지 못했습니다. 열린 파일을 닫고 저장 위치와 쓰기 권한을 확인하세요. 중간 저장 결과는 보존됩니다.") from None
    finally:
        book.close()


def _error_result(task, opts):
    # A bad row-specific item-8 policy must not prevent safe error recording.
    fallback = opts.to_dict()
    fallback.update(selection_mode="auto", selected_items=[], weights={})
    result = make_result(task, fallback)
    result.update(status="오류", options_fingerprint=options_fingerprint(opts),
                  selected_items=list(opts.selected_items), weights=dict(opts.weights) if opts.method == "impact" else {},
                  unresolved_reasons=["과제 분류에 실패했습니다. 과제별 부처 기준, 입력 또는 모델 연결 설정을 확인하세요."])
    return result


def run_excel(in_path, out_path, options, client_factory, *, resume=True, workers=3,
              make_forms=True, progress_cb=None, status_cb=None, cancel_event=None):
    """Process selected rows; return ordered results, counts and persistence status.

    ``client_factory()`` is called once per new task, inside its worker. Options
    are frozen before workers start; factories/clients/keys are never serialized.
    Both callbacks accept one string and may run on workers, so GUI callers must
    queue messages. Cancellation stops new work and preserves in-flight results.
    Global file failures raise BatchError; individual task/document failures do
    not discard successful classifications. Audit errors are reported separately.
    """
    if type(workers) is not int or not 1 <= workers <= 8:
        raise ValueError("동시 처리 수는 1~8 정수여야 합니다.")
    opts = normalize_options(options, allow_row_policy=True)
    source, output = Path(in_path).resolve(), Path(out_path).resolve()
    if source == output or (source.exists() and output.exists() and os.path.samefile(source, output)):
        raise BatchError("결과 저장 위치는 입력 엑셀과 달라야 합니다.")
    tasks, input_counts = _read_tasks(source)
    checkpoint = Path(str(output) + ".checkpoint.jsonl")
    audit = Path(str(output) + ".audit.jsonl")
    run_id = uuid4().hex
    forms = output.with_name(output.stem + "_HWPX") / run_id
    warnings = []

    def notify(callback, message):
        if callback is not None:
            callback(str(message))

    def warn(message):
        warnings.append(message)
        notify(progress_cb, message)
        notify(status_cb, message)

    if any(input_counts.values()):
        warn(f"입력 제외: 빈 과제명 {input_counts['blank']}건, 예시 {input_counts['example']}건, 중복 과제명 {input_counts['duplicate']}건. 중복은 첫 행을 유지합니다.")
    reused = _load_checkpoint(checkpoint, opts, tasks, warn) if resume else {}
    ordered_ids = [task_fingerprint(task) for task in tasks]
    saved = {}
    audit_status = "완료"
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        def initialize(temporary):
            # Keep every validated reusable classification in the initial atomic
            # snapshot, even if cancellation interrupts later document retries.
            with temporary.open("w", encoding="utf-8", newline="\n") as stream:
                for identity in ordered_ids:
                    if identity in reused:
                        stream.write(json.dumps(reused[identity], ensure_ascii=False, allow_nan=False,
                                                separators=(",", ":")) + "\n")
                stream.flush()
                os.fsync(stream.fileno())
        _atomic_write(checkpoint, initialize)
    except Exception:
        raise BatchError("중간 저장파일을 준비하지 못했습니다. 저장 위치와 쓰기 권한을 확인하세요. 기존 파일은 보존됩니다.") from None

    def ordered():
        return [saved[identity] for identity in ordered_ids if identity in saved]

    def persist(result):
        nonlocal audit_status
        encoded = json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n"
        try:
            _append_checkpoint(checkpoint, encoded)
        except Exception:
            raise BatchError("과제 결과의 중간 저장에 실패했습니다. 이전 저장 결과를 보존하고 처리를 중단합니다.") from None
        saved[result["task_fingerprint"]] = result
        if audit_status != "실패":
            try:
                event = {"schema_version": SCHEMA_VERSION, "run_id": run_id,
                         "timestamp": datetime.now(timezone.utc).isoformat(), "event": "batch_task_saved",
                         "task_fingerprint": result["task_fingerprint"], "options_fingerprint": result["options_fingerprint"],
                         "status": result["status"], "document_status": result["documents"]["status"]}
                with audit.open("a", encoding="utf-8", newline="\n") as stream:
                    stream.write(json.dumps(event, ensure_ascii=False, allow_nan=False) + "\n")
                    stream.flush()
                    os.fsync(stream.fileno())
            except Exception:
                audit_status = "실패"
                warn("감사기록 저장 실패: 분류 결과는 중간 저장되었으나 감사기록이 누락되었습니다. 저장 위치와 쓰기 권한을 확인하세요.")
        notify(status_cb, f"처리 {len(saved)}/{len(tasks)}건 · 분류 {result['status']} · 문서 {result['documents']['status']}")

    def documents(result):
        result = deepcopy(result)
        if make_forms:
            docs = result["documents"]
            if docs["status"] != "완료" or any(not Path(entry["path"]).is_file() for entry in docs["files"]):
                result = regenerate_documents(result, forms)
        else:
            result["documents"] = {"status": "미생성", "expected_files": _expected_files(result, opts), "files": [], "errors": []}
        return result

    # Reused results are checkpointed before starting any new model requests.
    for task in tasks:
        identity = task_fingerprint(task)
        if identity in reused:
            persist(documents(reused[identity]))

    def process(task):
        try:
            result = classify_task(task, opts, client_factory(), progress_callback=lambda message: notify(progress_cb, message))
            _validate_result(result, opts, task)
        except Exception:
            result = _error_result(task, opts)
        return documents(result)

    def cancelled():
        return cancel_event is not None and cancel_event.is_set()

    pending = iter(task for task in tasks if task_fingerprint(task) not in reused)
    interval = max(50, len(tasks) // 20)
    notify(progress_cb, f"대상 {len(tasks)}건 · 이전 완료 재사용 {len(reused)}건 · 동시 처리 {workers}건")
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="v14-batch") as pool:
        active = set()

        def fill():
            while len(active) < workers and not cancelled():
                task = next(pending, None)
                if task is None:
                    break
                active.add(pool.submit(process, task))

        fill()
        while active:
            finished, _ = wait(active, return_when=FIRST_COMPLETED)
            for future in finished:
                active.remove(future)
                persist(future.result())
                if len(saved) % interval == 0:
                    _save_workbook(output, ordered(), opts)
            fill()
    results = ordered()
    _save_workbook(output, results, opts)
    status = "취소" if cancelled() else "완료"
    counts = _counts(results)
    notify(status_cb, f"대량 처리 {status} · 정상 {counts['완료']} · 검토필요 {counts['검토필요']} · 후보제외 {counts['후보제외']} · 오류 {counts['오류']} · 문서 누락 {counts['documents_missing']} · 감사기록 {audit_status}")
    return {"status": status, "results": results, "counts": counts, "warnings": warnings,
            "input_counts": input_counts, "task_count": len(tasks), "reused_count": len(reused),
            "output_path": str(output), "checkpoint_path": str(checkpoint), "audit_path": str(audit),
            "audit_status": audit_status, "documents_dir": str(forms),
            "input_path": str(source), "input_fingerprints": ordered_ids,
            "options_fingerprint": options_fingerprint(opts), "selection_mode": opts.selection_mode}
