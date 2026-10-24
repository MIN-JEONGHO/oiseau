# -*- coding: utf-8 -*-
# security_grade_classifier_v5.py
# 연구개발과제 보안등급 분류 시뮬레이터
# 필요 패키지: pip install openai anthropic google-generativeai pandas openpyxl python-docx requests
#
# [v5 변경 사항]
#  - 등급분류 절차 개선: ①법적 기준 체크리스트로 7개 분류항목(부처 재량 제외) 중 해당 항목 선택
#    → ②선택된 항목에 대해서만 '유출 시 국가영향 수준 지표' 채점 → ③선택 항목 평균으로 등급 도출
#    (평균 4.5 이상: 보안과제 / 3.5 이상~4.5 미만: 민감과제 / 3.5 미만: 일반과제 우선 검토)
#  - 위원 구성 선택: 위원 1명(기술전문가) 또는 위원 3명(기술전문가, 연구보안 및 수출통제 전문가,
#    기술사업화 전문가) — 각 위원은 해당 전문가 관점에서 독립적으로 평가
#  - GUI 개선: 과제 입력창 확대, 결과를 보고서 형태로 정리하여 표시
# [v4] Ollama(로컬 LLM) 지원  [v3] 위원용 심의서·자가점검 양식 생성
#
# ★ 중요 안내 ★
# 1) 본 도구의 결과는 '보안등급분류위원회 심의를 위한 참고자료'이며,
#    최종 보안등급은 위원회 심의·의결 및 중앙행정기관 확정 절차를 따라야 합니다.
# 2) 등급 구간(보안≥4.5, 민감≥3.5)은 가이드라인(안)의 '예시' 기준으로 미확정 사항입니다.
# 3) 보안·민감과제 후보의 과제정보를 외부 상용 LLM API로 전송하는 행위 자체가
#    정보 유출 리스크가 될 수 있습니다. 실제 운영 시 내부망 모델 사용,
#    입력정보 최소화, API 제공사의 데이터 미학습 계약 여부 등을 반드시 확인하세요.

import os
import sys
import json
import time
import copy
import datetime
import threading
import traceback

import pandas as pd


def resource_path(rel_path):
    """스크립트로 실행 시/PyInstaller onefile exe로 실행 시 모두 동작하는 리소스 경로 반환."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, rel_path)

# =========================================================
# 1. 설정
# =========================================================
CONFIG = {
    # "openai" | "claude" | "gemini" 중 선택
    "provider": "claude",

    "openai": {
        "api_key": os.getenv("OPENAI_API_KEY", ""),
        "model": "gpt-4o",
    },
    "claude": {
        "api_key": os.getenv("ANTHROPIC_API_KEY", ""),
        "model": "claude-opus-4-8",
    },
    "gemini": {
        "api_key": os.getenv("GEMINI_API_KEY", ""),
        "model": "gemini-2.5-pro",
    },
    # 로컬 LLM (Ollama) — 과제정보가 외부로 나가지 않는 폐쇄망/내부망 운영에 적합
    # 사전 준비: ollama serve 실행 및 모델 다운로드 (예: ollama pull qwen2.5:14b)
    "ollama": {
        "api_key": "",  # 불필요 (로컬 서버)
        "model": os.getenv("OLLAMA_MODEL", "qwen2.5:14b"),
        "base_url": os.getenv("OLLAMA_BASE_URL", "http://localhost:11434"),
        "timeout": 600,  # 로컬 모델은 응답이 느릴 수 있어 여유 있게 설정
    },
    # 로컬 LLM (LM Studio) — LM Studio의 OpenAI 호환 로컬 서버 사용
    # 사전 준비: LM Studio 앱에서 모델 다운로드 후 '로컬 서버 시작'(기본 포트 1234)
    "lmstudio": {
        "api_key": "",  # 불필요 (로컬 서버)
        "model": os.getenv("LMSTUDIO_MODEL", "qwen2.5-14b-instruct"),
        "base_url": os.getenv("LMSTUDIO_BASE_URL", "http://localhost:1234"),
        "timeout": 600,
    },

    # 위원 구성: 1 또는 3 (각 위원은 아래 전문가 관점에서 '독립 호출'로 평가)
    "reviewer_panel": 3,
    "panel_personas": {
        1: ["기술분야 전문가"],
        3: ["기술분야 전문가", "연구보안·수출통제 전문가", "기술사업화 전문가"],
    },

    # 위원 평가 시 temperature (관점 다양성 확보용)
    "reviewer_temperature": 0.7,
    # 일반 판단(1,2,4단계) temperature
    "temperature": 0.2,

    # 등급 구간(가이드라인(안) 판정 가이드 기준)
    # 평균 4.5 이상: 보안과제 / 3.5 이상~4.5 미만: 민감과제 / 3.5 미만: 일반과제 우선 검토
    "thresholds": {"security": 4.5, "sensitive": 3.5},

    # API 재시도 설정
    "max_retries": 3,
    "retry_base_wait": 2.0,   # 지수 백오프: 2, 4, 8초...

    # 감사로그(판단근거 기록·관리) 파일 경로. None이면 기록 안 함.
    "audit_log": "분류_감사로그.jsonl",

    # 대량 분류 시 건별 대기(초) — rate limit 여유
    "batch_sleep": 1.0,
}

# 제공자별 모델 선택 목록(참고용) — GUI 콤보박스는 편집 가능하므로 목록에 없는 모델명도 직접 입력 가능
OPENAI_MODELS = [
    "gpt-4o",
    "gpt-4o-mini",
    "gpt-4.1",
    "gpt-4.1-mini",
    "gpt-4.1-nano",
    "o3",
    "o3-mini",
]
CLAUDE_MODELS = [
    "claude-opus-4-8",
    "claude-sonnet-5",
    "claude-haiku-4-5-20251001",
    "claude-fable-5",
]
GEMINI_MODELS = [
    "gemini-3.1-pro-preview",
    "gemini-3.5-flash",
    "gemini-3-flash-preview",
    "gemini-3.1-flash-lite-preview",
    "gemini-2.5-pro",
    "gemini-2.5-flash",
    "gemini-2.5-flash-lite",
]
OLLAMA_MODELS = [
    "qwen2.5:14b",
    "qwen2.5:32b",
    "llama3.1:8b",
    "llama3.1:70b",
    "gemma2:9b",
    "mistral:7b",
    "deepseek-r1:14b",
]
LMSTUDIO_MODELS = [
    "qwen2.5-14b-instruct",
    "llama-3.1-8b-instruct",
    "llama-3.2-3b-instruct",
    "gemma-2-9b-it",
    "mistral-7b-instruct-v0.3",
    "deepseek-r1-distill-qwen-14b",
]
MODEL_OPTIONS = {
    "openai": OPENAI_MODELS,
    "claude": CLAUDE_MODELS,
    "gemini": GEMINI_MODELS,
    "ollama": OLLAMA_MODELS,     # 실제 사용 가능 모델은 'ollama pull'로 설치한 목록에 따라 다름 — 직접 입력 가능
    "lmstudio": LMSTUDIO_MODELS,  # 실제 사용 가능 모델은 LM Studio에 다운로드한 목록에 따라 다름 — 직접 입력 가능
}

# 분류항목 7개 (부처 재량 항목 제외)
# — ①법적 기준 체크리스트 검토 대상이자 ②선택 시 국가영향 수준 채점 대상
IMPACT_ITEMS = [
    "군사·안보 활용성",
    "기술자립·공급망 전략성",
    "국가핵심기술",
    "전략물자",
    "기술 수준 및 희소성",
    "산업·경제적 파급력",
    "연구수행 관리 필요성",
]

DISCLAIMER = ("※ 본 결과는 보안등급분류위원회 심의를 위한 참고자료이며, "
              "최종 등급은 위원회 의결 및 중앙행정기관 확정에 따릅니다. "
              "(점수 구간은 가이드라인(안) 예시 기준)")

# 위원용 심의서(붙임2 평가양식) 분류항목 8개: (항목명, 분류기준 상세, 근거 법령)
# ※ 8번 항목은 국가영향 점수(IMPACT_ITEMS 7개) 산출 대상이 아닌 정성 검토 항목
FORM_ITEMS = [
    ("1. 군사·안보 활용성",
     "「방위사업법」 제3조제1호에 따른 방위력개선사업과 직접 관련된 과제 / "
     "기술의 활용 방식·적용 맥락상 민간 목적 외 국가안보·방위 분야 활용 가능성",
     "[보안] 영제45조1항1호 / [민감] 영제45조2항1호"),
    ("2. 기술자립·공급망 전략성",
     "외국 기술이전 거부로 국산화 추진 중인 기술 관련 / "
     "공급망 안정성·필수자원 확보·국제통상 대응 등 국가 차원의 전략적 관리 필요성",
     "[보안] 영제45조1항2호다 또는 1항3호 / [민감] 영제45조2항1호"),
    ("3. 국가핵심기술",
     "「산업기술의 유출방지 및 보호에 관한 법률」상 국가핵심기술 관련 "
     "(「국가첨단전략산업법」상 국가첨단전략기술 포함) / "
     "향후 국가핵심기술·첨단전략기술로 발전할 가능성",
     "[보안] 영제45조1항2호다 또는 1항3호 / [민감] 영제45조2항1호"),
    ("4. 전략물자",
     "「대외무역법」상 수출허가 등 제한이 필요한 기술 관련 과제",
     "[보안] 영제45조1항2호다 또는 1항3호 / [민감] 영제45조2항1호"),
    ("5. 기술 수준 및 희소성",
     "세계적 수준의 기술경쟁력 확보 가능성이 높거나 "
     "해외에서 독자적으로 재현·대체하기 어려운 희소 원천기술 포함",
     "[보안] 영제45조1항2호나 또는 1항3호 / [민감] 영제45조2항2호"),
    ("6. 산업·경제적 파급력",
     "미래 신산업·국가 주력산업과 밀접 연계되어 산업 생태계·시장 형성에 중대한 영향 / "
     "기술자립·글로벌 특허·국제표준 선점 기여 가능성",
     "[보안] 영제45조1항2호나 또는 1항3호 / [민감] 영제45조2항2호"),
    ("7. 연구수행 관리 필요성",
     "연구수행 과정 또는 연구성과의 활용·이전 과정에서 추가적인 관리가 필요하다고 인정되는 경우 "
     "(국제공동연구, 데이터·시료·장비 관리 등)",
     "[보안] 영제45조1항3호 / [민감] 영제45조2항3호"),
    ("8. (기타) 부처 정책상 관리 필요성",
     "중앙행정기관의 세부기준에 따라 보호 또는 추가관리가 필요하다고 인정되는 과제",
     "[보안] 영제45조1항3호 / [민감] 영제45조2항3호"),
]

# 연구자용 자가점검 체크리스트(붙임3) 9개 문항
SELF_CHECK_ITEMS = [
    "이 연구가 군사·방위력 개선과 직접적으로 관련되어 있습니까?",
    "이 연구가 해외에서 기술이전이 제한되거나 거부되어 국산화가 필요한 기술과 관련되어 있습니까?",
    "이 연구에 특정 국가·기업 의존도가 높아 공급망 안정성 확보가 중요한 기술이 포함되어 있습니까?",
    "연구 내용 유출 시 특정 산업(반도체, 배터리, 바이오 등)에 중대한 악영향을 미치게 될 것이라고 예상하십니까?",
    "이 연구 내용 유출 시 국제 협력, 외교, 통상 협상 등에 영향을 줄 가능성이 있습니까?",
    "기술패권 경쟁이 첨예한 분야이며 우리나라의 기술력이 세계 최고 수준이어서 적절한 관리가 필요하다고 보십니까?",
    "이 연구 내용이 '첨단전략기술, 국가핵심기술, 전략물자'와 관련성이 높은 기술이라고 생각하십니까?",
    "이 연구 관련 기술이 관계 법령에 의해 별도로 보호·육성되고 있습니까?",
    "이 연구가 사회적·윤리적 영향이 크거나 정부 차원의 관리 필요성이 있는 경우에 해당한다고 보십니까?",
]


# =========================================================
# 2. LLM 어댑터
# =========================================================
class LLMClient:
    def __init__(self, cfg):
        self.provider = cfg["provider"]
        self.cfg = cfg
        self._init_client()

    def _init_client(self):
        p = self.provider
        if p == "openai":
            from openai import OpenAI
            self.client = OpenAI(api_key=self.cfg["openai"]["api_key"])
            self.model = self.cfg["openai"]["model"]
        elif p == "claude":
            import anthropic
            self.client = anthropic.Anthropic(api_key=self.cfg["claude"]["api_key"])
            self.model = self.cfg["claude"]["model"]
        elif p == "gemini":
            import google.generativeai as genai
            genai.configure(api_key=self.cfg["gemini"]["api_key"])
            self.client = genai
            self.model = self.cfg["gemini"]["model"]
        elif p == "ollama":
            import requests
            self.client = requests
            self.model = self.cfg["ollama"]["model"]
            self.base_url = self.cfg["ollama"]["base_url"].rstrip("/")
            self.timeout = self.cfg["ollama"].get("timeout", 600)
            # 서버 연결 및 모델 존재 여부 사전 확인
            try:
                r = requests.get(f"{self.base_url}/api/tags", timeout=5)
                r.raise_for_status()
                names = [m.get("name", "") for m in r.json().get("models", [])]
                if names and not any(n == self.model or n.split(":")[0] == self.model
                                     for n in names):
                    raise RuntimeError(
                        f"Ollama에 '{self.model}' 모델이 없습니다. "
                        f"설치된 모델: {', '.join(names) or '없음'} "
                        f"(설치: ollama pull {self.model})")
            except RuntimeError:
                raise
            except Exception as e:
                raise RuntimeError(
                    f"Ollama 서버({self.base_url})에 연결할 수 없습니다. "
                    f"'ollama serve' 실행 여부를 확인하세요. ({e})")
        elif p == "lmstudio":
            import requests
            self.client = requests
            self.model = self.cfg["lmstudio"]["model"]
            self.base_url = self.cfg["lmstudio"]["base_url"].rstrip("/")
            self.timeout = self.cfg["lmstudio"].get("timeout", 600)
            # 서버 연결 및 모델 로드 여부 사전 확인(OpenAI 호환 /v1/models)
            try:
                r = requests.get(f"{self.base_url}/v1/models", timeout=5)
                r.raise_for_status()
                names = [m.get("id", "") for m in r.json().get("data", [])]
                if names and not any(n == self.model for n in names):
                    raise RuntimeError(
                        f"LM Studio에 '{self.model}' 모델이 로드되어 있지 않습니다. "
                        f"로드된 모델: {', '.join(names) or '없음'} "
                        f"(LM Studio 앱에서 모델을 다운로드·로드한 뒤 다시 시도하세요)")
            except RuntimeError:
                raise
            except Exception as e:
                raise RuntimeError(
                    f"LM Studio 서버({self.base_url})에 연결할 수 없습니다. "
                    f"LM Studio 앱의 '로컬 서버 시작' 여부를 확인하세요. ({e})")
        else:
            raise ValueError(f"지원하지 않는 provider: {p}")

    def _ask_once(self, system_prompt, user_prompt, temperature):
        p = self.provider
        if p == "openai":
            resp = self.client.chat.completions.create(
                model=self.model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                temperature=temperature,
                response_format={"type": "json_object"},
            )
            return resp.choices[0].message.content

        elif p == "claude":
            # assistant 프리필("{")로 JSON-only 응답 유도
            resp = self.client.messages.create(
                model=self.model,
                max_tokens=4000,
                temperature=temperature,
                system=system_prompt,
                messages=[
                    {"role": "user", "content": user_prompt},
                    {"role": "assistant", "content": "{"},
                ],
            )
            return "{" + resp.content[0].text

        elif p == "gemini":
            model = self.client.GenerativeModel(
                self.model,
                system_instruction=system_prompt,
                generation_config={
                    "temperature": temperature,
                    "response_mime_type": "application/json",
                },
            )
            resp = model.generate_content(user_prompt)
            return resp.text

        elif p == "ollama":
            resp = self.client.post(
                f"{self.base_url}/api/chat",
                json={
                    "model": self.model,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                    "stream": False,
                    "format": "json",  # JSON-only 응답 강제
                    "options": {"temperature": temperature},
                },
                timeout=self.timeout,
            )
            resp.raise_for_status()
            return resp.json().get("message", {}).get("content", "")

        elif p == "lmstudio":
            # LM Studio 로컬 서버는 OpenAI 호환 API 형식을 사용
            resp = self.client.post(
                f"{self.base_url}/v1/chat/completions",
                json={
                    "model": self.model,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                    "temperature": temperature,
                    "response_format": {"type": "json_object"},
                },
                timeout=self.timeout,
            )
            resp.raise_for_status()
            return resp.json()["choices"][0]["message"]["content"]

    def ask_json(self, system_prompt, user_prompt, temperature=None):
        """JSON dict를 반환. 파싱 실패·API 오류 시 지수 백오프 재시도."""
        if temperature is None:
            temperature = self.cfg.get("temperature", 0.2)
        retries = self.cfg.get("max_retries", 3)
        base = self.cfg.get("retry_base_wait", 2.0)
        last_err = None
        for attempt in range(retries + 1):
            try:
                raw = self._ask_once(system_prompt, user_prompt, temperature)
                return parse_json(raw)
            except Exception as e:
                last_err = e
                if attempt < retries:
                    time.sleep(base * (2 ** attempt))
        raise RuntimeError(f"LLM 호출/파싱 {retries + 1}회 실패: {last_err}")


def parse_json(text):
    """LLM 응답에서 JSON 객체를 안전하게 추출."""
    if not text:
        raise ValueError("빈 응답")
    text = text.strip()
    # 마크다운 펜스 제거 후 첫 '{' ~ 마지막 '}' 구간 파싱
    text = text.replace("```json", "```")
    if "```" in text:
        parts = [p for p in text.split("```") if "{" in p]
        if parts:
            text = parts[0]
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError(f"JSON 객체를 찾을 수 없음: {text[:200]}")
    return json.loads(text[start:end + 1])


# =========================================================
# 3. 감사로그 (위원회 운영기준: 주요 판단 근거 기록·관리)
# =========================================================
_audit_lock = threading.Lock()

def audit_log(event, payload):
    path = CONFIG.get("audit_log")
    if not path:
        return
    record = {
        "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
        "provider": CONFIG["provider"],
        "model": CONFIG[CONFIG["provider"]]["model"],
        "event": event,
        "payload": payload,
    }
    try:
        with _audit_lock:
            with open(path, "a", encoding="utf-8") as f:
                f.write(json.dumps(record, ensure_ascii=False) + "\n")
    except OSError:
        pass  # 로그 실패가 분류를 중단시키지 않도록


# =========================================================
# 4. 분류 프로세스 (가이드라인(안) 4단계)
# =========================================================
SYSTEM_PROMPT = """당신은 국가연구개발 보안등급분류위원회의 심의를 보조하는 전문가입니다.
「국가연구개발혁신법」 제21조제4항·제5항 및 같은 법 시행령 제45조제1항·제2항에 근거하여
연구개발과제의 보안등급 분류를 위한 참고 의견을 제시합니다.

[등급 정의]
- 보안과제(법 제21조제4항): 외부 유출 시 기술적·재산적 가치에 상당한 손실이 예상되거나
  국가안보를 위하여 별도로 높은 수준의 보안관리가 필요한 연구개발과제
- 민감과제(법 제21조제5항): 보안과제로 분류하지 아니하였으나 유출 시 국가에 미치는
  영향을 고려하여 추가적인 관리가 필요한 연구개발과제
- 일반과제: 위에 해당하지 않고 유출 시 국가에 미치는 영향이 크지 않아
  일반적 연구관리·기본 보안조치 하에서 수행 가능한 과제

[보안과제 분류기준(영 제45조제1항)]
① 방위력개선사업 관련 ② 외국 기술이전 거부로 국산화 추진 중인 기술
③ 보호 필요성이 인정되는 미래핵심기술 ④ 국가핵심기술
⑤ 대외무역법상 수출허가 등 제한 필요 기술 ⑥ 그 밖에 중앙행정기관장이 인정하는 과제

[민감과제 분류기준(영 제45조제2항) — 종합 고려]
① 국가안보·외교 등을 위한 전략적 필요성
② 기술적·경제적 가치 및 관련 산업의 성장 잠재력
③ 그 밖에 연구수행 과정 및 결과에 따른 추가적 관리 필요성

[판단 원칙]
- 기술의 우수성·성공 가능성이 아니라 '유출 시 국가에 미치는 영향과 보호 필요 수준'을 판단
- 단일 요건이 아닌 종합적 판단, 법적 기준 해당 여부와 정의 부합 여부를 함께 검토
- 판단이 불확실하면 불확실하다고 명시

반드시 지정된 JSON 형식으로만 답변하세요. JSON 외의 텍스트를 출력하지 마세요."""


def _task_block(task):
    return (f"과제명: {task.get('과제명', '')}\n"
            f"연구목표: {task.get('연구목표', '')}\n"
            f"연구내용: {task.get('연구내용', '')}\n"
            f"기대효과: {task.get('기대효과', '')}")


def step1_identify(llm, task):
    """1단계: 등급분류 대상(후보군) 식별"""
    prompt = f"""[과제정보]
{_task_block(task)}

[1단계: 등급분류 대상 식별]
다음 후보군 선별 기준에 비추어 이 과제가 보안등급 분류 검토 대상에 해당하는지 판단하세요.
- 혁신법 및 타 법령상 보호 대상 기술(국가전략기술, 국가핵심기술, 방위산업기술 등)
- 연구개발정보·성과의 비공개 필요성이 인정되는 경우
- 국제공동연구 과제(수출통제 기술·중요 연구자산 포함 여부)
- 성능목표·기술사양이 구체화된 응용·개발 단계 과제
- 그 밖에 국가적 중요성이 높은 과제

JSON 형식:
{{"is_candidate": true/false, "reason": "간단한 사유", "uncertainty": "판단이 불확실한 부분(없으면 빈 문자열)"}}"""
    return llm.ask_json(SYSTEM_PROMPT, prompt)


def step2_criteria(llm, task):
    """2단계: 법적 정의·기준 체크리스트 검토 — 7개 분류항목별 해당 여부 판단.
    '해당/일부/판단곤란' 항목이 3단계 국가영향 채점 대상으로 선택됨."""
    checklist = "\n".join(
        f"{i + 1}. {name.split('. ', 1)[-1]}: {crit} ({law})"
        for i, (name, crit, law) in enumerate(FORM_ITEMS[:7]))
    prompt = f"""[과제정보]
{_task_block(task)}

[2단계: 법적 정의·기준 체크리스트 검토]
아래 7개 분류항목 각각에 대해 이 과제의 해당 여부(해당/일부/미해당/판단곤란)와 근거를 제시하세요.
'판단곤란'은 제공된 정보만으로 판단할 수 없는 경우에만 사용하세요.
항목명은 아래에 제시된 명칭을 그대로 사용하세요.

{checklist}

JSON 형식:
{{"criteria": [{{"item": "항목명", "level": "해당|일부|미해당|판단곤란", "note": "근거"}}],
  "summary": "종합 검토의견",
  "needs_expert_review": true/false,
  "expert_review_reason": "전문 기술판정(무역안보관리원·산업기술보호협회 등) 필요 여부와 사유"}}"""
    return llm.ask_json(SYSTEM_PROMPT, prompt)


def select_items_from_criteria(s2):
    """2단계 결과에서 3단계 채점 대상 항목 추출('해당/일부/판단곤란')."""
    selected = []
    for c in s2.get("criteria", []):
        name = str(c.get("item", "")).strip()
        # 항목명 매칭(번호·공백 차이 허용)
        matched = next((it for it in IMPACT_ITEMS
                        if it == name or it in name or name in it), None)
        if matched and c.get("level") in ("해당", "일부", "판단곤란"):
            if matched not in selected:
                selected.append(matched)
    return selected


def _validate_scores(scores_dict, items, n):
    """점수 검증: 선택 항목 기준으로 누락 보정, 1~5 클램핑, 위원 수 맞춤."""
    fixed = {}
    for item in items:
        vals = scores_dict.get(item, [])
        if not isinstance(vals, list):
            vals = [vals]
        clean = []
        for v in vals:
            try:
                v = float(v)
            except (TypeError, ValueError):
                continue
            clean.append(min(5.0, max(1.0, v)))
        if not clean:
            clean = [3.0]  # 정보 부족 시 중간값 — 감사로그로 추적 가능
        clean = (clean + [clean[-1]] * n)[:n]
        fixed[item] = clean
    return fixed


def grade_by_avg(avg):
    """가이드라인(안) 판정 가이드에 따른 점수 기반 등급(우선 검토)."""
    th = CONFIG["thresholds"]
    if avg >= th["security"]:
        return "보안과제"      # 평균 4.5 이상: 보안과제 지정 필요성 우선 검토
    if avg >= th["sensitive"]:
        return "민감과제"      # 평균 3.5 이상~4.5 미만: 민감과제 지정 필요성 우선 검토
    return "일반과제"          # 평균 3.5 미만: 일반과제 우선 검토


def step3_impact(llm, task, selected_items, criteria_notes=None):
    """3단계: '선택된 분류항목에 대해서만' 유출 시 국가영향 수준 지표 채점.
    위원 구성(1명 또는 3명)에 따라 각 위원이 전문가 관점에서 독립 호출로 평가.
    평균은 반드시 파이썬에서 재계산."""
    panel = int(CONFIG.get("reviewer_panel", 3))
    personas = CONFIG["panel_personas"].get(panel, CONFIG["panel_personas"][3])
    n = len(personas)

    scale = """[유출 시 국가영향 수준 척도]
5점: 상당한 국가영향 — 국가 안보·외교·산업 등에 상당한 영향
4점: 높은 국가영향 — 보안과제 수준은 아니나 직접적 파급·민감 영향으로 별도 관리 필요
3점: 제한적 국가영향 — 일부 영향 가능하나 직접적 손실은 제한적
2점: 경미한 영향 — 국가 안보·외교·산업 등에 미치는 영향이 경미한 수준
1점: 영향 미미 — 공개·확산 시 국가 안보·외교·산업 등에 영향이 크지 않은 수준"""

    ref = ""
    if criteria_notes:
        ref = f"\n[2단계 법적 기준 검토 참고]\n{json.dumps(criteria_notes, ensure_ascii=False)}\n"

    scores = {item: [] for item in selected_items}
    notes = []
    for i, persona in enumerate(personas):
        prompt = f"""[과제정보]
{_task_block(task)}

당신은 보안등급분류위원회 위원이며, 다음 전문가 관점에서 독립적으로 평가합니다: {persona}
{ref}
[3단계: 유출 시 국가영향 수준 평가]
{scale}

이 과제는 2단계 법적 기준 체크리스트 검토 결과, 아래 분류항목이 검토 대상으로 선택되었습니다.
※ 반드시 아래 '선택된 항목'에 대해서만 평가하세요. 다른 항목은 평가하지 않습니다.

선택된 평가항목: {", ".join(selected_items)}

각 항목에 대해 1~5점(정수)으로 평가하고 항목별 근거를 간단히 제시하세요.

JSON 형식:
{{"scores": {{"항목명": 점수(정수)}}, "notes": {{"항목명": "근거"}}}}"""
        res = llm.ask_json(SYSTEM_PROMPT, prompt,
                           temperature=CONFIG.get("reviewer_temperature", 0.7))
        one = res.get("scores", {})
        for item in selected_items:
            # 항목명 유연 매칭
            val = one.get(item)
            if val is None:
                for k, v in one.items():
                    if item in str(k) or str(k) in item:
                        val = v
                        break
            scores[item].append(val if val is not None else 3)
        notes.append({"위원": f"위원{i + 1}({persona})", "근거": res.get("notes", {})})

    scores = _validate_scores(scores, selected_items, n)
    # ★ 평균은 LLM이 아닌 파이썬에서 계산 — '선택 항목'의 평균만 사용
    item_avg = {k: round(sum(v) / len(v), 2) for k, v in scores.items()}
    overall_avg = round(sum(item_avg.values()) / len(item_avg), 2) if item_avg else 0.0
    return {"위원구성": [f"위원{i + 1}: {p}" for i, p in enumerate(personas)],
            "선택항목": selected_items,
            "scores": scores, "item_avg": item_avg,
            "overall_avg": overall_avg, "notes": notes}


def step4_decide(llm, task, s2, s3):
    """4단계: 종합 의결(참고 의견)"""
    th = CONFIG["thresholds"]
    avg = s3["overall_avg"]
    rule_grade = grade_by_avg(avg)

    criteria_detail = json.dumps(s2.get("criteria", []), ensure_ascii=False)
    prompt = f"""[과제명] {task.get('과제명', '')}
[2단계 법적기준 체크리스트 결과] {criteria_detail}
[2단계 종합의견] {s2.get('summary', '')}
[3단계 채점 대상(선택 항목)] {", ".join(s3.get('선택항목', []))}
[3단계 선택 항목별 위원회 평균] {json.dumps(s3['item_avg'], ensure_ascii=False)}
[3단계 위원회 종합평균] {avg}
[판정 가이드] 평균 {th['security']} 이상: 보안과제 우선 검토 / {th['sensitive']} 이상~{th['security']} 미만: 민감과제 우선 검토 / {th['sensitive']} 미만: 일반과제 우선 검토
[점수 기반 우선 검토 등급] {rule_grade}

[4단계: 최종 등급 의견 제시]
법적 기준 해당 여부, 선택 항목의 국가영향 수준, 과제 특성을 종합하여 보안등급 의견을 제시하세요.
점수 구간은 우선 검토 참고 도구이며, 법 제21조의 보안·민감과제 '정의' 부합 여부를 함께 확인하세요.
점수 기반 등급과 다른 의견을 제시하는 경우 그 사유를 명확히 하세요.

JSON 형식:
{{"final_grade": "보안과제|민감과제|일반과제",
  "confidence": "높음|보통|낮음",
  "rationale": "의견 사유(법적 근거 조항 포함)",
  "key_factor": "등급 판단에 가장 중요하게 고려한 지표",
  "committee_discussion_points": "위원회 종합 토론 시 중점 논의가 필요한 사항",
  "reclassify_condition": "향후 재분류 검토 조건",
  "measures": "등급별 주요 보안조치 요지"}}"""
    result = llm.ask_json(SYSTEM_PROMPT, prompt)
    result["rule_grade_reference"] = rule_grade
    result["overall_avg"] = avg
    return result


def classify_one(llm, task, selected_override=None):
    """단일 과제 전체 프로세스 실행.
    selected_override: 사용자가 GUI에서 직접 선택한 분류항목 목록(없으면 2단계 결과로 자동 선택)."""
    out = {"과제명": task.get("과제명", ""), "면책문구": DISCLAIMER}
    audit_log("classify_start", {"task": task, "selected_override": selected_override})

    s1 = step1_identify(llm, task)
    out["1_식별"] = s1
    audit_log("step1", {"과제명": out["과제명"], "result": s1})

    if not s1.get("is_candidate", True):
        out["최종등급"] = "일반과제(후보 제외)"
        out["사유"] = s1.get("reason", "")
        audit_log("classify_end", {"과제명": out["과제명"], "최종등급": out["최종등급"]})
        return out

    s2 = step2_criteria(llm, task)
    audit_log("step2", {"과제명": out["과제명"], "result": s2})
    out["2_기준검토"] = s2

    # 채점 대상 항목 결정: 사용자 선택 우선, 없으면 2단계 결과로 자동 선택
    if selected_override:
        selected = [it for it in IMPACT_ITEMS if it in selected_override]
        out["항목선택방식"] = "사용자 직접 선택"
    else:
        selected = select_items_from_criteria(s2)
        out["항목선택방식"] = "2단계 체크리스트 자동 선택"
    out["선택항목"] = selected

    if not selected:
        # 7개 법적 기준 모두 미해당 → 국가영향 채점 생략, 일반과제 우선 검토
        out["최종등급"] = "일반과제"
        out["종합평균"] = "-"
        out["사유"] = ("2단계 법적 정의·기준 체크리스트 검토 결과 7개 분류항목 모두 미해당으로 "
                       "판단되어 국가영향 수준 채점을 생략하고 일반과제 우선 검토로 분류")
        audit_log("classify_end", {"과제명": out["과제명"], "최종등급": out["최종등급"]})
        return out

    criteria_notes = [c for c in s2.get("criteria", [])
                      if any(it in str(c.get("item", "")) for it in selected)]
    s3 = step3_impact(llm, task, selected, criteria_notes)
    audit_log("step3", {"과제명": out["과제명"], "result": s3})
    s4 = step4_decide(llm, task, s2, s3)
    audit_log("step4", {"과제명": out["과제명"], "result": s4})

    out["3_국가영향"] = s3
    out["4_의결의견"] = s4
    out["최종등급"] = s4.get("final_grade", "")
    out["신뢰도"] = s4.get("confidence", "")
    out["종합평균"] = s3["overall_avg"]
    audit_log("classify_end", {"과제명": out["과제명"], "최종등급": out["최종등급"]})
    return out


def format_report(res):
    """분류 결과를 사람이 읽기 좋은 보고서 텍스트로 정리(GUI 출력용)."""
    L = []
    bar = "━" * 46
    L.append(bar)
    L.append(f"  분류 결과: {res.get('과제명', '')}")
    L.append(bar)
    grade = res.get("최종등급", "-")
    avg = res.get("종합평균", "-")
    conf = res.get("신뢰도", "")
    line = f"  ▶ 등급 의견: {grade}"
    if conf:
        line += f"   (신뢰도: {conf})"
    L.append(line)
    if avg != "-":
        th = CONFIG["thresholds"]
        L.append(f"  ▶ 위원회 종합평균: {avg}점"
                 f"  [보안 ≥{th['security']} / 민감 {th['sensitive']}~{th['security']} / 일반 <{th['sensitive']}]")
    L.append("")

    s1 = res.get("1_식별", {})
    if s1:
        mark = "해당" if s1.get("is_candidate", True) else "미해당"
        L.append(f"[1단계] 분류대상 식별: {mark} — {s1.get('reason', '')}")

    s2 = res.get("2_기준검토", {})
    if s2:
        L.append(f"[2단계] 법적 정의·기준 체크리스트 ({res.get('항목선택방식', '')})")
        selected = res.get("선택항목", [])
        for c in s2.get("criteria", []):
            name = str(c.get("item", ""))
            chk = "■" if any(it in name or name in it for it in selected) else "□"
            L.append(f"  {chk} {name:<16s} [{c.get('level', '-')}] {c.get('note', '')}")
        if s2.get("needs_expert_review"):
            L.append(f"  ※ 전문 기술판정 검토 필요: {s2.get('expert_review_reason', '')}")

    s3 = res.get("3_국가영향", {})
    if s3:
        L.append("")
        L.append(f"[3단계] 유출 시 국가영향 채점 — 선택 항목만 평가 "
                 f"({' / '.join(s3.get('위원구성', []))})")
        for item in s3.get("선택항목", []):
            sc = s3.get("scores", {}).get(item, [])
            iavg = s3.get("item_avg", {}).get(item, "-")
            sc_txt = ", ".join(str(int(x)) if float(x).is_integer() else str(x) for x in sc)
            L.append(f"  · {item:<16s} 평균 {iavg}점  (위원별: {sc_txt})")
        L.append(f"  ⇒ 종합평균 {s3.get('overall_avg', '-')}점 → "
                 f"{grade_by_avg(s3.get('overall_avg', 0))} 지정 필요성 우선 검토")

    s4 = res.get("4_의결의견", {})
    L.append("")
    if s4:
        L.append("[4단계] 의결 의견 (AI 참고)")
        L.append(f"  · 사유: {s4.get('rationale', '')}")
        L.append(f"  · 핵심 고려 지표: {s4.get('key_factor', '')}")
        L.append(f"  · 위원회 논의 필요사항: {s4.get('committee_discussion_points', '')}")
        L.append(f"  · 재분류 검토 조건: {s4.get('reclassify_condition', '')}")
        L.append(f"  · 보안조치 요지: {s4.get('measures', '')}")
    elif res.get("사유"):
        L.append(f"[의견] {res.get('사유', '')}")
    L.append("")
    L.append(DISCLAIMER)
    L.append(bar)
    return "\n".join(L)




# =========================================================
# 5. 엑셀 대량 처리 (건별 저장 + 재개 기능)
# =========================================================
RESULT_COLUMNS = ["과제명", "상태", "최종등급(참고)", "신뢰도", "종합평균",
                  "채점항목(선택)", "항목선택방식", "핵심고려지표", "의결사유",
                  "위원회논의필요사항", "재분류조건", "면책문구", "상세(JSON)"]


def make_template(path="분류대상_양식.xlsx"):
    df = pd.DataFrame(columns=["과제명", "연구목표", "연구내용", "기대효과"])
    df.loc[0] = [
        "1.5kW급 태양전지용 GaN 전력반도체 기술개발",
        "고효율 GaN 전력반도체 소자·모듈 개발",
        "태양광 인버터용 화합물반도체 소자 설계·공정·실증",
        "전력변환 효율 향상 및 소자 국산화",
    ]
    df.to_excel(path, index=False)
    print(f"양식 생성 완료: {path}")
    return path


def _row_to_result(res):
    s4 = res.get("4_의결의견", {})
    return {
        "과제명": res.get("과제명", ""),
        "상태": "완료",
        "최종등급(참고)": res.get("최종등급", ""),
        "신뢰도": res.get("신뢰도", ""),
        "종합평균": res.get("종합평균", ""),
        "채점항목(선택)": ", ".join(res.get("선택항목", [])),
        "항목선택방식": res.get("항목선택방식", ""),
        "핵심고려지표": s4.get("key_factor", ""),
        "의결사유": s4.get("rationale", res.get("사유", "")),
        "위원회논의필요사항": s4.get("committee_discussion_points", ""),
        "재분류조건": s4.get("reclassify_condition", ""),
        "면책문구": DISCLAIMER,
        "상세(JSON)": json.dumps(res, ensure_ascii=False),
    }


def run_excel(in_path, out_path="분류결과.xlsx", resume=True,
              progress_cb=None, make_forms=False):
    """엑셀 대량 분류. 건별 즉시 저장, 기존 결과가 있으면 완료 건은 건너뜀.
    make_forms=True 시 과제별 위원용 심의서(docx)를 '<결과파일명>_심의서' 폴더에 생성."""
    df = pd.read_excel(in_path).fillna("")
    if "과제명" not in df.columns:
        raise ValueError("입력 엑셀에 '과제명' 컬럼이 필요합니다. '엑셀 양식 생성'을 이용하세요.")

    forms_dir = None
    if make_forms:
        forms_dir = os.path.splitext(out_path)[0] + "_심의서"
        os.makedirs(forms_dir, exist_ok=True)

    done = set()
    rows = []
    if resume and os.path.exists(out_path):
        try:
            prev = pd.read_excel(out_path).fillna("")
            for _, r in prev.iterrows():
                d = r.to_dict()
                if d.get("상태") == "완료":
                    done.add(str(d.get("과제명", "")))
                    rows.append(d)
        except Exception:
            pass  # 기존 파일 손상 시 새로 시작

    llm = LLMClient(CONFIG)
    total = len(df)
    for i, r in df.iterrows():
        task = {k: str(v).strip() for k, v in r.to_dict().items()}
        name = task.get("과제명", "")
        msg = f"[{i + 1}/{total}] {name}"
        if not name:
            continue
        if name in done:
            if progress_cb:
                progress_cb(msg + " → 기존 결과 존재, 건너뜀")
            continue
        if progress_cb:
            progress_cb(msg + " 분류 중...")
        try:
            res = classify_one(llm, task)
            rows.append(_row_to_result(res))
            if forms_dir:
                try:
                    fpath = os.path.join(forms_dir, _safe_filename(name) + "_심의서.docx")
                    export_review_form_docx(task, res, fpath)
                    if progress_cb:
                        progress_cb(f"    심의서 생성: {fpath}")
                except Exception as fe:
                    if progress_cb:
                        progress_cb(f"    [심의서 생성 실패] {fe}")
        except Exception as e:
            audit_log("error", {"과제명": name, "error": str(e),
                                "trace": traceback.format_exc()})
            rows.append({"과제명": name, "상태": "오류",
                         "의결사유": str(e), "면책문구": DISCLAIMER})
        # ★ 건별 즉시 저장 — 중단 시에도 결과 보존
        pd.DataFrame(rows).reindex(columns=RESULT_COLUMNS).to_excel(out_path, index=False)
        time.sleep(CONFIG.get("batch_sleep", 1.0))

    if progress_cb:
        progress_cb(f"완료: {out_path} 저장됨 (오류 건은 '상태' 컬럼 확인 후 재실행 시 자동 재시도)")
    return out_path


# =========================================================
# 6. 위원용 심의서(붙임2 양식 기반) Word 생성
# =========================================================
def _safe_filename(name, maxlen=60):
    bad = '\\/:*?"<>|\n\r\t'
    out = "".join(c for c in str(name) if c not in bad).strip()
    return (out or "무제")[:maxlen]


def _require_docx():
    try:
        from docx import Document
        return Document
    except ImportError:
        raise RuntimeError("python-docx가 설치되어 있지 않습니다: pip install python-docx")


def _impact_note_text(res, item_name):
    """3단계 결과에서 항목별 근거 텍스트 추출(단일/독립 위원 모드 모두 지원)."""
    s3 = res.get("3_국가영향", {})
    notes = s3.get("notes", {})
    if isinstance(notes, dict):
        return str(notes.get(item_name, ""))
    if isinstance(notes, list):  # 독립 위원 모드: 위원별 근거 일부 요약
        parts = []
        for rev in notes:
            reason = rev.get("근거", {})
            if isinstance(reason, dict) and reason.get(item_name):
                parts.append(f"위원{rev.get('위원', '')}: {reason[item_name]}")
        return " / ".join(parts[:3])
    return ""


def export_review_form_docx(task, res, path):
    """AI 사전검토 결과가 참고정보로 포함된 위원용 심의서(붙임2 양식 기반) 생성.
    '위원 판단'란과 서명란은 공란으로 두어 위원이 직접 기재."""
    Document = _require_docx()
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("연구개발과제 보안등급 심의서 (위원용)", level=1)
    doc.add_paragraph(
        "※ 본 심의는 유출 시 국가영향에 따른 보호수준을 결정하기 위한 절차이며, "
        "기술의 우수성 등을 평가하는 과제평가와는 다른 성격임")
    doc.add_paragraph(
        "※ 'AI 사전검토' 열은 심의 보조용 참고정보이며, '위원 판단'란은 위원이 직접 기재합니다. "
        + DISCLAIMER)

    # 과제 기본 정보
    doc.add_heading("과제 기본 정보", level=2)
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    rows_info = [
        ("세부사업명", task.get("세부사업명", "")),
        ("과제명", task.get("과제명", "")),
        ("연구목표", task.get("연구목표", "")),
        ("연구내용 요약", (task.get("연구내용", "") or "")[:300]),
    ]
    for i, (k, v) in enumerate(rows_info):
        t.cell(i, 0).text = k
        t.cell(i, 1).text = str(v)

    # AI 사전검토 요약
    s4 = res.get("4_의결의견", {})
    doc.add_heading("AI 사전검토 요약 (참고)", level=2)
    doc.add_paragraph(
        f"등급 의견: {res.get('최종등급', '-')}  |  "
        f"국가영향 종합평균: {res.get('종합평균', '-')}  |  "
        f"신뢰도: {res.get('신뢰도', '-')}")
    doc.add_paragraph(f"핵심 고려 지표: {s4.get('key_factor', '-')}")
    doc.add_paragraph(f"의견 사유: {s4.get('rationale', res.get('사유', '-'))}")

    # 분류항목별 검토표
    doc.add_heading("분류항목별 검토", level=2)
    item_avg = res.get("3_국가영향", {}).get("item_avg", {})
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for j, h in enumerate(["분류항목", "보안등급 분류 기준 상세",
                           "AI 사전검토 (참고)", "위원 판단\n(⑤~① 기재)", "근거 법령"]):
        hdr[j].text = h
    for name, criteria, law in FORM_ITEMS:
        row = tbl.add_row().cells
        row[0].text = name
        row[1].text = criteria
        plain = name.split(". ", 1)[-1].replace("(기타) ", "")
        selected = res.get("선택항목", list(item_avg.keys()))
        avg = item_avg.get(plain)
        if avg is not None:
            ai_txt = f"평균 {avg}점"
            note = _impact_note_text(res, plain)
            if note:
                ai_txt += f"\n{note}"
        elif plain in IMPACT_ITEMS and plain not in selected:
            ai_txt = "체크리스트 미해당 → 미채점"
        else:
            ai_txt = "정성 검토 항목"
        row[2].text = ai_txt
        row[3].text = "□ ⑤  □ ④  □ ③  □ ②  □ ①"
        row[4].text = law

    # 보안등급 결정
    doc.add_heading("보안등급 결정 (위원 기재)", level=2)
    ai_grade = res.get("최종등급", "")
    grade_rows = [
        ("보안과제", "국가안보 위한 별도 관리 필요 → 강력한 통제 및 비공개 관리"),
        ("민감과제", "국가전략에 잠재적 영향 줄 수 있어 추가 관리 필요 → 기관·연구자 단위 관리하되 개방 가능"),
        ("일반과제", "유출 시 국가안보 및 전략 등 영향 제한적 → 성과공개·확산 원칙"),
    ]
    gt = doc.add_table(rows=len(grade_rows), cols=3)
    gt.style = "Table Grid"
    for i, (g, desc) in enumerate(grade_rows):
        gt.cell(i, 0).text = g
        gt.cell(i, 1).text = "□"
        mark = "  ← AI 참고의견" if g and g in ai_grade else ""
        gt.cell(i, 2).text = desc + mark

    # 종합 의견 서술
    doc.add_heading("종합 의견 서술", level=2)
    doc.add_paragraph("등급 결정에 가장 중요하게 고려한 지표 및 이유:")
    doc.add_paragraph(f"(AI 참고: {s4.get('key_factor', '-')})")
    doc.add_paragraph("\n\n")
    doc.add_paragraph("위원회 종합 토론 시 중점 논의가 필요한 사항:")
    doc.add_paragraph(f"(AI 참고: {s4.get('committee_discussion_points', '-')})")
    doc.add_paragraph("\n\n")

    # 제출 확인
    doc.add_heading("제출 확인", level=2)
    st = doc.add_table(rows=3, cols=2)
    st.style = "Table Grid"
    st.cell(0, 0).text = "제출일시"
    st.cell(0, 1).text = "20    년     월     일"
    st.cell(1, 0).text = "평가위원 서명"
    st.cell(1, 1).text = "(서명)"
    st.cell(2, 0).text = "간사 확인"
    st.cell(2, 1).text = "(서명)"

    doc.save(path)
    return path


# =========================================================
# 7. 연구자용 자가점검 체크리스트(붙임3 기반)
# =========================================================
def selfcheck_judge(answers):
    """붙임3 판정 규칙: 하나라도 '예' 또는 '불확실'이면 보안·민감과제 후보 검토 필요."""
    if any(a == "예" for a in answers):
        return "보안·민감과제 후보로 검토 필요"
    if any(a == "불확실" for a in answers):
        return "보안·민감과제 후보로 검토 필요 (불확실 항목은 기관 연구보안 담당자 검토)"
    return "일반과제 후보 (모든 항목 '아니오')"


def selfcheck_auto(llm, task):
    """AI가 과제정보를 바탕으로 자가점검 9개 문항을 사전 점검(참고용)."""
    items = "\n".join(f"{i + 1}. {q}" for i, q in enumerate(SELF_CHECK_ITEMS))
    prompt = f"""[과제정보]
{_task_block(task)}

[연구자용 보안등급 자가점검 체크리스트 사전 점검]
아래 9개 문항 각각에 대해 '예/아니오/불확실' 중 하나로 답하고 근거를 제시하세요.
제공된 정보만으로 판단할 수 없으면 반드시 '불확실'로 답하세요.

{items}

JSON 형식:
{{"answers": [{{"no": 문항번호, "answer": "예|아니오|불확실", "reason": "근거"}}],
  "summary": "종합 의견"}}"""
    res = llm.ask_json(SYSTEM_PROMPT, prompt)
    answers = ["불확실"] * len(SELF_CHECK_ITEMS)
    reasons = [""] * len(SELF_CHECK_ITEMS)
    for a in res.get("answers", []):
        try:
            idx = int(a.get("no", 0)) - 1
        except (TypeError, ValueError):
            continue
        if 0 <= idx < len(SELF_CHECK_ITEMS):
            ans = str(a.get("answer", "불확실"))
            answers[idx] = ans if ans in ("예", "아니오", "불확실") else "불확실"
            reasons[idx] = str(a.get("reason", ""))
    result = {"answers": answers, "reasons": reasons,
              "summary": res.get("summary", ""),
              "judgment": selfcheck_judge(answers)}
    audit_log("selfcheck_auto", {"과제명": task.get("과제명", ""), "result": result})
    return result


def export_selfcheck_docx(task, answers, path, reasons=None, researcher=""):
    """붙임3 양식 기반 자가점검서(연구자용) Word 생성."""
    Document = _require_docx()
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("연구개발과제 보안등급 자가점검서 (연구자용)", level=1)
    doc.add_paragraph("※ (작성시기) ①연구개발기관 선정 단계 및 협약시기, "
                      "②보안등급 재분류를 원하는 경우, ③과제 최종평가 시기")
    doc.add_paragraph("※ 본 점검서는 보안등급 분류 또는 재분류 검토를 위한 참고자료이며, "
                      "보안등급은 보안등급분류위원회를 통해 최종 결정됨")
    doc.add_paragraph(f"과제명: {task.get('과제명', '')}")

    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for j, h in enumerate(["연번", "보안·민감과제 후보 검토 기준", "예", "아니오", "불확실"]):
        hdr[j].text = h
    for i, q in enumerate(SELF_CHECK_ITEMS):
        row = tbl.add_row().cells
        row[0].text = str(i + 1)
        qtext = q
        if reasons and i < len(reasons) and reasons[i]:
            qtext += f"\n(AI 참고 근거: {reasons[i]})"
        row[1].text = qtext
        ans = answers[i] if i < len(answers) else ""
        row[2].text = "■" if ans == "예" else "□"
        row[3].text = "■" if ans == "아니오" else "□"
        row[4].text = "■" if ans == "불확실" else "□"

    doc.add_paragraph(f"\n최종 확인 결과: {selfcheck_judge(answers)}")
    doc.add_paragraph("※ 상기 항목 중 하나라도 해당 또는 불확실한 경우, 보안·민감과제 후보로 검토될 수 있음")
    doc.add_paragraph("※ '불확실' 체크 사항은 기관 내 연구보안 담당자(또는 심의회의) 등에서 검토")
    doc.add_paragraph("\n본인은 위 모든 사항을 성실하게 검토하여 작성하였음을 확인합니다.")
    doc.add_paragraph("20    년     월     일")
    doc.add_paragraph(f"연구책임자(심사자)  {researcher}          (인)")
    doc.save(path)
    return path


# =========================================================
# 8. GUI (스레드 처리로 화면 멈춤 방지)
# =========================================================
def launch_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("연구개발과제 보안등급 분류 시뮬레이터 (위원회 심의 보조용) v5")
    root.geometry("900x920")
    try:
        root.iconbitmap(resource_path("book_icon.ico"))
    except Exception:
        pass

    busy = threading.Event()

    def log(msg):
        out_box.insert("end", msg + "\n")
        out_box.see("end")

    def log_threadsafe(msg):
        root.after(0, log, msg)

    # --- 안내 문구 ---
    tk.Label(root, text=DISCLAIMER, fg="darkred", wraplength=860,
             justify="left").pack(fill="x", padx=10, pady=(8, 0))
    tk.Label(root, text="※ 외부 API(openai/claude/gemini) 사용 시 과제정보가 외부로 전송됩니다. "
                        "비공개 정보를 다루는 경우 로컬 LLM(ollama/LM Studio) 사용을 권장합니다.",
             fg="darkred", wraplength=860, justify="left").pack(fill="x", padx=10)

    # --- API 설정 ---
    fr_api = ttk.LabelFrame(root, text="① API 설정")
    fr_api.pack(fill="x", padx=10, pady=6)

    ttk.Label(fr_api, text="LLM 제공자").grid(row=0, column=0, padx=5, pady=4)
    provider_var = tk.StringVar(value=CONFIG["provider"])
    cb = ttk.Combobox(fr_api, textvariable=provider_var,
                      values=["openai", "claude", "gemini", "ollama", "lmstudio"],
                      state="readonly", width=12)
    cb.grid(row=0, column=1, padx=5)

    ttk.Label(fr_api, text="API Key").grid(row=0, column=2, padx=5)
    key_var = tk.StringVar()
    key_entry = ttk.Entry(fr_api, textvariable=key_var, width=34, show="*")
    key_entry.grid(row=0, column=3, padx=5)

    # 모델명: 제공자별 추천 목록에서 선택하거나 직접 입력(편집 가능한 콤보박스)
    ttk.Label(fr_api, text="모델명").grid(row=1, column=0, padx=5, pady=4)
    model_var = tk.StringVar(value=CONFIG[CONFIG["provider"]]["model"])
    model_combo = ttk.Combobox(
        fr_api, textvariable=model_var,
        values=MODEL_OPTIONS.get(CONFIG["provider"], []),
        state="normal", width=28)
    model_combo.grid(row=1, column=1, columnspan=2, sticky="w")

    # 위원 구성 선택 (각 위원은 전문가 관점에서 독립 평가)
    panel_var = tk.IntVar(value=CONFIG.get("reviewer_panel", 3))
    fr_panel = ttk.LabelFrame(fr_api, text="위원 구성")
    fr_panel.grid(row=1, column=3, rowspan=2, padx=5, pady=2, sticky="w")
    ttk.Radiobutton(fr_panel, text="위원 1명 (기술전문가)",
                    value=1, variable=panel_var).pack(anchor="w")
    ttk.Radiobutton(fr_panel, text="위원 3명 (기술전문가, 연구보안 및\n수출통제 전문가, 기술사업화 전문가)",
                    value=3, variable=panel_var).pack(anchor="w")

    LOCAL_PROVIDERS = ("ollama", "lmstudio")

    ttk.Label(fr_api, text="서버 주소\n(로컬 전용)").grid(row=2, column=0, padx=5, pady=4)
    url_var = tk.StringVar(value=CONFIG["ollama"]["base_url"])
    url_entry = ttk.Entry(fr_api, textvariable=url_var, width=30)
    url_entry.grid(row=2, column=1, columnspan=2, sticky="w")
    local_lbl = ttk.Label(fr_api, text="", foreground="green")
    local_lbl.grid(row=3, column=1, columnspan=2, sticky="w")

    def on_provider_change(_=None):
        p = provider_var.get()
        model_var.set(CONFIG[p]["model"])
        model_combo["values"] = MODEL_OPTIONS.get(p, [])
        key_var.set(CONFIG[p].get("api_key", ""))
        if p in LOCAL_PROVIDERS:
            url_var.set(CONFIG[p]["base_url"])
            key_entry.config(state="disabled")
            url_entry.config(state="normal")
            local_lbl.config(text="✓ 로컬 실행 — 과제정보 외부 미전송")
        else:
            key_entry.config(state="normal")
            url_entry.config(state="disabled")
            local_lbl.config(text="")
    cb.bind("<<ComboboxSelected>>", on_provider_change)
    on_provider_change()

    def apply_config():
        p = provider_var.get()
        CONFIG["provider"] = p
        if key_var.get().strip():
            CONFIG[p]["api_key"] = key_var.get().strip()
        CONFIG[p]["model"] = model_var.get().strip()
        CONFIG["reviewer_panel"] = panel_var.get()
        if p in LOCAL_PROVIDERS:
            if url_var.get().strip():
                CONFIG[p]["base_url"] = url_var.get().strip()
        elif not CONFIG[p]["api_key"]:
            messagebox.showwarning("설정", "API Key가 비어 있습니다. 환경변수 또는 입력란을 확인하세요.")

    # --- 단건 입력 (입력창 확대) ---
    fr_one = ttk.LabelFrame(root, text="② 단건 과제 입력")
    fr_one.pack(fill="x", padx=10, pady=6)
    fields = {}
    heights = {"과제명": 2, "연구목표": 3, "연구내용": 5, "기대효과": 3}
    for i, name in enumerate(["과제명", "연구목표", "연구내용", "기대효과"]):
        ttk.Label(fr_one, text=name).grid(row=i, column=0, sticky="ne", padx=5, pady=3)
        t = tk.Text(fr_one, width=88, height=heights[name], wrap="word",
                    font=("맑은 고딕", 10))
        t.grid(row=i, column=1, padx=5, pady=3, sticky="we")
        fields[name] = t
    fr_one.columnconfigure(1, weight=1)

    # --- 분류항목 선택 (법적 정의·기준 체크리스트, 7개 항목) ---
    fr_items = ttk.LabelFrame(
        root, text="③ 분류항목 선택 — 선택된 항목에 대해서만 '유출 시 국가영향 수준' 채점")
    fr_items.pack(fill="x", padx=10, pady=6)
    auto_var = tk.BooleanVar(value=True)
    item_vars = {}

    def toggle_items():
        state = "disabled" if auto_var.get() else "normal"
        for cbx in item_checks:
            cbx.config(state=state)

    ttk.Checkbutton(
        fr_items, text="AI 자동 선택 (2단계 법적 기준 체크리스트 검토 결과로 자동 결정)",
        variable=auto_var, command=lambda: toggle_items()
    ).grid(row=0, column=0, columnspan=4, sticky="w", padx=5, pady=(2, 4))

    item_checks = []
    for i, item in enumerate(IMPACT_ITEMS):
        var = tk.BooleanVar(value=False)
        item_vars[item] = var
        cbx = ttk.Checkbutton(fr_items, text=f"{i + 1}. {item}", variable=var)
        cbx.grid(row=1 + i // 4, column=i % 4, sticky="w", padx=8, pady=1)
        item_checks.append(cbx)
    toggle_items()

    # --- 결과 출력 ---
    # ※ fr_out은 expand=True로 남는 공간을 채우는 영역이라 마지막에 pack해야 함.
    #   먼저 pack하면 버튼 행·크레딧이 창 높이를 넘어가 화면 밖으로 밀려나 보이지 않게 됨.
    fr_out = ttk.LabelFrame(root, text="④ 분류 결과")
    out_box = scrolledtext.ScrolledText(fr_out, width=96, height=16, wrap="word",
                                        font=("맑은 고딕", 10))
    out_box.pack(fill="both", expand=True, padx=4, pady=4)
    out_box.tag_configure("title", font=("맑은 고딕", 10, "bold"))
    json_var = tk.BooleanVar(value=False)

    # 하단 고정 요소(크레딧 → 버튼 행 순으로 pack)를 결과창보다 먼저 배치해 항상 보이게 함
    tk.Label(root, text="Powered by MIN", fg="gray",
             font=("", 9)).pack(side="bottom", pady=(0, 6))

    fr_btn = ttk.Frame(root)
    fr_btn.pack(side="bottom", fill="x", padx=10, pady=(0, 8))
    btn_single = ttk.Button(fr_btn, text="단건 분류 실행")
    btn_tmpl = ttk.Button(fr_btn, text="엑셀 양식 생성", command=lambda: (make_template(), log("양식 생성: 분류대상_양식.xlsx")))
    btn_batch = ttk.Button(fr_btn, text="엑셀 대량 분류")
    btn_form = ttk.Button(fr_btn, text="위원용 심의서 저장", state="disabled")
    btn_self = ttk.Button(fr_btn, text="자가점검(연구자용)")
    for b in (btn_single, btn_tmpl, btn_batch, btn_form, btn_self):
        b.pack(side="left", padx=5)
    forms_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(fr_btn, text="대량 분류 시 심의서 생성",
                    variable=forms_var).pack(side="left", padx=5)
    ttk.Checkbutton(fr_btn, text="상세 JSON 표시",
                    variable=json_var).pack(side="left", padx=5)

    fr_out.pack(fill="both", expand=True, padx=10, pady=6)

    # 최근 단건 분류 결과 보관 (심의서 생성용)
    last = {"task": None, "res": None}

    def set_busy(state):
        busy.set() if state else busy.clear()
        s = "disabled" if state else "normal"
        btn_single.config(state=s)
        btn_batch.config(state=s)
        btn_self.config(state=s)
        btn_form.config(state="disabled" if state or not last["res"] else "normal")

    def run_single():
        if busy.is_set():
            return
        apply_config()
        task = {k: v.get("1.0", "end").strip() for k, v in fields.items()}
        if not task["과제명"]:
            messagebox.showwarning("입력", "과제명을 입력하세요.")
            return
        # 분류항목 선택 방식 결정
        if auto_var.get():
            selected_override = None
        else:
            selected_override = [it for it, v in item_vars.items() if v.get()]
            if not selected_override:
                messagebox.showwarning(
                    "분류항목", "채점할 분류항목을 1개 이상 선택하거나 'AI 자동 선택'을 사용하세요.")
                return
        set_busy(True)
        panel = panel_var.get()
        log(f"\n=== {task['과제명']} 분류 시작 (위원 {panel}명"
            f"{', 항목 자동 선택' if selected_override is None else ', 항목 직접 선택'}) ===")

        def worker():
            try:
                llm = LLMClient(CONFIG)
                res = classify_one(llm, task, selected_override=selected_override)
                last["task"], last["res"] = task, res
                log_threadsafe(format_report(res))
                if json_var.get():
                    log_threadsafe("[상세 JSON]")
                    log_threadsafe(json.dumps(res, ensure_ascii=False, indent=2))
                log_threadsafe("→ '위원용 심의서 저장' 버튼으로 붙임2 양식(docx)을 생성할 수 있습니다.")

                # AI 자동 선택 결과를 체크박스에 반영(참고 표시)
                def reflect():
                    for it, v in item_vars.items():
                        v.set(it in res.get("선택항목", []))
                root.after(0, reflect)
            except Exception as e:
                log_threadsafe(f"[오류] {e}")
            finally:
                root.after(0, set_busy, False)
        threading.Thread(target=worker, daemon=True).start()

    def save_review_form():
        if not last["res"]:
            messagebox.showwarning("심의서", "먼저 단건 분류를 실행하세요.")
            return
        default = _safe_filename(last["task"].get("과제명", "심의서")) + "_심의서.docx"
        path = filedialog.asksaveasfilename(
            defaultextension=".docx", initialfile=default,
            filetypes=[("Word 문서", "*.docx")])
        if not path:
            return
        try:
            export_review_form_docx(last["task"], last["res"], path)
            log(f"위원용 심의서 저장 완료: {path}")
        except Exception as e:
            messagebox.showerror("심의서", str(e))

    # --- 자가점검(연구자용) 창 ---
    def open_selfcheck():
        task = {k: v.get("1.0", "end").strip() for k, v in fields.items()}
        win = tk.Toplevel(root)
        win.title("연구개발과제 보안등급 자가점검 (연구자용, 붙임3 기반)")
        win.geometry("880x640")
        tk.Label(win, text="※ 본 점검서는 참고자료이며, 보안등급은 보안등급분류위원회를 통해 최종 결정됩니다.",
                 fg="darkred", wraplength=840, justify="left").pack(fill="x", padx=10, pady=(8, 4))

        frame = ttk.Frame(win)
        frame.pack(fill="both", expand=True, padx=10)
        ans_vars = []
        for i, q in enumerate(SELF_CHECK_ITEMS):
            row = ttk.Frame(frame)
            row.pack(fill="x", pady=2)
            tk.Label(row, text=f"{i + 1}. {q}", wraplength=580,
                     justify="left", anchor="w").pack(side="left", fill="x", expand=True)
            var = tk.StringVar(value="")
            ans_vars.append(var)
            for opt in ("예", "아니오", "불확실"):
                ttk.Radiobutton(row, text=opt, value=opt, variable=var).pack(side="left", padx=2)

        result_lbl = tk.Label(win, text="", fg="blue", wraplength=840, justify="left")
        result_lbl.pack(fill="x", padx=10, pady=4)
        reasons_store = {"reasons": None}

        def do_judge(save=False):
            answers = [v.get() for v in ans_vars]
            if any(a == "" for a in answers):
                messagebox.showwarning("자가점검", "모든 문항에 응답해 주세요.", parent=win)
                return
            judgment = selfcheck_judge(answers)
            result_lbl.config(text=f"판정: {judgment}")
            audit_log("selfcheck_manual", {"과제명": task.get("과제명", ""),
                                           "answers": answers, "judgment": judgment})
            if save:
                default = _safe_filename(task.get("과제명", "자가점검")) + "_자가점검서.docx"
                path = filedialog.asksaveasfilename(
                    parent=win, defaultextension=".docx", initialfile=default,
                    filetypes=[("Word 문서", "*.docx")])
                if not path:
                    return
                try:
                    export_selfcheck_docx(task, answers, path,
                                          reasons=reasons_store["reasons"])
                    messagebox.showinfo("자가점검", f"저장 완료: {path}", parent=win)
                except Exception as e:
                    messagebox.showerror("자가점검", str(e), parent=win)

        def do_ai_check():
            if not task.get("과제명"):
                messagebox.showwarning("자가점검", "메인 화면의 '② 단건 과제 입력'에 과제정보를 먼저 입력하세요.",
                                       parent=win)
                return
            apply_config()
            result_lbl.config(text="AI 사전 점검 중... (참고용이며 연구자 확인 후 수정 필요)")

            def worker():
                try:
                    llm = LLMClient(CONFIG)
                    res = selfcheck_auto(llm, task)

                    def apply_result():
                        for i, a in enumerate(res["answers"]):
                            ans_vars[i].set(a)
                        reasons_store["reasons"] = res["reasons"]
                        result_lbl.config(
                            text=f"AI 사전 점검 완료(참고): {res['judgment']}\n"
                                 f"종합의견: {res.get('summary', '')}\n"
                                 f"※ 연구자가 각 항목을 직접 확인·수정한 뒤 판정·저장하세요.")
                    win.after(0, apply_result)
                except Exception as e:
                    win.after(0, lambda: result_lbl.config(text=f"[오류] {e}"))
            threading.Thread(target=worker, daemon=True).start()

        fb = ttk.Frame(win)
        fb.pack(fill="x", padx=10, pady=8)
        ttk.Button(fb, text="AI 자동 점검(참고)", command=do_ai_check).pack(side="left", padx=5)
        ttk.Button(fb, text="판정", command=lambda: do_judge(False)).pack(side="left", padx=5)
        ttk.Button(fb, text="판정 + 양식 저장(docx)", command=lambda: do_judge(True)).pack(side="left", padx=5)

    btn_form.config(command=save_review_form)
    btn_self.config(command=open_selfcheck)

    def run_batch():
        if busy.is_set():
            return
        apply_config()
        f = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if not f:
            return
        set_busy(True)
        log(f"\n=== 엑셀 대량 분류: {f} ===")

        def worker():
            try:
                run_excel(f, progress_cb=log_threadsafe, make_forms=forms_var.get())
            except Exception as e:
                log_threadsafe(f"[오류] {e}")
            finally:
                root.after(0, set_busy, False)
        threading.Thread(target=worker, daemon=True).start()

    btn_single.config(command=run_single)
    btn_batch.config(command=run_batch)

    root.mainloop()


# =========================================================
if __name__ == "__main__":
    launch_gui()
    # 코드로 직접 실행:
    # make_template()
    # run_excel("분류대상_양식.xlsx")
