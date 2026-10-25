"""Offline v16 bulk check: synthetic input, mock-model bulk run, HWPX check, analysis.

Nothing here calls a real model, the network or the Windows credential store.
Every model reply comes from tools/v16_mock_model.MockModel (one per task).

Public functions
    write_input(tasks, path)                  9-column 분류대상 workbook, string cells
    load_task_list(path)                      validated 모의과제_목록.json entries
    run_bulk(input_xlsx, task_list_json, run_dir)
        DesktopBridge.start_batch -> batch_v16.run_excel with injected mock client
        factory, in-memory key store and path selector; stops if run_dir exists
    check_documents(results, documents_dir)   HWPX files of every processed result
    analyze(run, checks)                      분석_결과.xlsx and 요약보고서.md
    ols(xs, ys)                               least squares slope/intercept/R²/n
    format_stats(xlsx_path, sheet, column_map) numbers only; never row text
    check_notice(task)                        in-memory notice review with the mock

Command line (run from the simulator folder)
    python tools/v16_bulk_check.py run <input.xlsx> <task_list.json> <run_dir>
    python tools/v16_bulk_check.py stats <xlsx> [--sheet NAME] [--map map.json]
    python tools/v16_bulk_check.py notice <input.xlsx>

Real task data rule (spec 4.2): format_stats reads cells only to count; its
result holds field labels from the given header mapping, fixed bullet/unit
labels and numbers. It never returns, prints or writes a row's text.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import json
import math
import os
from pathlib import Path
import re
import sys
import threading
import time
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

_TOOLS = Path(__file__).resolve().parent
_PROJECT = _TOOLS.parent
for _path in (_PROJECT, _TOOLS):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import batch_v16  # noqa: E402
from credential_store import KeyManager  # noqa: E402
from desktop_bridge_v16 import DesktopBridge  # noqa: E402
import hwpx_writer_v16 as hw  # noqa: E402
import notice_review_v16 as nr  # noqa: E402
import v16_domain as d  # noqa: E402
import v16_mock_model as mm  # noqa: E402


COLUMNS = tuple(d.TASK_COLUMNS) + tuple(d.OPTIONAL_TASK_COLUMNS)
FIELDS_12 = ('반도체·디스플레이', '이차전지', '첨단 모빌리티', '차세대 원자력', '첨단 바이오', '우주항공·해양',
             '수소', '사이버보안', '인공지능', '차세대 통신', '첨단로봇·제조', '양자')
EXCLUDED_FIELD = '제외확인'
EXPECTATIONS = ('등급', '검토필요', '제외')
LIST_KEYS = frozenset(('과제명', '분야', '고시대상', '기대', '노린조항'))
EXCEL_CELL_LIMIT = 32767
OUTPUT_NAME = 'v16_검토결과.xlsx'
DOCUMENTS_NAME = '의견서'
ANALYSIS_NAME = '분석_결과.xlsx'
REPORT_NAME = '요약보고서.md'
LOG_NAME = '실행기록.json'
FAKE_KEY = 'synthetic-offline-mock-key-not-real'
RUN_OPTIONS = dict(method='impact', num_reviewers=3, strictness='표준', selection_mode='auto',
                   ministry_policy='', provider=mm.PROVIDER, model=mm.MODEL)
WORKERS = 3
PHASES = ('facts', 'screen', 'compare', 'candidate', 'criteria', 'reviewer', 'summary')
STAGE_PHASES = frozenset(('candidate', 'criteria', 'reviewer', 'summary'))
ACTUAL_LABELS = {'완료': '등급', '검토필요': '검토필요', '후보제외': '후보제외', '오류': '오류'}
ACTUAL_COLUMNS = ('등급', '검토필요', '후보제외', '오류', '제외', '미처리')
DOC_STATUSES = ('완료', '일부실패', '실패', '미생성')
MIMETYPE = b'application/hwp+zip'
# Strings hwpx_writer_v16._notice_rows writes for the source appendix.
APPENDIX_TITLES = ('부록 1. 고시 대조 및 저장 기준본', '출처 및 버전·고시일·시행일·URL·SHA256')
CANDIDATE_MARK = '후보 조항: '
SOURCE_MARK = '출처: '
SPEC_COLUMN_MAP = {'사업명': '세부사업명', '연구개발과제명-국문': '과제명', '요약문_연구목표': '연구목표',
                   '요약문_연구내용': '연구내용', '요약문_기대효과': '기대효과'}
INPUT_CLEANED_EVENT = '입력 글의 엑셀 줄바꿈 표시·줄 끝 공백·연속 빈 줄 정리 후 검토'

NOTES = (
    '등급·후보는 모의 규칙(과제명 sha256 규칙·조항 낱말 일치)이 정했다. 실제 AI 판단이 아니다.',
    '호출 수는 고시 분할 규칙(요청 크기 100,000자, screen 조항 60개 한도)과 모의 후보 규칙에서 나온다.',
    '과제별 처리 시간은 이 PC에서 모의 클라이언트 생성부터 마지막 호출까지이며 의견서 쓰기는 빠졌다. 전체 실행 시간은 의견서 쓰기를 포함한다.',
    '회귀분석은 처리 규모·시간의 관계이고 모델 품질과 무관하다.',
    '의견서 HWPX는 ZIP·XML 구조와 파일 이름만 검사했다. 한글 프로그램으로 여는 확인은 미확인이다.',
    '모의과제는 합성 시험 과제이며 미정인 예시과제집의 수록 과제나 정답이 아니다.',
    '실제 AI로 큰 요청(최대 100,000자)을 보냈을 때의 동작·응답 크기·시간은 미확인이다.',
)

LAST_OUTCOME = None


# ---------------------------------------------------------------------------
# Input workbook and task list
# ---------------------------------------------------------------------------

def _header():
    return [name + ('(필수)' if name in batch_v16.REQUIRED_TEMPLATE_COLUMNS else '(선택)') for name in COLUMNS]


def _strings(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'


def write_input(tasks, path):
    """Write the 9-column 분류대상 sheet; every cell is a string (data_type 's').

    Text starting with '=' stays text. A value over Excel's cell limit or an
    unknown column is refused before anything is written.
    """
    rows = []
    for task in tasks:
        if not isinstance(task, dict) or set(task) - set(COLUMNS):
            raise ValueError('입력 과제는 9열 이름만 쓸 수 있습니다: ' + ', '.join(sorted(set(task) - set(COLUMNS))))
        row = []
        for column in COLUMNS:
            value = task.get(column, '')
            if not isinstance(value, str):
                raise ValueError(f'입력 칸은 문자열이어야 합니다: {column}')
            if len(value) > EXCEL_CELL_LIMIT:
                raise ValueError(f'엑셀 칸 한도({EXCEL_CELL_LIMIT}자)를 넘는 칸: {column}')
            row.append(value)
        rows.append(row)
    target = Path(path).resolve()
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = '분류대상'
        sheet.append(_header())
        for row in rows:
            sheet.append(row)
        _strings(sheet)
        alignment = Alignment(vertical='top', wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = alignment
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243B53')
        sheet.freeze_panes = 'A2'
        target.parent.mkdir(parents=True, exist_ok=True)
        book.save(target)
    finally:
        book.close()
    return str(target)


def _read_rows(path):
    """Raw input rows as dicts over the 9 columns (marks removed), in sheet order."""
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book['분류대상'] if '분류대상' in book.sheetnames else book.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header = [str(v or '').replace('(필수)', '').replace('(선택)', '').strip() for v in next(iterator, ())]
        rows = []
        for values in iterator:
            raw = {name: '' if value is None else str(value) for name, value in zip(header, values)}
            rows.append({column: raw.get(column, '') for column in COLUMNS})
        return rows
    finally:
        book.close()


def _clean_name(name):
    return d.clean_input_text(str(name))


def load_task_list(path):
    """Validate 모의과제_목록.json (spec 4.3) and return its entries."""
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    if not isinstance(data, list) or not data:
        raise ValueError('과제 목록은 비어 있지 않은 JSON 배열이어야 합니다.')
    seen = set()
    for index, entry in enumerate(data, 1):
        if not isinstance(entry, dict) or not {'과제명', '분야', '고시대상', '기대'} <= set(entry) <= LIST_KEYS:
            raise ValueError(f'과제 목록 {index}번 항목의 키가 맞지 않습니다.')
        if not isinstance(entry['과제명'], str):
            raise ValueError(f'과제 목록 {index}번 과제명은 문자열이어야 합니다.')
        if entry['분야'] not in FIELDS_12 + (EXCLUDED_FIELD,):
            raise ValueError(f'과제 목록 {index}번 분야가 12대 국가전략기술 또는 제외확인이 아닙니다.')
        if type(entry['고시대상']) is not bool or entry['기대'] not in EXPECTATIONS:
            raise ValueError(f'과제 목록 {index}번 고시대상 또는 기대 값이 맞지 않습니다.')
        targets = entry.get('노린조항', [])
        if not isinstance(targets, list) or not all(isinstance(t, str) and t for t in targets):
            raise ValueError(f'과제 목록 {index}번 노린조항은 조항 ID 목록이어야 합니다.')
        if entry['기대'] != '제외':
            name = _clean_name(entry['과제명'])
            if not name or name in seen:
                raise ValueError(f'과제 목록 {index}번 과제명이 비었거나 중복됩니다.')
            seen.add(name)
    return data


# ---------------------------------------------------------------------------
# Bulk run
# ---------------------------------------------------------------------------

class _MemoryBackend:
    """Dictionary credential backend; never touches the Windows store."""

    def __init__(self, values):
        self.values = dict(values)

    def get(self, provider):
        return self.values.get(provider)

    def set(self, provider, key):
        self.values[provider] = key

    def delete(self, provider):
        self.values.pop(provider, None)


def _unwrap(envelope):
    if not envelope.get('ok'):
        raise RuntimeError('브리지 호출 실패: ' + json.dumps(envelope.get('error'), ensure_ascii=False))
    return envelope['data']


def _client_stats(created, log):
    phases = Counter(entry['phase'] for entry in log)
    names = [entry['task_name'] for entry in log if entry['phase'] in STAGE_PHASES and entry['task_name']]
    names = names or [entry['task_name'] for entry in log if entry['task_name']]
    last = max((entry['timestamp'] for entry in log), default=created)
    return dict(task_name=Counter(names).most_common(1)[0][0] if names else None, calls=len(log),
                phases={phase: phases.get(phase, 0) for phase in PHASES},
                partial=sum(entry['partial'] for entry in log),
                system_chars=sum(entry['system_chars'] for entry in log),
                prompt_chars=sum(entry['prompt_chars'] for entry in log),
                response_chars=sum(entry['response_chars'] for entry in log),
                created=created, last_call=last, seconds=max(0.0, last - created))


def run_bulk(input_xlsx, task_list_json, run_dir, *, workers=WORKERS, catalogue_loader=None,
             poll_seconds=0.05, timeout=None):
    """Run the executable's batch path with the mock model (spec 4.6).

    Stops with FileExistsError when ``run_dir`` already exists, before any
    file or client is created. The task list is validated before run_dir is made.
    """
    run_dir = Path(run_dir).resolve()
    if run_dir.exists():
        raise FileExistsError(f'실행 폴더가 이미 있어 멈춤: {run_dir}')
    source = Path(input_xlsx).resolve()
    if not source.is_file():
        raise FileNotFoundError(f'입력 엑셀 없음: {source}')
    task_list = load_task_list(task_list_json)
    run_dir.mkdir(parents=True, exist_ok=False)
    output = run_dir / OUTPUT_NAME
    documents = run_dir / DOCUMENTS_NAME
    documents.mkdir()

    clients, lock, captured = [], threading.Lock(), {}

    def client_factory(options, key):
        client = mm.MockModel()
        with lock:
            clients.append((time.time(), client))
        return client

    def batch_runner(*args, **kwargs):
        report = batch_v16.run_excel(*args, **kwargs)
        captured['report'] = report
        return report

    selected = {'input_excel': source, 'output_excel': output, 'output_directory': documents}
    bridge = DesktopBridge(key_manager=KeyManager(_MemoryBackend({'claude': FAKE_KEY}), environ={}),
                           catalogue_loader=catalogue_loader, client_factory=client_factory,
                           path_selector=lambda purpose: selected[purpose], batch_runner=batch_runner)
    tokens = {purpose: _unwrap(bridge.select_path(purpose))['path_token'] for purpose in selected}
    request = dict(input_token=tokens['input_excel'], output_token=tokens['output_excel'],
                   output_dir_token=tokens['output_directory'], options=dict(RUN_OPTIONS),
                   workers=workers, resume=False, make_forms=True)
    started_at = datetime.now().astimezone().isoformat(timespec='seconds')
    begin = time.perf_counter()
    job_id = _unwrap(bridge.start_batch(request))['job_id']
    events, state, job_result = [], None, None
    while True:
        snapshot = _unwrap(bridge.get_job(job_id, events[-1]['sequence'] if events else 0))
        events.extend(snapshot['events'])
        state = snapshot['state']
        if state in ('done', 'cancelled', 'failed'):
            job_result = snapshot['result']
            break
        if timeout is not None and time.perf_counter() - begin > timeout:
            _unwrap(bridge.cancel_job(job_id))
            timeout = None
        time.sleep(poll_seconds)
    total = time.perf_counter() - begin

    stats = [_client_stats(created, client.call_log()) for created, client in clients]
    task_calls, unmatched, duplicated = {}, 0, []
    for entry in stats:
        name = entry['task_name']
        if name is None:
            unmatched += 1
        elif name in task_calls:
            duplicated.append(name)
        else:
            task_calls[name] = entry
    run = dict(run_dir=str(run_dir), input_path=str(source), output_path=str(output),
               documents_dir=str(documents), task_list=task_list, options=dict(RUN_OPTIONS),
               workers=workers, started_at=started_at, total_seconds=total, job_state=state,
               events=events, job_counts=(job_result or {}).get('counts'), report=captured.get('report'),
               clients=len(clients), task_calls=task_calls, unmatched_clients=unmatched,
               duplicated_clients=duplicated)
    log = {key: value for key, value in run.items() if key != 'report'}
    report = captured.get('report')
    if report is not None:
        log['report'] = {key: value for key, value in report.items() if key != 'results'}
    (run_dir / LOG_NAME).write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding='utf-8')
    return run


# ---------------------------------------------------------------------------
# HWPX document check (spec 4.7)
# ---------------------------------------------------------------------------

def _expected_names(result):
    stem = f"{hw._safe_name(result['task_name'], 20)}_{result['task_fingerprint'][:12]}_"
    names = []
    for index, reviewer in enumerate(result['reviewers']):
        rid = reviewer.get('reviewer_id')
        names.append(stem + f'에이전트_{index + 1}_{hw._safe_name(rid, 12)}_{hw._hash(rid, 8)}.hwpx')
    names.append(stem + '종합.hwpx')
    return names


def _section_texts(data):
    root = ElementTree.fromstring(data)
    return [element.text or '' for element in root.iter(f'{{{hw.HP}}}t')]


def _check_file(path, expected_name, notice_strings):
    record = dict(name=path.name, path=str(path), size=None, zip_ok=False, mimetype_ok=False, xml_ok=False,
                  name_ok=path.name == expected_name and '[모의]' in path.name, notice_ok=None, problems=[])
    if not record['name_ok']:
        record['problems'].append('파일 이름 규칙 불일치')
    if not path.is_file():
        record['problems'].append('파일 없음')
        return record
    record['size'] = os.path.getsize(path)
    try:
        with ZipFile(path) as archive:
            record['zip_ok'] = archive.testzip() is None
            names = archive.namelist()
            record['mimetype_ok'] = 'mimetype' in names and archive.read('mimetype') == MIMETYPE
            if hw.SECTION not in names:
                record['problems'].append('본문 XML 없음')
            else:
                for name in names:
                    if name.endswith('.xml') or name.endswith('.hpf'):
                        ElementTree.fromstring(archive.read(name))
                record['xml_ok'] = True
                if notice_strings:
                    text = '\n'.join(_section_texts(archive.read(hw.SECTION)))
                    missing = [s for s in notice_strings if s not in text]
                    record['notice_ok'] = not missing
                    if missing:
                        record['problems'].append('후보 조항 표시·출처 부록 누락 ' + str(len(missing)) + '건')
    except (BadZipFile, OSError, ElementTree.ParseError, KeyError):
        record['problems'].append('ZIP 또는 XML 열기 실패')
    if not record['zip_ok']:
        record['problems'].append('ZIP 검사 실패')
    if not record['mimetype_ok']:
        record['problems'].append('mimetype 불일치')
    return record


def _hwpx_on_disk(folder):
    # os.walk, not glob: names contain '['.
    found = []
    for root, _, files in os.walk(folder):
        found.extend(os.path.normcase(os.path.abspath(os.path.join(root, name)))
                     for name in files if name.lower().endswith('.hwpx'))
    return found


def check_documents(results, documents_dir=None):
    """Check every processed result's HWPX files (spec 4.7)."""
    tasks, problems, listed = [], [], []
    for result in results:
        docs, reviewers = result['documents'], result['reviewers']
        notice = result['notice_review']
        candidates = [cid for family in notice['families'] for cid in family['candidate_ids']]
        task = dict(task_name=result['task_name'], fingerprint=result['task_fingerprint'],
                    status=result['status'], document_status=docs['status'], expected_files=docs['expected_files'],
                    candidate_ids=candidates, files=[], total_size=0, problems=[])
        if not (len(docs['files']) == len(reviewers) + 1 == docs['expected_files']):
            task['problems'].append('의견서 수·에이전트 수+1·예정 수 불일치')
        if docs['status'] != '완료':
            task['problems'].append('문서 상태 ' + docs['status'])
        notice_strings = []
        if candidates:
            notice_strings = (list(APPENDIX_TITLES) + [CANDIDATE_MARK + cid for cid in candidates]
                              + [SOURCE_MARK + source['source_id'] for source in notice['sources']])
        expected = _expected_names(result)
        folders = set()
        for index, entry in enumerate(docs['files']):
            path = Path(entry['path'])
            listed.append(os.path.normcase(os.path.abspath(path)))
            folders.add(path.parent)
            record = _check_file(path, expected[index] if index < len(expected) else '', notice_strings)
            if entry['status'] != '완료':
                record['problems'].append('파일 상태 ' + entry['status'])
            kind = 'reviewer' if index < len(reviewers) else 'summary'
            if entry['kind'] != kind:
                record['problems'].append('문서 종류 불일치')
            record['kind'] = entry['kind']
            task['files'].append(record)
            task['total_size'] += record['size'] or 0
            task['problems'].extend(f"{record['name']}: {p}" for p in record['problems'])
        for folder in folders:
            if folder.is_dir():
                on_disk = {entry.name for entry in os.scandir(folder) if entry.is_file()}
                listed_here = {Path(e['path']).name for e in docs['files'] if Path(e['path']).parent == folder}
                if on_disk != listed_here:
                    task['problems'].append('결과 파일 목록과 폴더 파일 불일치')
            else:
                task['problems'].append('의견서 폴더 없음')
        task['ok'] = not task['problems']
        problems.extend(f"{task['task_name']}: {p}" for p in task['problems'])
        tasks.append(task)
    disk = _hwpx_on_disk(documents_dir) if documents_dir is not None else None
    if disk is not None:
        extra = set(disk) - set(listed)
        missing = set(listed) - set(disk)
        if extra:
            problems.append(f'결과에 없는 의견서 파일 {len(extra)}개')
        if missing:
            problems.append(f'의견서 폴더 밖이거나 없는 파일 {len(missing)}개')
    return dict(ok=not problems, problems=problems, tasks=tasks, listed_files=len(listed),
                disk_files=len(disk) if disk is not None else None,
                files_ok=sum(not r['problems'] for t in tasks for r in t['files']),
                total_size=sum(t['total_size'] for t in tasks))


# ---------------------------------------------------------------------------
# Statistics helpers
# ---------------------------------------------------------------------------

def ols(xs, ys):
    """Least squares y = slope*x + intercept; slope None when x does not vary."""
    xs, ys = [float(x) for x in xs], [float(y) for y in ys]
    if len(xs) != len(ys):
        raise ValueError('x와 y의 개수가 다릅니다.')
    n = len(xs)
    if n < 2:
        return dict(slope=None, intercept=None, r2=None, n=n)
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    if sxx == 0:
        return dict(slope=None, intercept=None, r2=None, n=n)
    slope = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx
    intercept = my - slope * mx
    total = sum((y - my) ** 2 for y in ys)
    residual = sum((y - (slope * x + intercept)) ** 2 for x, y in zip(xs, ys))
    r2 = (1.0 if residual == 0 else None) if total == 0 else 1 - residual / total
    return dict(slope=slope, intercept=intercept, r2=r2, n=n)


def _quantile(values, q):
    position = (len(values) - 1) * q
    low, high = math.floor(position), math.ceil(position)
    return values[low] + (values[high] - values[low]) * (position - low)


def describe(values):
    values = sorted(float(v) for v in values)
    if not values:
        return dict(n=0, min=None, p10=None, p25=None, median=None, p75=None, p90=None, max=None, mean=None)
    return dict(n=len(values), min=values[0], p10=_quantile(values, .1), p25=_quantile(values, .25),
                median=_quantile(values, .5), p75=_quantile(values, .75), p90=_quantile(values, .9),
                max=values[-1], mean=sum(values) / len(values))


# ---------------------------------------------------------------------------
# Format statistics of a real workbook: numbers only (spec 4.2)
# ---------------------------------------------------------------------------

_BULLET_SYMBOLS = ('-', '–', '•', '·', '○', '◦', 'ㅇ', '▪', '●', '■', '□', '◎', '*', '※', '▶', '►', '>',
                   '◇', '◆', '√', '✓', '◈', '▷')
_BULLET_PATTERNS = ((re.compile(r'\(\d+\)'), '(숫자)'), (re.compile(r'\d+\.(?!\d)'), '숫자.'),
                    (re.compile(r'\d+\)'), '숫자)'), (re.compile(r'[가-하]\.'), '가.'), (re.compile(r'[가-하]\)'), '가)'),
                    (re.compile(r'[①-⑳]'), '원문자'), (re.compile(r'[ⅰ-ⅻⅠ-Ⅻ]'), '로마숫자'))
_EXTRA_UNITS = ('%', '억원', '백만원', '천원', '만원', '원', '건', '명', '개', '종', '회', '편', '배', '년', '개월',
                '℃', '°C', 'V', 'kV', 'mAh', 'Ah', 'Wh', 'kWh', 'MWh', 'GWh', 'GB', 'TB', 'MB', 'dB', 'TOPS', 'fps')
_UNIT_LABELS = sorted(set(d._UNITS) | set(_EXTRA_UNITS), key=len, reverse=True)
_NUMBER_UNIT = re.compile(r'(?<![\w.])\d+(?:[.,]\d+)*\s?(' + '|'.join(re.escape(u) for u in _UNIT_LABELS)
                          + r')(?![A-Za-z가-힣])')


def _bullet(line):
    stripped = line.lstrip()
    if not stripped:
        return None
    for symbol in _BULLET_SYMBOLS:
        if stripped.startswith(symbol):
            return symbol
    for pattern, label in _BULLET_PATTERNS:
        if pattern.match(stripped):
            return label
    return None


def format_stats(xlsx_path, sheet, column_map, *, short_chars=10, header_search_rows=30):
    """Count format features of mapped columns; numbers and fixed labels only.

    ``column_map`` maps a header name to a field label (spec 4.2), e.g.
    {'연구개발과제명-국문': '과제명'}. ``sheet`` None picks the first sheet whose
    header row holds every mapped name. Rows are read only to count.
    """
    if not isinstance(column_map, dict) or not column_map:
        raise ValueError('칸 대응은 비어 있지 않은 사전이어야 합니다.')
    wanted = list(column_map)
    book = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        names = [sheet] if sheet is not None else list(book.sheetnames)
        if sheet is not None and sheet not in book.sheetnames:
            raise ValueError('시트 없음: ' + str(sheet))
        chosen = header_row = positions = None
        for name in names:
            for index, values in enumerate(book[name].iter_rows(max_row=header_search_rows, values_only=True), 1):
                header = [str(v).strip() if v is not None else '' for v in values]
                if all(w in header for w in wanted):
                    chosen, header_row = name, index
                    positions = {w: header.index(w) for w in wanted}
                    break
            if chosen is not None:
                break
        if chosen is None:
            raise ValueError('머리글을 찾지 못함: ' + ', '.join(wanted))
        collected = {label: [] for label in column_map.values()}
        rows = 0
        for values in book[chosen].iter_rows(min_row=header_row + 1, values_only=True):
            if values is None or all(v is None or str(v).strip() == '' for v in values):
                continue
            rows += 1
            for header_name, label in column_map.items():
                position = positions[header_name]
                value = values[position] if position < len(values) else None
                collected[label].append('' if value is None else str(value))
    finally:
        book.close()
    fields = {}
    for label, cells in collected.items():
        nonempty = [c for c in cells if c.strip()]
        lengths, cleaned_lengths, line_counts = [], [], []
        bullets, units = Counter(), Counter()
        marker = carriage = short = bullet_lines = total_lines = 0
        for cell in nonempty:
            lengths.append(len(cell))
            marker += '_x000D_' in cell
            carriage += '\r' in cell
            cleaned = d.clean_input_text(cell)
            cleaned_lengths.append(len(cleaned))
            short += len(cleaned) <= short_chars
            lines = [line for line in cleaned.split('\n') if line.strip()]
            line_counts.append(len(lines))
            total_lines += len(lines)
            for line in lines:
                shape = _bullet(line)
                if shape is not None:
                    bullets[shape] += 1
                    bullet_lines += 1
            for match in _NUMBER_UNIT.finditer(cleaned):
                units[match.group(1)] += 1
        count = len(nonempty)
        fields[label] = dict(
            cells=len(cells), nonempty=count, empty_ratio=(len(cells) - count) / len(cells) if cells else None,
            length=describe(lengths), cleaned_length=describe(cleaned_lengths),
            x000d_ratio=marker / count if count else None, carriage_return_ratio=carriage / count if count else None,
            lines=describe(line_counts), bullet_line_ratio=bullet_lines / total_lines if total_lines else None,
            bullets=dict(sorted(bullets.items())), units=dict(sorted(units.items())),
            number_unit_per_cell=sum(units.values()) / count if count else None,
            very_short_ratio=short / count if count else None, very_short_chars=short_chars)
    return dict(sheet=chosen, header_row=header_row, rows=rows, fields=fields)


# ---------------------------------------------------------------------------
# In-memory notice check (used to set 기대)
# ---------------------------------------------------------------------------

def check_notice(task, *, catalogue=None, chunk_chars=None):
    """Notice review only, in memory, with a fresh MockModel."""
    client = mm.MockModel()
    notice = nr.review_notices(task, client, catalogue=catalogue, chunk_chars=chunk_chars)
    return dict(candidate_ids=[cid for family in notice['families'] for cid in family['candidate_ids']],
                status=notice['status'],
                families={family['family']: family['status'] for family in notice['families']},
                comparisons=[dict(clause_id=c['clause_id'], current=c['current_assessment'],
                                  target=c['target_assessment']) for c in notice['comparisons']],
                blocking_reasons=list(notice['blocking_reasons']), calls=len(client.call_log()),
                partial_calls=sum(entry['partial'] for entry in client.call_log()))


# ---------------------------------------------------------------------------
# Analysis (spec 4.8)
# ---------------------------------------------------------------------------

def _excluded_reason(name):
    return 'blank' if not name else 'example' if name.startswith('(예시)') else 'duplicate'


def _judge(entry, record):
    """Spec 7 expectation rule; returns (match, reason)."""
    expect, status = entry['기대'], record['상태']
    candidates, notice = record['후보 조항 수'], record['고시 대조 상태']
    if expect == '등급' and entry['고시대상']:
        ok = status == '완료' and candidates >= 1 and notice == 'complete'
        return ok, '' if ok else f'고시대상 기대 등급: 상태 {status}, 후보 {candidates}, 고시 대조 {notice}'
    if expect == '등급':
        ok = candidates == 0 and notice == 'complete' and status in ('완료', '후보제외')
        return ok, '' if ok else f'일반 과제 기대 등급: 상태 {status}, 후보 {candidates}, 고시 대조 {notice}'
    if expect == '검토필요':
        ok = status == '검토필요' and (candidates >= 1 or not entry['고시대상'])
        return ok, '' if ok else f'기대 검토필요: 상태 {status}, 후보 {candidates}, 고시 대조 {notice}'
    return False, '처리되었지만 기대 제외'


def _count_lines(path):
    if not path.is_file():
        return None
    with path.open('r', encoding='utf-8') as stream:
        return sum(1 for line in stream if line.strip())


def _audit_saved(path):
    if not path.is_file():
        return None
    count = 0
    with path.open('r', encoding='utf-8') as stream:
        for line in stream:
            if line.strip() and json.loads(line).get('event') == 'batch_task_saved':
                count += 1
    return count


def _workbook_task_rows(path):
    if not path.is_file():
        return None
    book = load_workbook(path, read_only=True)
    try:
        rows = list(book['분류결과'].iter_rows(min_row=2, values_only=True))
        return sum(1 for row in rows if row and isinstance(row[-1], str) and row[-1].startswith('1/'))
    finally:
        book.close()


def _mean(values):
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else None


def analyze(run, checks, run_dir=None):
    """Build per-task records, tables and regressions; write the workbook and report."""
    run_dir = Path(run_dir or run['run_dir'])
    report = run['report'] or {}
    results = report.get('results', [])
    by_name = {result['task_name']: result for result in results}
    checks_by_name = {task['task_name']: task for task in checks['tasks']}
    task_list = run['task_list']
    entries = {_clean_name(e['과제명']): e for e in task_list if e['기대'] != '제외'}
    raw_rows = _read_rows(run['input_path'])
    raw_by_name = {}
    for row in raw_rows:
        name = _clean_name(row['과제명'])
        if name and not name.startswith('(예시)') and name not in raw_by_name:
            raw_by_name[name] = row

    records, mismatches = [], []
    for number, result in enumerate(results, 1):
        name = result['task_name']
        entry = entries.get(name)
        raw = raw_by_name.get(name, {column: '' for column in COLUMNS})
        snapshot = result['task_snapshot']
        calls = run['task_calls'].get(name)
        check = checks_by_name.get(name)
        notice = result['notice_review']
        candidate_ids = [cid for family in notice['families'] for cid in family['candidate_ids']]
        record = {'과제 번호': number, '분야': entry['분야'] if entry else '목록 없음',
                  '고시대상': entry['고시대상'] if entry else None, '기대': entry['기대'] if entry else None,
                  '과제명': name}
        for column in COLUMNS:
            record[f'{column} 원문 글자'] = len(raw[column])
            record[f'{column} 정리 후 글자'] = len(snapshot[column])
        record['입력 총 글자(정리 후)'] = sum(len(snapshot[column]) for column in COLUMNS)
        record['입력 정리'] = any(d.clean_input_text(raw[column]) != raw[column].strip() for column in COLUMNS)
        record['AI 호출 수'] = calls['calls'] if calls else None
        for phase in PHASES:
            record[f'호출 {phase}'] = calls['phases'][phase] if calls else None
        record['조각 호출'] = calls['partial'] if calls else None
        record['지시문 총 문자'] = calls['system_chars'] + calls['prompt_chars'] if calls else None
        record['요청 문자(사용자)'] = calls['prompt_chars'] if calls else None
        record['응답 문자'] = calls['response_chars'] if calls else None
        record['처리 시간(초)'] = calls['seconds'] if calls else None
        record['상태'] = result['status']
        record['등급'] = result['final_grade'] or ''
        record['후보 조항 수'] = len(candidate_ids)
        record['후보 조항'] = ', '.join(candidate_ids)
        record['고시 대조 상태'] = notice['status']
        record['문서 상태'] = result['documents']['status']
        record['의견서 수'] = len(check['files']) if check else 0
        record['의견서 총 크기(바이트)'] = check['total_size'] if check else None
        record['의견서 검사'] = '통과' if check and check['ok'] else '실패'
        targets = entry.get('노린조항') if entry else None
        record['노린조항 일치'] = (sorted(targets) == sorted(candidate_ids)) if targets else None
        if entry is None:
            match, reason = False, '과제 목록에 없음'
        else:
            match, reason = _judge(entry, record)
        record['기대 일치'] = match
        record['불일치 사유'] = reason
        if not match:
            mismatches.append(dict(과제명=name, 사유=reason))
        records.append(record)

    # Expected entries without a result, and expected exclusions.
    for name, entry in entries.items():
        if name not in by_name:
            mismatches.append(dict(과제명=name, 사유='결과 없음(미처리)'))
    input_counts = dict(report.get('input_counts') or {})
    excluded_entries = [e for e in task_list if e['기대'] == '제외']
    matrix = {expect: {column: 0 for column in ACTUAL_COLUMNS} for expect in EXPECTATIONS}
    for record in records:
        if record['기대'] in matrix:
            matrix[record['기대']][ACTUAL_LABELS[record['상태']]] += 1
    for name, entry in entries.items():
        if name not in by_name:
            matrix[entry['기대']]['미처리'] += 1
    remaining = dict(input_counts)
    for entry in excluded_entries:
        reason = _excluded_reason(_clean_name(entry['과제명']))
        if remaining.get(reason, 0) > 0:
            remaining[reason] -= 1
            matrix['제외']['제외'] += 1
        else:
            matrix['제외']['미처리'] += 1
            mismatches.append(dict(과제명=entry['과제명'], 사유='기대 제외였지만 제외 건수에 없음'))
    unexpected_excluded = sum(max(0, v) for v in remaining.values())
    if unexpected_excluded:
        mismatches.append(dict(과제명='', 사유=f'목록에 없는 제외 {unexpected_excluded}건'))

    status_counts = Counter(result['status'] for result in results)
    doc_counts = Counter(result['documents']['status'] for result in results)
    grade_counts = Counter(result['final_grade'] for result in results if result['status'] == '완료')
    by_field = []
    for field in FIELDS_12 + ('목록 없음',):
        group = [r for r in records if r['분야'] == field]
        if not group:
            continue
        by_field.append({'분야': field, '건수': len(group), '고시대상': sum(bool(r['고시대상']) for r in group),
                         **{status: sum(r['상태'] == status for r in group) for status in d.STATUSES},
                         '기대 일치': sum(r['기대 일치'] for r in group),
                         '평균 입력 글자': _mean([r['입력 총 글자(정리 후)'] for r in group]),
                         '평균 호출 수': _mean([r['AI 호출 수'] for r in group]),
                         '평균 처리 시간(초)': _mean([r['처리 시간(초)'] for r in group]),
                         '평균 후보 조항 수': _mean([r['후보 조항 수'] for r in group]),
                         '평균 의견서 크기(바이트)': _mean([r['의견서 총 크기(바이트)'] for r in group])})
    distributions = {label: describe([r[key] for r in records if r[key] is not None]) for label, key in (
        ('입력 총 글자(정리 후)', '입력 총 글자(정리 후)'), ('AI 호출 수', 'AI 호출 수'),
        ('지시문 총 문자', '지시문 총 문자'), ('처리 시간(초)', '처리 시간(초)'),
        ('의견서 총 크기(바이트)', '의견서 총 크기(바이트)'), ('후보 조항 수', '후보 조항 수'))}
    regression = {}
    for label, key in (('AI 호출 수', 'AI 호출 수'), ('지시문 총 문자 수', '지시문 총 문자'),
                       ('처리 시간(초)', '처리 시간(초)'), ('의견서 총 크기(바이트)', '의견서 총 크기(바이트)')):
        pairs = [(r['입력 총 글자(정리 후)'], r[key]) for r in records if r[key] is not None]
        regression[label] = ols([p[0] for p in pairs], [p[1] for p in pairs])
    comparison = []
    for label, flag in (('고시대상', True), ('일반', False)):
        group = [r for r in records if r['고시대상'] is flag]
        comparison.append({'구분': label, '건수': len(group),
                           '평균 후보 조항 수': _mean([r['후보 조항 수'] for r in group]),
                           '평균 호출 수': _mean([r['AI 호출 수'] for r in group]),
                           '평균 compare 호출': _mean([r['호출 compare'] for r in group]),
                           '평균 처리 시간(초)': _mean([r['처리 시간(초)'] for r in group])})
    output = Path(run['output_path'])
    persistence = dict(checkpoint_lines=_count_lines(Path(str(output) + '.checkpoint.jsonl')),
                       audit_saved=_audit_saved(Path(str(output) + '.audit.jsonl')),
                       workbook_rows=_workbook_task_rows(output))
    persistence['consistent'] = all(value == len(results) for value in persistence.values())
    raw_count = len(raw_rows)
    analysis = dict(
        records=records, mismatches=mismatches, matrix=matrix, excluded=input_counts,
        expected_excluded=len(excluded_entries), unexpected_excluded=unexpected_excluded,
        cleaned_tasks=sum(r['입력 정리'] for r in records), input_rows=raw_count, processed=len(results),
        status_counts={status: status_counts.get(status, 0) for status in d.STATUSES},
        document_counts={status: doc_counts.get(status, 0) for status in DOC_STATUSES},
        grade_counts={grade: grade_counts.get(grade, 0) for grade in d.GRADES},
        by_field=by_field, distributions=distributions, regression=regression, comparison=comparison,
        persistence=persistence, target_match=sum(r['노린조항 일치'] is True for r in records),
        target_mismatch=sum(r['노린조항 일치'] is False for r in records),
        candidate_excluded=sum(r['상태'] == '후보제외' and r['기대'] == '등급' for r in records),
        notices=sum(e['message'] == INPUT_CLEANED_EVENT for e in run['events']))
    _write_workbook(run_dir / ANALYSIS_NAME, run, checks, analysis)
    (run_dir / REPORT_NAME).write_text(_report(run, checks, analysis), encoding='utf-8')
    return analysis


def _summary_rows(run, checks, analysis):
    report = run['report'] or {}
    rows = [('실행 폴더', Path(run['run_dir']).name), ('실행 시작', run['started_at']),
            ('작업 상태', run['job_state']), ('실행 조건', json.dumps(run['options'], ensure_ascii=False)),
            ('동시 처리', run['workers']), ('전체 실행 시간(초, 의견서 쓰기 포함)', round(run['total_seconds'], 3)),
            ('입력 행 수(제목 제외)', analysis['input_rows']), ('처리 과제 수', analysis['processed']),
            ('제외 빈 과제명', analysis['excluded'].get('blank', 0)),
            ('제외 예시', analysis['excluded'].get('example', 0)),
            ('제외 중복', analysis['excluded'].get('duplicate', 0)),
            ('목록의 기대 제외 수', analysis['expected_excluded']),
            ('목록에 없는 제외', analysis['unexpected_excluded']),
            ('입력 정리 과제 수', analysis['cleaned_tasks']), ('입력 정리 알림 횟수', analysis['notices'])]
    rows += [(f'상태 {k}', v) for k, v in analysis['status_counts'].items()]
    rows += [(f'문서 상태 {k}', v) for k, v in analysis['document_counts'].items()]
    rows += [('의견서 예정 수', sum(r['documents']['expected_files'] for r in report.get('results', []))),
             ('의견서 결과 목록 수', checks['listed_files']), ('의견서 폴더 파일 수', checks['disk_files']),
             ('의견서 파일 검사 통과', checks['files_ok']), ('의견서 검사 결과', '통과' if checks['ok'] else '실패'),
             ('의견서 총 크기(바이트)', checks['total_size']),
             ('기대 일치', sum(r['기대 일치'] for r in analysis['records'])),
             ('불일치·누락 건수', len(analysis['mismatches'])),
             ('노린조항 일치', analysis['target_match']), ('노린조항 불일치', analysis['target_mismatch']),
             ('기대 등급 중 대상 제외(후보제외)', analysis['candidate_excluded']),
             ('중간 저장 줄 수', analysis['persistence']['checkpoint_lines']),
             ('감사 기록 과제 저장 수', analysis['persistence']['audit_saved']),
             ('결과 엑셀 과제 행 수', analysis['persistence']['workbook_rows']),
             ('저장 건수 일치', analysis['persistence']['consistent']),
             ('진행 이벤트 수', len(run['events'])), ('모의 클라이언트 수', run['clients']),
             ('과제와 연결 안 된 클라이언트', run['unmatched_clients']),
             ('같은 과제 중복 클라이언트', len(run['duplicated_clients'])),
             ('보고 경고 수', len(report.get('warnings', [])))]
    return rows


def _cell(value):
    if isinstance(value, float):
        return round(value, 6)
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _table(sheet, title, header, rows):
    if title:
        sheet.append([title])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
    sheet.append(list(header))
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_cell(value) for value in row])
    sheet.append([])


def _write_workbook(path, run, checks, analysis):
    book = Workbook()
    try:
        summary = book.active
        summary.title = '요약'
        _table(summary, '', ('항목', '값'), _summary_rows(run, checks, analysis))
        _table(summary, '해석 표시', ('내용',), [(note,) for note in NOTES])
        _table(summary, '불일치·누락', ('과제명', '사유'), [(m['과제명'], m['사유']) for m in analysis['mismatches']])
        tasks = book.create_sheet('과제별')
        header = list(analysis['records'][0]) if analysis['records'] else ['과제 번호', '과제명']
        _table(tasks, '', header, [[record[key] for key in header] for record in analysis['records']])
        stats = book.create_sheet('통계')
        total = analysis['processed'] or 1
        _table(stats, '상태별 건수·비율', ('상태', '건수', '비율'),
               [(k, v, v / total) for k, v in analysis['status_counts'].items()])
        _table(stats, '문서 상태별 건수·비율', ('문서 상태', '건수', '비율'),
               [(k, v, v / total) for k, v in analysis['document_counts'].items()])
        _table(stats, '기대와 실제 일치표(행: 기대, 열: 실제)', ('기대',) + ACTUAL_COLUMNS,
               [[expect] + [analysis['matrix'][expect][c] for c in ACTUAL_COLUMNS] for expect in EXPECTATIONS])
        _table(stats, '제외 건수(보고의 input_counts)', ('구분', '건수'),
               [('빈 과제명', analysis['excluded'].get('blank', 0)), ('예시', analysis['excluded'].get('example', 0)),
                ('중복', analysis['excluded'].get('duplicate', 0))])
        _table(stats, '입력 정리', ('구분', '건수'), [('입력 정리 과제 수', analysis['cleaned_tasks'])])
        if analysis['by_field']:
            keys = list(analysis['by_field'][0])
            _table(stats, '분야별 요약', keys, [[row[k] for k in keys] for row in analysis['by_field']])
        dist_keys = ('n', 'min', 'p10', 'p25', 'median', 'p75', 'p90', 'max', 'mean')
        _table(stats, '길이·호출 수·시간 분포', ('항목',) + dist_keys,
               [[label] + [values[k] for k in dist_keys] for label, values in analysis['distributions'].items()])
        _table(stats, '등급 분포(상태 완료)', ('등급', '건수'), list(analysis['grade_counts'].items()))
        regression = book.create_sheet('회귀')
        _table(regression, '최소제곱 회귀: 각 값 ~ 입력 총 글자 수(정리 후)', ('대상', '기울기', '절편', 'R²', 'n'),
               [(label, fit['slope'], fit['intercept'], fit['r2'], fit['n']) for label, fit in analysis['regression'].items()])
        keys = list(analysis['comparison'][0])
        _table(regression, '고시대상 대 일반 과제', keys, [[row[k] for k in keys] for row in analysis['comparison']])
        _table(regression, '해석 표시', ('내용',), [(note,) for note in NOTES])
        documents = book.create_sheet('의견서 검사')
        _table(documents, '', ('과제명', '종류', '파일 이름', '크기(바이트)', 'ZIP', 'mimetype', 'XML', '이름 규칙·[모의]',
                               '후보 조항·출처 부록', '결과', '문제'),
               [(task['task_name'], record.get('kind'), record['name'], record['size'], record['zip_ok'],
                 record['mimetype_ok'], record['xml_ok'], record['name_ok'],
                 '해당 없음' if record['notice_ok'] is None else record['notice_ok'],
                 '통과' if not record['problems'] else '실패', '; '.join(record['problems']))
                for task in checks['tasks'] for record in task['files']])
        _table(documents, '', ('전체 문제',), [(p,) for p in checks['problems']] or [('없음',)])
        _table(documents, '', ('비고',), [('한글 프로그램으로 여는 확인: 미확인',)])
        for sheet in book:
            _strings(sheet)
            sheet.column_dimensions['A'].width = 40
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
    finally:
        book.close()


def _md(value):
    if isinstance(value, float):
        value = round(value, 4)
    return str('' if value is None else value).replace('|', '\\|').replace('\n', ' ')


def _md_table(header, rows):
    lines = ['| ' + ' | '.join(_md(h) for h in header) + ' |', '|' + ' --- |' * len(header)]
    lines += ['| ' + ' | '.join(_md(v) for v in row) + ' |' for row in rows]
    return '\n'.join(lines)


def _report(run, checks, analysis):
    dist_keys = ('n', 'min', 'median', 'p90', 'max', 'mean')
    parts = [
        '# v16 대량 점검 요약보고서',
        '',
        '이 실행의 과제는 모두 **합성 시험 과제**이고 AI 응답은 모두 입력 기반 **모의 응답**이다. '
        '미정인 예시과제집의 수록 과제나 정답이 아니며, 실제 AI는 부르지 않았다.',
        '',
        '## 실행 조건',
        '',
        f"- 경로: DesktopBridge.start_batch → batch_v16.run_excel (모의 클라이언트, 메모리 키 저장소)",
        f"- 옵션: {json.dumps(run['options'], ensure_ascii=False)}, workers {run['workers']}, resume 끔, 의견서 만들기 켬",
        f"- 실행 시작 {run['started_at']}, 작업 상태 {run['job_state']}, 전체 실행 시간 {run['total_seconds']:.2f}초(의견서 쓰기 포함)",
        f"- 고시 대조 요청 크기 기본 {nr.DEFAULT_CHUNK_CHARS:,}자, screen 조항 한도 {nr.SCREEN_CLAUSE_LIMIT}",
        '',
        '## 건수',
        '',
        _md_table(('항목', '값'), _summary_rows(run, checks, analysis)),
        '',
        '## 상태·문서 상태',
        '',
        _md_table(('상태', '건수'), list(analysis['status_counts'].items())),
        '',
        _md_table(('문서 상태', '건수'), list(analysis['document_counts'].items())),
        '',
        '## 기대와 실제 (행: 기대, 열: 실제)',
        '',
        _md_table(('기대',) + ACTUAL_COLUMNS,
                  [[e] + [analysis['matrix'][e][c] for c in ACTUAL_COLUMNS] for e in EXPECTATIONS]),
        '',
        '### 불일치·누락',
        '',
        _md_table(('과제명', '사유'), [(m['과제명'], m['사유']) for m in analysis['mismatches']])
        if analysis['mismatches'] else '없음',
        '',
        '## 분야별 요약',
        '',
        _md_table(list(analysis['by_field'][0]), [list(r.values()) for r in analysis['by_field']])
        if analysis['by_field'] else '없음',
        '',
        '## 분포',
        '',
        _md_table(('항목',) + dist_keys,
                  [[label] + [v[k] for k in dist_keys] for label, v in analysis['distributions'].items()]),
        '',
        '## 등급 분포 (상태 완료)',
        '',
        _md_table(('등급', '건수'), list(analysis['grade_counts'].items())),
        '',
        '## 회귀분석 (최소제곱, 각 값 ~ 입력 총 글자 수)',
        '',
        _md_table(('대상', '기울기', '절편', 'R²', 'n'),
                  [(k, f['slope'], f['intercept'], f['r2'], f['n']) for k, f in analysis['regression'].items()]),
        '',
        _md_table(list(analysis['comparison'][0]), [list(r.values()) for r in analysis['comparison']]),
        '',
        '## 의견서 검사',
        '',
        f"- 결과 목록 파일 {checks['listed_files']}개, 의견서 폴더 파일 {checks['disk_files']}개, "
        f"파일 검사 통과 {checks['files_ok']}개, 전체 결과 {'통과' if checks['ok'] else '실패'}",
        f"- 문제: {'; '.join(checks['problems']) if checks['problems'] else '없음'}",
        '- 한글 프로그램으로 여는 확인: 미확인',
        '',
        '## 해석 표시',
        '',
    ]
    parts += ['- ' + note for note in NOTES]
    parts += ['', '## 미확인', '',
              '- 의견서를 한글 프로그램으로 여는 확인',
              '- 실제 AI로 큰 요청(최대 100,000자)을 보냈을 때의 동작·응답 크기·시간',
              '- 실제 AI 분류 품질, 실제 과제의 대량 처리', '']
    return '\n'.join(parts)


# ---------------------------------------------------------------------------
# Orchestration and command line
# ---------------------------------------------------------------------------

def run_all(input_xlsx, task_list_json, run_dir, **kwargs):
    global LAST_OUTCOME
    run = run_bulk(input_xlsx, task_list_json, run_dir, **kwargs)
    results = (run['report'] or {}).get('results', [])
    checks = check_documents(results, run['documents_dir'])
    analysis = analyze(run, checks)
    LAST_OUTCOME = dict(run=run, checks=checks, analysis=analysis)
    return LAST_OUTCOME


def _input_tasks(path):
    tasks, seen = [], set()
    for row in _read_rows(path):
        cleaned = {column: d.clean_input_text(value) for column, value in row.items()}
        name = cleaned['과제명']
        if not name or name.startswith('(예시)') or name in seen:
            continue
        seen.add(name)
        tasks.append(cleaned)
    return tasks


def main(argv=None):
    parser = argparse.ArgumentParser(prog='v16_bulk_check', description='v16 오프라인 대량 점검(모의 응답)')
    sub = parser.add_subparsers(dest='command', required=True)
    run_parser = sub.add_parser('run', help='대량 실행·의견서 검사·분석')
    run_parser.add_argument('input_xlsx')
    run_parser.add_argument('task_list_json')
    run_parser.add_argument('run_dir')
    run_parser.add_argument('--workers', type=int, default=WORKERS)
    stats_parser = sub.add_parser('stats', help='형식 통계(숫자만)')
    stats_parser.add_argument('xlsx')
    stats_parser.add_argument('--sheet', default=None)
    stats_parser.add_argument('--map', default=None, help='머리글→칸 대응 JSON 파일(기본: spec 4.2 대응)')
    notice_parser = sub.add_parser('notice', help='입력 엑셀 과제별 메모리 고시 대조(모의)')
    notice_parser.add_argument('input_xlsx')
    args = parser.parse_args(argv)

    if args.command == 'run':
        try:
            outcome = run_all(args.input_xlsx, args.task_list_json, args.run_dir, workers=args.workers)
        except FileExistsError as error:
            print(str(error), file=sys.stderr)
            return 1
        except (ValueError, FileNotFoundError) as error:
            print('실행 전 확인 실패: ' + str(error), file=sys.stderr)
            return 1
        run, checks, analysis = outcome['run'], outcome['checks'], outcome['analysis']
        print(f"작업 상태 {run['job_state']} · 처리 {analysis['processed']} · 제외 {json.dumps(analysis['excluded'])}"
              f" · 상태 {json.dumps(analysis['status_counts'], ensure_ascii=False)}")
        print(f"의견서 {checks['listed_files']}개 · 검사 {'통과' if checks['ok'] else '실패'} · "
              f"기대 불일치·누락 {len(analysis['mismatches'])}건 · 저장 건수 일치 {analysis['persistence']['consistent']}")
        print(f"결과: {Path(run['run_dir']) / ANALYSIS_NAME}, {Path(run['run_dir']) / REPORT_NAME}")
        good = (run['job_state'] == 'done' and checks['ok'] and not analysis['mismatches']
                and analysis['persistence']['consistent'])
        return 0 if good else 2
    if args.command == 'stats':
        mapping = json.loads(Path(args.map).read_text(encoding='utf-8')) if args.map else SPEC_COLUMN_MAP
        print(json.dumps(format_stats(args.xlsx, args.sheet, mapping), ensure_ascii=False, indent=2))
        return 0
    for task in _input_tasks(args.input_xlsx):
        found = check_notice(task)
        print(json.dumps(dict(과제명=task['과제명'], **found), ensure_ascii=False))
    return 0


if __name__ == '__main__':
    sys.exit(main())
