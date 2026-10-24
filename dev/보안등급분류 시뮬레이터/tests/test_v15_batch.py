"""Offline batch regressions: synthetic catalogues, workbooks and model clients."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import batch_v15 as batch
import v15_domain as d
from notice_catalogue_v15 import CatalogueError
from test_v15_engine import Client
from test_v15_notice_review import catalogue, fact, repack


def documents(result, output_dir):
    files = [dict(kind="reviewer", reviewer_id=r["reviewer_id"],
                  path=str(Path(output_dir) / (r["reviewer_id"] + ".hwpx")), status="완료", error="")
             for r in result["reviewers"]]
    files.append(dict(kind="summary", reviewer_id=None, path=str(Path(output_dir) / "summary.hwpx"), status="완료", error=""))
    return dict(status="완료", expected_files=len(files), files=files, errors=[])


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.input, self.output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.cat = catalogue()
        self.clients = []
        self.workbook([("합성 과제", "속도 20 GHz")])

    def workbook(self, rows, headers=("과제명", "현재 기술사양"), preferred=True):
        book = Workbook()
        if preferred:
            book.active.title = "작성안내"
            sheet = book.create_sheet("분류대상")
        else:
            sheet = book.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(self.input)
        book.close()

    def factory(self):
        client = Client(facts=[])
        self.clients.append(client)
        return client

    def run_batch(self, **kwargs):
        return batch.run_excel(self.input, self.output, kwargs.pop("options", {}),
            kwargs.pop("client_factory", self.factory), catalogue=kwargs.pop("catalogue", self.cat),
            make_forms=kwargs.pop("make_forms", False), **kwargs)

    def records(self):
        return [json.loads(line) for line in Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines()]

    def book(self):
        book = load_workbook(BytesIO(self.output.read_bytes()), data_only=False)
        self.addCleanup(book.close)
        return book

    def rewrite(self, change):
        record = self.records()[-1]
        change(record)
        Path(str(self.output) + ".checkpoint.jsonl").write_text(json.dumps(record, ensure_ascii=False) + "\n", encoding="utf-8")

    def test_nine_fields_and_historical_five_seven_columns(self):
        fields = d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS
        self.assertEqual(fields, ("세부사업명", "과제명", "연구목표", "연구내용", "기대효과",
                                  "추가 검토정보", "부처 세부기준", "현재 기술사양", "개발목표 사양"))
        for count in (5, 7, 9):
            with self.subTest(count=count):
                values = ["합성" + str(i) for i in range(count)]
                self.workbook([values], [" " + f + ("(필수)" if f == "과제명" else "(선택)") + " " for f in fields[:count]])
                result = self.run_batch(resume=False)["results"][0]
                self.assertEqual(list(result["task_snapshot"]), list(fields))
                self.assertEqual([result["task_snapshot"][f] for f in fields[:count]], values)

    def test_sheet_fallback_order_and_exclusions(self):
        self.workbook([("", ""), ("(예시)합성", ""), ("둘", ""), ("하나", ""), ("둘", "변경")], preferred=False)
        report = self.run_batch()
        self.assertEqual([r["task_name"] for r in report["results"]], ["둘", "하나"])
        self.assertEqual(report["input_counts"], dict(blank=1, example=1, duplicate=1))
        self.assertTrue(report["warnings"])

    def test_duplicate_and_missing_headers(self):
        for headers in (("과제명", "과제명(필수)"), ("과제명", "현재 기술사양", "현재 기술사양(선택)"), ("연구내용",)):
            self.workbook([("합성",) * len(headers)], headers)
            with self.subTest(headers=headers), self.assertRaises(ValueError):
                self.run_batch()
        self.assertFalse(self.clients)

    def test_template_has_nine_columns_and_examples(self):
        path = self.root / "template.xlsx"
        batch.make_template(path)
        book = load_workbook(BytesIO(path.read_bytes()))
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["작성안내", "분류대상"])
        self.assertEqual(book["분류대상"].max_column, 9)
        self.assertTrue(book["분류대상"]["B2"].value.startswith("(예시)"))

    def test_same_input_rejected_and_workers_boundaries(self):
        with self.assertRaises(batch.BatchError):
            batch.run_excel(self.input, self.input, {}, self.factory, catalogue=self.cat)
        for workers in (0, 9, True, 1.5):
            with self.subTest(workers=workers), self.assertRaises(ValueError):
                self.run_batch(workers=workers)
        for workers in (1, 8):
            self.assertEqual(self.run_batch(workers=workers, resume=False)["counts"]["완료"], 1)

    def test_invalid_options_and_actual_row_policy_fail_before_any_client(self):
        for options in (1, {"strictness": "엄격"}, {"comparison_rules_version": "old"},
                        {"selection_mode": "manual", "selected_items": [8]},
                        {"selection_mode": "manual", "selected_items": [1], "weights": {1: .8}}):
            with self.subTest(options=options), self.assertRaises(ValueError):
                self.run_batch(options=options)
        self.workbook([("정책있음", "합성 기준"), ("정책없음", "")], ("과제명", "부처 세부기준"))
        with self.assertRaises(ValueError):
            self.run_batch(options={"selection_mode": "manual", "selected_items": [8]})
        self.assertFalse(self.clients)
        self.assertFalse(Path(str(self.output) + ".checkpoint.jsonl").exists())

    def test_single_catalogue_load_and_same_identity_for_all_tasks(self):
        self.workbook([(str(i), "") for i in range(3)])
        real = batch.classify_task
        identities = []
        def observe(*args, **kwargs):
            identities.append(kwargs["catalogue"])
            return real(*args, **kwargs)
        with patch.object(batch, "load_catalogue", return_value=self.cat) as loader, patch.object(batch, "classify_task", side_effect=observe):
            self.run_batch(catalogue=None)
        self.assertEqual(loader.call_count, 1)
        self.assertEqual(len(identities), 3)
        self.assertTrue(all(c is self.cat for c in identities))

    def test_invalid_catalogue_raises_before_clients_and_storage(self):
        for cat in (object(), None, replace(self.cat, fingerprint="b" * 64), replace(self.cat, clauses=())):
            with patch.object(batch, "load_catalogue", side_effect=CatalogueError("합성")):
                with self.subTest(cat=cat), self.assertRaises(CatalogueError):
                    self.run_batch(catalogue=cat)
        self.assertFalse(self.clients)
        self.assertFalse(self.output.exists())

    def test_complete_reuse_auto_ignores_stale_settings_and_secrets(self):
        self.run_batch(options={"api_key": "SYNTHETIC-SECRET", "screen": object()})
        count = len(self.clients)
        report = self.run_batch(options={"selected_items": object(), "weights": "bad", "weight_tolerance": "bad"})
        self.assertEqual(report["reused_count"], 1)
        self.assertEqual(len(self.clients), count)
        self.assertNotIn("SECRET", Path(report["checkpoint_path"]).read_text(encoding="utf-8"))

    def test_every_input_field_invalidates_reuse(self):
        fields = d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS
        for field in fields:
            row = dict.fromkeys(fields, "") | {"과제명": "동일"}
            self.workbook([list(row.values())], fields)
            self.run_batch(resume=False)
            row[field] += "변경"
            self.workbook([list(row.values())], fields)
            with self.subTest(field=field):
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])

    def test_effective_options_and_baseline_invalidate_reuse(self):
        variants = [dict(method="checklist"), dict(provider="openai", model="gpt-6-luna"),
            dict(provider="ollama", model="synthetic-local"), dict(num_reviewers=1), dict(ministry_policy="합성"),
            dict(selection_mode="manual", selected_items=[1, 2], weights={1: .8, 2: 1.2}),
            dict(local_base_url="http://localhost:11434")]
        for options in variants:
            self.run_batch(resume=False)
            with self.subTest(options=options):
                self.assertEqual(self.run_batch(options=options)["reused_count"], 0)
        plain = self.cat.to_dict()
        plain["manifest"]["catalogue_id"] = "changed"
        self.run_batch(resume=False)
        self.assertEqual(self.run_batch(catalogue=repack(plain))["reused_count"], 0)

    def test_manual_weight_and_selection_change_invalidate_reuse(self):
        options = dict(selection_mode="manual", selected_items=[1, 2])
        self.run_batch(options=options)
        self.assertEqual(self.run_batch(options=options | {"weights": {1: .8, 2: 1.2}})["reused_count"], 0)
        self.assertEqual(self.run_batch(options=options | {"selected_items": [1]})["reused_count"], 0)

    def test_corrupt_legacy_duplicate_keys_and_nan_rejected(self):
        mutations = [lambda r: r.update(schema_version=1), lambda r: r.update(guideline_id="old"),
            lambda r: r.update(catalogue_fingerprint="a" * 64), lambda r: r.update(options_fingerprint="bad"),
            lambda r: r["notice_review"].update(sources=[]), lambda r: r["documents"].update(expected_files=1)]
        for change in mutations:
            self.run_batch(resume=False)
            self.rewrite(change)
            with self.subTest(change=change):
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])
        for raw in ('{"schema_version":2,"schema_version":2}\n', "NaN\n", "{broken\n", "\n"):
            Path(str(self.output) + ".checkpoint.jsonl").write_text(raw, encoding="utf-8")
            report = self.run_batch()
            self.assertEqual(report["reused_count"], 0)
            self.assertTrue(report["warnings"])

    def test_only_complete_status_reused_and_task_failures_are_safe(self):
        self.workbook([("실패", ""), ("정상", "")])
        calls = 0
        def factory():
            nonlocal calls
            calls += 1
            if calls == 1:
                raise RuntimeError("SYNTHETIC-SECRET")
            return self.factory()
        report = self.run_batch(workers=1, client_factory=factory)
        self.assertEqual([r["status"] for r in report["results"]], ["오류", "완료"])
        self.assertIsNone(report["results"][0]["final_grade"])
        self.assertEqual(self.run_batch()["reused_count"], 1)
        self.assertNotIn("SECRET", Path(report["checkpoint_path"]).read_text(encoding="utf-8"))

    def test_resume_false_replaces_old_records(self):
        self.run_batch()
        self.workbook([("변경", "")])
        self.run_batch(resume=False)
        self.assertEqual([r["task_name"] for r in self.records()], ["변경"])

    def test_distinct_classification_statuses_not_reused(self):
        excluded = self.run_batch(resume=False, client_factory=lambda: Client(candidate=False))
        self.assertEqual(excluded["results"][0]["status"], "후보제외")
        self.assertEqual(excluded["counts"]["후보제외"], 1)
        self.assertEqual(self.run_batch()["reused_count"], 0)
        pending = self.run_batch(resume=False, client_factory=lambda: Client(candidates=("advanced:1",), facts=[]))
        self.assertEqual(pending["results"][0]["status"], "검토필요")
        self.assertIsNone(pending["results"][0]["final_grade"])
        self.assertEqual(self.run_batch()["reused_count"], 0)

    def test_real_input_and_response_formulas_are_literal_text(self):
        self.workbook([("=SUM(1,2)", "+DDE()")])
        def change(data, response):
            if data.get("stage") == "summary":
                response["rationale"] = "=SUM(3,4)"
                response["discussion_points"] = "@SUM(5,6)"
            return response
        self.run_batch(client_factory=lambda: Client(facts=[], transform=change))
        sheet = self.book()["분류결과"]
        for address, expected in (("B2", "=SUM(1,2)"), ("J2", "=SUM(3,4)"), ("K2", "@SUM(5,6)")):
            self.assertEqual(sheet[address].value, expected)
            self.assertEqual(sheet[address].data_type, "s")

    def test_concurrency_preserves_order_and_coordinator_writes(self):
        self.workbook([(str(i), "") for i in range(18)])
        real = batch.classify_task
        coordinator = threading.get_ident()
        writes = []
        def delayed(task, *args, **kwargs):
            time.sleep((18 - int(task["과제명"])) % 4 * .002)
            return real(task, *args, **kwargs)
        append, save = batch._append_checkpoint, batch._save_workbook
        def record_append(*args):
            writes.append(threading.get_ident())
            return append(*args)
        def record_save(*args):
            writes.append(threading.get_ident())
            return save(*args)
        with patch.object(batch, "classify_task", side_effect=delayed), patch.object(batch, "_append_checkpoint", side_effect=record_append), patch.object(batch, "_save_workbook", side_effect=record_save):
            report = self.run_batch(workers=8)
        self.assertEqual([r["task_name"] for r in report["results"]], list(map(str, range(18))))
        self.assertEqual(len({id(c) for c in self.clients}), 18)
        self.assertEqual(set(writes), {coordinator})

    def test_options_frozen_before_client_factory(self):
        options = dict(num_reviewers=1)
        def factory():
            options["num_reviewers"] = 3
            return self.factory()
        report = self.run_batch(options=options, client_factory=factory)
        self.assertEqual(len(report["results"][0]["reviewers"]), 1)

    def test_cancel_preserves_inflight_and_stops_new_tasks(self):
        self.workbook([(str(i), "") for i in range(9)])
        event = threading.Event()
        def factory():
            event.set()
            return self.factory()
        report = self.run_batch(workers=1, client_factory=factory, cancel_event=event)
        self.assertEqual(report["status"], "취소")
        self.assertEqual(len(report["results"]), 1)
        self.assertEqual(len(self.records()), 1)

    def test_partial_document_failure_and_offline_regeneration(self):
        def partial(result, directory):
            docs = documents(result, directory)
            docs["files"][0].update(status="오류", error="합성 저장 실패")
            docs.update(status="일부실패", errors=["합성 저장 실패"])
            return docs
        with patch.object(batch, "export_documents", side_effect=partial):
            report = self.run_batch(make_forms=True, output_dir=self.root / "chosen")
        saved = report["results"][0]
        before = deepcopy(saved)
        self.assertEqual(saved["status"], "완료")
        self.assertEqual(report["counts"]["documents_failed"], 1)
        with patch.object(batch, "load_catalogue", side_effect=AssertionError("no catalogue")), patch.object(batch, "classify_task", side_effect=AssertionError("no model")), patch.object(batch, "export_documents", side_effect=documents):
            result = batch.regenerate_documents(saved, self.root / "retry")
        self.assertEqual(saved, before)
        self.assertEqual(result["documents"]["status"], "완료")
        self.assertEqual(Path(report["documents_dir"]), self.root / "chosen")

    def test_writer_exception_and_invalid_identity_are_safe(self):
        def wrong(result, directory):
            docs = documents(result, directory)
            docs["files"][0]["reviewer_id"] = "invalid"
            return docs
        for effect in (OSError("SECRET"), wrong):
            with patch.object(batch, "export_documents", side_effect=effect):
                report = self.run_batch(make_forms=True, resume=False)
            self.assertEqual(report["results"][0]["status"], "완료")
            self.assertEqual(report["results"][0]["documents"]["status"], "실패")
            self.assertNotIn("SECRET", str(report))

    def test_real_hwpx_export_and_saved_integrity(self):
        report = self.run_batch(make_forms=True, options={"num_reviewers": 1})
        self.assertEqual(report["counts"]["documents_completed"], 2)
        for entry in report["results"][0]["documents"]["files"]:
            self.assertTrue(Path(entry["path"]).is_file())
        broken = deepcopy(report["results"][0])
        broken["notice_review"]["sources"][0]["sha256"] = "b" * 64
        with self.assertRaises(ValueError):
            batch.regenerate_documents(broken, self.root / "bad")

    def test_1001_requested_exports_never_disabled(self):
        # Reuse an actually classified DTO shape while keeping 1001 input identities.
        base = self.run_batch(options={"num_reviewers": 1})["results"][0]
        self.workbook([(f"합성{i}", "") for i in range(1001)])
        def classify(task, opts, client, **kwargs):
            result = deepcopy(base)
            result.update(task_snapshot=d.task_snapshot(task), task_name=task["과제명"], task_fingerprint=d.task_fingerprint(task), options_fingerprint=d.options_fingerprint(opts))
            return result
        with patch.object(batch, "classify_task", side_effect=classify), patch.object(batch, "export_documents", side_effect=documents) as writer:
            report = self.run_batch(make_forms=True, workers=8, resume=False, options={"num_reviewers": 1})
        self.assertEqual(writer.call_count, 1001)
        self.assertEqual(report["counts"]["documents_completed"], 2002)

    def test_new_sheets_match_saved_candidates_conditions_sources_and_empty_families(self):
        report = self.run_batch(client_factory=lambda: Client(candidates=("advanced:1",), facts=[fact()]))
        notice = report["results"][0]["notice_review"]
        book = self.book()
        self.assertEqual(book.sheetnames, ["분류결과", "요약", "고시대조", "조건별근거", "고시기준본"])
        self.assertEqual([c.value for c in book["고시대조"][1]][:5], ["과제명", "과제지문", "범위", "범위상태", "후보상태"])
        self.assertEqual([c.value for c in book["조건별근거"][1]][:6], ["과제명", "과제지문", "기록종류", "대조ID", "조항ID", "조건ID"])
        self.assertEqual([c.value for c in book["고시기준본"][1]][:7], ["과제명", "과제지문", "기록종류", "기준본ID", "기준지문", "대조규칙버전", "출처ID"])
        def rows(name):
            values = list(book[name].values)
            return [dict(zip(values[0], row)) for row in values[1:]]
        comparison = [r for r in rows("고시대조") if r["대조ID"]][0]
        self.assertEqual(comparison["대조ID"], notice["comparisons"][0]["comparison_id"])
        self.assertEqual(comparison["현재결과"], notice["comparisons"][0]["current_assessment"])
        self.assertEqual(comparison["조항원문"], notice["clause_snapshots"][0]["text"])
        conditions = [r for r in rows("조건별근거") if r["기록종류"] == "조건"]
        self.assertEqual(len(conditions), len(notice["comparisons"][0]["conditions"]))
        self.assertEqual(conditions[0]["입력인용"], json.dumps([notice["facts"][0]], ensure_ascii=False, separators=(",", ":")))
        sources = [r for r in rows("고시기준본") if r["기록종류"] == "출처"]
        self.assertEqual({r["출처ID"] for r in sources}, {s["source_id"] for s in notice["sources"]})
        self.assertEqual(len([r for r in rows("고시대조") if r["후보상태"] == "후보 없음"]), 3)
        self.run_batch(resume=False)
        book2 = self.book()
        values = list(book2["고시대조"].values)
        empty = [dict(zip(values[0], r)) for r in values[1:]]
        self.assertEqual({r["범위"] for r in empty if r["후보상태"] == "후보 없음"}, set(d.FAMILIES))

    def test_every_string_cell_lossless_continuation_and_formula_text(self):
        report = self.run_batch()
        saved = deepcopy(report["results"][0])
        long = "=+@-한글😀" * 7000
        saved.update(rationale=long, discussion_points=long)
        saved["task_name"] = long
        # Workbook serializer is independently checked with long values in every
        # output column; DTO validation happens before callers reach it.
        with patch.object(batch, "_notice_rows") as notice_rows:
            def output(book, result):
                for name in ("고시대조", "조건별근거", "고시기준본"):
                    sheet = book[name]
                    batch._append_row(sheet, [long] * (sheet.max_column - 1))
            notice_rows.side_effect = output
            batch._save_workbook(self.output, [saved], batch.normalize_options({}))
        book = self.book()
        for name, columns in (("분류결과", (2, 10, 11)), ("고시대조", (1, 2)), ("조건별근거", (1, 2)), ("고시기준본", (1, 2))):
            for column in columns:
                self.assertEqual("".join(str(book[name].cell(row, column).value or "") for row in range(2, book[name].max_row + 1)), long)
        self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))

    def test_checkpoint_init_failure_preserves_previous(self):
        self.run_batch()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        previous = checkpoint.read_bytes()
        with patch.object(batch.os, "replace", side_effect=PermissionError("SECRET")):
            with self.assertRaises(batch.BatchError):
                self.run_batch(resume=False)
        self.assertEqual(checkpoint.read_bytes(), previous)

    def test_failed_append_rolls_back_and_locked_excel_keeps_previous(self):
        self.run_batch()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        before = checkpoint.read_bytes()
        with patch.object(batch.os, "fsync", side_effect=OSError("SECRET")):
            with self.assertRaises(OSError):
                batch._append_checkpoint(checkpoint, '{"bad":true}\n')
        self.assertEqual(checkpoint.read_bytes(), before)
        old_excel = self.output.read_bytes()
        replace = batch.os.replace
        def locked(source, target):
            if Path(target) == self.output:
                raise PermissionError("SECRET")
            return replace(source, target)
        with patch.object(batch.os, "replace", side_effect=locked), self.assertRaises(batch.BatchError):
            self.run_batch(resume=False)
        self.assertEqual(self.output.read_bytes(), old_excel)
        self.assertEqual(len(self.records()), 1)
        self.assertFalse(list(self.root.glob(".v15-*")))

    def test_audit_failure_separate_and_metadata_only(self):
        report = self.run_batch()
        events = Path(report["audit_path"]).read_text(encoding="utf-8")
        self.assertNotIn("합성 과제", events)
        self.assertNotIn("속도", events)
        Path(report["audit_path"]).unlink()
        Path(report["audit_path"]).mkdir()
        report = self.run_batch()
        self.assertEqual(report["audit_status"], "실패")
        self.assertEqual(report["counts"]["완료"], 1)
        self.assertTrue(report["warnings"])

    def generated(self, report, index=0):
        with patch.object(batch, "export_documents", side_effect=documents):
            return batch.regenerate_documents(report["results"][index], self.root / ("retry" + str(index)))

    def test_regeneration_persists_in_order_offline_and_reports_each_store(self):
        report = self.run_batch()
        regenerated = self.generated(report)
        events = []
        append, save = batch._append_checkpoint, batch._save_workbook
        def record(path, encoded):
            events.append("checkpoint" if str(path).endswith("checkpoint.jsonl") else "audit")
            return append(path, encoded)
        def workbook(*args):
            events.append("workbook")
            return save(*args)
        with patch.object(batch, "load_catalogue", side_effect=AssertionError("no live catalogue")), patch.object(batch, "classify_task", side_effect=AssertionError("no model")), patch.object(batch, "_append_checkpoint", side_effect=record), patch.object(batch, "_save_workbook", side_effect=workbook):
            updated = batch.persist_regenerated_documents(report, report["results"][0], regenerated)
        self.assertEqual(events, ["checkpoint", "workbook", "audit"])
        self.assertEqual(updated["persistence"]["status"], "완료")
        self.assertEqual(self.records()[-1], regenerated)
        self.assertEqual(self.book()["분류결과"]["M2"].value, "완료")

    def test_stale_A_then_B_report_rejected_before_any_store(self):
        self.workbook([("A", ""), ("B", "")])
        report = self.run_batch()
        a = batch.persist_regenerated_documents(report, report["results"][0], self.generated(report, 0))
        checkpoint = Path(report["checkpoint_path"])
        audit = Path(report["audit_path"])
        before = [p.read_bytes() for p in (checkpoint, self.output, audit)]
        b = batch.persist_regenerated_documents(report, report["results"][1], self.generated(report, 1))
        self.assertEqual([p.read_bytes() for p in (checkpoint, self.output, audit)], before)
        self.assertEqual([b["persistence"][k] for k in ("checkpoint", "workbook", "audit")], ["미시도"] * 3)
        self.assertTrue(b["persistence"]["warnings"])
        self.assertEqual(a["report"]["counts"]["documents_completed"], 4)
        # The returned report also keeps the latest A state when refusing B.
        self.assertEqual(b["report"]["results"][0]["documents"], a["result"]["documents"])
        fresh = batch.persist_regenerated_documents(a["report"], report["results"][1], self.generated(report, 1))
        self.assertEqual(fresh["persistence"]["status"], "완료")
        self.assertEqual(fresh["report"]["counts"]["documents_completed"], 8)

    def test_regeneration_independent_checkpoint_excel_audit_failures(self):
        for failing, expected in (("checkpoint", ("실패", "미시도", "완료")), ("workbook", ("완료", "실패", "완료")), ("audit", ("완료", "완료", "실패"))):
            report = self.run_batch(resume=False)
            generated = self.generated(report)
            append, save = batch._append_checkpoint, batch._save_workbook
            def append_failure(path, encoded):
                kind = "checkpoint" if str(path).endswith("checkpoint.jsonl") else "audit"
                if kind == failing:
                    raise OSError("SECRET")
                return append(path, encoded)
            def workbook_failure(*args):
                if failing == "workbook":
                    raise OSError("SECRET")
                return save(*args)
            with patch.object(batch, "_append_checkpoint", side_effect=append_failure), patch.object(batch, "_save_workbook", side_effect=workbook_failure):
                result = batch.persist_regenerated_documents(report, report["results"][0], generated)
            with self.subTest(failing=failing):
                self.assertEqual(tuple(result["persistence"][k] for k in ("checkpoint", "workbook", "audit")), expected)
                self.assertEqual(result["result"]["documents"]["status"], "완료")
                self.assertNotIn("SECRET", str(result))

    def test_regeneration_modified_input_options_or_classification_rejected(self):
        report = self.run_batch()
        generated = self.generated(report)
        changed = deepcopy(generated)
        changed["rationale"] += "변경"
        with self.assertRaises(batch.BatchError):
            batch.persist_regenerated_documents(report, report["results"][0], changed)
        for change in (lambda r: r.update(options_fingerprint="b" * 64), lambda r: r.update(catalogue_fingerprint="b" * 64)):
            altered = deepcopy(report)
            change(altered)
            self.assertEqual(batch.persist_regenerated_documents(altered, report["results"][0], generated)["persistence"]["checkpoint"], "미시도")
        self.workbook([("입력변경", "")])
        self.assertEqual(batch.persist_regenerated_documents(report, report["results"][0], generated)["persistence"]["checkpoint"], "미시도")


if __name__ == "__main__":
    unittest.main()
