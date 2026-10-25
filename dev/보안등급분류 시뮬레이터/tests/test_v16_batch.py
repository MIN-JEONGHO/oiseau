"""Offline v16 batch regressions: synthetic catalogues, workbooks and model clients."""
from contextlib import ExitStack
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import re
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

import batch_v16 as batch
import v16_domain as d
from notice_catalogue_v16 import CatalogueError
from test_v16_engine import Client
from test_v16_notice_review import (STYLE_ENDINGS, catalogue, catalogue_mapping, fact, normalise,
                                    repack, substitute)


# Spec 4.4 '입력 양식': the only exception to 'Excel wording stays v15'.
TEMPLATE_HEADER = ["세부사업명(선택)", "과제명(필수)", "연구목표(필수)", "연구내용(필수)", "기대효과(필수)",
                   "추가 검토정보(선택)", "부처 세부기준(선택)", "현재 기술사양(선택)", "개발목표 사양(선택)"]
# Earlier v16 template header (before 세부사업명 became optional); reading ignores marks.
OLD_TEMPLATE_HEADER = ["세부사업명(필수)"] + TEMPLATE_HEADER[1:]
TEMPLATE_EXAMPLES = [
    ["합성 세부사업 A", "(예시) 합성 과제 A", "시험용 목표", "실제 과제정보가 아닌 합성 내용", "시험용 효과",
     None, None, None, None],
    ["합성 세부사업 B", "(예시) 합성 과제 B", "시험용 목표", "실제 과제정보가 아닌 합성 내용", "시험용 효과",
     "합성 검토정보", "합성 부처 기준", "속도 5 GHz", "속도 20 GHz"],
]
TEMPLATE_GUIDE = [
    "연구보안 등급분류 입력 양식",
    "분류대상 시트에 과제별 한 행씩 입력",
    "과제명·연구목표·연구내용·기대효과 4개 항목 필수 입력, "
    "세부사업명과 추가 4개 항목(추가 검토정보·부처 세부기준·현재 기술사양·개발목표 사양) 선택 입력",
    "과제명이 빈 행과 '(예시)'로 시작하는 행은 검토에서 제외, 같은 과제명은 첫 행만 처리",
    "추가 검토정보는 실제 확인한 사실만, 부처 세부기준은 적용할 기준이 있을 때만 입력",
    "현재 기술사양과 개발목표 사양은 구분해 입력",
    "AI 에이전트 검토의견은 참고 의견이며 실제 심의위원회 심의와 최종 등급 확정을 대신하지 않음",
]
TEMPLATE_STYLE_PATTERN = re.compile("(?:" + "|".join(STYLE_ENDINGS) + r")[)\]」』’”'\"]*(?=$|[\s.!?,…:;])")


def documents(result, output_dir):
    files = [dict(kind="reviewer", reviewer_id=r["reviewer_id"],
                  path=str(Path(output_dir) / (r["reviewer_id"] + ".hwpx")), status="완료", error="")
             for r in result["reviewers"]]
    files.append(dict(kind="summary", reviewer_id=None, path=str(Path(output_dir) / "summary.hwpx"), status="완료", error=""))
    return dict(status="완료", expected_files=len(files), files=files, errors=[])


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.input, self.output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.cat = catalogue()
        self.clients = []
        self.workbook([("합성 과제", "속도 20 GHz")])

    def workbook(self, rows, headers=("과제명", "현재 기술사양"), preferred=True):
        book = Workbook()
        if preferred:
            book.active.title = "작성안내"
            sheet = book.create_sheet("분류대상")
        else:
            sheet = book.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(self.input)
        book.close()

    def factory(self):
        client = Client(facts=[])
        self.clients.append(client)
        return client

    def run_batch(self, **kwargs):
        return batch.run_excel(self.input, self.output, kwargs.pop("options", {}),
            kwargs.pop("client_factory", self.factory), catalogue=kwargs.pop("catalogue", self.cat),
            make_forms=kwargs.pop("make_forms", False), **kwargs)

    def records(self):
        return [json.loads(line) for line in Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines()]

    def book(self):
        book = load_workbook(BytesIO(self.output.read_bytes()), data_only=False)
        self.addCleanup(book.close)
        return book

    def rewrite(self, change):
        record = self.records()[-1]
        change(record)
        Path(str(self.output) + ".checkpoint.jsonl").write_text(json.dumps(record, ensure_ascii=False) + "\n", encoding="utf-8")

    def test_nine_fields_and_historical_five_seven_columns(self):
        fields = d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS
        self.assertEqual(fields, ("세부사업명", "과제명", "연구목표", "연구내용", "기대효과",
                                  "추가 검토정보", "부처 세부기준", "현재 기술사양", "개발목표 사양"))
        for count in (5, 7, 9):
            with self.subTest(count=count):
                values = ["합성" + str(i) for i in range(count)]
                self.workbook([values], [" " + f + ("(필수)" if f == "과제명" else "(선택)") + " " for f in fields[:count]])
                result = self.run_batch(resume=False)["results"][0]
                self.assertEqual(list(result["task_snapshot"]), list(fields))
                self.assertEqual([result["task_snapshot"][f] for f in fields[:count]], values)

    def test_sheet_fallback_order_and_exclusions(self):
        self.workbook([("", ""), ("(예시)합성", ""), ("둘", ""), ("하나", ""), ("둘", "변경")], preferred=False)
        report = self.run_batch()
        self.assertEqual([r["task_name"] for r in report["results"]], ["둘", "하나"])
        self.assertEqual(report["input_counts"], dict(blank=1, example=1, duplicate=1))
        self.assertTrue(report["warnings"])

    def test_duplicate_and_missing_headers(self):
        for headers in (("과제명", "과제명(필수)"), ("과제명", "현재 기술사양", "현재 기술사양(선택)"), ("연구내용",)):
            self.workbook([("합성",) * len(headers)], headers)
            with self.subTest(headers=headers), self.assertRaises(ValueError):
                self.run_batch()
        self.assertFalse(self.clients)

    def test_template_has_nine_columns_and_examples(self):
        path = self.root / "template.xlsx"
        batch.make_template(path)
        book = load_workbook(BytesIO(path.read_bytes()))
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["작성안내", "분류대상"])
        self.assertEqual(book["분류대상"].max_column, 9)
        self.assertTrue(book["분류대상"]["B2"].value.startswith("(예시)"))

    def template_cells(self):
        path = self.root / "template.xlsx"
        batch.make_template(path)
        book = load_workbook(BytesIO(path.read_bytes()))
        try:
            return path, {sheet.title: [list(row) for row in sheet.values] for sheet in book}
        finally:
            book.close()

    def test_template_header_marks_four_required_and_five_optional(self):
        _, cells = self.template_cells()
        self.assertEqual(cells["분류대상"][0], TEMPLATE_HEADER)
        self.assertEqual(len(TEMPLATE_HEADER), 9)
        self.assertEqual(TEMPLATE_HEADER[0], "세부사업명(선택)")
        self.assertEqual(batch.REQUIRED_TEMPLATE_COLUMNS, ("과제명", "연구목표", "연구내용", "기대효과"))
        # Marks are display only: column meaning and order stay as TASK_COLUMNS.
        self.assertEqual(d.TASK_COLUMNS, ("세부사업명", "과제명", "연구목표", "연구내용", "기대효과"))
        self.assertEqual([name + ("(필수)" if name in batch.REQUIRED_TEMPLATE_COLUMNS else "(선택)")
                          for name in d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS], TEMPLATE_HEADER)

    def test_template_example_rows_keep_current_cell_values(self):
        _, cells = self.template_cells()
        self.assertEqual(cells["분류대상"][1:], TEMPLATE_EXAMPLES)

    def test_template_guide_is_noun_style_and_carries_the_spec_facts(self):
        _, cells = self.template_cells()
        guide = cells["작성안내"]
        self.assertTrue(all(len(row) == 1 and isinstance(row[0], str) and row[0].strip() for row in guide), guide)
        self.assertEqual([row[0] for row in guide], TEMPLATE_GUIDE)
        for text in TEMPLATE_GUIDE:
            with self.subTest(text=text):
                self.assertEqual(TEMPLATE_STYLE_PATTERN.findall(text), [])
                for sentence in re.split(r"[.,·\n]\s*", text):
                    self.assertFalse(sentence.rstrip().endswith(STYLE_ENDINGS), sentence)
        joined = "\n".join(TEMPLATE_GUIDE)
        self.assertNotIn("기본 5개 항목", joined)
        for fact_text in ("과제별 한 행", "과제명·연구목표·연구내용·기대효과 4개 항목 필수 입력",
                          "세부사업명과 추가 4개 항목", "선택 입력",
                          "과제명이 빈 행", "'(예시)'로 시작하는 행", "첫 행만 처리", "실제 확인한 사실만",
                          "적용할 기준이 있을 때만", "구분해 입력", "참고 의견", "대신하지 않음"):
            self.assertIn(fact_text, joined)

    def test_template_style_check_rejects_sentence_endings(self):
        for text in ("과제명을 입력하세요", "제외합니다.", "입력해 주세요", "처리해요", "있어요", "좋아요",
                     "의견이에요", "예시예요", "할까요?"):
            with self.subTest(text=text):
                self.assertTrue(TEMPLATE_STYLE_PATTERN.findall(text))

    def test_saved_template_reads_as_two_examples_excluded_and_zero_tasks(self):
        path, _ = self.template_cells()
        tasks, excluded, cleaned = batch._read_tasks(path)
        self.assertEqual((tasks, excluded, cleaned), ([], dict(blank=0, example=2, duplicate=0), False))
        notices = []
        report = batch.run_excel(path, self.output, {}, self.factory, catalogue=self.cat, make_forms=False,
                                 notice_cb=notices.append)
        self.assertEqual(notices, [])
        self.assertEqual(report["input_counts"], dict(blank=0, example=2, duplicate=0))
        self.assertEqual((report["task_count"], report["results"]), (0, []))
        self.assertIn("입력 제외: 빈 과제명 0건, 예시 2건, 중복 과제명 0건. 중복은 첫 행 유지", report["warnings"])
        self.assertFalse(self.clients)

    def test_old_template_header_with_required_business_name_reads_the_same(self):
        rows = [["합성 세부사업", "합성 과제", "합성 목표", "합성 내용", "합성 효과", "", "", "속도 5 GHz", "속도 20 GHz"],
                [None, "합성 과제 둘", "합성 목표", "합성 내용", "합성 효과", None, None, None, None]]
        read = []
        for header in (OLD_TEMPLATE_HEADER, TEMPLATE_HEADER):
            self.workbook([list(TEMPLATE_EXAMPLES[0])] + rows, header)
            read.append(batch._read_tasks(self.input))
        self.assertEqual(read[0], read[1])
        tasks, excluded, cleaned = read[0]
        self.assertEqual([task["과제명"] for task in tasks], ["합성 과제", "합성 과제 둘"])
        self.assertEqual(tasks[0]["세부사업명"], "합성 세부사업")
        self.assertEqual(tasks[1]["세부사업명"], "")
        self.assertEqual((excluded, cleaned), (dict(blank=0, example=1, duplicate=0), False))

    def test_excel_marker_only_task_name_is_blank_and_batch_continues(self):
        self.workbook([("_x000D_", "속도 20 GHz"), (" _x000A__x0009_ ", ""), ("합성 과제", "속도 20 GHz")])
        tasks, excluded, _ = batch._read_tasks(self.input)
        self.assertEqual([task["과제명"] for task in tasks], ["합성 과제"])
        self.assertEqual(excluded, dict(blank=2, example=0, duplicate=0))
        report = self.run_batch()
        self.assertEqual([r["task_name"] for r in report["results"]], ["합성 과제"])
        self.assertEqual(report["input_counts"], dict(blank=2, example=0, duplicate=0))
        self.assertIn("입력 제외: 빈 과제명 2건, 예시 0건, 중복 과제명 0건. 중복은 첫 행 유지", report["warnings"])
        self.assertEqual(len(self.clients), 1)

    def test_names_equal_after_cleaning_count_second_as_duplicate(self):
        self.workbook([("합성_x000D_과제", "속도 20 GHz"), ("합성\n과제", "변경"),
                       ("합성 과제 둘", ""), ("합성 과제 둘_x000D_  ", "변경"), ("(예시)_x000D_합성", "")])
        tasks, excluded, cleaned = batch._read_tasks(self.input)
        self.assertEqual([task["과제명"] for task in tasks], ["합성\n과제", "합성 과제 둘"])
        self.assertEqual(tasks[0]["현재 기술사양"], "속도 20 GHz")
        self.assertEqual(excluded, dict(blank=0, example=1, duplicate=2))
        self.assertTrue(cleaned)
        report = self.run_batch()
        self.assertEqual([r["task_name"] for r in report["results"]], ["합성\n과제", "합성 과제 둘"])
        self.assertEqual(report["input_counts"], dict(blank=0, example=1, duplicate=2))

    def test_input_cleaned_notice_once_before_any_write_or_model_call(self):
        fields = list(d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS)
        dirty = ["합성 세부_x000D_사업", "합성 과제 A", "목표  \t\n둘째 줄", "내용\n\n\n\n끝", "효과",
                 "검토_x000A_정보", "", "속도 5 GHz \r\n", "속도 20 GHz"]
        self.workbook([dirty, ["", "합성 과제 B_x000D_", "", "", "", "", "", "", ""]], fields)
        tasks, excluded, cleaned = batch._read_tasks(self.input)
        self.assertTrue(cleaned)
        self.assertEqual(excluded, dict(blank=0, example=0, duplicate=0))
        self.assertEqual([task["과제명"] for task in tasks], ["합성 과제 A", "합성 과제 B"])
        self.assertEqual(tasks[0]["연구목표"], "목표\n둘째 줄")
        self.assertEqual(tasks[0]["연구내용"], "내용\n\n끝")
        events = []
        original_guard, original_write = batch._guard_previous_result, batch._atomic_write

        def guard(*args):
            events.append("guard")
            return original_guard(*args)

        def write(*args):
            events.append("write")
            return original_write(*args)

        def factory():
            events.append("client")
            return self.factory()
        with patch.object(batch, "_guard_previous_result", guard), patch.object(batch, "_atomic_write", write):
            report = self.run_batch(client_factory=factory, workers=1,
                                    notice_cb=lambda name: events.append(("notice", name)))
        self.assertEqual([event for event in events if isinstance(event, tuple)], [("notice", "input_cleaned")])
        self.assertEqual(events[:2], ["guard", ("notice", "input_cleaned")])
        self.assertIn("write", events)
        self.assertIn("client", events)
        # Screen only: no warning, no stored text, report structure unchanged.
        self.assertEqual(report["warnings"], [])
        self.assertEqual(report["input_counts"], dict(blank=0, example=0, duplicate=0))
        self.assertEqual(report["results"][0]["task_snapshot"]["세부사업명"], "합성 세부\n사업")
        stored = Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8")
        self.assertNotIn("input_cleaned", stored)
        self.assertNotIn("_x000D_", stored)

    def test_clean_input_never_sends_input_cleaned(self):
        self.workbook([("  합성 과제  ", " 속도 20 GHz "), ("합성 과제 둘", 20), ("합성 과제 셋", None)])
        self.assertFalse(batch._read_tasks(self.input)[2])
        notices = []
        report = self.run_batch(notice_cb=notices.append)
        self.assertEqual(notices, [])
        self.assertEqual(len(report["results"]), 3)
        # The default keeps working without a callback.
        self.assertEqual(len(self.run_batch()["results"]), 3)

    def test_only_excluded_or_non_task_columns_dirty_never_sends_input_cleaned(self):
        self.workbook([("", "속도_x000D_20"), ("_x000D_", "속도  \n\n\n20"), ("(예시) 합성_x000D_", "a_x000D_b"),
                       ("합성 과제", "속도 20 GHz", "메모_x000D_칸"), ("합성 과제", "속도_x000D_20", ""),
                       ("합성 과제_x000D_", "속도_x000A_20", "")],
                      ("과제명", "현재 기술사양", "메모"))
        tasks, excluded, cleaned = batch._read_tasks(self.input)
        self.assertEqual([task["과제명"] for task in tasks], ["합성 과제"])
        self.assertEqual(excluded, dict(blank=2, example=1, duplicate=2))
        self.assertFalse(cleaned)
        notices = []
        report = self.run_batch(notice_cb=notices.append)
        self.assertEqual(notices, [])
        self.assertEqual(report["input_counts"], dict(blank=2, example=1, duplicate=2))

    def test_template_formatting_matches_v15(self):
        import batch_v15
        books = {}
        for module in (batch_v15, batch):
            path = self.root / (module.__name__ + "-format.xlsx")
            module.make_template(path)
            books[module.__name__] = load_workbook(BytesIO(path.read_bytes()))
            self.addCleanup(books[module.__name__].close)

        def styles(sheet):
            def style(cell):
                return (cell.font.name, cell.font.size, cell.font.bold, cell.font.color and cell.font.color.rgb,
                        cell.fill.fill_type, cell.fill.fgColor.rgb, cell.alignment.vertical, cell.alignment.wrap_text,
                        cell.data_type)
            rows = list(sheet.iter_rows())
            return ({style(cell) for cell in rows[0]}, {style(cell) for row in rows[1:] for cell in row if cell.value is not None})

        old, new = books["batch_v15"], books["batch_v16"]
        self.assertEqual(old.sheetnames, new.sheetnames)
        for name in old.sheetnames:
            with self.subTest(sheet=name):
                self.assertEqual(old[name].freeze_panes, new[name].freeze_panes)
                self.assertEqual(old[name].auto_filter.ref, new[name].auto_filter.ref)
                self.assertEqual({k: v.width for k, v in old[name].column_dimensions.items()},
                                 {k: v.width for k, v in new[name].column_dimensions.items()})
                self.assertEqual(styles(old[name]), styles(new[name]))

    def test_same_input_rejected_and_workers_boundaries(self):
        with self.assertRaises(batch.BatchError):
            batch.run_excel(self.input, self.input, {}, self.factory, catalogue=self.cat)
        for workers in (0, 9, True, 1.5):
            with self.subTest(workers=workers), self.assertRaises(ValueError):
                self.run_batch(workers=workers)
        for workers in (1, 8):
            self.assertEqual(self.run_batch(workers=workers, resume=False)["counts"]["완료"], 1)

    def test_invalid_options_and_actual_row_policy_fail_before_any_client(self):
        for options in (1, {"strictness": "엄격"}, {"comparison_rules_version": "old"},
                        {"selection_mode": "manual", "selected_items": [8]},
                        {"selection_mode": "manual", "selected_items": [1], "weights": {1: .8}}):
            with self.subTest(options=options), self.assertRaises(ValueError):
                self.run_batch(options=options)
        self.workbook([("정책있음", "합성 기준"), ("정책없음", "")], ("과제명", "부처 세부기준"))
        with self.assertRaises(ValueError):
            self.run_batch(options={"selection_mode": "manual", "selected_items": [8]})
        self.assertFalse(self.clients)
        self.assertFalse(Path(str(self.output) + ".checkpoint.jsonl").exists())

    def test_single_catalogue_load_and_same_identity_for_all_tasks(self):
        self.workbook([(str(i), "") for i in range(3)])
        real = batch.classify_task
        identities = []
        def observe(*args, **kwargs):
            identities.append(kwargs["catalogue"])
            return real(*args, **kwargs)
        with patch.object(batch, "load_catalogue", return_value=self.cat) as loader, patch.object(batch, "classify_task", side_effect=observe):
            self.run_batch(catalogue=None)
        self.assertEqual(loader.call_count, 1)
        self.assertEqual(len(identities), 3)
        self.assertTrue(all(c is self.cat for c in identities))

    def test_invalid_catalogue_raises_before_clients_and_storage(self):
        for cat in (object(), None, replace(self.cat, fingerprint="b" * 64), replace(self.cat, clauses=())):
            with patch.object(batch, "load_catalogue", side_effect=CatalogueError("합성")):
                with self.subTest(cat=cat), self.assertRaises(CatalogueError):
                    self.run_batch(catalogue=cat)
        self.assertFalse(self.clients)
        self.assertFalse(self.output.exists())

    def test_complete_reuse_auto_ignores_stale_settings_and_secrets(self):
        self.run_batch(options={"api_key": "SYNTHETIC-SECRET", "screen": object()})
        count = len(self.clients)
        report = self.run_batch(options={"selected_items": object(), "weights": "bad", "weight_tolerance": "bad"})
        self.assertEqual(report["reused_count"], 1)
        self.assertEqual(len(self.clients), count)
        self.assertNotIn("SECRET", Path(report["checkpoint_path"]).read_text(encoding="utf-8"))

    def test_every_input_field_invalidates_reuse(self):
        fields = d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS
        for field in fields:
            row = dict.fromkeys(fields, "") | {"과제명": "동일"}
            self.workbook([list(row.values())], fields)
            self.run_batch(resume=False)
            row[field] += "변경"
            self.workbook([list(row.values())], fields)
            with self.subTest(field=field):
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])

    def test_effective_options_and_baseline_invalidate_reuse(self):
        variants = [dict(method="checklist"), dict(provider="openai", model="gpt-6-luna"),
            dict(provider="ollama", model="synthetic-local"), dict(num_reviewers=1), dict(ministry_policy="합성"),
            dict(selection_mode="manual", selected_items=[1, 2], weights={1: .8, 2: 1.2}),
            dict(local_base_url="http://localhost:11434")]
        for options in variants:
            self.run_batch(resume=False)
            with self.subTest(options=options):
                self.assertEqual(self.run_batch(options=options)["reused_count"], 0)
        plain = self.cat.to_dict()
        plain["manifest"]["catalogue_id"] = "changed"
        self.run_batch(resume=False)
        self.assertEqual(self.run_batch(catalogue=repack(plain))["reused_count"], 0)

    def test_manual_weight_and_selection_change_invalidate_reuse(self):
        options = dict(selection_mode="manual", selected_items=[1, 2])
        self.run_batch(options=options)
        self.assertEqual(self.run_batch(options=options | {"weights": {1: .8, 2: 1.2}})["reused_count"], 0)
        self.assertEqual(self.run_batch(options=options | {"selected_items": [1]})["reused_count"], 0)

    def test_corrupt_legacy_duplicate_keys_and_nan_rejected(self):
        # v16: an integer schema_version below 3 stops the run (spec 4.5,
        # PreviousResultProtectionTests); other damaged v16 records still warn.
        mutations = [lambda r: r.update(schema_version=4), lambda r: r.update(guideline_id="old"),
            lambda r: r.update(catalogue_fingerprint="a" * 64), lambda r: r.update(options_fingerprint="bad"),
            lambda r: r["notice_review"].update(sources=[]), lambda r: r["documents"].update(expected_files=1)]
        for change in mutations:
            self.run_batch(resume=False)
            self.rewrite(change)
            with self.subTest(change=change):
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])
        for raw in ('{"schema_version":3,"schema_version":3}\n', "NaN\n", "{broken\n", "\n"):
            Path(str(self.output) + ".checkpoint.jsonl").write_text(raw, encoding="utf-8")
            report = self.run_batch()
            self.assertEqual(report["reused_count"], 0)
            self.assertTrue(report["warnings"])

    def test_only_complete_status_reused_and_task_failures_are_safe(self):
        self.workbook([("실패", ""), ("정상", "")])
        calls = 0
        def factory():
            nonlocal calls
            calls += 1
            if calls == 1:
                raise RuntimeError("SYNTHETIC-SECRET")
            return self.factory()
        report = self.run_batch(workers=1, client_factory=factory)
        self.assertEqual([r["status"] for r in report["results"]], ["오류", "완료"])
        self.assertIsNone(report["results"][0]["final_grade"])
        self.assertEqual(self.run_batch()["reused_count"], 1)
        self.assertNotIn("SECRET", Path(report["checkpoint_path"]).read_text(encoding="utf-8"))

    def test_resume_false_replaces_old_records(self):
        self.run_batch()
        self.workbook([("변경", "")])
        self.run_batch(resume=False)
        self.assertEqual([r["task_name"] for r in self.records()], ["변경"])

    def test_distinct_classification_statuses_not_reused(self):
        excluded = self.run_batch(resume=False, client_factory=lambda: Client(candidate=False))
        self.assertEqual(excluded["results"][0]["status"], "후보제외")
        self.assertEqual(excluded["counts"]["후보제외"], 1)
        self.assertEqual(self.run_batch()["reused_count"], 0)
        pending = self.run_batch(resume=False, client_factory=lambda: Client(candidates=("advanced:1",), facts=[]))
        self.assertEqual(pending["results"][0]["status"], "검토필요")
        self.assertIsNone(pending["results"][0]["final_grade"])
        self.assertEqual(self.run_batch()["reused_count"], 0)

    def test_real_input_and_response_formulas_are_literal_text(self):
        self.workbook([("=SUM(1,2)", "+DDE()")])
        def change(data, response):
            if data.get("stage") == "summary":
                response["rationale"] = "=SUM(3,4)"
                response["discussion_points"] = "@SUM(5,6)"
            return response
        self.run_batch(client_factory=lambda: Client(facts=[], transform=change))
        sheet = self.book()["분류결과"]
        for address, expected in (("B2", "=SUM(1,2)"), ("J2", "=SUM(3,4)"), ("K2", "@SUM(5,6)")):
            self.assertEqual(sheet[address].value, expected)
            self.assertEqual(sheet[address].data_type, "s")

    def test_concurrency_preserves_order_and_coordinator_writes(self):
        self.workbook([(str(i), "") for i in range(18)])
        real = batch.classify_task
        coordinator = threading.get_ident()
        writes = []
        def delayed(task, *args, **kwargs):
            time.sleep((18 - int(task["과제명"])) % 4 * .002)
            return real(task, *args, **kwargs)
        append, save = batch._append_checkpoint, batch._save_workbook
        def record_append(*args):
            writes.append(threading.get_ident())
            return append(*args)
        def record_save(*args):
            writes.append(threading.get_ident())
            return save(*args)
        with patch.object(batch, "classify_task", side_effect=delayed), patch.object(batch, "_append_checkpoint", side_effect=record_append), patch.object(batch, "_save_workbook", side_effect=record_save):
            report = self.run_batch(workers=8)
        self.assertEqual([r["task_name"] for r in report["results"]], list(map(str, range(18))))
        self.assertEqual(len({id(c) for c in self.clients}), 18)
        self.assertEqual(set(writes), {coordinator})

    def test_options_frozen_before_client_factory(self):
        options = dict(num_reviewers=1)
        def factory():
            options["num_reviewers"] = 3
            return self.factory()
        report = self.run_batch(options=options, client_factory=factory)
        self.assertEqual(len(report["results"][0]["reviewers"]), 1)

    def test_cancel_preserves_inflight_and_stops_new_tasks(self):
        self.workbook([(str(i), "") for i in range(9)])
        event = threading.Event()
        def factory():
            event.set()
            return self.factory()
        report = self.run_batch(workers=1, client_factory=factory, cancel_event=event)
        self.assertEqual(report["status"], "취소")
        self.assertEqual(len(report["results"]), 1)
        self.assertEqual(len(self.records()), 1)

    def test_partial_document_failure_and_offline_regeneration(self):
        def partial(result, directory):
            docs = documents(result, directory)
            docs["files"][0].update(status="오류", error="합성 저장 실패")
            docs.update(status="일부실패", errors=["합성 저장 실패"])
            return docs
        with patch.object(batch, "export_documents", side_effect=partial):
            report = self.run_batch(make_forms=True, output_dir=self.root / "chosen")
        saved = report["results"][0]
        before = deepcopy(saved)
        self.assertEqual(saved["status"], "완료")
        self.assertEqual(report["counts"]["documents_failed"], 1)
        with patch.object(batch, "load_catalogue", side_effect=AssertionError("no catalogue")), patch.object(batch, "classify_task", side_effect=AssertionError("no model")), patch.object(batch, "export_documents", side_effect=documents):
            result = batch.regenerate_documents(saved, self.root / "retry")
        self.assertEqual(saved, before)
        self.assertEqual(result["documents"]["status"], "완료")
        self.assertEqual(Path(report["documents_dir"]), self.root / "chosen")

    def test_writer_exception_and_invalid_identity_are_safe(self):
        def wrong(result, directory):
            docs = documents(result, directory)
            docs["files"][0]["reviewer_id"] = "invalid"
            return docs
        for effect in (OSError("SECRET"), wrong):
            with patch.object(batch, "export_documents", side_effect=effect):
                report = self.run_batch(make_forms=True, resume=False)
            self.assertEqual(report["results"][0]["status"], "완료")
            self.assertEqual(report["results"][0]["documents"]["status"], "실패")
            self.assertNotIn("SECRET", str(report))

    def test_real_hwpx_export_and_saved_integrity(self):
        report = self.run_batch(make_forms=True, options={"num_reviewers": 1})
        self.assertEqual(report["counts"]["documents_completed"], 2)
        for entry in report["results"][0]["documents"]["files"]:
            self.assertTrue(Path(entry["path"]).is_file())
        broken = deepcopy(report["results"][0])
        broken["notice_review"]["sources"][0]["sha256"] = "b" * 64
        with self.assertRaises(ValueError):
            batch.regenerate_documents(broken, self.root / "bad")

    def test_1001_requested_exports_never_disabled(self):
        # Reuse an actually classified DTO shape while keeping 1001 input identities.
        base = self.run_batch(options={"num_reviewers": 1})["results"][0]
        self.workbook([(f"합성{i}", "") for i in range(1001)])
        def classify(task, opts, client, **kwargs):
            result = deepcopy(base)
            result.update(task_snapshot=d.task_snapshot(task), task_name=task["과제명"], task_fingerprint=d.task_fingerprint(task), options_fingerprint=d.options_fingerprint(opts))
            return result
        with patch.object(batch, "classify_task", side_effect=classify), patch.object(batch, "export_documents", side_effect=documents) as writer:
            report = self.run_batch(make_forms=True, workers=8, resume=False, options={"num_reviewers": 1})
        self.assertEqual(writer.call_count, 1001)
        self.assertEqual(report["counts"]["documents_completed"], 2002)

    def test_new_sheets_match_saved_candidates_conditions_sources_and_empty_families(self):
        report = self.run_batch(client_factory=lambda: Client(candidates=("advanced:1",), facts=[fact()]))
        notice = report["results"][0]["notice_review"]
        book = self.book()
        self.assertEqual(book.sheetnames, ["분류결과", "요약", "고시대조", "조건별근거", "고시기준본"])
        self.assertEqual([c.value for c in book["고시대조"][1]][:5], ["과제명", "과제지문", "범위", "범위상태", "후보상태"])
        self.assertEqual([c.value for c in book["조건별근거"][1]][:6], ["과제명", "과제지문", "기록종류", "대조ID", "조항ID", "조건ID"])
        self.assertEqual([c.value for c in book["고시기준본"][1]][:7], ["과제명", "과제지문", "기록종류", "기준본ID", "기준지문", "대조규칙버전", "출처ID"])
        def rows(name):
            values = list(book[name].values)
            return [dict(zip(values[0], row)) for row in values[1:]]
        comparison = [r for r in rows("고시대조") if r["대조ID"]][0]
        self.assertEqual(comparison["대조ID"], notice["comparisons"][0]["comparison_id"])
        self.assertEqual(comparison["현재결과"], notice["comparisons"][0]["current_assessment"])
        self.assertEqual(comparison["조항원문"], notice["clause_snapshots"][0]["text"])
        conditions = [r for r in rows("조건별근거") if r["기록종류"] == "조건"]
        self.assertEqual(len(conditions), len(notice["comparisons"][0]["conditions"]))
        self.assertEqual(conditions[0]["입력인용"], json.dumps([notice["facts"][0]], ensure_ascii=False, separators=(",", ":")))
        sources = [r for r in rows("고시기준본") if r["기록종류"] == "출처"]
        self.assertEqual({r["출처ID"] for r in sources}, {s["source_id"] for s in notice["sources"]})
        self.assertEqual(len([r for r in rows("고시대조") if r["후보상태"] == "후보 없음"]), 3)
        self.run_batch(resume=False)
        book2 = self.book()
        values = list(book2["고시대조"].values)
        empty = [dict(zip(values[0], r)) for r in values[1:]]
        self.assertEqual({r["범위"] for r in empty if r["후보상태"] == "후보 없음"}, set(d.FAMILIES))

    def test_every_string_cell_lossless_continuation_and_formula_text(self):
        report = self.run_batch()
        saved = deepcopy(report["results"][0])
        long = "=+@-한글😀" * 7000
        saved.update(rationale=long, discussion_points=long)
        saved["task_name"] = long
        # Workbook serializer is independently checked with long values in every
        # output column; DTO validation happens before callers reach it.
        with patch.object(batch, "_notice_rows") as notice_rows:
            def output(book, result):
                for name in ("고시대조", "조건별근거", "고시기준본"):
                    sheet = book[name]
                    batch._append_row(sheet, [long] * (sheet.max_column - 1))
            notice_rows.side_effect = output
            batch._save_workbook(self.output, [saved], batch.normalize_options({}))
        book = self.book()
        for name, columns in (("분류결과", (2, 10, 11)), ("고시대조", (1, 2)), ("조건별근거", (1, 2)), ("고시기준본", (1, 2))):
            for column in columns:
                self.assertEqual("".join(str(book[name].cell(row, column).value or "") for row in range(2, book[name].max_row + 1)), long)
        self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))

    def test_atomic_write_waits_out_a_briefly_held_target(self):
        # Antivirus, search indexing or an editor can hold a just-written file
        # for seconds; a short hold must not fail the whole batch.
        target = self.root / "held.xlsx"
        target.write_bytes(b"old")
        replace, failures = batch.os.replace, []

        def held(source, destination):
            if len(failures) < 6:
                failures.append(destination)
                raise PermissionError("held")
            return replace(source, destination)
        with patch.object(batch.os, "replace", side_effect=held), patch.object(batch.time, "sleep") as sleep:
            batch._atomic_write(target, lambda temporary: Path(temporary).write_bytes(b"new"))
        self.assertEqual(target.read_bytes(), b"new")
        delays = [call.args[0] for call in sleep.call_args_list]
        self.assertEqual(len(delays), 6)
        self.assertEqual(delays, sorted(delays))
        self.assertFalse(list(self.root.glob(".v16-*")))

    def test_atomic_write_gives_up_after_the_retry_budget(self):
        target = self.root / "locked.xlsx"
        target.write_bytes(b"old")
        with patch.object(batch.os, "replace", side_effect=PermissionError("SECRET")) as replace, \
                patch.object(batch.time, "sleep") as sleep, self.assertRaises(PermissionError):
            batch._atomic_write(target, lambda temporary: Path(temporary).write_bytes(b"new"))
        waited = sum(call.args[0] for call in sleep.call_args_list)
        self.assertGreaterEqual(waited, batch._REPLACE_RETRY_SECONDS - 2.0)
        self.assertLessEqual(waited, batch._REPLACE_RETRY_SECONDS)
        self.assertEqual(replace.call_count, len(sleep.call_args_list) + 1)
        self.assertEqual(target.read_bytes(), b"old")
        self.assertFalse(list(self.root.glob(".v16-*")))

    def test_atomic_write_does_not_retry_other_errors(self):
        target = self.root / "other.xlsx"
        with patch.object(batch.os, "replace", side_effect=OSError("disk")) as replace, \
                patch.object(batch.time, "sleep") as sleep, self.assertRaises(OSError):
            batch._atomic_write(target, lambda temporary: Path(temporary).write_bytes(b"new"))
        self.assertEqual(replace.call_count, 1)
        sleep.assert_not_called()

    def test_checkpoint_init_failure_preserves_previous(self):
        self.run_batch()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        previous = checkpoint.read_bytes()
        with patch.object(batch.os, "replace", side_effect=PermissionError("SECRET")), patch.object(batch.time, "sleep"):
            with self.assertRaises(batch.BatchError) as raised:
                self.run_batch(resume=False)
        # A v16 result is not a previous-version result; the write itself failed.
        self.assertNotIsInstance(raised.exception, batch.PreviousResultError)
        self.assertEqual(checkpoint.read_bytes(), previous)

    def test_failed_append_rolls_back_and_locked_excel_keeps_previous(self):
        self.run_batch()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        before = checkpoint.read_bytes()
        with patch.object(batch.os, "fsync", side_effect=OSError("SECRET")):
            with self.assertRaises(OSError):
                batch._append_checkpoint(checkpoint, '{"bad":true}\n')
        self.assertEqual(checkpoint.read_bytes(), before)
        old_excel = self.output.read_bytes()
        replace = batch.os.replace
        def locked(source, target):
            if Path(target) == self.output:
                raise PermissionError("SECRET")
            return replace(source, target)
        with patch.object(batch.os, "replace", side_effect=locked), patch.object(batch.time, "sleep"), \
                self.assertRaises(batch.BatchError):
            self.run_batch(resume=False)
        self.assertEqual(self.output.read_bytes(), old_excel)
        self.assertEqual(len(self.records()), 1)
        self.assertFalse(list(self.root.glob(".v16-*")))

    def test_audit_failure_separate_and_metadata_only(self):
        report = self.run_batch()
        events = Path(report["audit_path"]).read_text(encoding="utf-8")
        self.assertNotIn("합성 과제", events)
        self.assertNotIn("속도", events)
        Path(report["audit_path"]).unlink()
        Path(report["audit_path"]).mkdir()
        report = self.run_batch()
        self.assertEqual(report["audit_status"], "실패")
        self.assertEqual(report["counts"]["완료"], 1)
        self.assertTrue(report["warnings"])

    def generated(self, report, index=0):
        with patch.object(batch, "export_documents", side_effect=documents):
            return batch.regenerate_documents(report["results"][index], self.root / ("retry" + str(index)))

    def test_regeneration_persists_in_order_offline_and_reports_each_store(self):
        report = self.run_batch()
        regenerated = self.generated(report)
        events = []
        append, save = batch._append_checkpoint, batch._save_workbook
        def record(path, encoded):
            events.append("checkpoint" if str(path).endswith("checkpoint.jsonl") else "audit")
            return append(path, encoded)
        def workbook(*args):
            events.append("workbook")
            return save(*args)
        with patch.object(batch, "load_catalogue", side_effect=AssertionError("no live catalogue")), patch.object(batch, "classify_task", side_effect=AssertionError("no model")), patch.object(batch, "_append_checkpoint", side_effect=record), patch.object(batch, "_save_workbook", side_effect=workbook):
            updated = batch.persist_regenerated_documents(report, report["results"][0], regenerated)
        self.assertEqual(events, ["checkpoint", "workbook", "audit"])
        self.assertEqual(updated["persistence"]["status"], "완료")
        self.assertEqual(self.records()[-1], regenerated)
        self.assertEqual(self.book()["분류결과"]["M2"].value, "완료")

    def test_stale_A_then_B_report_rejected_before_any_store(self):
        self.workbook([("A", ""), ("B", "")])
        report = self.run_batch()
        a = batch.persist_regenerated_documents(report, report["results"][0], self.generated(report, 0))
        checkpoint = Path(report["checkpoint_path"])
        audit = Path(report["audit_path"])
        before = [p.read_bytes() for p in (checkpoint, self.output, audit)]
        b = batch.persist_regenerated_documents(report, report["results"][1], self.generated(report, 1))
        self.assertEqual([p.read_bytes() for p in (checkpoint, self.output, audit)], before)
        self.assertEqual([b["persistence"][k] for k in ("checkpoint", "workbook", "audit")], ["미시도"] * 3)
        self.assertTrue(b["persistence"]["warnings"])
        self.assertEqual(a["report"]["counts"]["documents_completed"], 4)
        # The returned report also keeps the latest A state when refusing B.
        self.assertEqual(b["report"]["results"][0]["documents"], a["result"]["documents"])
        fresh = batch.persist_regenerated_documents(a["report"], report["results"][1], self.generated(report, 1))
        self.assertEqual(fresh["persistence"]["status"], "완료")
        self.assertEqual(fresh["report"]["counts"]["documents_completed"], 8)

    def test_regeneration_independent_checkpoint_excel_audit_failures(self):
        for failing, expected in (("checkpoint", ("실패", "미시도", "완료")), ("workbook", ("완료", "실패", "완료")), ("audit", ("완료", "완료", "실패"))):
            report = self.run_batch(resume=False)
            generated = self.generated(report)
            append, save = batch._append_checkpoint, batch._save_workbook
            def append_failure(path, encoded):
                kind = "checkpoint" if str(path).endswith("checkpoint.jsonl") else "audit"
                if kind == failing:
                    raise OSError("SECRET")
                return append(path, encoded)
            def workbook_failure(*args):
                if failing == "workbook":
                    raise OSError("SECRET")
                return save(*args)
            with patch.object(batch, "_append_checkpoint", side_effect=append_failure), patch.object(batch, "_save_workbook", side_effect=workbook_failure):
                result = batch.persist_regenerated_documents(report, report["results"][0], generated)
            with self.subTest(failing=failing):
                self.assertEqual(tuple(result["persistence"][k] for k in ("checkpoint", "workbook", "audit")), expected)
                self.assertEqual(result["result"]["documents"]["status"], "완료")
                self.assertNotIn("SECRET", str(result))

    def test_regeneration_modified_input_options_or_classification_rejected(self):
        report = self.run_batch()
        generated = self.generated(report)
        changed = deepcopy(generated)
        changed["rationale"] += "변경"
        with self.assertRaises(batch.BatchError):
            batch.persist_regenerated_documents(report, report["results"][0], changed)
        for change in (lambda r: r.update(options_fingerprint="b" * 64), lambda r: r.update(catalogue_fingerprint="b" * 64)):
            altered = deepcopy(report)
            change(altered)
            self.assertEqual(batch.persist_regenerated_documents(altered, report["results"][0], generated)["persistence"]["checkpoint"], "미시도")
        self.workbook([("입력변경", "")])
        self.assertEqual(batch.persist_regenerated_documents(report, report["results"][0], generated)["persistence"]["checkpoint"], "미시도")


# ---------------------------------------------------------------------------
# v16 additions: 고시일 column, previous-version result protection, screen-only
# noun-style messages and the spec 6 "저장 문장 동일" parity with v15.
# ---------------------------------------------------------------------------
V15_CATALOGUE_COLUMNS = ["과제명", "과제지문", "기록종류", "기준본ID", "기준지문", "대조규칙버전", "출처ID", "범위",
    "고시명", "고시번호", "버전", "시행일", "URL", "해시", "상대경로", "수집시각", "검토상태", "처리범위", "검토메모", "저장원문"]
PROTECTION_MESSAGE = "기존 결과 파일 보호를 위해 중단. 새 결과 엑셀 경로 선택 필요"


def noun_style(message):
    return not message.rstrip(". ").endswith(STYLE_ENDINGS)


def sheet_rows(book, name):
    """Merge continuation rows ('k/n' in the last column) back into logical rows."""
    values = list(book[name].values)
    header, rows = list(values[0]), []
    assert header[-1] == "이어쓰기", header
    for raw in values[1:]:
        part, total = map(int, raw[-1].split("/"))
        cells = list(raw[:-1])
        if part == 1:
            rows.append(cells)
        else:
            current = rows[-1]
            for index, value in enumerate(cells):
                if isinstance(value, str):
                    current[index] = (current[index] or "") + value
                else:
                    assert value is None, (name, raw)
        assert 1 <= part <= total
    return header[:-1], rows


class CatalogueSheetV16Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name).resolve()
        self.input, self.output = root / "input.xlsx", root / "output.xlsx"
        book = Workbook()
        book.active.append(("과제명", "현재 기술사양"))
        book.active.append(("합성 과제", "속도 20 GHz"))
        book.save(self.input)
        book.close()

    def test_announcement_date_column_is_immediately_before_effective_date(self):
        expected = list(V15_CATALOGUE_COLUMNS)
        expected.insert(expected.index("시행일"), "고시일")
        self.assertEqual(batch.CATALOGUE_COLUMNS, expected)
        import batch_v15
        self.assertEqual(batch_v15.CATALOGUE_COLUMNS, V15_CATALOGUE_COLUMNS)
        self.assertEqual(batch.CATALOGUE_COLUMNS.index("고시일") + 1, batch.CATALOGUE_COLUMNS.index("시행일"))
        # The other sheets keep the v15 headers.
        self.assertEqual(batch.NOTICE_COLUMNS, batch_v15.NOTICE_COLUMNS)
        self.assertEqual(batch.CONDITION_COLUMNS, batch_v15.CONDITION_COLUMNS)
        self.assertEqual(batch.RESULT_COLUMNS, batch_v15.RESULT_COLUMNS)

    def test_saved_catalogue_rows_fill_every_shifted_column(self):
        report = batch.run_excel(self.input, self.output, {}, lambda: Client(candidates=("advanced:1",)),
                                 catalogue=catalogue(), make_forms=False)
        notice = report["results"][0]["notice_review"]
        book = load_workbook(BytesIO(self.output.read_bytes()), data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["분류결과", "요약", "고시대조", "조건별근거", "고시기준본"])
        header, rows = sheet_rows(book, "고시기준본")
        self.assertEqual(header, batch.CATALOGUE_COLUMNS)
        records = [dict(zip(header, row)) for row in rows]
        base = [r for r in records if r["기록종류"] == "기준본"]
        self.assertEqual(len(base), 1)
        base = base[0]
        for column in ("출처ID", "범위", "고시명", "고시번호", "버전", "고시일", "시행일", "URL", "해시", "상대경로", "수집시각"):
            self.assertIsNone(base[column], column)
        self.assertEqual(base["검토상태"], notice["status"])
        self.assertEqual(base["처리범위"], batch._json(notice["families"]))
        self.assertEqual(base["검토메모"], batch._json(notice["blocking_reasons"]))
        self.assertEqual(base["저장원문"], batch._json(notice["catalogue_manifest"]))
        sources = [r for r in records if r["기록종류"] == "출처"]
        self.assertEqual(len(sources), len(notice["sources"]))
        for row, source in zip(sources, notice["sources"]):
            with self.subTest(source=source["source_id"]):
                self.assertEqual(row["출처ID"], source["source_id"])
                self.assertEqual(row["고시명"], source["title"])
                self.assertEqual(row["고시번호"], source["notice_number"])
                self.assertEqual(row["버전"], source["announcement_version"])
                self.assertEqual(row["고시일"], source["announcement_date"])
                self.assertEqual(row["시행일"], source["effective_date"])
                self.assertEqual(row["URL"], source["official_url"])
                self.assertEqual(row["해시"], source["sha256"])
                self.assertEqual(row["상대경로"], source["relative_path"])
                self.assertEqual(row["수집시각"], source["retrieved_at"])
                self.assertEqual(row["검토상태"], source["review_status"])
                self.assertEqual(row["처리범위"], source["coverage"])
                self.assertEqual(row["검토메모"], source["review_note"])
                self.assertEqual(json.loads(row["저장원문"]), source)
                self.assertLess(row["고시일"], "9999")
                self.assertLessEqual(row["고시일"], row["시행일"])


class PreviousResultProtectionTests(unittest.TestCase):
    """Spec 4.5: a selected non-v16 result is never overwritten, resume or not."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.input, self.output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        self.audit = Path(str(self.output) + ".audit.jsonl")
        book = Workbook()
        book.active.title = "작성안내"
        sheet = book.create_sheet("분류대상")
        sheet.append(("과제명", "현재 기술사양"))
        sheet.append(("합성 과제 A", "속도 20 GHz"))
        sheet.append(("합성 과제 B", "속도 20 GHz"))
        book.save(self.input)
        book.close()
        self.clients = []

    def factory(self):
        client = Client(candidates=("advanced:1",))
        self.clients.append(client)
        return client

    def v15_run(self, output=None):
        import batch_v15
        return batch_v15.run_excel(self.input, output or self.output, {}, lambda: Client(candidates=("advanced:1",)),
                                   catalogue=catalogue(v15=True, catalogue_id="synthetic-v15"), make_forms=False)

    def v16_run(self, **kwargs):
        return batch.run_excel(self.input, self.output, {}, self.factory, catalogue=catalogue(),
                               make_forms=kwargs.pop("make_forms", False), **kwargs)

    def snapshot(self):
        return {path.name: (path.read_bytes() if path.is_file() else None)
                for path in (self.checkpoint, self.output, self.audit)}

    def entries(self):
        return sorted(path.name for path in self.root.iterdir())

    def assert_stopped(self, resume):
        before, listing, clients = self.snapshot(), self.entries(), len(self.clients)
        messages = []
        names = ("_atomic_write", "_append_checkpoint", "_save_workbook", "export_documents", "classify_task")
        with ExitStack() as stack:
            guards = {name: stack.enter_context(patch.object(batch, name, Mock(side_effect=AssertionError("write attempted"))))
                      for name in names}
            with self.assertRaises(batch.PreviousResultError) as raised:
                batch.run_excel(self.input, self.output, {}, self.factory, catalogue=catalogue(), resume=resume,
                                make_forms=True, progress_cb=messages.append, status_cb=messages.append)
        for name, guard in guards.items():
            self.assertEqual(guard.call_count, 0, name)
        self.assertIsInstance(raised.exception, batch.BatchError)
        self.assertEqual(str(raised.exception), PROTECTION_MESSAGE)
        self.assertEqual(batch.PREVIOUS_RESULT_MESSAGE, PROTECTION_MESSAGE)
        self.assertTrue(noun_style(str(raised.exception)))
        # Checkpoint, workbook and audit are byte-identical; nothing else appeared.
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(self.entries(), listing)
        # Zero new client factories means zero model calls.
        self.assertEqual(len(self.clients), clients)
        self.assertEqual(messages, [])

    def test_v15_checkpoint_workbook_and_audit_are_preserved(self):
        self.v15_run()
        self.assertTrue(all(value is not None for value in self.snapshot().values()))
        self.assertTrue(all(json.loads(line)["schema_version"] == 2
                            for line in self.checkpoint.read_text(encoding="utf-8").splitlines()))
        for resume in (True, False):
            with self.subTest(resume=resume):
                self.assert_stopped(resume)

    def test_previous_result_stop_never_sends_input_cleaned(self):
        self.v15_run()
        book = Workbook()
        book.active.title = "작성안내"
        sheet = book.create_sheet("분류대상")
        sheet.append(("과제명", "현재 기술사양"))
        sheet.append(("합성 과제 A_x000D_", "속도_x000D_20 GHz"))
        book.save(self.input)
        book.close()
        self.assertTrue(batch._read_tasks(self.input)[2])
        notices = []
        for resume in (True, False):
            with self.subTest(resume=resume), self.assertRaises(batch.PreviousResultError):
                batch.run_excel(self.input, self.output, {}, self.factory, catalogue=catalogue(), resume=resume,
                                make_forms=False, notice_cb=notices.append)
        self.assertEqual(notices, [])

    def test_v15_checkpoint_only(self):
        self.v15_run()
        self.output.unlink()
        for resume in (True, False):
            with self.subTest(resume=resume):
                self.assert_stopped(resume)

    def test_v15_result_workbook_only(self):
        self.v15_run()
        self.checkpoint.unlink()
        for resume in (True, False):
            with self.subTest(resume=resume):
                self.assert_stopped(resume)
        self.audit.unlink()
        for resume in (True, False):
            with self.subTest(resume=resume, audit=False):
                self.assert_stopped(resume)

    def test_one_older_record_in_a_v16_checkpoint_stops_even_with_v16_workbook(self):
        self.v16_run()
        other = self.root / "v15" / "other.xlsx"
        other.parent.mkdir()
        self.v15_run(other)
        old = Path(str(other) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines()[0]
        with self.checkpoint.open("a", encoding="utf-8", newline="\n") as stream:
            stream.write(old + "\n")
        for resume in (True, False):
            with self.subTest(resume=resume):
                self.assert_stopped(resume)

    def test_v14_and_minimal_older_checkpoint_records(self):
        for raw in ('{"schema_version":1}\n', '{"schema_version":2,"task_name":"합성"}\n',
                    '{"schema_version":0}\n{"broken\n', '{"schema_version":-1}\n',
                    '{"schema_version":2,"schema_version":2}\n'):
            self.checkpoint.write_text(raw, encoding="utf-8")
            for resume in (True, False):
                with self.subTest(raw=raw, resume=resume):
                    self.assert_stopped(resume)

    def test_v14_workbook_without_catalogue_sheet(self):
        book = Workbook()
        book.active.title = "분류결과"
        book.active.append(["연번", "과제명", "상태", "최종등급(참고)"])
        book.active.append([1, "합성", "완료", "민감과제"])
        book.create_sheet("요약").append(["구분", "건수", "비율"])
        book.save(self.output)
        book.close()
        self.checkpoint.write_text('{"schema_version":1}\n', encoding="utf-8")
        self.audit.write_text('{"schema_version":1,"event":"batch_task_saved"}\n', encoding="utf-8")
        for resume in (True, False):
            with self.subTest(resume=resume):
                self.assert_stopped(resume)
        self.checkpoint.unlink()
        for resume in (True, False):
            with self.subTest(resume=resume, checkpoint=False):
                self.assert_stopped(resume)

    def test_unreadable_result_file(self):
        for data in (b"not an xlsx", b"PK\x03\x04broken", b""):
            self.output.write_bytes(data)
            for resume in (True, False):
                with self.subTest(data=data, resume=resume):
                    self.assert_stopped(resume)

    def test_other_excel_files_are_not_v16_results(self):
        variants = []
        book = Workbook()
        book.active.append(["과제명", "현재 기술사양"])
        book.active.append(["합성", "속도"])
        variants.append(book)
        book = Workbook()
        book.active.title = "고시기준본"
        book.active.append(V15_CATALOGUE_COLUMNS + ["이어쓰기"])
        variants.append(book)
        book = Workbook()
        book.active.title = "고시기준본"
        variants.append(book)
        book = Workbook()
        book.active.title = "다른시트"
        book.active.append(["고시일"])
        book.create_sheet("고시기준본").append(["과제명", "시행일"])
        variants.append(book)
        for index, book in enumerate(variants):
            book.save(self.output)
            book.close()
            for resume in (True, False):
                with self.subTest(variant=index, resume=resume):
                    self.assert_stopped(resume)

    def test_new_output_path_proceeds_like_v15(self):
        self.assertTrue(all(value is None for value in self.snapshot().values()))
        report = self.v16_run()
        self.assertEqual(report["counts"]["완료"], 2)
        self.assertTrue(all(value is not None for value in self.snapshot().values()))
        records = [json.loads(line) for line in self.checkpoint.read_text(encoding="utf-8").splitlines()]
        self.assertEqual([r["schema_version"] for r in records], [3, 3])
        book = load_workbook(BytesIO(self.output.read_bytes()), read_only=True)
        self.addCleanup(book.close)
        self.assertIn("고시일", next(book["고시기준본"].iter_rows(max_row=1, values_only=True)))

    def test_normal_v16_resume_and_restart_work_as_in_v15(self):
        self.v16_run()
        calls = len(self.clients)
        report = self.v16_run()
        self.assertEqual(report["reused_count"], 2)
        self.assertEqual(len(self.clients), calls)
        self.assertFalse(report["warnings"])
        report = self.v16_run(resume=False)
        self.assertEqual(report["reused_count"], 0)
        self.assertEqual(len(self.clients), calls + 2)
        records = self.checkpoint.read_text(encoding="utf-8").splitlines()
        self.assertEqual(len(records), 2)

    def test_damaged_v16_checkpoint_still_warns_and_restarts(self):
        self.v16_run()
        for raw in ("{broken\n", "NaN\n", '{"schema_version":"2"}\n', '{"schema_version":true}\n',
                    '{"schema_version":2.0}\n', "[2]\n", "﻿\n", '{"schema_version":3}\n'):
            self.checkpoint.write_text(raw, encoding="utf-8")
            with self.subTest(raw=raw):
                report = self.v16_run()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])
                self.assertTrue(all(noun_style(message) for message in report["warnings"]))
        self.checkpoint.write_bytes(b"\xff\xfe invalid utf-8\n")
        report = self.v16_run()
        self.assertEqual(report["reused_count"], 0)
        self.assertTrue(report["warnings"])

    def test_v16_checkpoint_without_workbook_resumes(self):
        self.v16_run()
        self.output.unlink()
        calls = len(self.clients)
        report = self.v16_run()
        self.assertEqual(report["reused_count"], 2)
        self.assertEqual(len(self.clients), calls)


class ScreenMessageV16Tests(unittest.TestCase):
    """Spec 4.8: only screen-only strings change; saved sentences stay v15."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()

    def workbook(self, path, rows):
        book = Workbook()
        book.active.append(("과제명", "현재 기술사양"))
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()

    def test_progress_status_and_warnings_are_noun_style_and_never_stored(self):
        source, output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.workbook(source, [("합성", "속도 20 GHz"), ("", ""), ("(예시) 합성", ""), ("합성", "중복")])
        batch.run_excel(source, output, {}, lambda: Client(candidates=("advanced:1",)), catalogue=catalogue(), make_forms=False)
        messages = []
        report = batch.run_excel(source, output, {"num_reviewers": 1}, lambda: Client(candidates=("advanced:1",)),
            catalogue=catalogue(), make_forms=False, progress_cb=messages.append, status_cb=messages.append)
        self.assertTrue(report["warnings"])
        self.assertGreaterEqual(len(set(messages)), 4)
        stored = "".join(Path(str(output) + suffix).read_text(encoding="utf-8")
                         for suffix in (".checkpoint.jsonl", ".audit.jsonl"))
        book = load_workbook(BytesIO(output.read_bytes()))
        self.addCleanup(book.close)
        stored += json.dumps([list(sheet.values) for sheet in book], ensure_ascii=False, default=str)
        for message in set(messages) | set(report["warnings"]):
            with self.subTest(message=message):
                self.assertTrue(noun_style(message), message)
                self.assertNotIn(message, stored)
        self.assertIn("입력 제외: 빈 과제명 1건, 예시 1건, 중복 과제명 1건. 중복은 첫 행 유지", report["warnings"])

    def test_corrupt_checkpoint_and_audit_failure_warnings_are_noun_style(self):
        source, output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.workbook(source, [("합성", "속도 20 GHz")])
        batch.run_excel(source, output, {}, lambda: Client(), catalogue=catalogue(), make_forms=False)
        Path(str(output) + ".checkpoint.jsonl").write_text("{broken\n", encoding="utf-8")
        audit = Path(str(output) + ".audit.jsonl")
        audit.unlink()
        audit.mkdir()
        report = batch.run_excel(source, output, {}, lambda: Client(), catalogue=catalogue(), make_forms=False)
        self.assertEqual(len(report["warnings"]), 2)
        self.assertTrue(all(noun_style(message) for message in report["warnings"]), report["warnings"])

    def test_regeneration_persistence_warnings_are_noun_style(self):
        source, output = self.root / "input.xlsx", self.root / "output.xlsx"
        self.workbook(source, [("A", "속도 20 GHz"), ("B", "속도 20 GHz")])
        report = batch.run_excel(source, output, {}, lambda: Client(), catalogue=catalogue(), make_forms=False)

        def generated(index):
            with patch.object(batch, "export_documents", side_effect=documents):
                return batch.regenerate_documents(report["results"][index], self.root / ("retry" + str(index)))
        batch.persist_regenerated_documents(report, report["results"][0], generated(0))
        stale = batch.persist_regenerated_documents(report, report["results"][1], generated(1))
        warnings = list(stale["persistence"]["warnings"])
        for failing in ("checkpoint", "workbook", "audit"):
            fresh = batch.run_excel(source, output, {}, lambda: Client(), catalogue=catalogue(), make_forms=False, resume=False)
            append, save = batch._append_checkpoint, batch._save_workbook

            def append_failure(path, encoded, failing=failing):
                kind = "checkpoint" if str(path).endswith("checkpoint.jsonl") else "audit"
                if kind == failing:
                    raise OSError("SECRET")
                return append(path, encoded)

            def workbook_failure(*args, failing=failing):
                if failing == "workbook":
                    raise OSError("SECRET")
                return save(*args)
            with patch.object(batch, "export_documents", side_effect=documents):
                regenerated = batch.regenerate_documents(fresh["results"][0], self.root / ("fail-" + failing))
            with patch.object(batch, "_append_checkpoint", side_effect=append_failure), \
                    patch.object(batch, "_save_workbook", side_effect=workbook_failure):
                warnings += batch.persist_regenerated_documents(fresh, fresh["results"][0], regenerated)["persistence"]["warnings"]
        self.assertEqual(len(warnings), 4)
        self.assertEqual(len(set(warnings)), 4)
        for message in warnings:
            self.assertTrue(noun_style(message), message)

    def test_persisted_program_sentences_and_template_are_v15_bytes(self):
        import batch_v15
        for module in (batch_v15, batch):
            path = self.root / (module.__name__ + "-template.xlsx")
            module.make_template(path)
        def cells(name):
            book = load_workbook(BytesIO((self.root / (name + "-template.xlsx")).read_bytes()))
            try:
                return [[list(row) for row in sheet.values] for sheet in book]
            finally:
                book.close()
        # Spec 4.4 exception: only the guide text and the header marks differ from v15.
        old, new = cells("batch_v15"), cells("batch_v16")
        self.assertEqual(len(old), len(new))
        self.assertEqual([row[0] for row in new[0]], TEMPLATE_GUIDE)
        self.assertEqual(new[1][0], TEMPLATE_HEADER)
        strip = lambda header: [value.replace("(필수)", "").replace("(선택)", "") for value in header]
        self.assertEqual(strip(old[1][0]), strip(new[1][0]))
        self.assertEqual(old[1][1:], new[1][1:])
        self.assertEqual(new[1][1:], TEMPLATE_EXAMPLES)
        source = self.root / "input.xlsx"
        self.workbook(source, [("합성", "속도 20 GHz")])

        def failing_factory():
            raise RuntimeError("SECRET")
        stored = {}
        for module, cat in ((batch_v15, catalogue(v15=True)), (batch, catalogue())):
            output = self.root / (module.__name__ + "-out.xlsx")
            with patch.object(module, "export_documents", side_effect=OSError("SECRET")):
                report = module.run_excel(source, output, {"num_reviewers": 1}, failing_factory, catalogue=cat,
                                          make_forms=True, output_dir=self.root / (module.__name__ + "-forms"))
            result = report["results"][0]
            stored[module.__name__] = (result["unresolved_reasons"], result["documents"]["errors"])
        self.assertEqual(stored["batch_v15"], stored["batch_v16"])
        self.assertEqual(stored["batch_v16"], (["과제 분류에 실패했습니다. 입력 또는 모델 연결 설정을 확인하세요."],
                                               ["HWPX 문서 저장에 실패했습니다. 저장 위치와 템플릿을 확인하세요."]))


RUN_COMPONENT = re.compile(r"\d{8}_\d{6}_[0-9a-f]{12}")


class V15SavedTextParityTests(unittest.TestCase):
    """Spec 6 '저장 문장 동일': same batch and mock responses through v15 and v16."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / "input.xlsx"
        book = Workbook()
        book.active.title = "분류대상"
        book.active.append(("과제명", "현재 기술사양", "부처 세부기준"))
        for name in ("합성 과제 A", "합성 과제 B", "합성 과제 C", "합성 과제 D"):
            book.active.append((name, "속도 20 GHz", "합성 부처 기준" if name.endswith("A") else ""))
        book.save(self.source)
        book.close()

    @staticmethod
    def factory(complete=("advanced:1", "core:1"), pending=("advanced:1",)):
        """Deterministic per-task clients: 완료, 후보제외, 검토필요, 오류 (workers=1)."""
        calls = []
        long_text = "합성 장문 근거 =+@-한글😀 " * 1400

        def summary(data, out):
            if data.get("stage") == "summary":
                out["rationale"] = long_text
            return out

        def build():
            calls.append(None)
            index = len(calls)
            if index == 1:
                return Client(candidates=complete, transform=summary)
            if index == 2:
                return Client(candidate=False)
            if index == 3:
                return Client(candidates=pending, facts=[])
            raise RuntimeError("SECRET")
        return build

    def run_version(self, module, cat, name, options, candidates):
        folder = self.root / name
        folder.mkdir()
        output = folder / "output.xlsx"
        report = module.run_excel(self.source, output, options, self.factory(*candidates), catalogue=cat, workers=1,
                                  make_forms=True, output_dir=folder / "forms")
        return folder, output, report

    def comparable(self, folder, mapping):
        def text(value):
            value = substitute(value, mapping)
            return RUN_COMPONENT.sub("<run>", value.replace(str(folder), "<root>"))

        def walk(value):
            if isinstance(value, str):
                return text(value)
            if isinstance(value, list):
                return [walk(item) for item in value]
            if isinstance(value, dict):
                return {key: walk(item) for key, item in value.items()}
            return value
        return walk, text

    def records(self, output):
        return [json.loads(line) for line in Path(str(output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines()]

    def sheets(self, output, text, v16):
        book = load_workbook(BytesIO(output.read_bytes()), data_only=False)
        try:
            self.assertEqual(book.sheetnames, ["분류결과", "요약", "고시대조", "조건별근거", "고시기준본"])
            merged = {}
            for name in book.sheetnames:
                header, rows = sheet_rows(book, name)
                if name == "고시기준본":
                    if v16:
                        position = header.index("고시일")
                        self.assertEqual(header[position + 1], "시행일")
                        header = header[:position] + header[position + 1:]
                        rows = [row[:position] + row[position + 1:] for row in rows]
                    self.assertEqual(header, V15_CATALOGUE_COLUMNS)
                    kind, raw = header.index("기록종류"), header.index("저장원문")
                    for row in rows:
                        if row[kind] == "기준본":
                            row[raw] = "<catalogue_manifest>"
                        else:
                            source = json.loads(row[raw])
                            source.pop("announcement_date", None)
                            row[raw] = json.dumps(source, ensure_ascii=False)
                merged[name] = (header, [[text(cell) if isinstance(cell, str) else cell for cell in row] for row in rows])
            return merged
        finally:
            book.close()

    def compare(self, old_module, old_cat, new_cat, options, candidates=()):
        from test_v16_hwpx import comparable_hwpx, hwpx_rows
        old_folder, old_output, old_report = self.run_version(old_module, old_cat, "v15", options, candidates)
        new_folder, new_output, new_report = self.run_version(batch, new_cat, "v16", options, candidates)
        old_records, new_records = self.records(old_output), self.records(new_output)
        self.assertEqual([r["schema_version"] for r in old_records], [2] * len(old_records))
        self.assertEqual([r["schema_version"] for r in new_records], [3] * len(new_records))
        self.assertEqual(len(old_records), len(new_records))
        mapping = catalogue_mapping(old_cat, new_cat)
        for old, new in zip(old_records, new_records):
            self.assertEqual(old["task_fingerprint"], new["task_fingerprint"])
            mapping[old["options_fingerprint"]] = new["options_fingerprint"]
        old_walk, old_text = self.comparable(old_folder, mapping)
        new_walk, new_text = self.comparable(new_folder, {})
        statuses = [r["status"] for r in new_records]
        for old, new in zip(old_records, new_records):
            old_saved, new_saved = normalise(old_walk(old)), normalise(new_walk(new))
            old_saved.pop("schema_version")
            new_saved.pop("schema_version")
            self.assertEqual(old_saved, new_saved)
        self.assertTrue(all("announcement_date" in source for r in new_records if r["notice_review"]["catalogue_manifest"]
                            for source in r["notice_review"]["sources"]))
        old_sheets = self.sheets(old_output, old_text, False)
        new_sheets = self.sheets(new_output, new_text, True)
        for name in old_sheets:
            with self.subTest(sheet=name):
                self.assertEqual(old_sheets[name], new_sheets[name])
        old_files = [entry for r in old_records for entry in r["documents"]["files"]]
        new_files = [entry for r in new_records for entry in r["documents"]["files"]]
        self.assertEqual(len(old_files), len(new_files))
        self.assertGreater(len(new_files), 0)
        for old, new in zip(old_files, new_files):
            self.assertEqual((old["status"], new["status"]), ("완료", "완료"))
            old_rows = comparable_hwpx(hwpx_rows(old["path"]), mapping, v16=False)
            new_rows = comparable_hwpx(hwpx_rows(new["path"]), {}, v16=True)
            self.assertEqual(old_rows, new_rows)
        old_events = [json.loads(line) for line in Path(str(old_output) + ".audit.jsonl").read_text(encoding="utf-8").splitlines()]
        new_events = [json.loads(line) for line in Path(str(new_output) + ".audit.jsonl").read_text(encoding="utf-8").splitlines()]
        drop = ("schema_version", "run_id", "timestamp")
        self.assertEqual([{k: substitute(v, mapping) if isinstance(v, str) else v for k, v in e.items() if k not in drop} for e in old_events],
                         [{k: v for k, v in e.items() if k not in drop} for e in new_events])
        self.assertEqual({e["schema_version"] for e in new_events}, {3})
        return statuses, old_sheets

    def test_synthetic_batch_saves_identical_text_after_spec_mapping(self):
        import batch_v15
        statuses, sheets = self.compare(batch_v15, catalogue(v15=True, catalogue_id="synthetic-v15"),
                                        catalogue(catalogue_id="synthetic-v16"), {"num_reviewers": 3})
        self.assertEqual(statuses, ["완료", "후보제외", "검토필요", "오류"])
        header, rows = sheets["분류결과"]
        self.assertTrue(any(len(row[header.index("AI 종합 검토의견")] or "") > batch.CELL_CHUNK_SIZE for row in rows))

    # v15 comparison means "same request size -> same requests": v16 runs at the
    # former 24,000-character size without the screen clause limit (spec 4.1).
    @patch("notice_review_v16.SCREEN_CLAUSE_LIMIT", None)
    @patch("notice_review_v16.DEFAULT_CHUNK_CHARS", 24000)
    def test_installed_catalogues_with_long_manifests_save_identical_text(self):
        import batch_v15
        import notice_catalogue_v15
        from notice_catalogue_v16 import load_catalogue
        old_cat, new_cat = notice_catalogue_v15.load_catalogue(), load_catalogue()
        self.assertEqual(new_cat.manifest["catalogue_id"], "kr-notices-v16-1")
        candidate = old_cat.by_family["advanced"][0]["clause_id"]
        self.compare(batch_v15, old_cat, new_cat, {"num_reviewers": 1}, ((candidate,), (candidate,)))
        # The longer v16 manifest needs its own continuation rows (merged above).
        for folder in ("v15", "v16"):
            book = load_workbook(BytesIO((self.root / folder / "output.xlsx").read_bytes()), read_only=True)
            try:
                parts = [row[-1] for row in book["고시기준본"].iter_rows(min_row=2, values_only=True)]
            finally:
                book.close()
            self.assertTrue(any(part != "1/1" for part in parts), folder)


if __name__ == "__main__":
    unittest.main()
