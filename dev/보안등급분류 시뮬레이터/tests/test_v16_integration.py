"""Synthetic real engine/batch/document integration through the v16 desktop API.

The first two classes are the v15 integration checks retargeted to v16. The
rest are v16 release-level checks (spec 6):

- the full bridge pipeline (single and batch) sends v15-identical system
  prompts and v15-equal user prompts after the 기준본 파생값 substitution,
  also when only v16 receives task and policy text that needs cleaning;
- every screen-only Python string across bridge, host, batch, classification
  and notice review is a noun-ending 개조식 (exception texts and persisted
  sentences intentionally stay as in v15 and are not scanned);
- no v16 runtime module imports a v15/v14 module except llm_client_v14.
"""
import ast
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import batch_v16
from credential_store import KeyManager
import desktop_bridge_v16
from desktop_bridge_v16 import DesktopBridge
import security_grade_classifier_v16 as host
import v16_domain as d
from test_v16_engine import Client, RecordingClient
from test_v16_notice_review import (STYLE_ENDINGS, catalogue, catalogue_mapping, fact, normalise,
                                    substitute)


ROOT = Path(__file__).resolve().parents[1]


class MemoryBackend:
    def __init__(self):
        self.values = {}

    def get(self, provider):
        return self.values.get(provider)

    def set(self, provider, key):
        self.values[provider] = key

    def delete(self, provider):
        self.values.pop(provider, None)


def finish(bridge, started):
    if not started["ok"]:
        raise AssertionError(started)
    job_id = started["data"]["job_id"]
    deadline = time.monotonic() + 120
    while time.monotonic() < deadline:
        reply = bridge.get_job(job_id, 0)
        if not reply["ok"]:
            raise AssertionError(reply)
        if reply["data"]["state"] in ("done", "cancelled", "failed"):
            return reply["data"]
        time.sleep(.01)
    raise AssertionError("합성 작업 시간 초과")


class DesktopIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source, self.output = self.root / "in.xlsx", self.root / "out.xlsx"
        book = Workbook()
        book.active.append(["과제명", "현재 기술사양"])
        for name in ("합성 A", "합성 B"):
            book.active.append([name, "속도 20 GHz"])
        book.save(self.source)
        book.close()
        self.cat = catalogue()
        self.loader_calls = 0
        self.clients = []
        self.paths = dict(input_excel=self.source, output_excel=self.output,
                          output_directory=self.root / "documents")

        def loader():
            self.loader_calls += 1
            return self.cat

        def factory(opts, key):
            self.assertEqual(key, "synthetic-secret")
            client = Client()
            self.clients.append(client)
            return client

        self.bridge = DesktopBridge(key_manager=KeyManager(MemoryBackend(), environ={}),
            catalogue_loader=loader, client_factory=factory,
            path_selector=lambda purpose: self.paths[purpose])
        self.assertTrue(self.bridge.save_key("claude", "synthetic-secret")["ok"])
        self.tokens = {purpose: self.bridge.select_path(purpose)["data"]["path_token"]
                       for purpose in self.paths}

    def batch(self):
        return finish(self.bridge, self.bridge.start_batch(dict(
            input_token=self.tokens["input_excel"], output_token=self.tokens["output_excel"],
            output_dir_token=None, options={"num_reviewers": 1}, workers=2,
            resume=False, make_forms=False)))

    def export(self, job_id, result):
        return finish(self.bridge, self.bridge.start_export(dict(job_id=job_id,
            task_fingerprint=result["task_fingerprint"],
            output_dir_token=self.tokens["output_directory"])))

    def latest(self):
        latest = {}
        for line in Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines():
            result = json.loads(line)
            latest[result["task_fingerprint"]] = result
        return latest

    def test_real_batch_A_B_A_exports_preserve_both_records_without_models(self):
        initial = self.batch()
        self.assertEqual(initial["state"], "done")
        a, b = initial["result"]["results"]
        self.assertEqual([a["status"], b["status"]], ["완료", "완료"])
        calls = len(self.clients), self.loader_calls
        with patch("notice_catalogue_v16.load_catalogue", side_effect=AssertionError("export catalogue")), \
                patch("llm_client_v14.LLMClient", side_effect=AssertionError("export model")), \
                patch("llm_json_v16.TolerantLLMClient", side_effect=AssertionError("export model")), \
                patch.object(self.bridge._keys, "resolve", side_effect=AssertionError("export credentials")):
            first = self.export(initial["job_id"], a)
            second = self.export(initial["job_id"], b)
            third = self.export(first["job_id"], first["result"])
        self.assertEqual((len(self.clients), self.loader_calls), calls)
        for item in (first, second, third):
            self.assertEqual(item["state"], "done")
            self.assertEqual(item["result"]["schema_version"], 3)
            self.assertEqual(item["result"]["documents"]["status"], "완료")
            d.validate_result(item["result"])
            for entry in item["result"]["documents"]["files"]:
                self.assertTrue(Path(entry["path"]).is_file())
        latest = self.latest()
        self.assertEqual(latest[a["task_fingerprint"]], third["result"])
        self.assertEqual(latest[b["task_fingerprint"]], second["result"])
        book = load_workbook(BytesIO(self.output.read_bytes()))
        self.addCleanup(book.close)
        self.assertIn("고시기준본", book.sheetnames)
        header = next(book["고시기준본"].iter_rows(min_row=1, max_row=1, values_only=True))
        self.assertEqual(header.index("고시일") + 1, header.index("시행일"))

    def test_checkpoint_failure_keeps_durable_context_for_retry(self):
        initial = self.batch()
        a, b = initial["result"]["results"]
        before = deepcopy(self.latest())
        append = batch_v16._append_checkpoint

        def fail_checkpoint(path, encoded):
            if str(path).endswith(".checkpoint.jsonl"):
                raise OSError("synthetic secret path failure")
            return append(path, encoded)

        with patch.object(batch_v16, "_append_checkpoint", side_effect=fail_checkpoint):
            failed = self.export(initial["job_id"], a)
        self.assertEqual(failed["result"]["documents"]["status"], "완료")
        self.assertEqual(self.latest(), before)
        self.assertTrue(any(event["stage"] == "persistence" and "실패" in event["message"]
                            for event in failed["events"]))
        self.assertEqual(failed["events"][-1]["stage"], "persistence")
        self.assertIn("실패", failed["events"][-1]["message"])
        self.assertNotIn("synthetic secret path failure", json.dumps(failed, ensure_ascii=False))
        retried = self.export(failed["job_id"], failed["result"])
        other = self.export(initial["job_id"], b)
        self.assertEqual(retried["state"], "done")
        self.assertEqual(other["state"], "done")
        latest = self.latest()
        self.assertEqual(latest[a["task_fingerprint"]], retried["result"])
        self.assertEqual(latest[b["task_fingerprint"]], other["result"])

    def test_real_batch_progress_adapter_exposes_only_fixed_stages_and_messages(self):
        job = self.batch()
        allowed = {"preflight", "facts", "notices", "candidate", "criteria", "agents", "summary", "documents", "persistence"}
        fixed = set(desktop_bridge_v16._MESSAGES.values()) | set(desktop_bridge_v16._EVENTS.values())
        self.assertTrue(job["events"])
        for event in job["events"]:
            self.assertIn(event["stage"], allowed)
            self.assertIn(event["message"], fixed)
            self.assertNotIn("{'stage':", event["message"])
            self.assertNotIn("합성 A", event["message"])
            self.assertNotIn("합성 B", event["message"])
        self.assertEqual(job["events"][-1]["stage"], "persistence")


class HostBoundaryIntegrationTests(unittest.TestCase):
    def test_selfcheck_imports_real_delayed_sdks_without_network_or_client_creation(self):
        import socket
        import openai
        import anthropic
        code = """
import json
import socket
def forbidden(*args, **kwargs):
    raise AssertionError('SDK import made a network request')
socket.socket.connect = forbidden
socket.create_connection = forbidden
import security_grade_classifier_v16 as host
print(json.dumps(host._selfcheck_sdks()))
"""
        fresh = subprocess.run([sys.executable, "-X", "utf8", "-B", "-c", code],
            cwd=Path(host.__file__).parent, capture_output=True, text=True, timeout=60)
        self.assertEqual(fresh.returncode, 0, fresh.stderr)
        fresh_result = json.loads(fresh.stdout)
        with patch.object(socket.socket, "connect", side_effect=AssertionError("network")), \
                patch.object(socket, "create_connection", side_effect=AssertionError("network")), \
                patch.object(openai.OpenAI, "__init__", side_effect=AssertionError("client")), \
                patch.object(anthropic.Anthropic, "__init__", side_effect=AssertionError("client")):
            result = host._selfcheck_sdks()
        self.assertEqual(result["openai"]["version"], "2.47.0")
        self.assertEqual(result["anthropic"]["version"], "1.3.0")
        self.assertEqual(result["openai"]["client_class"], "OpenAI")
        self.assertEqual(result["anthropic"]["client_class"], "Anthropic")
        self.assertEqual(result, fresh_result)

    def test_sdk_import_diagnostics_are_exclusive_to_selfcheck_branch(self):
        tree = ast.parse(Path(host.__file__).read_text(encoding="utf-8"))
        guarded_calls = []
        all_calls = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "_selfcheck_sdks":
                all_calls.append(node)
            if isinstance(node, ast.If) and isinstance(node.test, ast.Name) and node.test.id == "checking":
                guarded_calls.extend(child for statement in node.body for child in ast.walk(statement)
                    if isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
                    and child.func.id == "_selfcheck_sdks")
        self.assertEqual(len(all_calls), 1)
        self.assertEqual(guarded_calls, all_calls)

    def test_native_message_allowlist_rejects_dunder_builtins_and_malformed_packets(self):
        server = host.LocalAssetServer()
        self.addCleanup(server.close)
        self.assertEqual(len(host.PUBLIC_METHODS), 12)
        self.assertEqual(len(set(host.PUBLIC_METHODS)), 12)
        self.assertEqual(host.PUBLIC_METHODS[-1], "save_input_template")
        self.assertEqual(tuple(host.PUBLIC_METHODS), tuple(desktop_bridge_v16.PUBLIC_METHODS))
        for name in host.PUBLIC_METHODS:
            packet = json.dumps([name, "[]", "1234567890"])
            self.assertTrue(host._allow_native_message(server, server.url, packet))
            self.assertFalse(host._allow_native_message(server, "https://example.invalid/", packet))
        forbidden = ("__str__", "__class__.__base__.__subclasses__", "_keys.resolve", "_keys._backend.get",
                     "_paths.clear", "_clause_details.get", "_get_catalogue", "_pywebviewAlert", "console",
                     "pywebviewStateUpdate", "pywebviewStateDelete", "pywebviewMoveWindow",
                     "pywebviewEventHandler", "pywebviewAsyncCallback")
        for name in forbidden:
            with self.subTest(name=name):
                self.assertFalse(host._allow_native_message(server, server.url, json.dumps([name, "[]", "safe-id"])))
        for packet in ("null", "not JSON", "[]", json.dumps([{}, "[]", "safe-id"]),
                       json.dumps(["get_initial_state", {}, "safe-id"]),
                       json.dumps(["get_initial_state", "{}", "safe-id"]),
                       json.dumps(["get_initial_state", "[NaN]", "safe-id"]),
                       json.dumps(["get_initial_state", '[{"value":1,"value":2}]', "safe-id"]),
                       json.dumps(["get_initial_state", "[]", 'bad"id']),
                       json.dumps(["get_initial_state", "[]", "x" * 129])):
            with self.subTest(packet=packet):
                self.assertFalse(host._allow_native_message(server, server.url, packet))

    def test_host_bridge_and_screen_mock_name_the_same_twelve_functions(self):
        """Follow-up spec 4.5/6: host allowlist, bridge PUBLIC_METHODS and the UI mock agree."""
        from tests.test_v16_frontend import BRIDGE_FUNCTIONS
        host_names, bridge_names = tuple(host.PUBLIC_METHODS), tuple(desktop_bridge_v16.PUBLIC_METHODS)
        self.assertEqual(host_names, bridge_names)
        self.assertEqual(len(host_names), 12)
        self.assertEqual(len(BRIDGE_FUNCTIONS), 12)
        self.assertEqual(set(BRIDGE_FUNCTIONS), set(host_names))
        self.assertEqual(set(BRIDGE_FUNCTIONS), set(bridge_names))
        for name in bridge_names:
            with self.subTest(name=name):
                self.assertTrue(callable(getattr(DesktopBridge, name, None)))

    def test_actual_static_server_serves_only_local_asset_snapshot(self):
        server = host.LocalAssetServer().start()
        self.addCleanup(server.close)
        self.assertEqual(server.root.name, "ui_v16")
        results = host._verify_static_server(server)
        self.assertTrue(results)
        self.assertTrue(all(entry["status"] in (403, 404) for entry in results))
        import http.client
        from urllib.parse import urlsplit
        for name in ("index.html", "styles.css", "app.js"):
            url = server.url.replace("index.html", name)
            self.assertTrue(server.allowed(url))
            connection = http.client.HTTPConnection(server.authority, timeout=3)
            try:
                connection.request("GET", urlsplit(url).path)
                response = connection.getresponse()
                self.assertEqual(response.status, 200)
                self.assertIn("default-src 'none'", response.getheader("Content-Security-Policy"))
                self.assertEqual(response.read(), (server.root / name).read_bytes())
            finally:
                connection.close()
        for url in ("https://example.invalid/", "file:///C:/Windows/win.ini", server.origin + "/credential_store.py",
                    server.origin + "/resources_v16/notices/manifest.json"):
            self.assertFalse(server.allowed(url))

    def test_runtime_absence_emits_fixed_guidance_before_keys_and_legacy_import(self):
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "runtime-absent.json"
            with patch.object(host, "require_webview2", side_effect=host.HostError(host.RUNTIME_MESSAGE)), \
                    patch("credential_store.WindowsCredentialBackend", side_effect=AssertionError("real credentials")):
                self.assertEqual(host.run_app(self_check_path=report_path), 1)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertFalse(report["pass"])
            self.assertEqual(report["failures"], [host.RUNTIME_MESSAGE])
            self.assertEqual(report["fake_credential_operations"], [])
            self.assertTrue(report["profile_cleaned"])
            import builtins
            original_import = builtins.__import__

            def missing_webview(name, *args, **kwargs):
                if name == "webview":
                    raise ModuleNotFoundError("synthetic hidden import diagnostic")
                return original_import(name, *args, **kwargs)

            with patch.object(host, "require_webview2", return_value="155.0.0.0"), \
                    patch("builtins.__import__", side_effect=missing_webview), \
                    patch("credential_store.WindowsCredentialBackend", side_effect=AssertionError("real credentials")):
                self.assertEqual(host.run_app(self_check_path=report_path), 1)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["failures"], [host.INITIALIZATION_MESSAGE])
            self.assertEqual(report["fake_credential_operations"], [])
            self.assertNotIn("synthetic hidden import diagnostic", report_path.read_text(encoding="utf-8"))
        guard = host._NoLegacyRenderer()
        for module in ("webview.platforms.mshtml", "webview.platforms.cef"):
            with self.assertRaisesRegex(ImportError, "WebView2 Runtime"):
                guard.find_spec(module)
        self.assertIsNone(guard.find_spec("webview.platforms.edgechromium"))


# ---------------------------------------------------------------------------
# Spec 6 '지시문 동일' through the whole bridge pipeline (v15 bridge vs v16).
# ---------------------------------------------------------------------------
class _Pipeline:
    """One bridge (v15 or v16) with a recording client per model connection."""

    def __init__(self, module, loader, root, client_options):
        self.root = Path(root)
        self.root.mkdir()
        self.clients = []
        self.client_options = list(client_options)
        self.paths = {}

        def factory(opts, key):
            if key != "synthetic-secret":
                raise AssertionError("unexpected key")
            options = self.client_options[min(len(self.clients), len(self.client_options) - 1)]
            client = RecordingClient(**options)
            self.clients.append(client)
            return client

        self.bridge = module.DesktopBridge(key_manager=KeyManager(MemoryBackend(), environ={}),
            catalogue_loader=loader, client_factory=factory,
            path_selector=lambda purpose: self.paths[purpose])
        reply = self.bridge.save_key("claude", "synthetic-secret")
        if not reply["ok"]:
            raise AssertionError(reply)

    def token(self, purpose, path):
        self.paths[purpose] = path
        reply = self.bridge.select_path(purpose)
        if not reply["ok"]:
            raise AssertionError(reply)
        return reply["data"]["path_token"]

    def single(self, task, options):
        return finish(self.bridge, self.bridge.start_single(dict(task=task, options=options, output_dir_token=None)))

    def batch(self, source, options):
        return finish(self.bridge, self.bridge.start_batch(dict(
            input_token=self.token("input_excel", source),
            output_token=self.token("output_excel", self.root / "output.xlsx"),
            output_dir_token=None, options=options, workers=1, resume=False, make_forms=False)))


class BridgePromptParityTests(unittest.TestCase):
    """Same synthetic tasks through desktop_bridge_v15 and desktop_bridge_v16."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()

    def pipelines(self, name, old_loader, new_loader, client_options):
        import desktop_bridge_v15
        old = _Pipeline(desktop_bridge_v15, old_loader, self.root / (name + "-v15"), client_options)
        new = _Pipeline(desktop_bridge_v16, new_loader, self.root / (name + "-v16"), client_options)
        return old, new

    def assert_prompt_parity(self, old, new, mapping):
        """System prompts byte-identical; user prompts equal after the spec 6 mapping."""
        self.assertEqual(len(old.clients), len(new.clients))
        self.assertTrue(new.clients)
        raw_differences = 0
        systems = set()
        for old_client, new_client in zip(old.clients, new.clients):
            self.assertEqual(len(old_client.wire), len(new_client.wire))
            self.assertTrue(new_client.wire)
            for (old_system, old_user, old_temperature), (new_system, new_user, new_temperature) in zip(
                    old_client.wire, new_client.wire):
                self.assertEqual(old_system.encode("utf-8"), new_system.encode("utf-8"))
                self.assertEqual(old_temperature, new_temperature)
                self.assertEqual(normalise(json.loads(substitute(old_user, mapping))), normalise(json.loads(new_user)))
                raw_differences += old_user != new_user
                systems.add(new_system)
        # The mapping was actually exercised: catalogue-derived values reached
        # the requests, and both notice-review and engine prompts were compared.
        self.assertGreater(raw_differences, 0)
        self.assertGreaterEqual(len(systems), 2)
        return sum(len(client.wire) for client in new.clients)

    def assert_saved_equal(self, old_result, new_result, mapping):
        self.assertEqual((old_result["schema_version"], new_result["schema_version"]), (2, 3))
        d.validate_result(new_result)
        old_saved = normalise(json.loads(substitute(json.dumps(old_result, ensure_ascii=False), mapping)))
        new_saved = normalise(deepcopy(new_result))
        for saved in (old_saved, new_saved):
            saved.pop("schema_version")
        self.assertEqual(old_saved, new_saved)

    def assert_screen_events(self, job):
        fixed = set(desktop_bridge_v16._MESSAGES.values()) | set(desktop_bridge_v16._EVENTS.values())
        self.assertTrue(job["events"])
        for event in job["events"]:
            self.assertIn(event["message"], fixed)
            self.assertEqual(style_violations(event["message"]), [])

    def test_single_flow_sends_v15_prompts_through_the_bridge(self):
        target_task = {"과제명": "합성 단건 목표", "현재 기술사양": "속도 5 GHz", "개발목표 사양": "속도 20 GHz",
                       "부처 세부기준": "합성 부처 기준"}
        scenarios = (
            ("impact_three", {"과제명": "합성 단건", "현재 기술사양": "속도 20 GHz"}, {"num_reviewers": 3},
             dict(candidates=("advanced:1", "core:1"))),
            ("checklist_policy_target", target_task, {"num_reviewers": 1, "method": "checklist"},
             dict(candidates=("advanced:1",), chosen=(1, 8),
                  facts=[fact("속도 5 GHz", number="5"), fact(field="개발목표 사양", basis="target", identity="t1")])),
            ("manual_weights", {"과제명": "합성 직접 선정", "현재 기술사양": "속도 20 GHz"},
             {"num_reviewers": 1, "selection_mode": "manual", "selected_items": [1, 2], "weights": {"1": .8, "2": 1.2}},
             dict(candidates=("advanced:1",))),
        )
        for name, task, options, client_options in scenarios:
            with self.subTest(scenario=name):
                old_cat = catalogue(v15=True, catalogue_id="synthetic-v15")
                new_cat = catalogue(catalogue_id="synthetic-v16")
                old, new = self.pipelines(name, lambda: old_cat, lambda: new_cat, [client_options])
                old_job, new_job = old.single(task, options), new.single(task, options)
                self.assertEqual((old_job["state"], new_job["state"]), ("done", "done"))
                self.assertEqual(old_job["result"]["status"], new_job["result"]["status"])
                mapping = catalogue_mapping(old_cat, new_cat)
                mapping[old_job["result"]["options_fingerprint"]] = new_job["result"]["options_fingerprint"]
                self.assert_prompt_parity(old, new, mapping)
                self.assert_saved_equal(old_job["result"], new_job["result"], mapping)
                self.assert_screen_events(new_job)

    def test_batch_flow_sends_v15_prompts_through_the_bridge(self):
        source = self.root / "input.xlsx"
        book = Workbook()
        book.active.title = "분류대상"
        book.active.append(("과제명", "현재 기술사양", "개발목표 사양", "부처 세부기준"))
        book.active.append(("합성 배치 A", "속도 20 GHz", "", "합성 부처 기준"))
        book.active.append(("합성 배치 B", "속도 20 GHz", "", ""))
        book.active.append(("합성 배치 C", "속도 5 GHz", "속도 20 GHz", ""))
        book.active.append(("", "", "", ""))
        book.active.append(("(예시) 합성", "", "", ""))
        book.save(source)
        book.close()
        clients = [dict(candidates=("advanced:1", "core:1")), dict(candidate=False),
                   dict(candidates=("advanced:1",), facts=[])]
        for name, options in (("impact", {"num_reviewers": 3}), ("checklist", {"num_reviewers": 1, "method": "checklist"})):
            with self.subTest(scenario=name):
                old_cat = catalogue(v15=True, catalogue_id="synthetic-v15")
                new_cat = catalogue(catalogue_id="synthetic-v16")
                old, new = self.pipelines(name, lambda: old_cat, lambda: new_cat, clients)
                old_job, new_job = old.batch(source, options), new.batch(source, options)
                self.assertEqual((old_job["state"], new_job["state"]), ("done", "done"))
                old_results, new_results = old_job["result"]["results"], new_job["result"]["results"]
                self.assertEqual([r["status"] for r in new_results], ["완료", "후보제외", "검토필요"])
                self.assertEqual(old_job["result"]["counts"], new_job["result"]["counts"])
                self.assertEqual(len(old_job["result"]["warnings"]), len(new_job["result"]["warnings"]))
                mapping = catalogue_mapping(old_cat, new_cat)
                for old_result, new_result in zip(old_results, new_results):
                    self.assertEqual(old_result["task_fingerprint"], new_result["task_fingerprint"])
                    mapping[old_result["options_fingerprint"]] = new_result["options_fingerprint"]
                self.assertEqual(len(old.clients), 3)
                self.assert_prompt_parity(old, new, mapping)
                for old_result, new_result in zip(old_results, new_results):
                    self.assert_saved_equal(old_result, new_result, mapping)
                self.assert_screen_events(new_job)

    # Spec 6 'v15 대조': text needing cleaning goes to v16 only; v15 gets the
    # same text already passed through v16_domain.clean_input_text.
    MESSY_TASK = {
        "과제명": "합성 정리 단건  \t",
        "연구목표": "목표 1_x000D_\n목표 2  \n\n\n\n\n목표 3\t",
        "연구내용": "측정 20_x000D_30 구간\r\n  들여쓴 줄 \t\n끝_x000a_정리",
        "기대효과": "효과_x0009_정리_x000D_",
        "현재 기술사양": "속도 5 GHz \t\n참고 줄",
        "개발목표 사양": "속도 20 GHz_x000D_\n",
        "부처 세부기준": "합성 과제 기준_x000D_\n둘째 줄  \n\n\n\n셋째 줄",
    }
    MESSY_POLICY = "합성 공통 기준_x000D_\n둘째 줄 \t\n\n\n\n\n셋째 줄  "

    @staticmethod
    def cleaned(task):
        return {name: d.clean_input_text(value) for name, value in task.items()}

    def notice_count(self, job, name):
        text = desktop_bridge_v16._EVENTS[name]
        return sum(event["message"] == text for event in job["events"])

    def test_single_flow_with_messy_input_matches_v15_with_cleaned_input(self):
        clean_task = self.cleaned(self.MESSY_TASK)
        clean_policy = d.clean_input_text(self.MESSY_POLICY)
        # The fixture really needs cleaning: every field but the name changes beyond strip().
        for name, value in self.MESSY_TASK.items():
            if name != "과제명":
                self.assertNotEqual(clean_task[name], value.strip(), name)
        self.assertNotEqual(clean_policy, self.MESSY_POLICY.strip())
        self.assertEqual(clean_task["연구내용"], "측정 20\n30 구간\n  들여쓴 줄\n끝\n정리")
        scenarios = (
            ("messy_checklist_policy", {"num_reviewers": 1, "method": "checklist"},
             dict(candidates=("advanced:1",), chosen=(1, 8),
                  facts=[fact("속도 5 GHz", number="5"), fact(field="개발목표 사양", basis="target", identity="t1")])),
            ("messy_impact_three", {"num_reviewers": 3},
             dict(candidates=("advanced:1", "core:1"),
                  facts=[fact("속도 5 GHz", number="5"), fact(field="개발목표 사양", basis="target", identity="t1")])),
        )
        for name, options, client_options in scenarios:
            with self.subTest(scenario=name):
                old_cat = catalogue(v15=True, catalogue_id="synthetic-v15")
                new_cat = catalogue(catalogue_id="synthetic-v16")
                old, new = self.pipelines(name, lambda: old_cat, lambda: new_cat, [client_options])
                old_job = old.single(clean_task, options | {"ministry_policy": clean_policy})
                new_job = new.single(dict(self.MESSY_TASK), options | {"ministry_policy": self.MESSY_POLICY})
                self.assertEqual((old_job["state"], new_job["state"]), ("done", "done"))
                self.assertEqual(old_job["result"]["status"], new_job["result"]["status"])
                self.assertEqual(old_job["result"]["task_fingerprint"], new_job["result"]["task_fingerprint"])
                self.assertEqual(new_job["result"]["task_snapshot"], d.task_snapshot(clean_task))
                mapping = catalogue_mapping(old_cat, new_cat)
                mapping[old_job["result"]["options_fingerprint"]] = new_job["result"]["options_fingerprint"]
                self.assert_prompt_parity(old, new, mapping)
                self.assert_saved_equal(old_job["result"], new_job["result"], mapping)
                self.assert_screen_events(new_job)
                self.assertEqual(self.notice_count(new_job, "input_cleaned"), 1)
                self.assertEqual(self.notice_count(new_job, "response_trimmed"), 0)
                wire = json.dumps([client.wire for client in new.clients], ensure_ascii=False)
                for marker in ("_x000D_", "_x000a_", "_x0009_", "목표 2  "):
                    self.assertNotIn(marker, wire)
                self.assertIn("측정 20\\\\n30 구간", wire)
                saved = json.dumps(new_job["result"], ensure_ascii=False)
                for text in desktop_bridge_v16._EVENTS.values():
                    self.assertNotIn(text, saved)

    def test_batch_flow_with_messy_input_matches_v15_with_cleaned_input(self):
        header = ("과제명", "연구목표", "연구내용", "기대효과", "현재 기술사양", "개발목표 사양", "부처 세부기준")
        messy_rows = [
            tuple(self.MESSY_TASK[name] if name != "과제명" else "합성 정리 배치 A_x000D_" for name in header),
            ("합성 정리 배치 B  ", "목표\t\n\n\n\n끝", "", "", "속도 20 GHz", "", ""),
            ("_x000D_", "제외 행", "", "", "", "", ""),
            ("합성 정리 배치 B_x000D_\n", "정리 후 중복", "", "", "", "", ""),
        ]
        new_source, old_source = self.root / "messy-input.xlsx", self.root / "cleaned-input.xlsx"
        for path, rows in ((new_source, messy_rows), (old_source, None)):
            book = Workbook()
            book.active.title = "분류대상"
            book.active.append(header)
            if rows is None:
                # What v16 actually reads back, passed through the v16 cleaning.
                read = load_workbook(new_source, read_only=True)
                try:
                    rows = [tuple(d.clean_input_text("" if value is None else str(value)) for value in values)
                            for values in list(read["분류대상"].iter_rows(values_only=True))[1:]]
                finally:
                    read.close()
                self.assertEqual([row[0] for row in rows],
                                 ["합성 정리 배치 A", "합성 정리 배치 B", "", "합성 정리 배치 B"])
            for row in rows:
                book.active.append(row)
            book.save(path)
            book.close()
        read = load_workbook(new_source, read_only=True)
        try:
            raw = json.dumps(list(read["분류대상"].iter_rows(values_only=True)), ensure_ascii=False)
        finally:
            read.close()
        self.assertIn("_x000D_", raw)  # the v16 file really carries the Excel markers
        clients = [dict(candidates=("advanced:1",), chosen=(1, 8),
                        facts=[fact("속도 5 GHz", number="5"), fact(field="개발목표 사양", basis="target", identity="t1")]),
                   dict(candidates=("advanced:1", "core:1"))]
        options = {"num_reviewers": 1, "method": "checklist"}
        old_cat = catalogue(v15=True, catalogue_id="synthetic-v15")
        new_cat = catalogue(catalogue_id="synthetic-v16")
        old, new = self.pipelines("messy-batch", lambda: old_cat, lambda: new_cat, clients)
        old_job = old.batch(old_source, options | {"ministry_policy": d.clean_input_text(self.MESSY_POLICY)})
        new_job = new.batch(new_source, options | {"ministry_policy": self.MESSY_POLICY})
        self.assertEqual((old_job["state"], new_job["state"]), ("done", "done"))
        old_results, new_results = old_job["result"]["results"], new_job["result"]["results"]
        self.assertEqual(len(new_results), 2)
        self.assertEqual(old_job["result"]["counts"], new_job["result"]["counts"])
        self.assertEqual(len(old_job["result"]["warnings"]), len(new_job["result"]["warnings"]))
        mapping = catalogue_mapping(old_cat, new_cat)
        for old_result, new_result in zip(old_results, new_results):
            self.assertEqual(old_result["task_fingerprint"], new_result["task_fingerprint"])
            mapping[old_result["options_fingerprint"]] = new_result["options_fingerprint"]
        self.assertEqual(len(old.clients), 2)
        self.assert_prompt_parity(old, new, mapping)
        for old_result, new_result in zip(old_results, new_results):
            self.assert_saved_equal(old_result, new_result, mapping)
        self.assert_screen_events(new_job)
        self.assertEqual(self.notice_count(new_job, "input_cleaned"), 1)
        self.assertEqual(self.notice_count(new_job, "response_trimmed"), 0)
        # Screen-only notices never reach the stored checkpoint, audit or workbook.
        output = new.root / "output.xlsx"
        stored = "".join(Path(str(output) + suffix).read_text(encoding="utf-8")
                         for suffix in (".checkpoint.jsonl", ".audit.jsonl"))
        self.assertNotIn("_x000D_", stored)
        book = load_workbook(BytesIO(output.read_bytes()))
        self.addCleanup(book.close)
        stored += json.dumps([list(sheet.values) for sheet in book], ensure_ascii=False, default=str)
        for text in desktop_bridge_v16._EVENTS.values():
            self.assertNotIn(text, stored)

    # v15 comparison means "same request size -> same requests": v16 runs at the
    # former 24,000-character size without the screen clause limit (spec 4.1).
    @patch("notice_review_v16.SCREEN_CLAUSE_LIMIT", None)
    @patch("notice_review_v16.DEFAULT_CHUNK_CHARS", 24000)
    def test_installed_catalogues_single_flow_sends_v15_prompts(self):
        """resources_v15/notices vs resources_v16/notices via each bridge's default loader."""
        import notice_catalogue_v15
        from notice_catalogue_v16 import load_catalogue
        old_cat, new_cat = notice_catalogue_v15.load_catalogue(), load_catalogue()
        self.assertEqual(new_cat.manifest["catalogue_id"], "kr-notices-v16-1")
        self.assertNotEqual(old_cat.fingerprint, new_cat.fingerprint)
        candidate = new_cat.by_family["advanced"][0]["clause_id"]
        old, new = self.pipelines("installed", None, None, [dict(candidates=(candidate,))])
        task = {"과제명": "합성 설치 기준본", "현재 기술사양": "속도 20 GHz"}
        old_job, new_job = old.single(task, {"num_reviewers": 1}), new.single(task, {"num_reviewers": 1})
        self.assertEqual((old_job["state"], new_job["state"]), ("done", "done"))
        mapping = catalogue_mapping(old_cat, new_cat)
        mapping[old_job["result"]["options_fingerprint"]] = new_job["result"]["options_fingerprint"]
        calls = self.assert_prompt_parity(old, new, mapping)
        self.assertGreater(calls, 3)
        self.assert_saved_equal(old_job["result"], new_job["result"], mapping)
        sources = new_job["result"]["notice_review"]["sources"]
        self.assertTrue(sources and all("announcement_date" in source for source in sources))


# ---------------------------------------------------------------------------
# Spec 4.8/6 문체 검사 over every screen-only Python string.
# ---------------------------------------------------------------------------
STYLE_PATTERN = re.compile("(?:" + "|".join(STYLE_ENDINGS) + r")[)\]」』’”'\"]*(?=$|[\s.!?,…:;·])")
HANGUL = re.compile("[가-힣]")


def style_violations(text):
    return [match.group(0) for match in STYLE_PATTERN.finditer(text)]


def _tree(module_name):
    return ast.parse((ROOT / (module_name + ".py")).read_text(encoding="utf-8-sig"))


def _literal(node, names):
    """The text of a string expression; placeholders stand for runtime values."""
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    if isinstance(node, ast.JoinedStr):
        return "".join(part.value if isinstance(part, ast.Constant) else "0" for part in node.values)
    if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Add):
        left, right = _literal(node.left, names), _literal(node.right, names)
        return None if left is None or right is None else left + right
    if isinstance(node, ast.Name) and isinstance(names.get(node.id), str):
        return names[node.id]
    return None


def bridge_screen_strings(module):
    return {"_MESSAGES." + key: value for key, value in module._MESSAGES.items()} | \
           {"_EVENTS." + key: value for key, value in getattr(module, "_EVENTS", {}).items()}


def host_screen_strings(module):
    """Title, startup guidance and HostError texts (shown in the startup message box)."""
    names = {name: getattr(module, name) for name in ("TITLE", "RUNTIME_MESSAGE", "INITIALIZATION_MESSAGE")}
    found = dict(names)
    # Default names shown in the native save dialogs (v16 follow-up adds the template).
    for name in ("SAVE_FILENAME", "TEMPLATE_FILENAME"):
        if isinstance(getattr(module, name, None), str):
            found[name] = getattr(module, name)
    for node in ast.walk(_tree(module.__name__)):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "HostError":
            for argument in node.args:
                text = _literal(argument, names)
                if text is not None:
                    found["HostError:" + text] = text
        if isinstance(node, ast.Assign) and any(isinstance(target, ast.Attribute) and target.attr == "failure"
                                               for target in node.targets):
            text = _literal(node.value, names)
            if text is not None and HANGUL.search(text):
                found["failure:" + text] = text
    return found


def batch_screen_strings(module):
    """warn(...) and notify(callback, ...) texts plus the previous-result stop."""
    found = {}
    for node in ast.walk(_tree(module.__name__)):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in ("warn", "notify"):
            position = 0 if node.func.id == "warn" else 1
            if len(node.args) > position:
                text = _literal(node.args[position], {})
                if text is not None and HANGUL.search(text):
                    found[node.func.id + ":" + text] = text
    if hasattr(module, "PREVIOUS_RESULT_MESSAGE"):
        found["PREVIOUS_RESULT_MESSAGE"] = module.PREVIOUS_RESULT_MESSAGE
    return found


def progress_label_strings(classification_module, notice_module):
    """classify_task's screen-only labels and review_notices' emit() messages."""
    found = {}
    for node in ast.walk(_tree(classification_module.__name__)):
        if isinstance(node, ast.FunctionDef) and node.name == "classify_task":
            for child in ast.walk(node):
                if isinstance(child, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "labels" for t in child.targets):
                    for keyword in child.value.keywords:
                        found["labels." + keyword.arg] = _literal(keyword.value, {})
    for node in ast.walk(_tree(notice_module.__name__)):
        if isinstance(node, ast.FunctionDef) and node.name == "emit":
            for child in ast.walk(node):
                if isinstance(child, ast.Constant) and isinstance(child.value, str) and HANGUL.search(child.value):
                    found["emit:" + child.value] = child.value
    return found


class ScreenStringStyleTests(unittest.TestCase):
    """Only screen strings are scanned; exception texts and persisted sentences stay v15."""

    def collect(self, version):
        import importlib
        modules = {name: importlib.import_module(name.format(v=version)) for name in (
            "desktop_bridge_v{v}", "security_grade_classifier_v{v}", "batch_v{v}", "classification_v{v}",
            "notice_review_v{v}")}
        bridge, host_module, batch_module, engine, notices = modules.values()
        return dict(bridge=bridge_screen_strings(bridge), host=host_screen_strings(host_module),
                    batch=batch_screen_strings(batch_module), labels=progress_label_strings(engine, notices))

    def test_every_v16_screen_only_python_string_is_noun_style(self):
        groups = self.collect(16)
        self.assertEqual(len(groups["bridge"]), len(desktop_bridge_v16._MESSAGES) + len(desktop_bridge_v16._EVENTS))
        self.assertGreaterEqual(len(groups["host"]), 9)
        self.assertGreaterEqual(len(groups["batch"]), 10)
        self.assertEqual(set(groups["labels"].values()), {"등급분류 대상 여부 검토", "검토항목 확인", "AI 에이전트 독립 검토",
                                                          "표준 종합 의견 작성", "입력 근거 확인 중", "고시 조건 대조 중"})
        self.assertEqual(groups["batch"]["PREVIOUS_RESULT_MESSAGE"], "기존 결과 파일 보호를 위해 중단. 새 결과 엑셀 경로 선택 필요")
        self.assertEqual(desktop_bridge_v16._EVENTS["previous_result"], batch_v16.PREVIOUS_RESULT_MESSAGE)
        # Follow-up spec 4.4/4.5: the protected-path refusal is scanned, and the
        # screen shows exactly this bridge text; the save dialog default is scanned.
        from tests.test_v16_frontend import PROTECTED_MESSAGE
        self.assertEqual(groups["bridge"]["_EVENTS.template_protected"], PROTECTED_MESSAGE)
        self.assertEqual(set(desktop_bridge_v16._MESSAGES), {"validation", "catalogue", "credential", "busy",
                                                              "not_found", "io", "runtime"})
        self.assertEqual(groups["host"]["TEMPLATE_FILENAME"], "v16_일괄검토_입력양식.xlsx")
        for group, strings in groups.items():
            for name, text in strings.items():
                with self.subTest(group=group, name=name):
                    self.assertIsInstance(text, str)
                    self.assertTrue(text.strip())
                    self.assertEqual(style_violations(text), [])
                    for sentence in re.split(r"[.,·\n]\s*", text):
                        self.assertFalse(sentence.rstrip().endswith(STYLE_ENDINGS), sentence)

    def test_the_same_scan_finds_v15_sentence_endings(self):
        """The extractors reach real strings: v15 counterparts still end in 니다/세요."""
        groups = self.collect(15)
        for group in ("bridge", "host", "batch", "labels"):
            with self.subTest(group=group):
                self.assertTrue(groups[group])
                self.assertTrue(any(style_violations(text) for text in groups[group].values() if text), group)

    def test_exception_and_persisted_sentences_keep_v15_wording(self):
        """Not scanned on purpose: these are stored or raised exactly as in v15."""
        import batch_v15
        import classification_v15
        stored = ("과제 분류에 실패했습니다. 입력 또는 모델 연결 설정을 확인하세요.",
                  "HWPX 문서 저장에 실패했습니다. 저장 위치와 템플릿을 확인하세요.")
        for module in (batch_v15, batch_v16):
            source = Path(module.__file__).read_text(encoding="utf-8")
            for sentence in stored:
                with self.subTest(module=module.__name__, sentence=sentence):
                    self.assertIn(sentence, source)
        self.assertEqual(classification_v15._MODEL_ERROR, __import__("classification_v16")._MODEL_ERROR)
        self.assertEqual(classification_v15._NOTICE_PENDING, __import__("classification_v16")._NOTICE_PENDING)
        for text in batch_screen_strings(batch_v16).values():
            self.assertNotIn(text, stored)


# ---------------------------------------------------------------------------
# Import boundary: v16 runtime never loads v15/v14 modules (llm_client_v14 only).
# ---------------------------------------------------------------------------
V16_RUNTIME = ("security_grade_classifier_v16", "desktop_bridge_v16", "v16_domain", "notice_catalogue_v16",
               "notice_review_v16", "classification_v16", "batch_v16", "hwpx_writer_v16", "llm_json_v16")
REUSED = ("llm_client_v14", "credential_store")
LEGACY_NAME = re.compile(r"(?:^|_)v(?:[0-9]|1[0-5])(?:\.5)?(?:_test)?$|^v1[0-5]_|_claude$|^embedded_key$")


def _imported_names(tree):
    names = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            names.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            names.add(node.module)
        elif isinstance(node, ast.Call):
            function = node.func
            called = function.attr if isinstance(function, ast.Attribute) else function.id if isinstance(function, ast.Name) else ""
            if called in ("import_module", "__import__") and node.args and isinstance(node.args[0], ast.Constant):
                names.add(node.args[0].value)
    return {name.split(".")[0] for name in names if isinstance(name, str)}


class ImportBoundaryTests(unittest.TestCase):
    def test_legacy_name_pattern(self):
        for name in ("v15_domain", "batch_v15", "hwpx_writer_v15", "security_grade_classifier_v13_test",
                     "security_grade_classifier_v7.5", "v14_domain", "classification_v14", "llm_client_v14",
                     "security_grade_classifier_v12_claude", "embedded_key"):
            self.assertTrue(LEGACY_NAME.search(name), name)
        for name in V16_RUNTIME + ("credential_store",):
            self.assertFalse(LEGACY_NAME.search(name), name)

    def test_v16_runtime_sources_import_only_v16_and_reused_modules(self):
        local_modules = {path.stem for path in ROOT.glob("*.py")} | {"hwpx_writer", "reference_lists", "excel_viewer"}
        allowed = set(V16_RUNTIME) | set(REUSED)
        for module in V16_RUNTIME + REUSED:
            names = _imported_names(_tree(module))
            with self.subTest(module=module):
                legacy = {name for name in names if LEGACY_NAME.search(name) and name != "llm_client_v14"}
                self.assertEqual(legacy, set())
                self.assertTrue({name for name in names if name in local_modules} <= allowed,
                                {name for name in names if name in local_modules} - allowed)

    def test_importing_the_v16_runtime_loads_no_legacy_module(self):
        code = """
import json, sys
for name in %r:
    __import__(name)
print(json.dumps(sorted(sys.modules)))
""" % (V16_RUNTIME,)
        fresh = subprocess.run([sys.executable, "-X", "utf8", "-B", "-c", code], cwd=ROOT,
                               capture_output=True, text=True, timeout=60)
        self.assertEqual(fresh.returncode, 0, fresh.stderr)
        loaded = {name.split(".")[0] for name in json.loads(fresh.stdout)}
        local = {path.stem for path in ROOT.glob("*.py")}
        self.assertTrue(set(V16_RUNTIME) <= loaded)
        self.assertEqual({name for name in loaded if name in local} - set(V16_RUNTIME) - set(REUSED), set())
        self.assertEqual({name for name in loaded if LEGACY_NAME.search(name)} - {"llm_client_v14"}, set())


if __name__ == "__main__":
    unittest.main()
