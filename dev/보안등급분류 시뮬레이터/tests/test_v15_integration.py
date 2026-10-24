"""Synthetic real engine/batch/document integration through the desktop API."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import batch_v15
from credential_store import KeyManager
from desktop_bridge_v15 import DesktopBridge
import security_grade_classifier_v15 as host
import v15_domain as d
from test_v15_engine import Client
from test_v15_notice_review import catalogue


class MemoryBackend:
    def __init__(self):
        self.values = {}

    def get(self, provider):
        return self.values.get(provider)

    def set(self, provider, key):
        self.values[provider] = key

    def delete(self, provider):
        self.values.pop(provider, None)


def finish(bridge, started):
    if not started["ok"]:
        raise AssertionError(started)
    job_id = started["data"]["job_id"]
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        reply = bridge.get_job(job_id, 0)
        if not reply["ok"]:
            raise AssertionError(reply)
        if reply["data"]["state"] in ("done", "cancelled", "failed"):
            return reply["data"]
        time.sleep(.01)
    raise AssertionError("합성 작업 시간 초과")


class DesktopIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source, self.output = self.root / "in.xlsx", self.root / "out.xlsx"
        book = Workbook()
        book.active.append(["과제명", "현재 기술사양"])
        for name in ("합성 A", "합성 B"):
            book.active.append([name, "속도 20 GHz"])
        book.save(self.source)
        book.close()
        self.cat = catalogue()
        self.loader_calls = 0
        self.clients = []
        self.paths = dict(input_excel=self.source, output_excel=self.output,
                          output_directory=self.root / "documents")

        def loader():
            self.loader_calls += 1
            return self.cat

        def factory(opts, key):
            self.assertEqual(key, "synthetic-secret")
            client = Client()
            self.clients.append(client)
            return client

        self.bridge = DesktopBridge(key_manager=KeyManager(MemoryBackend(), environ={}),
            catalogue_loader=loader, client_factory=factory,
            path_selector=lambda purpose: self.paths[purpose])
        self.assertTrue(self.bridge.save_key("claude", "synthetic-secret")["ok"])
        self.tokens = {purpose: self.bridge.select_path(purpose)["data"]["path_token"]
                       for purpose in self.paths}

    def batch(self):
        return finish(self.bridge, self.bridge.start_batch(dict(
            input_token=self.tokens["input_excel"], output_token=self.tokens["output_excel"],
            output_dir_token=None, options={"num_reviewers": 1}, workers=2,
            resume=False, make_forms=False)))

    def export(self, job_id, result):
        return finish(self.bridge, self.bridge.start_export(dict(job_id=job_id,
            task_fingerprint=result["task_fingerprint"],
            output_dir_token=self.tokens["output_directory"])))

    def latest(self):
        latest = {}
        for line in Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines():
            result = json.loads(line)
            latest[result["task_fingerprint"]] = result
        return latest

    def test_real_batch_A_B_A_exports_preserve_both_records_without_models(self):
        initial = self.batch()
        self.assertEqual(initial["state"], "done")
        a, b = initial["result"]["results"]
        self.assertEqual([a["status"], b["status"]], ["완료", "완료"])
        calls = len(self.clients), self.loader_calls
        with patch("notice_catalogue_v15.load_catalogue", side_effect=AssertionError("export catalogue")), \
                patch("llm_client_v14.LLMClient", side_effect=AssertionError("export model")), \
                patch.object(self.bridge._keys, "resolve", side_effect=AssertionError("export credentials")):
            first = self.export(initial["job_id"], a)
            second = self.export(initial["job_id"], b)
            third = self.export(first["job_id"], first["result"])
        self.assertEqual((len(self.clients), self.loader_calls), calls)
        for item in (first, second, third):
            self.assertEqual(item["state"], "done")
            self.assertEqual(item["result"]["schema_version"], 2)
            self.assertEqual(item["result"]["documents"]["status"], "완료")
            d.validate_result(item["result"])
            for entry in item["result"]["documents"]["files"]:
                self.assertTrue(Path(entry["path"]).is_file())
        latest = self.latest()
        self.assertEqual(latest[a["task_fingerprint"]], third["result"])
        self.assertEqual(latest[b["task_fingerprint"]], second["result"])
        book = load_workbook(BytesIO(self.output.read_bytes()))
        self.addCleanup(book.close)
        self.assertIn("고시기준본", book.sheetnames)

    def test_checkpoint_failure_keeps_durable_context_for_retry(self):
        initial = self.batch()
        a, b = initial["result"]["results"]
        before = deepcopy(self.latest())
        append = batch_v15._append_checkpoint

        def fail_checkpoint(path, encoded):
            if str(path).endswith(".checkpoint.jsonl"):
                raise OSError("synthetic secret path failure")
            return append(path, encoded)

        with patch.object(batch_v15, "_append_checkpoint", side_effect=fail_checkpoint):
            failed = self.export(initial["job_id"], a)
        self.assertEqual(failed["result"]["documents"]["status"], "완료")
        self.assertEqual(self.latest(), before)
        self.assertTrue(any(event["stage"] == "persistence" and "실패" in event["message"]
                            for event in failed["events"]))
        self.assertEqual(failed["events"][-1]["stage"], "persistence")
        self.assertIn("실패", failed["events"][-1]["message"])
        retried = self.export(failed["job_id"], failed["result"])
        other = self.export(initial["job_id"], b)
        self.assertEqual(retried["state"], "done")
        self.assertEqual(other["state"], "done")
        latest = self.latest()
        self.assertEqual(latest[a["task_fingerprint"]], retried["result"])
        self.assertEqual(latest[b["task_fingerprint"]], other["result"])

    def test_real_batch_progress_adapter_exposes_only_fixed_stages_and_messages(self):
        job = self.batch()
        allowed = {"preflight", "facts", "notices", "candidate", "criteria", "agents", "summary", "documents", "persistence"}
        self.assertTrue(job["events"])
        for event in job["events"]:
            self.assertIn(event["stage"], allowed)
            self.assertNotIn("{'stage':", event["message"])
            self.assertNotIn("합성 A", event["message"])
            self.assertNotIn("합성 B", event["message"])
        self.assertEqual(job["events"][-1]["stage"], "persistence")


class HostBoundaryIntegrationTests(unittest.TestCase):
    def test_selfcheck_imports_real_delayed_sdks_without_network_or_client_creation(self):
        import socket
        import subprocess
        import sys
        import openai
        import anthropic
        code = """
import json
import socket
def forbidden(*args, **kwargs):
    raise AssertionError('SDK import made a network request')
socket.socket.connect = forbidden
socket.create_connection = forbidden
import security_grade_classifier_v15 as host
print(json.dumps(host._selfcheck_sdks()))
"""
        fresh = subprocess.run([sys.executable, "-X", "utf8", "-B", "-c", code],
            cwd=Path(host.__file__).parent, capture_output=True, text=True, timeout=30)
        self.assertEqual(fresh.returncode, 0, fresh.stderr)
        fresh_result = json.loads(fresh.stdout)
        with patch.object(socket.socket, "connect", side_effect=AssertionError("network")), \
                patch.object(socket, "create_connection", side_effect=AssertionError("network")), \
                patch.object(openai.OpenAI, "__init__", side_effect=AssertionError("client")), \
                patch.object(anthropic.Anthropic, "__init__", side_effect=AssertionError("client")):
            result = host._selfcheck_sdks()
        self.assertEqual(result["openai"]["version"], "2.47.0")
        self.assertEqual(result["anthropic"]["version"], "1.3.0")
        self.assertEqual(result["openai"]["client_class"], "OpenAI")
        self.assertEqual(result["anthropic"]["client_class"], "Anthropic")
        self.assertEqual(result, fresh_result)

    def test_sdk_import_diagnostics_are_exclusive_to_selfcheck_branch(self):
        import ast
        tree = ast.parse(Path(host.__file__).read_text(encoding="utf-8"))
        guarded_calls = []
        all_calls = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "_selfcheck_sdks":
                all_calls.append(node)
            if isinstance(node, ast.If) and isinstance(node.test, ast.Name) and node.test.id == "checking":
                guarded_calls.extend(child for statement in node.body for child in ast.walk(statement)
                    if isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
                    and child.func.id == "_selfcheck_sdks")
        self.assertEqual(len(all_calls), 1)
        self.assertEqual(guarded_calls, all_calls)

    def test_native_message_allowlist_rejects_dunder_builtins_and_malformed_packets(self):
        server = host.LocalAssetServer()
        self.addCleanup(server.close)
        for name in host.PUBLIC_METHODS:
            packet = json.dumps([name, "[]", "1234567890"])
            self.assertTrue(host._allow_native_message(server, server.url, packet))
            self.assertFalse(host._allow_native_message(server, "https://example.invalid/", packet))
        forbidden = ("__str__", "__class__.__base__.__subclasses__", "_keys.resolve", "_keys._backend.get",
                     "_paths.clear", "_pywebviewAlert", "console", "pywebviewStateUpdate",
                     "pywebviewStateDelete", "pywebviewMoveWindow", "pywebviewEventHandler", "pywebviewAsyncCallback")
        for name in forbidden:
            with self.subTest(name=name):
                self.assertFalse(host._allow_native_message(server, server.url, json.dumps([name, "[]", "safe-id"])))
        for packet in ("null", "not JSON", "[]", json.dumps([{}, "[]", "safe-id"]),
                       json.dumps(["get_initial_state", {}, "safe-id"]),
                       json.dumps(["get_initial_state", "{}", "safe-id"]),
                       json.dumps(["get_initial_state", "[NaN]", "safe-id"]),
                       json.dumps(["get_initial_state", '[{"value":1,"value":2}]', "safe-id"]),
                       json.dumps(["get_initial_state", "[]", 'bad"id']),
                       json.dumps(["get_initial_state", "[]", "x" * 129])):
            with self.subTest(packet=packet):
                self.assertFalse(host._allow_native_message(server, server.url, packet))

    def test_actual_static_server_serves_only_local_asset_snapshot(self):
        server = host.LocalAssetServer().start()
        self.addCleanup(server.close)
        results = host._verify_static_server(server)
        self.assertTrue(results)
        self.assertTrue(all(entry["status"] in (403, 404) for entry in results))
        import http.client
        from urllib.parse import urlsplit
        for name in ("index.html", "styles.css", "app.js"):
            url = server.url.replace("index.html", name)
            self.assertTrue(server.allowed(url))
            connection = http.client.HTTPConnection(server.authority, timeout=3)
            try:
                connection.request("GET", urlsplit(url).path)
                response = connection.getresponse()
                self.assertEqual(response.status, 200)
                self.assertIn("default-src 'none'", response.getheader("Content-Security-Policy"))
                self.assertEqual(response.read(), (server.root / name).read_bytes())
            finally:
                connection.close()
        for url in ("https://example.invalid/", "file:///C:/Windows/win.ini", server.origin + "/credential_store.py"):
            self.assertFalse(server.allowed(url))

    def test_runtime_absence_emits_fixed_guidance_before_keys_and_legacy_import(self):
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "runtime-absent.json"
            with patch.object(host, "require_webview2", side_effect=host.HostError(host.RUNTIME_MESSAGE)), \
                    patch("credential_store.WindowsCredentialBackend", side_effect=AssertionError("real credentials")):
                self.assertEqual(host.run_app(self_check_path=report_path), 1)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertFalse(report["pass"])
            self.assertEqual(report["failures"], [host.RUNTIME_MESSAGE])
            self.assertEqual(report["fake_credential_operations"], [])
            self.assertTrue(report["profile_cleaned"])
            import builtins
            original_import = builtins.__import__

            def missing_webview(name, *args, **kwargs):
                if name == "webview":
                    raise ModuleNotFoundError("synthetic hidden import diagnostic")
                return original_import(name, *args, **kwargs)

            with patch.object(host, "require_webview2", return_value="155.0.0.0"), \
                    patch("builtins.__import__", side_effect=missing_webview), \
                    patch("credential_store.WindowsCredentialBackend", side_effect=AssertionError("real credentials")):
                self.assertEqual(host.run_app(self_check_path=report_path), 1)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["failures"], [host.INITIALIZATION_MESSAGE])
            self.assertEqual(report["fake_credential_operations"], [])
            self.assertNotIn("synthetic hidden import diagnostic", report_path.read_text(encoding="utf-8"))
        guard = host._NoLegacyRenderer()
        for module in ("webview.platforms.mshtml", "webview.platforms.cef"):
            with self.assertRaisesRegex(ImportError, "WebView2 Runtime"):
                guard.find_spec(module)
        self.assertIsNone(guard.find_spec("webview.platforms.edgechromium"))


if __name__ == "__main__":
    unittest.main()
