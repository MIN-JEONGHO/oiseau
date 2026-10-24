"""Synthetic workbooks, isolated clients and temporary output only."""
import copy
import importlib.util
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import batch_v14 as batch
from v14_domain import ITEM_NAMES

_spec = importlib.util.spec_from_file_location("batch_engine_fixture", Path(__file__).with_name("test_v14_engine.py"))
_fixture = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_fixture)
FakeClient = _fixture.FakeClient


def fake_documents(result, output_dir):
    files = [{"kind": "reviewer", "reviewer_id": r["reviewer_id"],
              "path": str(Path(output_dir) / (r["reviewer_id"] + ".hwpx")), "status": "완료", "error": ""}
             for r in result["reviewers"]]
    files.append({"kind": "summary", "reviewer_id": None, "path": str(Path(output_dir) / "summary.hwpx"),
                  "status": "완료", "error": ""})
    return {"status": "완료", "expected_files": len(files), "files": files, "errors": []}


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.input = self.root / "input.xlsx"
        self.output = self.root / "output.xlsx"
        self.clients = []
        self.workbook([("가상 과제", "내용")])

    def workbook(self, rows, headers=("과제명", "연구내용"), preferred=True):
        book = Workbook()
        if preferred:
            book.active.title = "작성안내"
            sheet = book.create_sheet("분류대상")
        else:
            sheet = book.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(self.input)
        book.close()

    def factory(self):
        client = FakeClient()
        self.clients.append(client)
        return client

    def run_batch(self, **kwargs):
        return batch.run_excel(self.input, self.output, kwargs.pop("options", {}),
                               kwargs.pop("client_factory", self.factory), make_forms=kwargs.pop("make_forms", False), **kwargs)

    def records(self):
        return [json.loads(line) for line in Path(str(self.output) + ".checkpoint.jsonl").read_text(encoding="utf-8").splitlines()]

    def test_atomic_write_retries_transient_permission_without_rewriting(self):
        self.output.write_bytes(b'previous result')
        written = []
        def writer(path):
            written.append(path)
            path.write_bytes(b'new result')
        replace = batch.os.replace
        attempts = []
        def transient(source, target):
            attempts.append(Path(source))
            if len(attempts) == 1:
                raise PermissionError('synthetic transient permission')
            replace(source, target)
        with patch.object(batch.os, 'replace', side_effect=transient), patch('time.sleep') as sleep:
            batch._atomic_write(self.output, writer)
        self.assertEqual(self.output.read_bytes(), b'new result')
        self.assertEqual(len(written), 1)
        self.assertEqual(attempts, [written[0], written[0]])
        self.assertFalse(written[0].exists())
        sleep.assert_called_once_with(.05)

    def test_atomic_write_preserves_previous_on_persistent_or_other_errors(self):
        for exception, count, delays in ((PermissionError('synthetic lock'), 3, [.05, .1]),
                                          (OSError('synthetic other failure'), 1, [])):
            with self.subTest(exception=type(exception).__name__):
                self.output.write_bytes(b'previous result')
                written = []
                def writer(path):
                    written.append(path)
                    path.write_bytes(b'new result')
                with patch.object(batch.os, 'replace', side_effect=exception) as replace, patch('time.sleep') as sleep:
                    with self.assertRaises(type(exception)):
                        batch._atomic_write(self.output, writer)
                self.assertEqual(replace.call_count, count)
                self.assertEqual([call.args[0] for call in sleep.call_args_list], delays)
                self.assertEqual(len(written), 1)
                self.assertEqual(self.output.read_bytes(), b'previous result')
                self.assertFalse(written[0].exists())
                self.assertEqual(list(self.root.glob('.v14-*.tmp')), [])

    def rewrite_record(self, mutate):
        records = self.records()
        mutate(records[0])
        Path(str(self.output) + ".checkpoint.jsonl").write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in records) + "\n", encoding="utf-8")

    def test_input_contract_and_literal_excel_text(self):
        self.workbook([("", "빈행", ""), ("(예시) 과제", "예시", ""),
                       ("=SUM(1,2)", "첫내용", "사업1"), ("=SUM(1,2)", "중복", "사업2"),
                       ("둘째", "", "")], (" 과제명 (필수) ", "연구내용(선택)", "세부사업명"))
        original = self.input.read_bytes()
        result = self.run_batch(workers=2)
        self.assertEqual(self.input.read_bytes(), original)
        self.assertEqual([r["task_name"] for r in result["results"]], ["=SUM(1,2)", "둘째"])
        self.assertEqual(result["results"][0]["task_snapshot"]["연구내용"], "첫내용")
        self.assertEqual(result["input_counts"], {"blank": 1, "example": 1, "duplicate": 1})
        self.assertTrue(result["warnings"])
        with_workbook = load_workbook(self.output)
        self.addCleanup(with_workbook.close)
        self.assertEqual(with_workbook.sheetnames, ["분류결과", "요약"])
        sheet = with_workbook["분류결과"]
        self.assertEqual(sheet["B2"].value, "=SUM(1,2)")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual([c.value for c in sheet[1]][:11], batch.RESULT_COLUMNS[:11])
        self.assertIn("문서상태", [c.value for c in sheet[1]])

    def test_excel_agent_labels_and_average_formats_preserve_raw_values(self):
        from v14_domain import make_result, normalize_options
        values = (3.875, 3.499999, 4.49999999999, 4.5)
        results = []
        for index, value in enumerate(values):
            result = make_result({'과제명': f'합성 평균 {index}'}, {})
            result['overall_avg'] = value
            results.append(result)
        batch._save_workbook(self.output, results, normalize_options({}))
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        sheet = book['분류결과']
        self.assertIn('AI 종합 검토의견', [cell.value for cell in sheet[1]])
        self.assertNotIn('의결사유', [cell.value for cell in sheet[1]])
        for row, value, number_format in zip(range(2, 6), values, ('0.00', '0.000000', '0.00000000000', '0.00')):
            self.assertEqual(sheet.cell(row, 6).value, value)
            self.assertEqual(sheet.cell(row, 6).data_type, 'n')
            self.assertEqual(sheet.cell(row, 6).number_format, number_format)

    def test_regenerated_batch_documents_persist_without_model_calls(self):
        from io import BytesIO
        with patch('batch_v14.export_documents', side_effect=OSError('synthetic document failure')):
            report = self.run_batch(make_forms=True)
        original = copy.deepcopy(report['results'][0])
        call_count = sum(len(client.calls) for client in self.clients)
        regenerated = batch.regenerate_documents(original, self.root / 'new-documents')
        outcome = batch.persist_regenerated_documents(report, original, regenerated)
        self.assertEqual(sum(len(client.calls) for client in self.clients), call_count)
        self.assertEqual(outcome['persistence']['status'], '완료')
        latest = self.records()[-1]
        self.assertEqual(latest['documents'], regenerated['documents'])
        self.assertEqual({k: v for k, v in latest.items() if k != 'documents'},
                         {k: v for k, v in original.items() if k != 'documents'})
        self.assertEqual(outcome['report']['counts']['documents_missing'], 0)
        book = load_workbook(BytesIO(self.output.read_bytes()))
        self.addCleanup(book.close)
        self.assertEqual(book['분류결과']['M2'].value, '완료')
        self.assertEqual(book['분류결과']['O2'].value, 4)
        self.assertEqual(book['분류결과']['Q2'].value, 0)
        for entry in regenerated['documents']['files']:
            self.assertIn(entry['path'], book['분류결과']['R2'].value)
        audit = Path(report['audit_path']).read_text(encoding='utf-8').splitlines()
        event = json.loads(audit[-1])
        self.assertEqual(event['event'], 'batch_documents_regenerated')
        self.assertEqual(event['workbook_status'], '완료')
        self.assertNotIn('task_snapshot', event)
        self.assertNotIn('documents', event)
        self.assertEqual(report['results'][0], original)

    def test_regenerated_batch_metadata_failures_are_separate_and_recoverable(self):
        for store in ('checkpoint', 'workbook', 'audit'):
            with self.subTest(store=store):
                report = self.run_batch(make_forms=False, resume=False)
                original = copy.deepcopy(report['results'][0])
                regenerated = batch.regenerate_documents(original, self.root / ('docs-' + store))
                previous_checkpoint = Path(report['checkpoint_path']).read_bytes()
                previous_excel = self.output.read_bytes()
                append = batch._append_checkpoint
                replace = batch.os.replace
                def fail_append(path, encoded):
                    if (store == 'checkpoint' and str(path) == report['checkpoint_path']) or (store == 'audit' and str(path) == report['audit_path']):
                        raise OSError('synthetic secret must not surface')
                    return append(path, encoded)
                def fail_replace(source, target):
                    if store == 'workbook' and Path(target) == self.output:
                        raise PermissionError('synthetic locked file')
                    return replace(source, target)
                with patch('batch_v14._append_checkpoint', side_effect=fail_append), patch('batch_v14.os.replace', side_effect=fail_replace):
                    outcome = batch.persist_regenerated_documents(report, original, regenerated)
                self.assertEqual(outcome['persistence'][store], '실패')
                self.assertTrue(outcome['persistence']['warnings'])
                self.assertNotIn('synthetic secret', str(outcome))
                self.assertEqual(outcome['result']['documents']['status'], '완료')
                self.assertEqual(outcome['result']['overall_avg'], original['overall_avg'])
                if store == 'checkpoint':
                    self.assertEqual(Path(report['checkpoint_path']).read_bytes(), previous_checkpoint)
                    self.assertEqual(self.output.read_bytes(), previous_excel)
                    self.assertEqual(outcome['persistence']['workbook'], '미시도')
                elif store == 'workbook':
                    self.assertEqual(self.output.read_bytes(), previous_excel)
                    self.assertEqual(self.records()[-1]['documents'], regenerated['documents'])
                else:
                    self.assertEqual(outcome['report']['audit_status'], '실패')

    def test_regenerated_batch_rejects_changed_source_options_and_classification(self):
        report = self.run_batch()
        original = copy.deepcopy(report['results'][0])
        regenerated = batch.regenerate_documents(original, self.root / 'rejected-docs')
        old_checkpoint = Path(report['checkpoint_path']).read_bytes()
        changed = copy.deepcopy(regenerated)
        changed['rationale'] += ' changed classification'
        with self.assertRaises(batch.BatchError):
            batch.persist_regenerated_documents(report, original, changed)
        for change in ('source', 'options'):
            self.workbook([('가상 과제', '내용')])
            invalid_report = copy.deepcopy(report)
            if change == 'source':
                self.workbook([('changed input', 'changed contents')])
            else:
                invalid_report['options_fingerprint'] = '0' * 64
            outcome = batch.persist_regenerated_documents(invalid_report, original, regenerated)
            self.assertEqual(outcome['persistence']['status'], '실패')
            self.assertEqual(outcome['persistence']['workbook'], '미시도')
            self.assertEqual(Path(report['checkpoint_path']).read_bytes(), old_checkpoint)

    def test_only_task_column_required_and_first_sheet_fallback(self):
        self.workbook([("단일",)], ("과제명",), preferred=False)
        result = self.run_batch()
        self.assertEqual(result["results"][0]["task_snapshot"]["연구목표"], "")
        self.workbook([("자료",)], ("연구내용",))
        with self.assertRaisesRegex(batch.BatchError, "과제명"):
            self.run_batch()

    def test_make_template_roundtrip_optional_headers_and_synthetic_examples(self):
        template = self.root / "synthetic_template.xlsx"
        batch.make_template(template)
        book = load_workbook(template)
        try:
            self.assertEqual(book.sheetnames, ["작성안내", "분류대상"])
            sheet = book["분류대상"]
            self.assertEqual(len(list(sheet[1])), 7)
            self.assertEqual(sheet["B1"].value, "과제명(필수)")
            self.assertTrue(all(sheet.cell(row, 2).value.startswith("(예시)") for row in (2, 3)))
            sheet.append(["합성 세부사업", "합성 신규 과제", "", "시험용 연구", "", "추가 근거", "합성 부처 기준"])
            book.save(template)
        finally:
            book.close()
        report = batch.run_excel(template, self.output, {}, self.factory, make_forms=False)
        self.assertEqual(report["input_counts"]["example"], 2)
        self.assertEqual([r["task_name"] for r in report["results"]], ["합성 신규 과제"])
        self.assertEqual(report["results"][0]["task_snapshot"]["추가 검토정보"], "추가 근거")

    def test_resume_reuses_only_same_complete_result_and_no_secrets(self):
        self.run_batch(options={"api_key": "SYNTHETIC-DO-NOT-SAVE"})
        first = len(self.clients)
        result = self.run_batch()
        self.assertEqual(len(self.clients), first)
        self.assertEqual(result["reused_count"], 1)
        self.assertNotIn("SYNTHETIC", Path(result["checkpoint_path"]).read_text(encoding="utf-8"))
        self.assertNotIn("api_key", Path(result["checkpoint_path"]).read_text(encoding="utf-8"))
        self.assertNotIn("options", self.records()[0])

    def test_every_input_field_invalidates_resume(self):
        headers = ("과제명", "세부사업명", "연구목표", "연구내용", "기대효과", "추가 검토정보", "부처 세부기준")
        for index in range(1, len(headers)):
            with self.subTest(field=headers[index]):
                self.workbook([("동일", "", "", "", "", "", "")], headers)
                self.run_batch(resume=False)
                row = ["동일"] + [""] * 6
                row[index] = "합성 변경"
                self.workbook([row], headers)
                self.assertEqual(self.run_batch()["reused_count"], 0)

    def test_every_public_option_invalidates_resume(self):
        variants = [dict(method="checklist"), dict(provider="openai", model="gpt-6-luna"),
                    dict(provider="ollama", model="synthetic-local"), dict(strictness="엄격"),
                    dict(num_reviewers=1), dict(ministry_policy="합성 기준"),
                    dict(selection_mode="manual", selected_items=ITEM_NAMES[:2], weights={ITEM_NAMES[0]: 1.2, ITEM_NAMES[1]: .8}),
                    dict(weight_tolerance=0), dict(local_base_url="http://localhost:11434")]
        for options in variants:
            with self.subTest(options=options):
                self.run_batch(resume=False)
                self.assertEqual(self.run_batch(options=options)["reused_count"], 0)

    def test_corrupt_legacy_and_structurally_invalid_checkpoint_warn(self):
        mutations = [lambda r: r.update(schema_version=True),
                     lambda r: r.update(schema_version=99), lambda r: r.update(guideline_id="old"),
                     lambda r: r.update(task_snapshot=[]), lambda r: r.update(status=["완료"]),
                     lambda r: r.update(reviewers=[]), lambda r: r["reviewers"][0].update(status="미실시"),
                     lambda r: r["reviewers"][0].update(scores={}),
                     lambda r: r["reviewers"][0].update(reviewer_score=True),
                     lambda r: r["criteria"][0].update(law_refs="invalid"),
                     lambda r: r.update(final_grade="오류"), lambda r: r.update(options_fingerprint="bad"),
                     lambda r: r.update(documents={}), lambda r: r.update(task_name="unrelated")]
        for mutate in mutations:
            with self.subTest(mutation=mutate):
                self.run_batch(resume=False)
                self.rewrite_record(mutate)
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])
        for contents in ("{broken\n", '{"상태":"완료","과제명":"가상 과제"}\n', "[]\n", "NaN\n"):
            Path(str(self.output) + ".checkpoint.jsonl").write_text(contents, encoding="utf-8")
            report = self.run_batch()
            self.assertEqual(report["reused_count"], 0)
            self.assertTrue(report["warnings"])

    def test_resume_false_replaces_old_checkpoint(self):
        self.run_batch()
        self.workbook([("새 과제", "새 내용")])
        self.run_batch(resume=False)
        self.assertEqual([r["task_name"] for r in self.records()], ["새 과제"])

    def test_duplicate_json_keys_are_rejected(self):
        self.run_batch()
        path = Path(str(self.output) + ".checkpoint.jsonl")
        contents = path.read_text(encoding="utf-8")
        path.write_text(contents.replace('"schema_version":1', '"schema_version":99,"schema_version":1'), encoding="utf-8")
        report = self.run_batch()
        self.assertEqual(report["reused_count"], 0)
        self.assertTrue(report["warnings"])

    def test_concurrent_isolated_clients_checkpoint_and_output_order(self):
        self.workbook([(str(i), "합성") for i in range(30)])
        real = batch.classify_task
        def delayed(task, options, client, progress_callback=None):
            time.sleep((30 - int(task["과제명"])) % 4 * .001)
            return real(task, options, client, progress_callback)
        with patch.object(batch, "classify_task", side_effect=delayed):
            report = self.run_batch(workers=8)
        self.assertEqual(len({id(c) for c in self.clients}), 30)
        self.assertEqual([r["task_name"] for r in report["results"]], list(map(str, range(30))))
        self.assertEqual(len(self.records()), 30)
        self.assertEqual(self.run_batch()["reused_count"], 30)

    def test_options_frozen_before_worker_and_row_policy_validated(self):
        options = {"strictness": "표준"}
        def mutate_factory():
            options["strictness"] = "엄격"
            return self.factory()
        report = self.run_batch(options=options, client_factory=mutate_factory)
        self.assertEqual(self.clients[0].calls[-1][1]["strictness"], "표준")
        self.workbook([("정책있음", "합성 기준"), ("정책없음", "")], ("과제명", "부처 세부기준"))
        report = self.run_batch(options={"selection_mode": "manual", "selected_items": [ITEM_NAMES[7]]}, resume=False)
        self.assertEqual([r["status"] for r in report["results"]], ["완료", "오류"])

    def test_task_error_continues_and_grade_not_fabricated(self):
        self.workbook([("첫째", ""), ("둘째", "")])
        count = 0
        def factory():
            nonlocal count
            count += 1
            if count == 1:
                raise RuntimeError("SYNTHETIC SECRET should not appear")
            return self.factory()
        report = self.run_batch(client_factory=factory, workers=1)
        self.assertEqual([r["status"] for r in report["results"]], ["오류", "완료"])
        self.assertIsNone(report["results"][0]["final_grade"])
        self.assertNotIn("SECRET", Path(report["checkpoint_path"]).read_text(encoding="utf-8"))
        self.assertEqual(self.run_batch()["reused_count"], 1)

    def test_distinct_classification_statuses_and_summary(self):
        responses = [FakeClient(), FakeClient({"candidate": {"is_candidate": False, "reason": "명확한 제외", "uncertainty": ""}}),
                     FakeClient({"candidate": {"is_candidate": None, "reason": "정보 부족", "uncertainty": "미확인"}}),
                     FakeClient({"candidate": RuntimeError("synthetic")})]
        self.workbook([(str(i), "") for i in range(4)])
        report = self.run_batch(client_factory=lambda: responses.pop(0), workers=1)
        self.assertEqual([r["status"] for r in report["results"]], ["완료", "후보제외", "검토필요", "오류"])
        self.assertEqual([report["counts"][s] for s in ("완료", "후보제외", "검토필요", "오류")], [1] * 4)
        self.assertEqual(report["counts"]["documents_missing"], 13)

    def test_partial_document_failure_retry_does_not_call_model_or_mutate(self):
        def partial(result, output_dir):
            docs = fake_documents(result, output_dir)
            docs["files"][0].update(status="오류", error="합성 저장 실패")
            docs.update(status="일부실패", errors=["합성 저장 실패"])
            return docs
        with patch.object(batch, "export_documents", side_effect=partial):
            report = self.run_batch(make_forms=True)
        saved = report["results"][0]
        self.assertEqual(saved["status"], "완료")
        self.assertEqual(report["counts"]["documents_failed"], 1)
        before = copy.deepcopy(saved)
        calls = sum(len(c.calls) for c in self.clients)
        with patch.object(batch, "export_documents", side_effect=fake_documents):
            regenerated = batch.regenerate_documents(saved, self.root / "retry")
        self.assertEqual(saved, before)
        self.assertEqual(regenerated["documents"]["status"], "완료")
        self.assertEqual(sum(len(c.calls) for c in self.clients), calls)

    def test_writer_exception_preserves_classification_and_missing_counts(self):
        with patch.object(batch, "export_documents", side_effect=OSError("secret path")):
            report = self.run_batch(make_forms=True)
        self.assertEqual(report["results"][0]["status"], "완료")
        self.assertEqual(report["results"][0]["documents"]["status"], "실패")
        self.assertEqual(report["counts"]["documents_failed"], 4)
        self.assertEqual(report["counts"]["documents_missing"], 4)

    def test_document_generation_defaults_on(self):
        with patch.object(batch, "export_documents", side_effect=fake_documents) as writer:
            report = batch.run_excel(self.input, self.output, {}, self.factory)
        self.assertEqual(writer.call_count, 1)
        self.assertEqual(report["results"][0]["documents"]["status"], "완료")

    def test_document_regeneration_in_resume_never_reclassifies(self):
        self.run_batch()
        count = len(self.clients)
        with patch.object(batch, "export_documents", side_effect=fake_documents) as writer:
            report = self.run_batch(make_forms=True)
        self.assertEqual(writer.call_count, 1)
        self.assertEqual(len(self.clients), count)
        self.assertEqual(report["reused_count"], 1)

    def test_checkpoint_rejects_document_count_identity_and_persona_mismatch(self):
        def summary_only(result):
            result["documents"] = fake_documents(result, self.root)
            result["documents"]["files"] = result["documents"]["files"][-1:]
            result["documents"]["expected_files"] = 1
        def wrong_reviewer(result):
            result["documents"] = fake_documents(result, self.root)
            result["documents"]["files"][0]["reviewer_id"] = "reviewer-99"
        def missing_summary(result):
            result["documents"] = fake_documents(result, self.root)
            result["documents"]["files"][-1].update(kind="reviewer", reviewer_id="reviewer-4")
        mutations = [summary_only, wrong_reviewer, missing_summary,
                     lambda r: r["documents"].update(expected_files=1),
                     lambda r: r["reviewers"][0].update(persona=""),
                     lambda r: r["reviewers"][0].update(persona="AI 기술사업화 전문가")]
        for mutate in mutations:
            with self.subTest(mutation=mutate):
                self.run_batch(resume=False)
                self.rewrite_record(mutate)
                report = self.run_batch()
                self.assertEqual(report["reused_count"], 0)
                self.assertTrue(report["warnings"])

    def test_writer_wrong_document_identity_is_failure_without_model_retry(self):
        saved = self.run_batch()["results"][0]
        before = copy.deepcopy(saved)
        calls = sum(len(client.calls) for client in self.clients)
        def wrong_identity(result, directory):
            docs = fake_documents(result, directory)
            docs["files"][0]["reviewer_id"] = "reviewer-99"
            return docs
        with patch.object(batch, "export_documents", side_effect=wrong_identity):
            result = batch.regenerate_documents(saved, self.root / "retry")
        self.assertEqual(result["status"], "완료")
        self.assertEqual(result["documents"]["status"], "실패")
        self.assertEqual(result["documents"]["expected_files"], 4)
        self.assertEqual(saved, before)
        self.assertEqual(sum(len(client.calls) for client in self.clients), calls)

    def test_document_identity_contract_for_off_and_early_results(self):
        for count in (1, 3):
            with self.subTest(count=count):
                off = self.run_batch(options={"num_reviewers": count}, resume=False)
                self.assertEqual(off["results"][0]["documents"],
                                 {"status": "미생성", "expected_files": count + 1, "files": [], "errors": []})
                self.assertEqual(self.run_batch(options={"num_reviewers": count})["reused_count"], 1)
                for candidate in (False, None):
                    factory = lambda: FakeClient({"candidate": {"is_candidate": candidate, "reason": "합성 후보 검토", "uncertainty": "" if candidate is False else "정보 부족"}})
                    off = self.run_batch(options={"num_reviewers": count}, client_factory=factory, resume=False)
                    saved = off["results"][0]
                    self.assertEqual(saved["reviewers"], [])
                    self.assertEqual(saved["documents"]["expected_files"], 1 if candidate is False else count + 1)
                    with patch.object(batch, "export_documents", side_effect=fake_documents):
                        regenerated = batch.regenerate_documents(saved, self.root / "early")
                    self.assertEqual(regenerated["documents"]["expected_files"], 1)
                    self.assertEqual([(f["kind"], f["reviewer_id"]) for f in regenerated["documents"]["files"]], [("summary", None)])
                    batch._validate_result(regenerated)

    def test_1001_requested_documents_never_disabled(self):
        self.workbook([(f"합성{i}", "") for i in range(1001)])
        with patch.object(batch, "export_documents", side_effect=fake_documents) as writer:
            report = self.run_batch(make_forms=True, workers=8)
        self.assertEqual(writer.call_count, 1001)
        self.assertEqual(report["counts"]["documents_expected"], 4004)
        self.assertEqual(report["counts"]["documents_completed"], 4004)

    def test_cancellation_stops_new_tasks_and_preserves_completed(self):
        self.workbook([(str(i), "") for i in range(9)])
        event = threading.Event()
        def cancel_factory():
            event.set()
            return self.factory()
        report = self.run_batch(workers=1, client_factory=cancel_factory, cancel_event=event)
        self.assertEqual(report["status"], "취소")
        self.assertEqual(len(report["results"]), 1)
        self.assertEqual(len(self.records()), 1)
        self.assertEqual(len(self.clients), 1)

    def test_locked_excel_atomic_failure_keeps_previous_file_and_checkpoint(self):
        self.run_batch()
        previous = self.output.read_bytes()
        real_replace = batch.os.replace
        def locked(source, target):
            if Path(target) == self.output:
                raise PermissionError("synthetic locked file")
            return real_replace(source, target)
        with patch.object(batch.os, "replace", side_effect=locked):
            with self.assertRaisesRegex(batch.BatchError, "엑셀"):
                self.run_batch(resume=False)
        self.assertEqual(self.output.read_bytes(), previous)
        self.assertEqual(len(self.records()), 1)
        self.assertFalse(list(self.root.glob(".v14-*")))

    def test_checkpoint_atomic_init_failure_keeps_previous(self):
        self.run_batch()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        before = checkpoint.read_bytes()
        with patch.object(batch.os, "replace", side_effect=PermissionError("synthetic")):
            with self.assertRaises(batch.BatchError):
                self.run_batch(resume=False)
        self.assertEqual(checkpoint.read_bytes(), before)

    def test_checkpoint_append_failure_preserves_last_good_record(self):
        self.workbook([("첫째", ""), ("둘째", "")])
        real_fsync = batch.os.fsync
        calls = 0
        def fail_second_task(descriptor):
            nonlocal calls
            calls += 1
            # Initial snapshot, first checkpoint, first audit, second checkpoint.
            if calls == 4:
                raise OSError("synthetic disk failure")
            return real_fsync(descriptor)
        with patch.object(batch.os, "fsync", side_effect=fail_second_task):
            with self.assertRaisesRegex(batch.BatchError, "중간 저장"):
                self.run_batch(workers=1)
        self.assertEqual([r["task_name"] for r in self.records()], ["첫째"])
        self.assertEqual(self.run_batch()["reused_count"], 1)

    def test_audit_failure_visible_and_no_legacy_log(self):
        audit_path = Path(str(self.output) + ".audit.jsonl")
        audit_path.mkdir()
        messages = []
        report = self.run_batch(status_cb=messages.append, progress_cb=messages.append)
        self.assertEqual(report["audit_status"], "실패")
        self.assertTrue(any("감사" in message and "실패" in message for message in messages))
        self.assertTrue(report["warnings"])
        self.assertEqual(report["counts"]["완료"], 1)
        self.assertTrue(all(isinstance(message, str) for message in messages))

    def test_input_not_overwritten_and_invalid_worker_count(self):
        with self.assertRaisesRegex(batch.BatchError, "입력"):
            batch.run_excel(self.input, self.input, {}, self.factory)
        for workers in (0, 9, True, 1.5):
            with self.subTest(workers=workers), self.assertRaises(ValueError):
                self.run_batch(workers=workers)


if __name__ == "__main__":
    unittest.main()
