"""Offline tests for the v16 bulk check tool; synthetic tasks and temp folders only.

Every task here is invented test text (no real project content). The installed
catalogue bundle is used read-only; every model reply comes from the
input-based mock model. No network, key store or real model is involved.
"""
import contextlib
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / 'tools'))

import batch_v16
import hwpx_writer_v16 as hw
import notice_review_v16 as nr
import v16_domain as d
from notice_catalogue_v16 import load_catalogue
import v16_bulk_check as bc


CATALOGUE = load_catalogue()
_, BY_ID = nr._preflight(CATALOGUE)
TARGET_CLAUSE = 'core:item:029'
UNSUPPORTED_CLAUSE = 'core:item:035'

GOAL = ('마을 공동체 텃밭의 운영 기록과 참여자 설문을 묶어 운영 개선 방향을 찾는 연구이다. '
        '참여 주민과 운영 담당자 사이의 소통 방식을 정리하고 계절별 참여 흐름을 살핀다.\n'
        '- 참여자 의견을 주제별로 나누어 정리하는 절차 마련\n'
        '- 계절별 참여 흐름을 해마다 같은 방식으로 기록하는 양식 마련')
CONTENT = ('1차년도에는 운영 기록 정리 방식과 설문 문항을 마련하고 시범 텃밭 세 곳에서 의견을 모은다.\n'
           '2차년도에는 시범 결과를 바탕으로 운영 기준을 다듬고 안내서를 만든다.\n'
           '○ 운영 기록 정리 방식 마련\n'
           '○ 설문 문항 개발과 시범 운영\n'
           '○ 안내서 작성과 여러 텃밭 시범 적용')
EFFECT = ('주민 의견이 텃밭 운영에 꾸준히 반영되는 방식이 자리 잡는다. '
          '운영 담당자는 같은 기준으로 해마다 참여 흐름을 비교할 수 있다.')


def generic(name, **extra):
    task = {'세부사업명': '[모의] 생활연구 지원사업', '과제명': name,
            '연구목표': GOAL, '연구내용': CONTENT, '기대효과': EFFECT}
    task.update(extra)
    return task


def spec_lines(clause):
    lines, seen = [], set()
    for leaf in d.condition_leaves(clause['condition']):
        if leaf['type'] == 'unsupported' or leaf['subject'] in seen:
            continue
        seen.add(leaf['subject'])
        lines.append(f"{leaf['subject']} {leaf['threshold']}{leaf['unit']}" if leaf['type'] == 'numeric'
                     else leaf['subject'])
    return '\n'.join(lines)


def targeted(clause_id, name):
    clause = BY_ID[clause_id]
    title = re.sub(r'\s+', ' ', clause['title']).strip()
    lines = spec_lines(clause)
    return {'세부사업명': '[모의] 산업기술 고도화사업', '과제명': name,
            '연구목표': f'{title}을 확보한다.\n시험 장비와 공정 조건을 정리한다.',
            '연구내용': '시제품을 만들고 시험 결과를 정리한다.\n- 단계별 시험 절차 마련',
            '기대효과': '관련 산업의 시험 기반을 넓힌다.',
            '현재 기술사양': lines, '개발목표 사양': lines}


NAMES = ['[모의] 텃밭 운영 기록 정리 연구', '[모의] 텃밭 참여 흐름 비교 연구',
         '[모의] 텃밭 안내서 시범 적용 연구', '[모의] 자동 판정 조항 대상 합성 과제',
         '[모의] 판정 불가 조항 대상 합성 과제']
TASKS = [
    generic(NAMES[0]),
    generic(NAMES[1], 연구내용='=합계 칸처럼 보이는 글로 시작하는 연구내용이다._x000D_\n' + CONTENT),
    generic(NAMES[2], 연구목표=GOAL.replace('\n', '_x000D_\n') + '   \n\n\n\n끝.'),
    targeted(TARGET_CLAUSE, NAMES[3]),
    targeted(UNSUPPORTED_CLAUSE, NAMES[4]),
]
EXCLUDED = [{'과제명': '_x000D_'}, generic(NAMES[0]), {'과제명': '(예시) 합성 과제', '연구목표': '시험'}]
TASK_LIST = (
    [dict(과제명=NAMES[i], 분야='인공지능', 고시대상=False, 기대='등급') for i in range(3)]
    + [dict(과제명=NAMES[3], 분야='첨단로봇·제조', 고시대상=True, 기대='등급', 노린조항=[TARGET_CLAUSE]),
       dict(과제명=NAMES[4], 분야='첨단로봇·제조', 고시대상=True, 기대='검토필요', 노린조항=[UNSUPPORTED_CLAUSE])]
    + [dict(과제명='_x000D_', 분야='제외확인', 고시대상=False, 기대='제외'),
       dict(과제명=NAMES[0], 분야='제외확인', 고시대상=False, 기대='제외'),
       dict(과제명='(예시) 합성 과제', 분야='제외확인', 고시대상=False, 기대='제외')])


def hwpx_files(folder):
    found = []
    for root, _, files in os.walk(folder):
        found.extend(os.path.join(root, name) for name in files if name.endswith('.hwpx'))
    return found


class WriteInputTests(unittest.TestCase):
    def test_layout_header_marks_and_string_cells(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / '입력.xlsx'
            tasks = [dict(TASKS[1]), {'과제명': '[모의] 빈 칸 과제', '현재 기술사양': '=1+1'}]
            bc.write_input(tasks, path)
            template = Path(temp) / '양식.xlsx'
            batch_v16.make_template(template)
            expected_header = [c.value for c in load_workbook(template)['분류대상'][1]]
            book = load_workbook(path)
            self.assertIn('분류대상', book.sheetnames)
            sheet = book['분류대상']
            rows = list(sheet.iter_rows())
            self.assertEqual([c.value for c in rows[0]], expected_header)
            self.assertEqual(len(rows), 3)
            columns = d.TASK_COLUMNS + d.OPTIONAL_TASK_COLUMNS
            for row, task in zip(rows[1:], tasks):
                self.assertEqual(len(row), 9)
                for cell, column in zip(row, columns):
                    self.assertEqual(cell.value or '', task.get(column, ''))
                    if cell.value is not None:
                        self.assertEqual(cell.data_type, 's')
            self.assertTrue(rows[1][3].value.startswith('=합계'))
            self.assertEqual(rows[2][7].value, '=1+1')
            # The program reads the same rows back (cleaned text, no formula).
            tasks_read, excluded, cleaned = batch_v16._read_tasks(path)
            self.assertEqual(len(tasks_read), 2)
            self.assertTrue(cleaned)
            self.assertTrue(tasks_read[0]['연구내용'].startswith('=합계'))

    def test_too_long_cell_and_unknown_column_are_refused(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaises(ValueError):
                bc.write_input([{'과제명': '[모의] 긴 칸', '연구내용': '가' * 32768}], Path(temp) / 'a.xlsx')
            with self.assertRaises(ValueError):
                bc.write_input([{'과제명': '[모의] 모르는 칸', '사업명': 'x'}], Path(temp) / 'b.xlsx')
            self.assertFalse((Path(temp) / 'a.xlsx').exists())


class OlsTests(unittest.TestCase):
    def test_straight_line_and_known_values(self):
        fit = bc.ols([1, 2, 3, 4, 5], [5, 7, 9, 11, 13])
        self.assertAlmostEqual(fit['slope'], 2.0)
        self.assertAlmostEqual(fit['intercept'], 3.0)
        self.assertAlmostEqual(fit['r2'], 1.0)
        self.assertEqual(fit['n'], 5)
        fit = bc.ols([1, 2, 3, 4], [2, 4, 5, 4])
        self.assertAlmostEqual(fit['slope'], 0.7)
        self.assertAlmostEqual(fit['intercept'], 2.0)
        self.assertAlmostEqual(fit['r2'], 1 - 2.3 / 4.75)
        self.assertEqual(fit['n'], 4)

    def test_degenerate_inputs(self):
        fit = bc.ols([3, 3, 3], [1, 2, 3])
        self.assertIsNone(fit['slope'])
        self.assertEqual(fit['n'], 3)
        self.assertIsNone(bc.ols([], [])['slope'])
        with self.assertRaises(ValueError):
            bc.ols([1, 2], [1])


class FormatStatsTests(unittest.TestCase):
    SECRET_WORDS = ('가람누리솔빛', '다온해솔마루', '비밀고유번호사칠', '라온하제그루', '보람찬미리내',
                    '새라온결', '아라온담')

    def make(self, path):
        book = Workbook()
        notice = book.active
        notice.title = '유의사항'
        notice.append(['신청 업무 외 사용 금지 안내 문장'])
        sheet = book.create_sheet('과제')
        sheet.append(['2024년 합성 목록'])
        sheet.append(['과제고유번호', '사업명', '연구개발과제명-국문', '요약문_연구목표', '요약문_연구내용', '요약문_기대효과'])
        sheet.append(['비밀고유번호사칠', '가람누리솔빛 사업', '다온해솔마루 과제',
                      '○ 라온하제그루 출력 20 kW 달성_x000D_\n○ 효율 95 % 확보', '- 보람찬미리내 1단계\n- 2단계 3건',
                      '짧음'])
        sheet.append(['비밀고유번호사칠', '가람누리솔빛 사업', '새라온결 과제', '1. 아라온담 목표\n2. 두 번째 목표',
                      None, '기대 효과 문장으로 충분히 긴 글을 둔다 500 억원 규모'])
        book.save(path)

    def test_numbers_only_and_no_row_text(self):
        mapping = {'사업명': '세부사업명', '연구개발과제명-국문': '과제명', '요약문_연구목표': '연구목표',
                   '요약문_연구내용': '연구내용', '요약문_기대효과': '기대효과'}
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / '합성.xlsx'
            self.make(path)
            out = io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
                stats = bc.format_stats(path, '과제', mapping)
                auto = bc.format_stats(path, None, mapping)
            self.assertEqual(out.getvalue(), '')
            self.assertEqual(stats, auto)
            text = json.dumps(stats, ensure_ascii=False)
            for word in self.SECRET_WORDS:
                self.assertNotIn(word, text)
            self.assertEqual(stats['sheet'], '과제')
            self.assertEqual(stats['header_row'], 2)
            self.assertEqual(stats['rows'], 2)
            goal = stats['fields']['연구목표']
            self.assertEqual(goal['nonempty'], 2)
            self.assertAlmostEqual(goal['x000d_ratio'], 0.5)
            self.assertEqual(goal['bullets'].get('○'), 2)
            self.assertEqual(goal['bullets'].get('숫자.'), 2)
            self.assertEqual(goal['units'].get('kW'), 1)
            self.assertEqual(goal['units'].get('%'), 1)
            self.assertEqual(goal['lines']['max'], 2)
            content = stats['fields']['연구내용']
            self.assertEqual(content['nonempty'], 1)
            self.assertAlmostEqual(content['empty_ratio'], 0.5)
            self.assertEqual(content['bullets'].get('-'), 2)
            self.assertEqual(content['units'].get('건'), 1)
            effect = stats['fields']['기대효과']
            self.assertAlmostEqual(effect['very_short_ratio'], 0.5)
            self.assertEqual(effect['units'].get('억원'), 1)
            self.assertIsInstance(goal['length']['median'], (int, float))
            for value in stats['fields'].values():
                for key in value['bullets']:
                    self.assertLessEqual(len(key), 4)

    def test_missing_header_names_only_names(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / '합성.xlsx'
            self.make(path)
            with self.assertRaises(ValueError) as caught:
                bc.format_stats(path, '과제', {'없는머리글': '과제명'})
            for word in self.SECRET_WORDS:
                self.assertNotIn(word, str(caught.exception))


class CheckNoticeTests(unittest.TestCase):
    def test_generic_task_has_no_candidates_and_is_complete(self):
        found = bc.check_notice(generic('[모의] 메모리 확인 일반 과제'), catalogue=CATALOGUE)
        self.assertEqual(found['candidate_ids'], [])
        self.assertEqual(found['status'], 'complete')
        self.assertEqual(found['comparisons'], [])
        self.assertGreater(found['calls'], 0)

    def test_targeted_task_reports_assessments(self):
        found = bc.check_notice(targeted(TARGET_CLAUSE, '[모의] 메모리 확인 대상 과제'), catalogue=CATALOGUE)
        self.assertEqual(found['candidate_ids'], [TARGET_CLAUSE])
        self.assertEqual(found['status'], 'complete')
        self.assertEqual(len(found['comparisons']), 1)
        comparison = found['comparisons'][0]
        self.assertEqual(comparison['clause_id'], TARGET_CLAUSE)
        self.assertIn(comparison['current'], ('satisfied', 'not_satisfied'))
        self.assertIn(comparison['target'], ('satisfied', 'not_satisfied'))


class RunDirTests(unittest.TestCase):
    def test_existing_run_dir_stops_before_any_work(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / '입력.xlsx'
            bc.write_input([TASKS[0]], source)
            listing = root / '목록.json'
            listing.write_text(json.dumps(TASK_LIST[:1], ensure_ascii=False), encoding='utf-8')
            run_dir = root / '실행1'
            run_dir.mkdir()
            (run_dir / '기존.txt').write_text('keep', encoding='utf-8')
            with self.assertRaises(FileExistsError):
                bc.run_bulk(source, listing, run_dir)
            self.assertEqual(sorted(os.listdir(run_dir)), ['기존.txt'])
            out = io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
                code = bc.main(['run', str(source), str(listing), str(run_dir)])
            self.assertNotEqual(code, 0)
            self.assertEqual(sorted(os.listdir(run_dir)), ['기존.txt'])

    def test_bad_task_list_stops_before_creating_run_dir(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / '입력.xlsx'
            bc.write_input([TASKS[0]], source)
            listing = root / '목록.json'
            listing.write_text(json.dumps([dict(과제명=NAMES[0], 분야='없는분야', 고시대상=False, 기대='등급')],
                                          ensure_ascii=False), encoding='utf-8')
            with self.assertRaises(ValueError):
                bc.run_bulk(source, listing, root / '실행1')
            self.assertFalse((root / '실행1').exists())


class BulkRunTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        root = Path(cls.temp.name)
        cls.source = root / '입력_합성5.xlsx'
        bc.write_input(TASKS[:1] + EXCLUDED[:1] + TASKS[1:3] + EXCLUDED[1:] + TASKS[3:], cls.source)
        cls.listing = root / '모의과제_목록.json'
        cls.listing.write_text(json.dumps(TASK_LIST, ensure_ascii=False), encoding='utf-8')
        cls.run_dir = root / 'v16_test' / '실행1'
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            cls.code = bc.main(['run', str(cls.source), str(cls.listing), str(cls.run_dir)])
        cls.stdout = out.getvalue()
        cls.outcome = bc.LAST_OUTCOME

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_counts_and_persistence(self):
        self.assertEqual(self.code, 0, self.stdout)
        run = self.outcome['run']
        self.assertEqual(run['job_state'], 'done')
        report = run['report']
        self.assertEqual(report['task_count'], 5)
        self.assertEqual(report['input_counts'], {'blank': 1, 'example': 1, 'duplicate': 1})
        self.assertEqual(len(report['results']), 5)
        self.assertEqual(report['counts']['오류'], 0)
        output = self.run_dir / 'v16_검토결과.xlsx'
        self.assertTrue(output.is_file())
        checkpoint = Path(str(output) + '.checkpoint.jsonl')
        audit = Path(str(output) + '.audit.jsonl')
        lines = [json.loads(line) for line in checkpoint.read_text(encoding='utf-8').splitlines()]
        self.assertEqual(len(lines), 5)
        events = [json.loads(line) for line in audit.read_text(encoding='utf-8').splitlines()]
        self.assertEqual(sum(e['event'] == 'batch_task_saved' for e in events), 5)
        formula = next(r for r in lines if r['task_name'] == NAMES[1])
        self.assertTrue(formula['task_snapshot']['연구내용'].startswith('=합계'))
        self.assertTrue((self.run_dir / '의견서').is_dir())
        self.assertTrue(run['events'])
        self.assertEqual(sum(e['message'] == '입력 글의 엑셀 줄바꿈 표시·줄 끝 공백·연속 빈 줄 정리 후 검토'
                             for e in run['events']), 1)
        self.assertGreater(run['total_seconds'], 0)
        self.assertEqual(set(run['task_calls']), set(NAMES))
        for name in NAMES:
            calls = run['task_calls'][name]
            self.assertGreater(calls['calls'], 0)
            self.assertGreaterEqual(calls['seconds'], 0)
            self.assertGreater(calls['prompt_chars'], 0)
        self.assertEqual(run['unmatched_clients'], 0)

    def test_document_checks(self):
        checks = self.outcome['checks']
        self.assertTrue(checks['ok'], checks['problems'])
        results = self.outcome['run']['report']['results']
        expected = sum(len(r['reviewers']) + 1 for r in results)
        disk = hwpx_files(self.run_dir / '의견서')
        self.assertEqual(len(disk), expected)
        self.assertEqual(checks['disk_files'], expected)
        self.assertEqual(checks['listed_files'], expected)
        by_name = {task['task_name']: task for task in checks['tasks']}
        for result in results:
            task = by_name[result['task_name']]
            self.assertTrue(task['ok'], task['problems'])
            self.assertEqual(len(task['files']), len(result['reviewers']) + 1)
            for record, entry in zip(task['files'], result['documents']['files']):
                name = Path(entry['path']).name
                self.assertEqual(record['name'], name)
                self.assertIn('[모의]', name)
                self.assertTrue(name.startswith(hw._safe_name(result['task_name'], 20) + '_'
                                                + result['task_fingerprint'][:12] + '_'))
                self.assertTrue(record['zip_ok'] and record['xml_ok'] and record['name_ok'])
                self.assertEqual(record['size'], os.path.getsize(entry['path']))
                with ZipFile(entry['path']) as archive:
                    self.assertEqual(archive.read('mimetype'), b'application/hwp+zip')
        targeted_task = by_name[NAMES[3]]
        self.assertTrue(targeted_task['candidate_ids'])
        self.assertTrue(all(record['notice_ok'] for record in targeted_task['files']))

    def test_check_detects_missing_file(self):
        results = json.loads(json.dumps(self.outcome['run']['report']['results'], ensure_ascii=False))
        broken = results[0]
        broken['documents']['files'][0]['path'] += '.없음'
        checks = bc.check_documents(results, self.run_dir / '의견서')
        self.assertFalse(checks['ok'])
        self.assertTrue(checks['problems'])

    def test_expectations_and_analysis_outputs(self):
        analysis = self.outcome['analysis']
        records = {r['과제명']: r for r in analysis['records']}
        self.assertEqual(set(records), set(NAMES))
        self.assertEqual(records[NAMES[3]]['상태'], '완료')
        self.assertGreaterEqual(records[NAMES[3]]['후보 조항 수'], 1)
        self.assertEqual(records[NAMES[4]]['상태'], '검토필요')
        for name in NAMES[:3]:
            self.assertEqual(records[name]['후보 조항 수'], 0)
            self.assertEqual(records[name]['고시 대조 상태'], 'complete')
        self.assertTrue(records[NAMES[1]]['입력 정리'])
        self.assertTrue(records[NAMES[2]]['입력 정리'])
        self.assertFalse(records[NAMES[0]]['입력 정리'])
        self.assertTrue(all(r['기대 일치'] for r in analysis['records']), analysis['mismatches'])
        self.assertEqual(analysis['mismatches'], [])
        self.assertEqual(analysis['excluded'], {'blank': 1, 'example': 1, 'duplicate': 1})
        self.assertEqual(analysis['expected_excluded'], 3)
        self.assertEqual(analysis['cleaned_tasks'], 2)
        self.assertEqual(set(analysis['regression']),
                         {'AI 호출 수', '지시문 총 문자 수', '처리 시간(초)', '의견서 총 크기(바이트)'})
        for fit in analysis['regression'].values():
            self.assertEqual(fit['n'], 5)
        workbook = self.run_dir / '분석_결과.xlsx'
        book = load_workbook(workbook)
        self.assertEqual(book.sheetnames, ['요약', '과제별', '통계', '회귀', '의견서 검사'])
        self.assertEqual(book['과제별'].max_row, 6)
        summary_text = ' '.join(str(c.value) for row in book['요약'].iter_rows() for c in row if c.value)
        self.assertIn('모의 규칙', summary_text)
        self.assertIn('모델 품질과 무관', summary_text)
        report = (self.run_dir / '요약보고서.md').read_text(encoding='utf-8')
        for phrase in ('합성 시험 과제', '모의 규칙', '의견서 쓰기는 빠졌', '모델 품질과 무관', '한글 프로그램',
                       '미확인', '100,000자', '예시과제집'):
            self.assertIn(phrase, report)


if __name__ == '__main__':
    unittest.main()
