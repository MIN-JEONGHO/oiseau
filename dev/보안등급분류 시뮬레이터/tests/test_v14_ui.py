"""Withdrawn Tk checks only: synthetic inputs, fake credentials and clients."""
import copy
import gc
import importlib.util
from pathlib import Path
import tempfile
import threading
import time
import tkinter as tk
import unittest
from unittest.mock import patch

from credential_store import KeyManager
from v14_domain import ITEM_NAMES, make_result
from security_grade_classifier_v14 import ClassifierApp, self_check


class MemoryStore:
    def __init__(self):
        self.keys = {}

    def get(self, provider):
        return self.keys.get(provider)

    def set(self, provider, key):
        self.keys[provider] = key

    def delete(self, provider):
        self.keys.pop(provider, None)


class UITests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.store = MemoryStore()
        self.app = ClassifierApp(self.root, key_manager=KeyManager(self.store, {}),
                                 client_factory=lambda **kw: object())

    def tearDown(self):
        self.app.close()
        # Destroyed Tcl interpreters must be finalized on this main thread.
        # A later worker's GC must not inherit cycles from a prior test root.
        self.app = None
        self.root = None
        gc.collect()

    def drain(self):
        deadline = time.monotonic() + 4
        while self.app.busy and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.01)
        self.root.update()
        self.assertFalse(self.app.busy)

    def task(self):
        self.app.task_fields['과제명'].insert('1.0', '합성 UI 과제')

    def test_provider_separation_apply_delete_reload(self):
        self.app.key_entry_var.set('synthetic-claude-key')
        self.app.apply_settings()
        self.assertEqual(self.store.keys['claude'], 'synthetic-claude-key')
        self.assertEqual(self.app.key_entry_var.get(), '')
        self.app.key_entry_var.set('unsaved-must-clear')
        self.app.provider_var.set('openai')
        self.app.on_provider_change()
        self.assertEqual(self.app.key_entry_var.get(), '')
        self.assertEqual(self.app.model_var.get(), 'gpt-6-luna')
        self.app.key_entry_var.set('synthetic-openai-key')
        self.app.apply_settings()
        self.app.provider_var.set('claude')
        self.app.on_provider_change()
        self.assertIn('저장', self.app.key_status_var.get())
        self.app.delete_key()
        self.assertNotIn('claude', self.store.keys)
        self.assertIn('openai', self.store.keys)
        self.assertEqual(KeyManager(self.store, {}).resolve('openai').source, 'saved')

    def test_agent_count_and_result_labels(self):
        self.assertEqual(self.app.reviewers_var.get(), '3개')
        count_combo = self.app.readonly_controls[0]
        self.assertEqual(tuple(count_combo.cget('values')), ('1개', '3개'))
        self.app.reviewers_var.set('1개')
        self.assertEqual(self.app.snapshot_options().num_reviewers, 1)
        self.app.reviewers_var.set('3개')
        self.assertEqual(self.app.snapshot_options().num_reviewers, 3)
        overview = self.app.detail_notebook.tabs()[0]
        self.assertEqual(self.app.detail_notebook.tab(overview, 'text'), 'AI 종합 검토의견')
        pending, labels = [self.root], []
        while pending:
            widget = pending.pop()
            pending.extend(widget.winfo_children())
            if 'text' in widget.keys():
                labels.append(widget.cget('text'))
        self.assertIn('AI 에이전트 검토', labels)
        self.assertIn('AI 에이전트 수', labels)
        self.assertIn('종합평균', labels)
        self.assertNotIn('가상 위원 수', labels)

    def test_checklist_and_conditional_eighth(self):
        self.app.selection_var.set('manual')
        self.app.refresh_method()
        self.assertIn('disabled', self.app.item_checks[7].state())
        self.app.policy_text.insert('1.0', '합성 정책 척도')
        self.app.refresh_method()
        self.assertNotIn('disabled', self.app.item_checks[7].state())
        self.app.method_var.set('checklist')
        self.app.refresh_method()
        self.assertTrue(all('disabled' in widget.state() for widget in self.app.weight_entries))
        self.assertTrue(all(v.get() for v in self.app.item_vars[:7]))
        options = self.app.snapshot_options()
        self.assertEqual(options.method, 'checklist')
        self.assertEqual(options.weights, ())

    def test_validation_and_cancelled_dialogs(self):
        self.app.selection_var.set('manual')
        self.app.weight_vars[0].set('1.2')
        with self.assertRaises(ValueError):
            self.app.snapshot_options()
        self.app.result = make_result({'과제명': '합성'}, {})
        before = self.app.status_var.get()
        with patch('security_grade_classifier_v14.filedialog.askdirectory', return_value=''):
            self.app.export_documents()
        self.assertEqual(self.app.status_var.get(), before)
        self.assertFalse(self.app.busy)

    def test_async_snapshot_busy_and_main_thread_queue(self):
        self.task()
        self.app.key_entry_var.set('synthetic-only')
        gate = threading.Event()
        observed = []
        main_ident = threading.get_ident()
        original_after = self.root.after
        def checked_after(*args, **kwargs):
            self.assertEqual(threading.get_ident(), main_ident)
            return original_after(*args, **kwargs)
        self.root.after = checked_after
        def classify(task, options, client, progress_callback=None):
            gate.wait(2)
            observed.append((copy.deepcopy(task), options.method, threading.get_ident()))
            progress_callback('합성 처리 중')
            result = make_result(task, options)
            result['status'] = '검토필요'
            result['unresolved_reasons'] = ['합성 부족 사유']
            return result
        self.app.classify_func = classify
        self.app.start_single()
        self.assertTrue(self.app.busy)
        self.assertIn('disabled', self.app.run_button.state())
        self.assertIn('disabled', self.app.apply_button.state())
        self.app.method_var.set('checklist')
        self.app.start_single()
        gate.set()
        self.drain()
        self.assertEqual(len(observed), 1)
        self.assertEqual(observed[0][1], 'impact')
        self.assertNotEqual(observed[0][2], main_ident)
        self.assertEqual(self.app.card_vars['status'].get(), '검토필요')
        self.assertEqual(self.app.card_vars['average'].get(), '미산출')
        self.assertIn('합성 부족 사유', self.app.result_text.get('1.0', 'end'))

    def test_document_retry_preserves_result_without_client(self):
        self.app.result = make_result({'과제명': '합성'}, {})
        calls = []
        def documents(result, folder):
            calls.append(folder)
            result = copy.deepcopy(result)
            result['documents'] = {'status': '완료', 'expected_files': 1,
                                   'files': [{'status': '완료'}], 'errors': []}
            return result
        self.app.document_func = documents
        self.app.client_factory = lambda **kw: self.fail('No model for document retry')
        with tempfile.TemporaryDirectory() as directory:
            with patch('security_grade_classifier_v14.filedialog.askdirectory', return_value=directory):
                self.app.export_documents()
                self.drain()
            self.app.retry_documents()
            self.drain()
        self.assertEqual(len(calls), 2)
        self.assertEqual(self.app.result['status'], '검토필요')

    def test_selected_batch_document_retry_durably_updates_original_run(self):
        import batch_v14 as batch
        from io import BytesIO
        import json
        from openpyxl import Workbook, load_workbook
        spec = importlib.util.spec_from_file_location('v14_ui_retry_fakes', Path(__file__).with_name('test_v14_engine.py'))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        clients = []
        def factory():
            client = module.FakeClient()
            clients.append(client)
            return client
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            input_path, output_path = folder / 'input.xlsx', folder / 'output.xlsx'
            book = Workbook()
            book.active.append(['과제명', '연구내용'])
            book.active.append(['합성 GUI 문서 재시도', '합성 연구'])
            book.save(input_path)
            book.close()
            options = {'selection_mode': 'manual', 'selected_items': [ITEM_NAMES[0]], 'weights': {ITEM_NAMES[0]: 1}}
            with patch('batch_v14.export_documents', side_effect=OSError('synthetic document failure')):
                report = batch.run_excel(input_path, output_path, options, factory, workers=1)
            before = sum(len(client.calls) for client in clients)
            self.app.render_batch(report)
            self.app.batch_list.selection_set('0')
            self.app.select_batch_result()
            # Later GUI settings must not become the retry's run context.
            self.app.provider_var.set('openai')
            self.app.method_var.set('checklist')
            self.app.input_path_var.set('not-the-original-input.xlsx')
            self.app.client_factory = lambda **kwargs: self.fail('No model call for document retry')
            with patch('security_grade_classifier_v14.filedialog.askdirectory', return_value=str(folder / 'new-documents')):
                self.app.export_documents()
            self.drain()
            self.assertEqual(sum(len(client.calls) for client in clients), before)
            self.assertIn('결과 기록 완료', self.app.status_var.get())
            self.assertIn('엑셀 완료', self.app.status_var.get())
            self.assertEqual(self.app.batch_report['counts']['documents_missing'], 0)
            latest = json.loads(Path(report['checkpoint_path']).read_text(encoding='utf-8').splitlines()[-1])
            self.assertEqual(latest['documents'], self.app.result['documents'])
            self.assertTrue(all(Path(entry['path']).is_file() for entry in latest['documents']['files']))
            updated_book = load_workbook(BytesIO(output_path.read_bytes()))
            self.addCleanup(updated_book.close)
            self.assertEqual(updated_book['분류결과']['H2'].value, '직접 선택')
            self.assertEqual(updated_book['분류결과']['M2'].value, '완료')
            self.assertEqual(updated_book['분류결과']['Q2'].value, 0)
            self.assertIn(latest['documents']['files'][0]['path'], updated_book['분류결과']['R2'].value)
            event = json.loads(Path(report['audit_path']).read_text(encoding='utf-8').splitlines()[-1])
            self.assertEqual(event['event'], 'batch_documents_regenerated')

    def test_batch_document_persistence_error_keeps_generated_files_visible(self):
        result = make_result({'과제명': '합성 기록 실패'}, {})
        report = {'status': '완료', 'results': [result], 'counts': {'완료': 0, '검토필요': 1, '후보제외': 0, '오류': 0,
                  'documents_expected': 4, 'documents_completed': 0, 'documents_failed': 0, 'documents_missing': 4},
                  'warnings': [], 'input_counts': {}, 'output_path': 'synthetic.xlsx', 'audit_status': '완료', 'documents_dir': ''}
        self.app.render_batch(report)
        self.app.batch_list.selection_set('0')
        self.app.select_batch_result()
        def generate(saved, folder):
            saved['rationale'] = 'synthetic changed classification must not replace original'
            saved['documents'] = {'status': '완료', 'expected_files': 1,
                                  'files': [{'status': '완료', 'path': str(Path(folder) / 'generated.hwpx')}], 'errors': []}
            return saved
        def fail_persistence(*args):
            raise OSError('synthetic raw persistence error')
        self.app.document_func = generate
        self.app.persist_documents_func = fail_persistence
        with tempfile.TemporaryDirectory() as directory:
            self.app.document_dir = directory
            self.app.retry_documents()
            self.drain()
        self.assertIn('문서 생성 완료', self.app.status_var.get())
        self.assertIn('결과 기록 실패', self.app.status_var.get())
        self.assertNotIn('synthetic raw', self.app.status_var.get())
        self.assertIn('완료 1/1', self.app.card_vars['documents'].get())
        self.assertEqual(self.app.result['status'], '검토필요')
        self.assertEqual(self.app.result['rationale'], result['rationale'])
        self.assertEqual(self.app.batch_report['results'][0]['documents']['status'], '미생성')

    def test_real_engine_fake_client_async_reviewer_details(self):
        spec = importlib.util.spec_from_file_location('v14_ui_engine_fakes', Path(__file__).with_name('test_v14_engine.py'))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        fake = module.FakeClient()
        self.app.client_factory = lambda **kwargs: fake
        self.task()
        self.app.key_entry_var.set('synthetic-engine-key')
        self.app.start_single()
        self.drain()
        self.assertEqual(self.app.result['status'], '완료')
        self.assertEqual(len(self.app.detail_notebook.tabs()), 4)
        self.assertEqual([self.app.detail_notebook.tab(tab, 'text') for tab in self.app.detail_notebook.tabs()[1:]],
                         ['에이전트 1 검토의견', '에이전트 2 검토의견', '에이전트 3 검토의견'])
        self.assertEqual(self.app.card_vars['grade'].get(), '민감과제')
        self.assertEqual(len(fake.calls), 6)
        self.assertNotIn('synthetic-engine-key', str(self.app.result))
        self.app.method_var.set('checklist')
        self.app.refresh_method()
        self.app.start_single()
        self.drain()
        self.assertEqual(self.app.card_vars['average'].get(), '해당 없음')

    def test_async_batch_cancel_counts_and_selected_result_retry(self):
        self.app.key_entry_var.set('synthetic-batch-key')
        self.app.input_path_var.set('synthetic-input.xlsx')
        self.app.output_path_var.set('synthetic-output.xlsx')
        seen = []
        started = threading.Event()
        def batch(input_path, output_path, options, factory, **kwargs):
            seen.append((options.method, kwargs['workers'], kwargs['make_forms']))
            factory()
            started.set()
            kwargs['cancel_event'].wait(2)
            result = make_result({'과제명': '합성 대량 과제'}, options)
            result['documents'] = {'status': '일부실패', 'expected_files': 4, 'files': [{'status': '완료', 'path': 'synthetic-document.hwpx'}], 'errors': ['합성 문서 실패']}
            return {'status': '취소', 'results': [result], 'counts': {'완료': 0, '검토필요': 1, '후보제외': 0, '오류': 0,
                    'documents_expected': 4, 'documents_completed': 1, 'documents_failed': 3, 'documents_missing': 3},
                    'audit_status': '실패', 'input_counts': {'중복': 2}, 'warnings': ['합성 감사기록 저장 실패'],
                    'output_path': output_path, 'documents_dir': '', 'reused_count': 0}
        self.app.batch_func = batch
        self.app.start_batch()
        self.assertTrue(started.wait(2))
        self.app.method_var.set('checklist')
        self.app.cancel_batch()
        self.drain()
        self.assertEqual(seen, [('impact', 3, True)])
        self.assertIn('누락 3', self.app.batch_status_var.get())
        self.assertIn('감사기록: 실패', self.app.batch_status_var.get())
        self.app.batch_list.selection_set('0')
        self.app.select_batch_result()
        self.assertEqual(self.app.result['task_name'], '합성 대량 과제')
        self.assertEqual(self.app.card_vars['status'].get(), '검토필요')

    def test_unexpected_exception_never_displays_raw_secret(self):
        self.task()
        self.app.key_entry_var.set('synthetic-error-key')
        def failed(**kwargs):
            raise RuntimeError('Authorization synthetic-error-key full raw exception')
        self.app.client_factory = failed
        self.app.start_single()
        self.drain()
        self.assertEqual(self.app.card_vars['status'].get(), '오류')
        self.assertNotIn('synthetic-error-key', self.app.status_var.get())
        self.assertNotIn('Authorization', self.app.result_text.get('1.0', 'end'))

    def test_storage_failure_and_environment_are_visible_without_key_text(self):
        class FailedStore(MemoryStore):
            def set(self, provider, key):
                raise OSError('synthetic-should-not-display')
        self.app.key_manager = KeyManager(FailedStore(), {'OPENAI_API_KEY': 'synthetic-env-key'})
        self.app.key_entry_var.set('synthetic-session-key')
        self.app.apply_settings()
        self.assertIn('이번 실행', self.app.key_status_var.get())
        self.assertIn('저장하지 못', self.app.key_status_var.get())
        self.assertNotIn('synthetic-', self.app.key_status_var.get())
        self.app.provider_var.set('openai')
        self.app.on_provider_change()
        self.assertIn('환경변수', self.app.key_status_var.get())
        self.assertEqual(self.app.key_entry_var.get(), '')
        self.app.delete_key()
        self.assertIn('환경변수', self.app.key_status_var.get())

    def test_small_window_and_font_scale_layout(self):
        for scale in (1.0, 1.25, 1.5):
            self.app.set_font_scale(scale)
            self.root.geometry('900x650')
            self.root.update_idletasks()
            self.assertEqual(self.root.state(), 'withdrawn')
            self.assertLessEqual(self.app.footer.winfo_reqwidth(), 900)
            self.assertLessEqual(self.app.header.winfo_reqheight() + self.app.footer.winfo_reqheight(), 300)
            self.assertEqual(self.app.settings_scroll.canvas.cget('yscrollcommand') != '', True)

    def test_self_check_fake_store_only(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.json'
            self.assertEqual(self_check(path), 0)
            self.assertTrue(path.exists())

    def test_close_while_busy_requests_cancel_and_preserves_window(self):
        self.app._set_busy(True)
        self.app.close()
        self.assertTrue(self.app.cancel_event.is_set())
        self.assertFalse(self.app.closed)
        self.assertTrue(self.root.winfo_exists())
        self.assertEqual(self.root.state(), 'withdrawn')
        self.assertIn('즉시 중단할 수 없습니다', self.app.status_var.get())
        self.app._set_busy(False)


if __name__ == '__main__':
    unittest.main()
