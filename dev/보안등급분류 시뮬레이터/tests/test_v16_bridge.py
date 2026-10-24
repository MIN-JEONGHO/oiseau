"""Offline v16 desktop bridge contract checks; no native store, dialog or API access.

The pywebview dispatcher check that needs the host module (bind_public_api)
lives in tests/test_v16_packaging.py with the v16 host.
"""
import ast
from collections import Counter
from contextlib import nullcontext
from copy import deepcopy
import json
import os
from pathlib import Path
import re
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import batch_v16
import classification_v16 as engine
from credential_store import KeyManager
import desktop_bridge_v16 as bridge_module
from desktop_bridge_v16 import DesktopBridge
from notice_catalogue_v16 import CatalogueError, load_catalogue
import v16_domain as domain
from test_v16_domain import APPENDIX_A_DEFINITIONS, APPENDIX_A_PRINCIPLES
from test_v16_engine import Client
from test_v16_notice_review import STYLE_ENDINGS, catalogue, repack


PUBLIC_METHODS = {
    "get_initial_state", "select_path", "save_key", "delete_key", "start_single",
    "start_batch", "get_job", "cancel_job", "start_export",
    "search_notice_clauses", "get_notice_clause", "save_input_template",
}
# Spec 4.4/4.5: the one failure whose message is not _MESSAGES[code].
TEMPLATE_PROTECTED = "검토에 쓰는 파일 보호를 위해 저장 중단. 다른 파일 이름 선택 필요"
ERROR_CODES = {"validation", "catalogue", "credential", "busy", "not_found", "io", "runtime"}
STAGES = {"preflight", "facts", "notices", "candidate", "criteria", "agents", "summary", "documents", "persistence"}
TERMINAL = {"done", "cancelled", "failed"}
SECRET = "synthetic-bridge-secret-never-render"
TASK = {"과제명": "합성 브리지 과제", "현재 기술사양": "속도 20 GHz"}
# Spec 4.1: notice number -> (title, announcement date, effective date).
NOTICES = {
    "산업통상자원부고시 제2025-71호": ("국가첨단전략기술 지정 등에 관한 고시", "2025-05-12", "2025-05-18"),
    "산업통상부고시 제2025-2호": ("국가핵심기술 지정 등에 관한 고시", "2025-10-02", "2025-10-02"),
    "산업통상부고시 제2026-101호": ("전략물자수출입고시", "2026-08-31", "2026-09-01"),
}
NOTICE_DATES = {number: dates for number, (_, *dates) in NOTICES.items()}
SEARCH_ITEM_KEYS = {"clause_id", "family", "source_id", "notice_number", "code", "title", "field"}
CLAUSE_KEYS = {"clause_id", "family", "source_id", "notice_number", "notice_title", "announcement_date",
               "effective_date", "official_url", "code", "title", "field", "source_anchor", "text"}
STYLE_PATTERN = re.compile("(?:" + "|".join(STYLE_ENDINGS) + r")[)\]」』’”'\"]*(?=$|[\s.!?,…:;])")
HANGUL = re.compile("[가-힣]")


def style_violations(text):
    return [match.group(0) for match in STYLE_PATTERN.finditer(text)]


def distinct_url_catalogue():
    """Synthetic catalogue whose sources each have their own official URL."""
    plain = catalogue().to_dict()
    for source in plain["manifest"]["sources"]:
        source["official_url"] = "https://example.gov.test/" + source["source_id"]
    return repack(plain)


def unvalidated_preflight(cat):
    """Stand-in for notice_review_v16._preflight that skips domain validation."""
    plain = json.loads(json.dumps(cat.to_dict(), ensure_ascii=False))
    return plain["manifest"], {clause["clause_id"]: clause for clause in plain["clauses"]}


class FakeBackend:
    def __init__(self):
        self.values = {}
        self.fail = set()
        self.calls = []

    def _check(self, operation, provider):
        self.calls.append((operation, provider))
        if operation in self.fail:
            raise RuntimeError("native credential error " + SECRET)

    def get(self, provider):
        self._check("get", provider)
        return self.values.get(provider)

    def set(self, provider, key):
        self._check("set", provider)
        self.values[provider] = key

    def delete(self, provider):
        self._check("delete", provider)
        self.values.pop(provider, None)


def regenerated(result, output_dir):
    result = deepcopy(result)
    files = [dict(kind="reviewer", reviewer_id=r["reviewer_id"],
                  path=str(Path(output_dir) / (r["reviewer_id"] + ".hwpx")), status="완료", error="")
             for r in result["reviewers"]]
    files.append(dict(kind="summary", reviewer_id=None, path=str(Path(output_dir) / "summary.hwpx"),
                      status="완료", error=""))
    result["documents"] = dict(status="완료", expected_files=len(files), files=files, errors=[])
    return result


class BridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.input = self.root / "input.xlsx"
        self.input.write_bytes(b"synthetic selection; batch runner is injected")
        self.output = self.root / "output.xlsx"
        self.directory = self.root / "documents"
        self.directory.mkdir()
        self.template = self.root / "v16_일괄검토_입력양식.xlsx"
        self.selected = {"input_excel": self.input, "output_excel": self.output,
                         "output_directory": self.directory, "input_template": self.template}
        self.backend = FakeBackend()
        self.backend.values["claude"] = SECRET
        self.cat = catalogue()
        self.clients = []
        self.catalogue_calls = []
        self.client_keys = []
        self.bridges = []
        self.bridge = self.make_bridge()

    def load_catalogue(self):
        self.catalogue_calls.append(True)
        return self.cat

    def make_client(self, options, key):
        self.client_keys.append(key)
        client = Client(facts=[])
        self.clients.append(client)
        return client

    def make_bridge(self, **overrides):
        kwargs = dict(key_manager=KeyManager(self.backend, environ={}),
                      catalogue_loader=self.load_catalogue, client_factory=self.make_client,
                      path_selector=lambda purpose: self.selected[purpose])
        kwargs.update(overrides)
        bridge = DesktopBridge(**kwargs)
        self.bridges.append(bridge)
        return bridge

    def envelope(self, response, *, ok=True, code=None, protected=False):
        self.assertIsInstance(response, dict)
        self.assertEqual(set(response), {"ok", "data", "error"})
        self.assertIs(response["ok"], ok)
        json.dumps(response, ensure_ascii=False, allow_nan=False)
        self.assertNotIn(SECRET, json.dumps(response, ensure_ascii=False))
        if ok:
            self.assertIsNone(response["error"])
            return response["data"]
        self.assertIsNone(response["data"])
        self.assertEqual(set(response["error"]), {"code", "message"})
        self.assertIn(response["error"]["code"], ERROR_CODES)
        self.assertTrue(response["error"]["message"])
        if code:
            self.assertEqual(response["error"]["code"], code)
        # Only the fixed message for the code: never exception text or a stack.
        # The single exception is save_input_template's protected-path refusal
        # (spec 4.5): code validation with its own fixed _EVENTS message.
        if protected:
            self.assertEqual(response["error"]["code"], "validation")
            self.assertEqual(response["error"]["message"], TEMPLATE_PROTECTED)
            self.assertEqual(response["error"]["message"], bridge_module._EVENTS["template_protected"])
        else:
            self.assertEqual(response["error"]["message"], bridge_module._MESSAGES[response["error"]["code"]])
        return response["error"]

    def token(self, purpose, bridge=None):
        data = self.envelope((bridge or self.bridge).select_path(purpose))
        self.assertIs(data["cancelled"], False)
        self.assertIsInstance(data["path_token"], str)
        self.assertTrue(data["path_token"])
        self.assertIsInstance(data["display_name"], str)
        self.assertTrue(data["display_name"])
        self.assertNotIn(str(self.root), data["path_token"])
        return data["path_token"]

    def single_request(self, **updates):
        return dict(task=deepcopy(TASK), options={"num_reviewers": 1}, output_dir_token=None) | updates

    def batch_request(self, bridge=None, **updates):
        return dict(input_token=self.token("input_excel", bridge), output_token=self.token("output_excel", bridge),
                    output_dir_token=None, options={"num_reviewers": 1}, workers=1,
                    resume=False, make_forms=False) | updates

    def wait_job(self, job_id, bridge=None):
        bridge = bridge or self.bridge
        deadline = time.monotonic() + 6
        while time.monotonic() < deadline:
            data = self.envelope(bridge.get_job(job_id, 0))
            self.assertEqual(data["job_id"], job_id)
            if data["state"] in TERMINAL:
                return data
            time.sleep(.005)
        self.fail("synthetic bridge job did not finish")

    def start_single(self, request=None, bridge=None):
        bridge = bridge or self.bridge
        data = self.envelope(bridge.start_single(request or self.single_request()))
        self.assertEqual(set(data), {"job_id"})
        self.assertIsInstance(data["job_id"], str)
        return data["job_id"]

    def result(self, task=None):
        return engine.classify_task(task or TASK, {"num_reviewers": 1}, Client(facts=[]), catalogue=self.cat)

    def test_only_twelve_public_methods_and_safe_arity_envelopes(self):
        actual = {name for name in dir(self.bridge) if not name.startswith("_") and callable(getattr(self.bridge, name))}
        self.assertEqual(actual, PUBLIC_METHODS)
        valid_args = {
            "get_initial_state": (), "select_path": ("input_excel",),
            "save_key": ("claude", SECRET), "delete_key": ("claude",),
            "start_single": ({},), "start_batch": ({},), "get_job": ("unknown", 0),
            "cancel_job": ("unknown",), "start_export": ({},),
            "search_notice_clauses": ({},), "get_notice_clause": ("advanced:1",),
            "save_input_template": ({},),
        }
        self.assertEqual(set(valid_args), PUBLIC_METHODS)
        for name, args in valid_args.items():
            method = getattr(self.bridge, name)
            with self.subTest(method=name, case="extra positional"):
                self.envelope(method(*args, None), ok=False, code="validation")
            with self.subTest(method=name, case="unknown keyword"):
                self.envelope(method(untrusted=True), ok=False, code="validation")
            if args:
                with self.subTest(method=name, case="missing positional"):
                    self.envelope(method(*args[:-1]), ok=False, code="validation")

    def test_initial_state_is_renderer_shape_without_credentials(self):
        data = self.envelope(self.bridge.get_initial_state())
        self.assertEqual(set(data), {"defaults", "criteria", "credential_status", "catalogue", "guideline"})
        self.assertEqual(data["defaults"]["model"], "claude-sonnet-5-5")
        self.assertEqual(data["credential_status"], {"claude": True, "openai": False})
        self.assertEqual(len(data["criteria"]), 8)
        self.assertEqual({entry["id"] for entry in data["criteria"]}, set(range(1, 9)))
        self.assertEqual(data["catalogue"]["status"], "ready")
        self.assertEqual(data["catalogue"]["fingerprint"], self.cat.fingerprint)
        self.assertEqual(len(data["catalogue"]["sources"]), 4)
        self.assertEqual(set(data["catalogue"]), {"catalogue_id", "fingerprint", "status", "sources", "families",
                                                  "notices", "notice_titles"})
        for source in data["catalogue"]["sources"]:
            self.assertEqual(source["announcement_date"], NOTICE_DATES[source["notice_number"]][0])
        self.assertEqual({entry["family"] for entry in data["catalogue"]["families"]}, set(domain.FAMILIES))
        for entry in data["catalogue"]["families"]:
            self.assertEqual((entry["total_clauses"], entry["interpreted_clauses"], entry["unsupported_clauses"]), (1, 1, 0))
        self.assertEqual(data["guideline"]["id"], domain.GUIDELINE_ID)
        for field in ("definitions", "principles"):
            self.assertTrue(data["guideline"][field])
        data["defaults"]["num_reviewers"] = 900
        self.assertNotEqual(self.envelope(self.bridge.get_initial_state())["defaults"]["num_reviewers"], 900)

    def test_credential_provider_isolation_and_safe_failures(self):
        self.envelope(self.bridge.save_key("openai", "synthetic-openai"))
        self.assertEqual(self.backend.values["openai"], "synthetic-openai")
        self.envelope(self.bridge.delete_key("openai"))
        self.assertNotIn("openai", self.backend.values)
        self.assertEqual(self.backend.values["claude"], SECRET)
        for args in (("ollama", SECRET), (None, SECRET), ({}, SECRET), (["claude"], SECRET),
                     ("openai", None), ("openai", 7), ("openai", "***")):
            with self.subTest(args=(args[0], type(args[1]).__name__)):
                self.envelope(self.bridge.save_key(*args), ok=False, code="validation")
        self.backend.fail.add("set")
        self.envelope(self.bridge.save_key("openai", SECRET), ok=False, code="credential")
        self.backend.fail.add("delete")
        self.envelope(self.bridge.delete_key("claude"), ok=False, code="credential")
        self.envelope(self.bridge.delete_key("lmstudio"), ok=False, code="validation")

    def test_new_bridge_reuses_saved_provider_key_and_reports_read_failure(self):
        self.envelope(self.bridge.save_key("openai", "synthetic-restart-openai"))
        fresh = self.make_bridge()
        data = self.envelope(fresh.get_initial_state())
        self.assertEqual(data["credential_status"], {"claude": True, "openai": True})
        self.envelope(fresh.delete_key("openai"))
        newest = self.make_bridge()
        self.assertEqual(self.envelope(newest.get_initial_state())["credential_status"], {"claude": True, "openai": False})
        self.backend.fail.add("get")
        failing = self.make_bridge()
        self.envelope(failing.get_initial_state(), ok=False, code="credential")
        self.envelope(failing.start_single(self.single_request()), ok=False, code="credential")
        self.assertFalse(self.clients)

    def test_failed_save_session_key_remains_usable_with_explicit_warning(self):
        self.backend.values.clear()
        self.backend.fail.add("set")
        self.envelope(self.bridge.save_key("claude", SECRET), ok=False, code="credential")
        self.assertNotIn("claude", self.backend.values)
        job = self.wait_job(self.start_single())
        self.assertEqual(job["state"], "done")
        self.assertEqual(job["result"]["status"], "완료")
        self.assertEqual(self.client_keys, [SECRET])
        warnings = " ".join(event["message"] for event in job["events"] if event["stage"] == "preflight")
        for text in ("저장", "이번 실행", "다음 실행"):
            self.assertIn(text, warnings)
        self.assertTrue("실패" in warnings or "못했습니다" in warnings)
        fresh = self.make_bridge()
        self.envelope(fresh.start_single(self.single_request()), ok=False, code="credential")
        self.assertEqual(self.client_keys, [SECRET])

    def test_environment_key_is_used_with_source_warning_and_never_saved(self):
        self.backend.values.clear()
        bridge = self.make_bridge(key_manager=KeyManager(self.backend, environ={"ANTHROPIC_API_KEY": SECRET}))
        self.assertEqual(self.envelope(bridge.get_initial_state())["credential_status"], {"claude": False, "openai": False})
        job = self.wait_job(self.start_single(bridge=bridge), bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(self.client_keys, [SECRET])
        self.assertEqual(self.backend.values, {})
        self.assertFalse(any(operation == "set" for operation, _ in self.backend.calls))
        self.assertTrue(any(event["stage"] == "preflight" and "환경변수" in event["message"] for event in job["events"]))

    def test_catalogue_aggregate_recurses_unsupported_conditions(self):
        supported = self.cat.to_dict()["clauses"][0]["condition"]
        unsupported = dict(type="unsupported", predicate_id="unknown", subject="분류 경계",
                           text="합성 미해석 조건", source_anchor="advanced.txt#1", reason="시험 미해석")
        mixed = catalogue(dict(type="all", children=[supported, dict(type="not", child=unsupported)]))
        bridge = self.make_bridge(catalogue_loader=lambda: mixed)
        families = {entry["family"]: entry for entry in self.envelope(bridge.get_initial_state())["catalogue"]["families"]}
        self.assertEqual((families["advanced"]["total_clauses"], families["advanced"]["interpreted_clauses"],
                          families["advanced"]["unsupported_clauses"]), (1, 0, 1))
        self.assertEqual((families["core"]["interpreted_clauses"], families["core"]["unsupported_clauses"]), (1, 0))

    def test_path_selection_cancellation_and_purpose_validation(self):
        first = self.token("input_excel")
        self.selected["input_excel"] = None
        data = self.envelope(self.bridge.select_path("input_excel"))
        self.assertEqual(data, {"cancelled": True, "path_token": None, "display_name": None})
        for purpose in (None, 1, "input", "../../escape", {}):
            with self.subTest(purpose=purpose):
                self.envelope(self.bridge.select_path(purpose), ok=False, code="validation")
        output = self.token("output_excel")
        directory = self.token("output_directory")
        for invalid in (first, output, str(self.directory), "forged-token", 5):
            with self.subTest(invalid=invalid):
                self.envelope(self.bridge.start_single(self.single_request(output_dir_token=invalid)), ok=False, code="validation")
        self.assertNotEqual(first, output)
        self.assertNotEqual(output, directory)

    def test_dialog_errors_are_safe_and_tokens_are_bridge_local(self):
        token = self.token("output_directory")
        other = self.make_bridge()
        self.envelope(other.start_single(self.single_request(output_dir_token=token)), ok=False, code="validation")

        def bad_selector(purpose):
            raise OSError("dialog diagnostic " + SECRET)

        bridge = self.make_bridge(path_selector=bad_selector)
        self.envelope(bridge.select_path("input_excel"), ok=False, code="io")

    def test_requests_reject_types_unknown_keys_and_raw_paths_before_model(self):
        for request in (None, [], "request", {}, self.single_request(task={}), self.single_request(task={"과제명": 7}),
                        self.single_request(task=TASK | {"arbitrary": "injected"}), self.single_request(options=[]),
                        self.single_request(options={"api_key": SECRET}), self.single_request(options={"num_reviewers": True}),
                        self.single_request() | {"output_dir": str(self.directory)}):
            with self.subTest(request=type(request).__name__):
                self.envelope(self.bridge.start_single(request), ok=False, code="validation")
        for request in ({}, self.batch_request(workers=True), self.batch_request(workers=9), self.batch_request(resume=1),
                        self.batch_request(make_forms="false"), self.batch_request(input_token=str(self.input)),
                        self.batch_request(output_token=self.token("input_excel")), self.batch_request() | {"input_path": str(self.input)}):
            with self.subTest(request=request):
                self.envelope(self.bridge.start_batch(request), ok=False, code="validation")
        for job_id, sequence in ((None, 0), (1, 0), ({}, 0), ("unknown", -1), ("unknown", True), ("unknown", .5)):
            with self.subTest(job_id=job_id, sequence=sequence):
                self.envelope(self.bridge.get_job(job_id, sequence), ok=False, code="validation")
        self.envelope(self.bridge.get_job("unknown", 0), ok=False, code="not_found")
        self.envelope(self.bridge.cancel_job("unknown"), ok=False, code="not_found")
        self.assertFalse(self.clients)

    def test_single_result_events_cursor_and_read_copy(self):
        request = self.single_request()
        job_id = self.start_single(request)
        request["task"]["과제명"] = "renderer changed after start"
        job = self.wait_job(job_id)
        self.assertEqual(job["state"], "done")
        result = job["result"]
        self.assertEqual(result["schema_version"], 3)
        self.assertEqual(result["task_name"], TASK["과제명"])
        self.assertEqual(result["documents"]["status"], "미생성")
        self.assertEqual(self.client_keys, [SECRET])
        self.assertTrue(job["events"])
        sequences = [event["sequence"] for event in job["events"]]
        self.assertEqual(sequences, sorted(set(sequences)))
        self.assertGreater(sequences[0], 0)
        for event in job["events"]:
            self.assertEqual(set(event), {"sequence", "stage", "message", "completed", "total"})
            self.assertIn(event["stage"], STAGES)
            self.assertIsInstance(event["message"], str)
        self.assertEqual(self.envelope(self.bridge.get_job(job_id, sequences[-1]))["events"], [])
        middle = sequences[len(sequences) // 2]
        self.assertTrue(all(event["sequence"] > middle for event in self.envelope(self.bridge.get_job(job_id, middle))["events"]))
        result["task_name"] = "mutated DTO"
        job["events"][0]["message"] = "mutated event"
        retained = self.envelope(self.bridge.get_job(job_id, 0))
        self.assertEqual(retained["result"]["task_name"], TASK["과제명"])
        self.assertNotEqual(retained["events"][0]["message"], "mutated event")
        self.assertEqual(self.envelope(self.bridge.cancel_job(job_id))["state"], "done")

    def test_busy_cancel_retains_inflight_result_then_accepts_next_job(self):
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)

        def classify(task, options, client, progress_callback=None, *, catalogue=None):
            entered.set()
            if not release.wait(4):
                raise AssertionError("test did not release synthetic classifier")
            return engine.classify_task(task, options, client, progress_callback=progress_callback, catalogue=catalogue)

        bridge = self.make_bridge(classify=classify)
        batch_request = self.batch_request(bridge)
        job_id = self.start_single(bridge=bridge)
        self.assertTrue(entered.wait(2))
        self.envelope(bridge.start_single(self.single_request()), ok=False, code="busy")
        self.envelope(bridge.start_batch(batch_request), ok=False, code="busy")
        self.assertEqual(self.envelope(bridge.cancel_job(job_id))["state"], "cancelling")
        release.set()
        job = self.wait_job(job_id, bridge)
        self.assertEqual(job["state"], "cancelled")
        self.assertEqual(job["result"]["task_name"], TASK["과제명"])
        next_id = self.start_single(bridge=bridge)
        self.assertNotEqual(next_id, job_id)
        self.assertEqual(self.wait_job(next_id, bridge)["state"], "done")
        self.assertEqual(self.envelope(bridge.get_job(job_id, 0))["state"], "cancelled")

    def test_late_callback_cannot_change_terminal_or_next_job_events(self):
        callbacks = []
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)

        def classify(task, options, client, progress_callback=None, *, catalogue=None):
            callbacks.append(progress_callback)
            if len(callbacks) == 2:
                entered.set()
                if not release.wait(4):
                    raise AssertionError("test did not release next synthetic classifier")
            return engine.classify_task(task, options, client, progress_callback=progress_callback, catalogue=catalogue)

        bridge = self.make_bridge(classify=classify)
        first_id = self.start_single(bridge=bridge)
        first = self.wait_job(first_id, bridge)
        self.assertEqual(first["state"], "done")
        second_id = self.start_single(bridge=bridge)
        self.assertTrue(entered.wait(2))
        second = self.envelope(bridge.get_job(second_id, 0))
        self.assertEqual(second["state"], "running")
        callbacks[0](dict(stage="candidate", message="late callback " + SECRET, completed=77, total=99))
        self.assertEqual(self.envelope(bridge.get_job(first_id, 0))["events"], first["events"])
        self.assertEqual(self.envelope(bridge.get_job(second_id, 0))["events"], second["events"])
        release.set()
        self.assertEqual(self.wait_job(second_id, bridge)["state"], "done")

    def test_catalogue_and_worker_exceptions_are_safe(self):
        def bad_catalogue():
            raise CatalogueError("catalogue diagnostic " + SECRET)

        bridge = self.make_bridge(catalogue_loader=bad_catalogue)
        self.envelope(bridge.get_initial_state(), ok=False, code="catalogue")
        self.envelope(bridge.start_single(self.single_request()), ok=False, code="catalogue")

        def bad_classify(*args, **kwargs):
            raise RuntimeError("model diagnostic " + SECRET)

        bridge = self.make_bridge(classify=bad_classify)
        job = self.wait_job(self.start_single(bridge=bridge), bridge)
        self.assertEqual(job["state"], "failed")
        self.assertIsNone(job["result"])
        self.assertTrue(job["events"])
        self.assertNotIn(SECRET, json.dumps(job, ensure_ascii=False))

    def test_export_uses_retained_dto_with_zero_model_or_catalogue_calls(self):
        regenerated_calls = []

        def regenerate(saved, output_dir):
            regenerated_calls.append(deepcopy(saved))
            return regenerated(saved, output_dir)

        bridge = self.make_bridge(regenerator=regenerate)
        source_id = self.start_single(bridge=bridge)
        original = self.wait_job(source_id, bridge)["result"]
        calls = len(self.clients), len(self.catalogue_calls)
        token = self.token("output_directory", bridge)
        request = dict(job_id=source_id, task_fingerprint=original["task_fingerprint"], output_dir_token=token)
        for invalid in ({}, request | {"task_fingerprint": "forged"}, request | {"result": original},
                        request | {"output_dir_token": str(self.directory)}, request | {"job_id": "missing"}):
            self.envelope(bridge.start_export(invalid), ok=False)
        export_id = self.envelope(bridge.start_export(request))["job_id"]
        self.assertNotEqual(export_id, source_id)
        exported = self.wait_job(export_id, bridge)["result"]
        self.assertEqual(exported["schema_version"], 3)
        self.assertEqual(exported["documents"]["status"], "완료")
        self.assertEqual({k: v for k, v in exported.items() if k != "documents"},
                         {k: v for k, v in original.items() if k != "documents"})
        second_id = self.envelope(bridge.start_export(request | {"job_id": export_id}))["job_id"]
        self.assertEqual(self.wait_job(second_id, bridge)["result"]["task_fingerprint"], original["task_fingerprint"])
        self.assertEqual((len(self.clients), len(self.catalogue_calls)), calls)
        self.assertEqual(len(regenerated_calls), 2)

    def test_export_refuses_tampered_retained_python_result(self):
        calls = []

        def regenerate(saved, output_dir):
            calls.append(deepcopy(saved))
            return regenerated(saved, output_dir)

        bridge = self.make_bridge(regenerator=regenerate)
        token = self.token("output_directory", bridge)
        source_id = self.start_single(bridge=bridge)
        original = self.wait_job(source_id, bridge)["result"]
        bridge._jobs[source_id].result["task_name"] = "tampered Python retained result"
        request = dict(job_id=source_id, task_fingerprint=original["task_fingerprint"], output_dir_token=token)
        self.envelope(bridge.start_export(request), ok=False, code="validation")
        self.assertEqual(calls, [])

    def test_document_failure_preserves_classification_and_reports_document_event(self):
        def fail_documents(saved, output_dir):
            saved = deepcopy(saved)
            saved["documents"] = dict(status="실패", expected_files=len(saved["reviewers"]) + 1,
                                      files=[], errors=["합성 의견서 저장 실패"])
            return saved

        bridge = self.make_bridge(regenerator=fail_documents)
        token = self.token("output_directory", bridge)
        job = self.wait_job(self.start_single(self.single_request(output_dir_token=token), bridge), bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(job["result"]["status"], "완료")
        self.assertEqual(job["result"]["documents"]["status"], "실패")
        self.assertTrue(any(event["stage"] == "documents" and "실패" in event["message"] for event in job["events"]))

    def batch_report(self, results):
        counts = {status: sum(result["status"] == status for result in results) for status in domain.STATUSES}
        return dict(status="완료", results=deepcopy(results), counts=counts, warnings=[],
                    input_counts={"blank": 0, "example": 0, "duplicate": 0}, task_count=len(results), reused_count=0,
                    output_path=str(self.output), input_path=str(self.input),
                    checkpoint_path=str(self.output) + ".checkpoint.jsonl", audit_path=str(self.output) + ".audit.jsonl",
                    documents_dir=str(self.directory), audit_status="완료", input_fingerprints=[r["task_fingerprint"] for r in results],
                    options_fingerprint=results[0]["options_fingerprint"], selection_mode="auto", catalogue_fingerprint=self.cat.fingerprint)

    def test_batch_runner_receives_frozen_context_and_renderer_report_shape(self):
        seen = []
        report = self.batch_report([self.result()])

        def runner(in_path, out_path, options, client_factory, **kwargs):
            seen.append((in_path, out_path, options, kwargs))
            client_factory()
            kwargs["status_cb"]("합성 저장 상태")
            return deepcopy(report)

        bridge = self.make_bridge(batch_runner=runner)
        request = self.batch_request(bridge, workers=3, resume=True, make_forms=True)
        job_id = self.envelope(bridge.start_batch(request))["job_id"]
        request["options"]["num_reviewers"] = 3
        job = self.wait_job(job_id, bridge)
        self.assertEqual(job["state"], "done")
        wire = job["result"]
        for field in ("results", "counts", "paths", "warnings", "persistence"):
            self.assertIn(field, wire)
        self.assertEqual(wire["results"][0]["schema_version"], 3)
        self.assertEqual(wire["paths"]["input_excel"], str(self.input))
        self.assertEqual(wire["paths"]["output_excel"], str(self.output))
        source, output, options, kwargs = seen[0]
        self.assertEqual((Path(source), Path(output)), (self.input, self.output))
        self.assertEqual(options.num_reviewers, 1)
        self.assertEqual((kwargs["workers"], kwargs["resume"], kwargs["make_forms"]), (3, True, True))
        self.assertIs(kwargs["catalogue"], self.cat)
        self.assertIsInstance(kwargs["cancel_event"], threading.Event)
        self.assertNotIn(SECRET, repr(options))

    def test_batch_cancel_passes_event_and_retains_completed_rows(self):
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)
        result = self.result()
        report = self.batch_report([result])
        report["task_count"] = 3
        seen = []

        def runner(in_path, out_path, options, client_factory, **kwargs):
            seen.append(kwargs["cancel_event"])
            entered.set()
            if not release.wait(4):
                raise AssertionError("test did not release synthetic batch runner")
            self.assertTrue(kwargs["cancel_event"].is_set())
            return deepcopy(report)

        bridge = self.make_bridge(batch_runner=runner)
        job_id = self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"]
        self.assertTrue(entered.wait(2))
        self.assertEqual(self.envelope(bridge.cancel_job(job_id))["state"], "cancelling")
        release.set()
        job = self.wait_job(job_id, bridge)
        self.assertEqual(job["state"], "cancelled")
        self.assertTrue(seen[0].is_set())
        self.assertEqual(job["result"]["results"], [result])
        self.assertEqual(job["result"]["counts"]["cancelled"], 2)

    def test_batch_export_chains_share_latest_durable_report(self):
        first, second = self.result(), self.result(TASK | {"과제명": "합성 두번째"})
        base_report = self.batch_report([first, second])
        seen = []

        def persist(report, saved, replacement):
            seen.append(deepcopy(report))
            updated = deepcopy(report)
            updated["results"] = [deepcopy(replacement) if r["task_fingerprint"] == saved["task_fingerprint"] else r
                                  for r in updated["results"]]
            return dict(result=deepcopy(replacement), report=updated,
                        persistence=dict(status="완료", checkpoint="완료", workbook="완료", audit="완료", warnings=[]))

        bridge = self.make_bridge(batch_runner=lambda *args, **kwargs: deepcopy(base_report),
                                  regenerator=regenerated, persister=persist)
        batch_id = self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"]
        self.wait_job(batch_id, bridge)
        calls = len(self.clients), len(self.catalogue_calls)
        output_tokens = []
        for suffix in ("first", "second", "third"):
            directory = self.root / suffix
            directory.mkdir()
            self.selected["output_directory"] = directory
            output_tokens.append(self.token("output_directory", bridge))

        def export(source_id, result, token):
            job_id = self.envelope(bridge.start_export(dict(job_id=source_id,
                task_fingerprint=result["task_fingerprint"], output_dir_token=token)))["job_id"]
            return job_id, self.wait_job(job_id, bridge)["result"]

        first_export, first_updated = export(batch_id, first, output_tokens[0])
        _, second_updated = export(batch_id, second, output_tokens[1])
        _, third_updated = export(first_export, first_updated, output_tokens[2])
        self.assertEqual(len(seen), 3)
        self.assertEqual(seen[1]["results"][0]["documents"], first_updated["documents"])
        self.assertEqual(seen[2]["results"][0]["documents"], first_updated["documents"])
        self.assertEqual(seen[2]["results"][1]["documents"], second_updated["documents"])
        self.assertNotEqual(third_updated["documents"], first_updated["documents"])
        self.assertEqual((len(self.clients), len(self.catalogue_calls)), calls)

    def test_failed_checkpoint_export_does_not_advance_shared_report(self):
        original = self.result()
        base_report = self.batch_report([original])
        seen = []

        def persist(report, saved, replacement):
            seen.append(deepcopy(report))
            updated = deepcopy(report)
            updated["results"] = [deepcopy(replacement)]
            return dict(result=deepcopy(replacement), report=updated,
                        persistence=dict(status="실패", checkpoint="실패", workbook="미시도", audit="미시도",
                                         warnings=["합성 체크포인트 실패"]))

        bridge = self.make_bridge(batch_runner=lambda *args, **kwargs: deepcopy(base_report),
                                  regenerator=regenerated, persister=persist)
        batch_id = self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"]
        self.wait_job(batch_id, bridge)
        token = self.token("output_directory", bridge)
        source_id = batch_id
        for _ in range(2):
            source_id = self.envelope(bridge.start_export(dict(job_id=source_id,
                task_fingerprint=original["task_fingerprint"], output_dir_token=token)))["job_id"]
            job = self.wait_job(source_id, bridge)
            self.assertEqual(job["result"]["documents"]["status"], "완료")
            self.assertTrue(any(event["stage"] == "persistence" for event in job["events"]))
        self.assertEqual(len(seen), 2)
        self.assertEqual(seen[1]["results"][0]["documents"], original["documents"])

    # --- v16: initial state notices, title table, -body rule, guideline -------

    def test_initial_state_notices_titles_and_body_source_url(self):
        self.cat = distinct_url_catalogue()
        catalogue_data = self.envelope(self.make_bridge().get_initial_state())["catalogue"]
        self.assertEqual(catalogue_data["notice_titles"], {number: title for number, (title, *_) in NOTICES.items()})
        notices = catalogue_data["notices"]
        self.assertEqual([notice["notice_number"] for notice in notices], list(NOTICES))
        expected = {
            "산업통상자원부고시 제2025-71호": (["advanced"], ["advanced-body"], "advanced-body"),
            "산업통상부고시 제2025-2호": (["core"], ["core-body"], "core-body"),
            "산업통상부고시 제2026-101호": (["dual_use", "military"], ["strategic-body", "military-annex"], "strategic-body"),
        }
        for notice in notices:
            number = notice["notice_number"]
            with self.subTest(notice=number):
                self.assertEqual(list(notice), ["notice_number", "title", "announcement_date", "effective_date",
                                                "families", "source_ids", "official_url"])
                families, source_ids, body = expected[number]
                self.assertEqual((notice["title"], notice["announcement_date"], notice["effective_date"]), NOTICES[number])
                self.assertEqual(notice["families"], families)
                self.assertEqual(notice["source_ids"], source_ids)
                self.assertEqual(notice["official_url"], "https://example.gov.test/" + body)

    def test_body_source_rule_and_unknown_notice_are_catalogue_errors(self):
        def renamed(old, new):
            plain = distinct_url_catalogue().to_dict()
            for source in plain["manifest"]["sources"]:
                if source["source_id"] == old:
                    source["source_id"] = new
            for clause in plain["clauses"]:
                if clause["source_id"] == old:
                    clause["source_id"] = new
            return repack(plain)

        def unknown_notice():
            plain = distinct_url_catalogue().to_dict()
            for source in plain["manifest"]["sources"]:
                if source["notice_number"] == "산업통상부고시 제2025-2호":
                    source["notice_number"] = "산업통상부고시 제2099-1호"
            return repack(plain)

        broken = {"two bodies": renamed("military-annex", "military-body"),
                  "no body": renamed("core-body", "core-annex"), "unknown notice": unknown_notice()}
        for unvalidated in (False, True):
            for name, cat in broken.items():
                stand_in = patch.object(bridge_module, "_preflight", unvalidated_preflight) if unvalidated else nullcontext()
                with self.subTest(case=name, bridge_rule_only=unvalidated), stand_in:
                    bridge = self.make_bridge(catalogue_loader=lambda cat=cat: cat)
                    self.envelope(bridge.get_initial_state(), ok=False, code="catalogue")
                    self.envelope(bridge.search_notice_clauses({}), ok=False, code="catalogue")
                    self.envelope(bridge.get_notice_clause("advanced:1"), ok=False, code="catalogue")
        # The unvalidated stand-in itself accepts a sound catalogue.
        with patch.object(bridge_module, "_preflight", unvalidated_preflight):
            bridge = self.make_bridge(catalogue_loader=distinct_url_catalogue)
            self.assertEqual(len(self.envelope(bridge.get_initial_state())["catalogue"]["notices"]), 3)

    def test_guideline_is_display_shape_not_prompt_constants(self):
        guideline = self.envelope(self.bridge.get_initial_state())["guideline"]
        self.assertEqual(list(guideline), ["id", "definitions", "principles"])
        self.assertEqual(guideline, domain.guideline_display())
        self.assertEqual(guideline["id"], domain.GUIDELINE_ID)
        self.assertEqual(guideline["definitions"], APPENDIX_A_DEFINITIONS)
        self.assertEqual(guideline["principles"], APPENDIX_A_PRINCIPLES)
        self.assertEqual([entry["heading"] for entry in guideline["definitions"]],
                         ["법 제21조제4항(보안과제)", "법 제21조제5항(민감과제)", "시행령 제45조제1항(보안과제 분류 기준)",
                          "시행령 제45조제2항(민감과제 분류 기준)", "일반과제(가이드라인)"])
        self.assertEqual([entry["title"] for entry in guideline["principles"]],
                         ["국가 차원의 보호 필요성에 비례한 종합적 판단", "협력적･전문적 의사결정",
                          "예측 가능성 및 절차적 안정성 원칙", "과제 단위 분류"])
        for entry in guideline["definitions"]:
            self.assertEqual(list(entry), ["heading", "text"])
        for entry in guideline["principles"]:
            self.assertEqual(list(entry), ["title", "details"])
            self.assertTrue(entry["details"] and all(type(detail) is str for detail in entry["details"]))
        self.assertEqual(guideline["definitions"][4]["text"],
                         "보안과제 또는 민감과제에 해당하지 않으며, 통상적인 연구관리 및 기본 보안관리 조치 하에서 수행 가능한 연구개발과제")
        self.assertEqual(guideline["definitions"][3]["text"].split("\n")[1:],
                         ["1. 국가안보･외교 등을 위한 전략적 필요성", "2. 기술적･재산적 가치",
                          "3. 그 밖에 연구개발과제의 수행 과정 및 결과의 유출 시 국가 전체에 미치는 영향의 정도"])
        wire = json.dumps(guideline, ensure_ascii=False)
        self.assertNotIn("대체값으로 사용하지 않는다", wire)
        for text in list(domain.LEGAL_DEFINITIONS.values()) + list(domain.CLASSIFICATION_PRINCIPLES):
            self.assertNotIn(text, wire)

    # --- v16: search_notice_clauses / get_notice_clause --------------------

    def test_search_defaults_items_without_text_and_screen_request_shape(self):
        data = self.envelope(self.bridge.search_notice_clauses({}))
        self.assertEqual(list(data), ["total", "offset", "limit", "items"])
        self.assertEqual((data["total"], data["offset"], data["limit"]), (4, 0, 50))
        self.assertEqual([item["clause_id"] for item in data["items"]], [c["clause_id"] for c in self.cat.clauses])
        for item in data["items"]:
            self.assertEqual(list(item), ["clause_id", "family", "source_id", "notice_number", "code", "title", "field"])
            self.assertNotIn("text", item)
        self.assertEqual(data["items"][2], dict(clause_id="dual_use:1", family="dual_use", source_id="strategic-body",
                         notice_number="산업통상부고시 제2026-101호", code="SYN", title="합성 속도", field="합성"))
        screen = self.envelope(self.bridge.search_notice_clauses(dict(query="", family=None, offset=0, limit=50)))
        self.assertEqual(screen, data)
        self.assertEqual(self.envelope(self.bridge.search_notice_clauses(dict(query="SYN", family="military"))),
                         dict(total=1, offset=0, limit=50, items=[data["items"][3]]))
        # Mutating a returned item never changes the retained index.
        data["items"][0]["title"] = "renderer changed"
        self.assertEqual(self.envelope(self.bridge.search_notice_clauses({}))["items"][0]["title"], "합성 속도")

    def test_search_rejects_every_invalid_request(self):
        invalid = [None, [], "속도", 5, True, ("query",), {"unknown": 1}, {"query": "", "extra": None},
                   {"query": None}, {"query": 5}, {"query": ["속도"]}, {"query": b"speed"},
                   {"query": "가" * 201}, {"query": " " * 201},
                   {"family": "Advanced"}, {"family": "dual-use"}, {"family": ""}, {"family": 1}, {"family": True},
                   {"family": ["core"]},
                   {"offset": -1}, {"offset": True}, {"offset": False}, {"offset": 1.0}, {"offset": "0"}, {"offset": None},
                   {"limit": 0}, {"limit": 101}, {"limit": -5}, {"limit": True}, {"limit": False}, {"limit": 50.0},
                   {"limit": "50"}, {"limit": None}]
        for request in invalid:
            with self.subTest(request=request):
                self.envelope(self.bridge.search_notice_clauses(request), ok=False, code="validation")
        for request, total in (({"query": "가" * 200}, 0), ({"query": " " * 200}, 4), ({"limit": 1}, 4),
                               ({"limit": 100}, 4), ({"family": None}, 4), ({"offset": 10 ** 12}, 4)):
            with self.subTest(valid=request):
                self.assertEqual(self.envelope(self.bridge.search_notice_clauses(request))["total"], total)

    def test_get_notice_clause_record_validation_and_not_found(self):
        self.cat = distinct_url_catalogue()
        bridge = self.make_bridge()
        data = self.envelope(bridge.get_notice_clause("military:1"))
        self.assertEqual(set(data), CLAUSE_KEYS)
        clause = next(c for c in self.cat.clauses if c["clause_id"] == "military:1")
        for key in ("clause_id", "family", "source_id", "code", "title", "field", "source_anchor", "text"):
            self.assertEqual(data[key], clause[key])
        self.assertEqual((data["notice_number"], data["notice_title"], data["announcement_date"], data["effective_date"]),
                         ("산업통상부고시 제2026-101호", "전략물자수출입고시", "2026-08-31", "2026-09-01"))
        # The clause's own source URL, not the notice's body-source URL.
        self.assertEqual(data["official_url"], "https://example.gov.test/military-annex")
        for invalid in (None, 5, True, ["military:1"], {"clause_id": "military:1"}, b"military:1"):
            with self.subTest(invalid=invalid):
                self.envelope(bridge.get_notice_clause(invalid), ok=False, code="validation")
        for missing in ("missing", "", "MILITARY:1", " military:1"):
            with self.subTest(missing=missing):
                self.envelope(bridge.get_notice_clause(missing), ok=False, code="not_found")

    def test_lookups_are_read_only_and_safe_on_catalogue_failure(self):
        bridge = self.make_bridge()
        self.envelope(bridge.search_notice_clauses({}))
        before = (list(self.backend.calls), len(self.catalogue_calls), dict(bridge._jobs), dict(bridge._paths),
                  bridge._active, list(self.clients), list(self.client_keys), sorted(self.root.rglob("*")))
        for _ in range(3):
            self.envelope(bridge.search_notice_clauses(dict(query="속도", family="core", offset=0, limit=10)))
            self.envelope(bridge.get_notice_clause("core:1"))
            self.envelope(bridge.get_notice_clause("missing"), ok=False, code="not_found")
            self.envelope(bridge.search_notice_clauses({"limit": 0}), ok=False, code="validation")
        after = (list(self.backend.calls), len(self.catalogue_calls), dict(bridge._jobs), dict(bridge._paths),
                 bridge._active, list(self.clients), list(self.client_keys), sorted(self.root.rglob("*")))
        self.assertEqual(after, before)
        self.assertEqual(self.backend.calls, [])
        self.assertEqual(len(self.catalogue_calls), 1)
        for error in (CatalogueError("catalogue diagnostic " + SECRET), RuntimeError("loader diagnostic " + SECRET),
                      OSError("disk diagnostic " + SECRET)):
            def failing(error=error):
                raise error
            with self.subTest(error=type(error).__name__):
                broken = self.make_bridge(catalogue_loader=failing)
                self.envelope(broken.search_notice_clauses({}), ok=False, code="catalogue")
                self.envelope(broken.get_notice_clause("advanced:1"), ok=False, code="catalogue")
        self.assertEqual(self.backend.calls, [])

    def test_lookups_work_while_a_job_is_running_without_touching_it(self):
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)

        def classify(task, options, client, progress_callback=None, *, catalogue=None):
            entered.set()
            if not release.wait(4):
                raise AssertionError("test did not release synthetic classifier")
            return engine.classify_task(task, options, client, progress_callback=progress_callback, catalogue=catalogue)

        bridge = self.make_bridge(classify=classify)
        job_id = self.start_single(bridge=bridge)
        self.assertTrue(entered.wait(2))
        running = self.envelope(bridge.get_job(job_id, 0))
        self.assertEqual(running["state"], "running")
        self.envelope(bridge.start_single(self.single_request()), ok=False, code="busy")
        found = self.envelope(bridge.search_notice_clauses(dict(query="합성 속도", family=None, offset=1, limit=2)))
        self.assertEqual((found["total"], [item["clause_id"] for item in found["items"]]), (4, ["core:1", "dual_use:1"]))
        self.assertEqual(self.envelope(bridge.get_notice_clause("dual_use:1"))["text"], "속도 10 GHz 이상")
        self.envelope(bridge.get_notice_clause("missing"), ok=False, code="not_found")
        self.envelope(bridge.search_notice_clauses({"family": "space"}), ok=False, code="validation")
        self.assertEqual(self.envelope(bridge.get_job(job_id, 0)), running)
        self.assertEqual(bridge._active, job_id)
        release.set()
        job = self.wait_job(job_id, bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(job["result"]["task_name"], TASK["과제명"])

    # --- v16: protected previous result (spec 4.5) -------------------------

    def assert_previous_result_stop(self, job):
        self.assertEqual(job["state"], "failed")
        self.assertIsNone(job["result"])
        messages = [event["message"] for event in job["events"]]
        self.assertEqual(job["events"][-1]["stage"], "preflight")
        self.assertEqual(messages[-1], batch_v16.PREVIOUS_RESULT_MESSAGE)
        self.assertEqual(messages[-1], "기존 결과 파일 보호를 위해 중단. 새 결과 엑셀 경로 선택 필요")
        self.assertNotIn(bridge_module._MESSAGES["io"], messages)
        self.assertNotIn(SECRET, json.dumps(job, ensure_ascii=False))

    def test_previous_result_error_keeps_io_and_emits_dedicated_event(self):
        def runner(*args, **kwargs):
            raise batch_v16.PreviousResultError()

        bridge = self.make_bridge(batch_runner=runner)
        job_id = self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"]
        self.assert_previous_result_stop(self.wait_job(job_id, bridge))
        self.assertIsNone(bridge._active)
        self.assertTrue(issubclass(batch_v16.PreviousResultError, batch_v16.BatchError))

        # Any other BatchError still maps to the generic io message.
        def other(*args, **kwargs):
            raise batch_v16.BatchError("batch diagnostic " + SECRET)

        bridge = self.make_bridge(batch_runner=other)
        job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"], bridge)
        self.assertEqual(job["state"], "failed")
        self.assertEqual(job["events"][-1]["message"], bridge_module._MESSAGES["io"])
        self.assertNotIn(batch_v16.PREVIOUS_RESULT_MESSAGE, [event["message"] for event in job["events"]])
        self.assertNotIn(SECRET, json.dumps(job, ensure_ascii=False))

    def test_previous_result_stop_ends_failed_even_after_cancel_request(self):
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)

        def runner(*args, **kwargs):
            entered.set()
            if not release.wait(4):
                raise AssertionError("test did not release synthetic batch runner")
            raise batch_v16.PreviousResultError()

        bridge = self.make_bridge(batch_runner=runner)
        job_id = self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"]
        self.assertTrue(entered.wait(2))
        self.assertEqual(self.envelope(bridge.cancel_job(job_id))["state"], "cancelling")
        release.set()
        self.assert_previous_result_stop(self.wait_job(job_id, bridge))

    def test_real_batch_runner_protects_non_v16_result_through_bridge(self):
        book = Workbook()
        book.active.title = "분류대상"
        book.active.append(("과제명", "현재 기술사양"))
        book.active.append(("합성 보호 과제", "속도 20 GHz"))
        book.save(self.input)
        book.close()
        previous = Workbook()
        previous.active.title = "분류결과"
        previous.active.append(("연번", "과제명"))
        previous.save(self.output)
        previous.close()
        checkpoint = Path(str(self.output) + ".checkpoint.jsonl")
        audit = Path(str(self.output) + ".audit.jsonl")
        output_bytes, listing = self.output.read_bytes(), sorted(self.root.rglob("*"))
        for resume in (True, False):
            with self.subTest(resume=resume):
                bridge = self.make_bridge()
                request = self.batch_request(bridge, resume=resume, make_forms=True)
                job = self.wait_job(self.envelope(bridge.start_batch(request))["job_id"], bridge)
                self.assert_previous_result_stop(job)
                self.assertEqual(self.output.read_bytes(), output_bytes)
                self.assertFalse(checkpoint.exists() or audit.exists())
                self.assertEqual(sorted(self.root.rglob("*")), listing)
                self.assertEqual(self.clients, [])

    # --- v16: input template download (spec 4.4, 4.5) -----------------------

    def assert_template_workbook(self, path):
        """Two sheets and nine columns; header and guide wording belong to batch_v16."""
        book = load_workbook(path)
        try:
            self.assertEqual(book.sheetnames, ["작성안내", "분류대상"])
            self.assertTrue(any(cell.value for row in book["작성안내"] for cell in row))
            sheet = book["분류대상"]
            self.assertEqual(sheet.max_column, 9)
            header = [cell.value for cell in sheet[1]]
            self.assertEqual(len(header), 9)
            self.assertTrue(all(type(value) is str and value.strip() for value in header))
        finally:
            book.close()

    def template_token(self, path, bridge=None):
        self.selected["input_template"] = path
        return self.token("input_template", bridge)

    def save_template(self, token, bridge=None):
        return (bridge or self.bridge).save_input_template({"output_token": token})

    def files(self):
        return {path: path.read_bytes() for path in self.root.rglob("*") if path.is_file()}

    def assert_protected_refusal(self, bridge, path):
        before, listing = self.files(), sorted(self.root.rglob("*"))
        token = self.template_token(path, bridge)
        response = self.save_template(token, bridge)
        self.envelope(response, ok=False, protected=True)
        self.assertNotIn(str(self.root), json.dumps(response, ensure_ascii=False))
        # The existing file (and every sibling) is left byte-identical.
        self.assertEqual(self.files(), before)
        self.assertEqual(sorted(self.root.rglob("*")), listing)

    def test_input_template_saves_two_sheet_nine_column_workbook(self):
        token = self.token("input_template")
        response = self.save_template(token)
        data = self.envelope(response)
        self.assertEqual(data, {"saved": True, "display_name": self.template.name})
        self.assertNotIn(str(self.root), json.dumps(response, ensure_ascii=False))
        self.assert_template_workbook(self.template)
        # The saved template is never selected as the input workbook (spec 4.4).
        self.assertEqual([purpose for purpose, _ in self.bridge._paths.values()], ["input_template"])
        # An ordinary existing file is replaced after the native overwrite confirmation.
        ordinary = self.root / "기존 파일.XLSX"
        ordinary.write_bytes(b"synthetic ordinary file")
        data = self.envelope(self.save_template(self.template_token(ordinary)))
        self.assertEqual(data, {"saved": True, "display_name": "기존 파일.XLSX"})
        self.assert_template_workbook(ordinary)
        self.assertEqual(sorted(path.name for path in self.root.iterdir()),
                         sorted(["documents", "input.xlsx", "기존 파일.XLSX", self.template.name]))
        self.assertFalse(self.clients)
        self.assertEqual(self.backend.calls, [])

    def test_input_template_cancel_is_not_an_error_and_writes_nothing(self):
        before = self.files()
        for cancelled in (None, ""):
            with self.subTest(cancelled=cancelled):
                self.selected["input_template"] = cancelled
                data = self.envelope(self.bridge.select_path("input_template"))
                self.assertEqual(data, {"cancelled": True, "path_token": None, "display_name": None})
        self.assertEqual(self.bridge._paths, {})
        self.assertEqual(self.files(), before)
        self.assertFalse(self.template.exists())

    def test_input_template_path_rules_and_purpose_bound_tokens(self):
        (self.root / "folder.xlsx").mkdir()
        for name, path in {"xls": self.root / "양식.xls", "csv": self.root / "양식.csv",
                           "no suffix": self.root / "양식", "xlsx then tmp": self.root / "양식.xlsx.tmp",
                           "directory": self.root / "folder.xlsx", "directory without suffix": self.directory,
                           "missing parent": self.root / "없는 폴더" / "양식.xlsx"}.items():
            with self.subTest(case=name):
                self.selected["input_template"] = path
                self.envelope(self.bridge.select_path("input_template"), ok=False, code="validation")
        self.assertEqual(self.bridge._paths, {})
        self.assertFalse((self.root / "없는 폴더").exists())
        template = self.template_token(self.template)
        others = {purpose: self.token(purpose) for purpose in ("input_excel", "output_excel", "output_directory")}
        for name, request in {**{purpose: {"output_token": token} for purpose, token in others.items()},
                              "forged": {"output_token": "forged-token"}, "raw path": {"output_token": str(self.template)},
                              "none": {"output_token": None}, "number": {"output_token": 5}, "empty": {"output_token": ""},
                              "bare token": template, "empty request": {}, "list": [template],
                              "extra key": {"output_token": template, "path": str(self.template)},
                              "wrong key": {"input_token": template}}.items():
            with self.subTest(request=name):
                self.envelope(self.bridge.save_input_template(request), ok=False, code="validation")
        # Template tokens are bridge-local and usable for no other purpose.
        self.envelope(self.save_template(template, self.make_bridge()), ok=False, code="validation")
        self.envelope(self.bridge.start_single(self.single_request(output_dir_token=template)), ok=False, code="validation")
        self.envelope(self.bridge.start_batch(self.batch_request(input_token=template)), ok=False, code="validation")
        self.envelope(self.bridge.start_batch(self.batch_request(output_token=template)), ok=False, code="validation")
        self.assertFalse(self.template.exists())
        self.assertFalse(self.clients)
        # An unknown purpose never falls through to the output-directory rule.
        with patch.object(bridge_module, "_PURPOSES", bridge_module._PURPOSES | {"mystery"}):
            self.selected["mystery"] = self.directory
            self.envelope(self.bridge.select_path("mystery"), ok=False, code="validation")
        self.assertNotIn("mystery", {purpose for purpose, _ in self.bridge._paths.values()})

    def test_input_template_refuses_paths_selected_as_input_or_output(self):
        self.input.write_bytes(b"synthetic selected input bytes")
        self.output.write_bytes(b"synthetic previous result bytes")
        missing_output = self.root / "아직 없는 결과.xlsx"
        self.token("input_excel")
        self.token("output_excel")
        self.selected["output_excel"] = missing_output
        self.token("output_excel")
        for path in (self.input, self.output, missing_output):
            with self.subTest(path=path.name):
                self.assert_protected_refusal(self.bridge, path)
        self.assertFalse(missing_output.exists())
        if os.name == "nt":
            with self.subTest(path="case variant of the input"):
                self.assert_protected_refusal(self.bridge, self.root / "INPUT.XLSX")
        # The rule is per app session: a bridge that never selected the path for
        # review treats it as an ordinary file (the native dialog confirms overwrite).
        fresh = self.make_bridge()
        self.envelope(self.save_template(self.template_token(self.output, fresh), fresh))
        self.assert_template_workbook(self.output)
        self.assertEqual(self.input.read_bytes(), b"synthetic selected input bytes")

    def test_input_template_refuses_result_workbooks_with_checkpoint_or_audit(self):
        bridge = self.make_bridge()
        for suffix in (".checkpoint.jsonl", ".audit.jsonl"):
            result = self.root / ("이전 결과" + suffix.split(".")[1] + ".xlsx")
            result.write_bytes(b"synthetic result workbook " + suffix.encode())
            Path(str(result) + suffix).write_bytes(b'{"schema_version": 3}\n')
            with self.subTest(sibling=suffix):
                self.assert_protected_refusal(bridge, result)
            # The workbook itself need not exist yet: the sibling alone protects it.
            orphan = self.root / ("결과 없음" + suffix.split(".")[1] + ".xlsx")
            Path(str(orphan) + suffix).write_bytes(b"{}\n")
            with self.subTest(sibling=suffix, workbook="missing"):
                self.assert_protected_refusal(bridge, orphan)
                self.assertFalse(orphan.exists())

    def test_input_template_is_busy_while_a_job_runs(self):
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)

        def classify(task, options, client, progress_callback=None, *, catalogue=None):
            entered.set()
            if not release.wait(4):
                raise AssertionError("test did not release synthetic classifier")
            return engine.classify_task(task, options, client, progress_callback=progress_callback, catalogue=catalogue)

        bridge = self.make_bridge(classify=classify)
        # Path selection is blocked during a job, so the token is taken first.
        token = self.template_token(self.template, bridge)
        job_id = self.start_single(bridge=bridge)
        self.assertTrue(entered.wait(2))
        self.envelope(bridge.select_path("input_template"), ok=False, code="busy")
        self.envelope(self.save_template(token, bridge), ok=False, code="busy")
        self.assertFalse(self.template.exists())
        release.set()
        self.assertEqual(self.wait_job(job_id, bridge)["state"], "done")
        self.envelope(self.save_template(token, bridge))
        self.assert_template_workbook(self.template)

    def test_input_template_write_holds_the_bridge_lock(self):
        bridge = self.make_bridge()
        token = self.template_token(self.template, bridge)
        original, observed = batch_v16.make_template, []

        def probe():
            acquired = bridge._lock.acquire(timeout=.05)
            observed.append(acquired)
            if acquired:
                bridge._lock.release()

        def writer(path):
            thread = threading.Thread(target=probe)
            thread.start()
            thread.join(2)
            observed.append(bridge._active)
            return original(path)

        with patch.object(batch_v16, "make_template", writer):
            self.envelope(self.save_template(token, bridge))
        # No job can start between the busy check and the write.
        self.assertEqual(observed, [False, None])
        self.assert_template_workbook(self.template)

    def test_input_template_write_failures_are_safe_io(self):
        token = self.token("input_template")
        for error in (batch_v16.BatchError("template diagnostic " + SECRET), OSError("disk diagnostic " + SECRET),
                      PermissionError("locked diagnostic " + SECRET)):
            with self.subTest(error=type(error).__name__), patch.object(batch_v16, "make_template", side_effect=error):
                response = self.save_template(token)
                self.envelope(response, ok=False, code="io")
                self.assertNotIn("diagnostic", json.dumps(response, ensure_ascii=False))
        self.assertFalse(self.template.exists())
        if os.name == "nt":
            # A real failure: an open file cannot be replaced on Windows.
            self.template.write_bytes(b"synthetic open template")
            with self.template.open("rb"):
                response = self.save_template(token)
            self.envelope(response, ok=False, code="io")
            self.assertNotIn("저장하지 못했습니다", json.dumps(response, ensure_ascii=False))
            self.assertEqual(self.template.read_bytes(), b"synthetic open template")
            self.assertEqual([path.name for path in self.root.iterdir() if path.name.startswith(".v16-")], [])

    # --- v16: screen-only message style -------------------------------------

    def test_emitted_events_come_from_the_fixed_noun_style_tables(self):
        allowed = set(bridge_module._MESSAGES.values()) | set(bridge_module._EVENTS.values())
        self.backend.values.clear()
        environment = KeyManager(self.backend, environ={"ANTHROPIC_API_KEY": SECRET})
        bridge = self.make_bridge(key_manager=environment, regenerator=regenerated)
        token = self.token("output_directory", bridge)
        job = self.wait_job(self.start_single(self.single_request(output_dir_token=token), bridge), bridge)
        self.assertEqual(job["state"], "done")
        report = self.batch_report([self.result()])
        bridge = self.make_bridge(key_manager=environment, batch_runner=lambda *args, **kwargs: deepcopy(report),
                                  regenerator=regenerated)
        batch_job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"], bridge)
        events = job["events"] + batch_job["events"]
        self.assertTrue(len({event["stage"] for event in events}) >= 4)
        for event in events:
            with self.subTest(message=event["message"]):
                self.assertIn(event["message"], allowed)
                self.assertEqual(style_violations(event["message"]), [])

    # --- v16 보완: screen-only input_cleaned / response_trimmed notices --------

    def notices(self, job, name):
        return [event for event in job["events"] if event["message"] == NOTICE_TEXTS[name]]

    def assert_notice_progress(self, job):
        """Each notice copies the stage, completed and total of the event before it."""
        events = job["events"]
        for index, event in enumerate(events):
            if event["message"] in NOTICE_TEXTS.values():
                self.assertGreater(index, 0)
                previous = events[index - 1]
                self.assertEqual((event["stage"], event["completed"], event["total"]),
                                 (previous["stage"], previous["completed"], previous["total"]))

    def assert_no_notice_text(self, value):
        dumped = json.dumps(value, ensure_ascii=False)
        for text in NOTICE_TEXTS.values():
            self.assertNotIn(text, dumped)

    def test_notice_texts_are_fixed_event_entries(self):
        for name, text in NOTICE_TEXTS.items():
            with self.subTest(name=name):
                self.assertEqual(bridge_module._EVENTS[name], text)
                self.assertEqual(style_violations(text), [])
                self.assertNotIn(text, bridge_module._MESSAGES.values())

    def test_messy_single_input_emits_input_cleaned_once_and_result_stays_clean(self):
        messy = TASK | {"연구목표": "목표 1_x000D_\n목표 2  ", "연구내용": "내용\r\n\n\n\n끝\t",
                        "기대효과": "효과_x000A_정리"}
        job = self.wait_job(self.start_single(self.single_request(task=messy)))
        self.assertEqual(job["state"], "done")
        self.assertEqual(len(self.notices(job, "input_cleaned")), 1)
        self.assertEqual(self.notices(job, "response_trimmed"), [])
        self.assert_notice_progress(job)
        self.assertEqual(job["result"]["task_snapshot"]["연구목표"], "목표 1\n목표 2")
        self.assert_no_notice_text(job["result"])
        self.assert_no_notice_text(self.bridge._jobs[job["job_id"]].result)

    def test_clean_or_only_space_padded_single_input_emits_no_notice(self):
        for task in (TASK, TASK | {"과제명": "  합성 브리지 과제  ", "연구내용": "\n내용\n"}):
            with self.subTest(task=task):
                job = self.wait_job(self.start_single(self.single_request(task=task)))
                self.assertEqual(job["state"], "done")
                self.assertEqual(self.notices(job, "input_cleaned"), [])
                self.assertEqual(self.notices(job, "response_trimmed"), [])
                self.assertTrue(all(event["message"] in bridge_module._EVENTS.values() for event in job["events"]))

    def test_only_messy_common_policy_emits_one_input_cleaned(self):
        cases = ((" 합성 기준  ", 0), ("합성 기준_x000D_\n세부  \n\n\n\n끝", 1), (None, 0))
        for policy, expected in cases:
            with self.subTest(policy=policy):
                request = self.single_request(options={"num_reviewers": 1, "ministry_policy": policy})
                job = self.wait_job(self.start_single(request))
                self.assertEqual(job["state"], "done")
                self.assertEqual(len(self.notices(job, "input_cleaned")), expected)
                self.assert_no_notice_text(job["result"])
        # Messy task text and messy policy together still notify once.
        request = self.single_request(task=TASK | {"연구목표": "목표_x000D_"},
                                      options={"num_reviewers": 1, "ministry_policy": "기준  \n끝"})
        job = self.wait_job(self.start_single(request))
        self.assertEqual(len(self.notices(job, "input_cleaned")), 1)

    def test_batch_policy_cleaned_and_runner_notices_emit_once_across_threads(self):
        report = self.batch_report([self.result()])
        seen = []

        def runner(in_path, out_path, options, client_factory, **kwargs):
            notice = kwargs["notice_cb"]
            seen.append(notice)
            kwargs["progress_cb"](dict(stage="criteria", completed=2, total=5))
            barrier = threading.Barrier(8)

            def worker():
                barrier.wait(2)
                for name in ("input_cleaned", "input_cleaned", "unknown", "response_trimmed"):
                    notice(name)
            threads = [threading.Thread(target=worker) for _ in range(8)]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join(4)
            return deepcopy(report)

        bridge = self.make_bridge(batch_runner=runner)
        options = {"num_reviewers": 1, "ministry_policy": "기준_x000D_\n끝"}
        job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge, options=options)))["job_id"], bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(len(seen), 1)
        cleaned = self.notices(job, "input_cleaned")
        self.assertEqual(len(cleaned), 1)
        self.assertEqual(self.notices(job, "response_trimmed"), [])
        self.assert_notice_progress(job)
        self.assertTrue(all(event["message"] in bridge_module._EVENTS.values() for event in job["events"]))
        self.assert_no_notice_text(job["result"])
        self.assert_no_notice_text(bridge._jobs[job["job_id"]].context.report)

    def test_batch_runner_notice_copies_last_progress_and_clean_batch_has_none(self):
        report = self.batch_report([self.result()])

        def runner(in_path, out_path, options, client_factory, **kwargs):
            kwargs["progress_cb"](dict(stage="criteria", completed=2, total=5))
            kwargs["notice_cb"]("input_cleaned")
            return deepcopy(report)

        bridge = self.make_bridge(batch_runner=runner)
        job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"], bridge)
        cleaned = self.notices(job, "input_cleaned")
        self.assertEqual([(e["stage"], e["completed"], e["total"]) for e in cleaned], [("criteria", 2, 5)])
        bridge = self.make_bridge(batch_runner=lambda *args, **kwargs: deepcopy(report))
        job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"], bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(self.notices(job, "input_cleaned"), [])

    def test_notice_without_events_uses_preflight_zero_and_respects_job_guards(self):
        bridge = self.make_bridge()
        job = bridge_module._Job("synthetic-notice-job")
        bridge._jobs[job.job_id] = job
        bridge._notice(job, "input_cleaned")
        bridge._notice(job, "input_cleaned")
        bridge._notice(job, "unknown")
        self.assertEqual(job.events, [dict(sequence=1, stage="preflight", message=NOTICE_TEXTS["input_cleaned"],
                                           completed=0, total=0)])
        job.state = "done"
        bridge._notice(job, "response_trimmed")
        stranger = bridge_module._Job("synthetic-unregistered-job")
        bridge._notice(stranger, "input_cleaned")
        self.assertEqual(len(job.events), 1)
        self.assertEqual(stranger.events, [])

    def test_growing_trimmed_counter_emits_one_response_trimmed(self):
        for grow, expected in ((True, 1), (False, 0)):
            with self.subTest(grow=grow):
                clients = []

                def factory(options, key):
                    client = TrimmingClient(grow=grow, facts=[])
                    clients.append(client)
                    return client

                bridge = self.make_bridge(client_factory=factory)
                job = self.wait_job(self.start_single(bridge=bridge), bridge)
                self.assertEqual(job["state"], "done")
                self.assertGreater(clients[0].trimmed_responses if grow else len(clients[0].engine_calls), 1)
                self.assertEqual(len(self.notices(job, "response_trimmed")), expected)
                self.assertEqual(self.notices(job, "input_cleaned"), [])
                self.assert_notice_progress(job)
                self.assert_no_notice_text(job["result"])
        # The default fake Client has no trimmed_responses attribute at all.
        self.assertFalse(hasattr(Client(facts=[]), "trimmed_responses"))
        job = self.wait_job(self.start_single())
        self.assertEqual(self.notices(job, "response_trimmed"), [])

    def test_batch_clients_each_trimmed_still_emit_one_response_trimmed(self):
        report = self.batch_report([self.result()])

        def runner(in_path, out_path, options, client_factory, **kwargs):
            clients = [client_factory() for _ in range(4)]

            def worker(client):
                for _ in range(3):
                    client.ask_json("system", "{}")
            threads = [threading.Thread(target=worker, args=(client,)) for client in clients]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join(4)
            return deepcopy(report)

        bridge = self.make_bridge(batch_runner=runner, client_factory=lambda options, key: CountingClient())
        job = self.wait_job(self.envelope(bridge.start_batch(self.batch_request(bridge)))["job_id"], bridge)
        self.assertEqual(job["state"], "done")
        self.assertEqual(len(self.notices(job, "response_trimmed")), 1)
        self.assert_notice_progress(job)

    def test_cancellable_client_keeps_cancel_and_model_checks_and_ignores_odd_counters(self):
        opts = domain.normalize_options({"num_reviewers": 1})
        calls = []
        client = CountingClient()
        wrapped = bridge_module._CancellableClient(client, opts, threading.Event(), lambda: calls.append(True))
        wrapped.ask_json("system", "{}")
        wrapped.ask_json("system", "{}")
        self.assertEqual(len(calls), 2)  # the once-per-job rule lives in the bridge helper
        odd = CountingClient()
        odd.trimmed_responses = "not a count"
        bridge_module._CancellableClient(odd, opts, threading.Event(), lambda: calls.append(True)).ask_json("s", "{}")
        self.assertEqual(len(calls), 2)
        cancel = threading.Event()
        cancel.set()
        with self.assertRaises(RuntimeError):
            bridge_module._CancellableClient(CountingClient(), opts, cancel, calls.append).ask_json("s", "{}")
        drifting = CountingClient()
        drifting.model = "other-model"
        with self.assertRaises(RuntimeError):
            bridge_module._CancellableClient(drifting, opts, threading.Event(), calls.append).ask_json("s", "{}")
        self.assertEqual(len(calls), 2)

    def test_default_client_factory_builds_tolerant_client_without_network(self):
        import llm_json_v16
        opts = domain.normalize_options({"num_reviewers": 1})
        client = bridge_module._client(opts, "synthetic-construct-only-key")
        self.assertIsInstance(client, llm_json_v16.TolerantLLMClient)
        self.assertEqual((client.provider, client.model, client.trimmed_responses), (opts.provider, opts.model, 0))


# Spec 4.3: the two screen-only notices, copied here so a wording change is deliberate.
NOTICE_TEXTS = {
    "input_cleaned": "입력 글의 엑셀 줄바꿈 표시·줄 끝 공백·연속 빈 줄 정리 후 검토",
    "response_trimmed": "AI 응답 앞뒤 설명문 제외 후 JSON 결과만 사용",
}


class TrimmingClient(Client):
    """Fake engine client whose counter grows on every call when grow is set."""

    def __init__(self, *, grow=True, **kwargs):
        super().__init__(**kwargs)
        self.grow = grow
        self.trimmed_responses = 0

    def ask_json(self, system_prompt, user_prompt, temperature=None):
        value = super().ask_json(system_prompt, user_prompt, temperature)
        if self.grow:
            self.trimmed_responses += 1
        return value


class CountingClient:
    """Minimal synthetic client: each call counts as a trimmed response."""
    provider = "claude"
    model = domain.DEFAULT_MODELS["claude"]

    def __init__(self):
        self.trimmed_responses = 0

    def ask_json(self, system_prompt, user_prompt, temperature=None):
        if type(self.trimmed_responses) is int:
            self.trimmed_responses += 1
        else:
            self.trimmed_responses += "+"
        return {}


class BridgeMessageStyleTests(unittest.TestCase):
    """Spec 4.8: every fixed error and event literal is a noun-ending 개조식."""

    def test_style_check_detects_v15_sentence_endings(self):
        import desktop_bridge_v15
        for message in desktop_bridge_v15._MESSAGES.values():
            with self.subTest(message=message):
                self.assertTrue(style_violations(message))
        for good in ("진행 중인 작업 있음. 완료 또는 취소 후 다시 시작", "입력 또는 실행 설정 확인 후 다시 시도"):
            self.assertEqual(style_violations(good), [])

    def test_fixed_messages_and_event_table_are_noun_style(self):
        self.assertEqual(set(bridge_module._MESSAGES), {"validation", "catalogue", "credential", "busy", "not_found",
                                                         "io", "runtime"})
        self.assertIn("진행 중", bridge_module._MESSAGES["busy"])
        self.assertEqual(bridge_module._EVENTS["previous_result"], batch_v16.PREVIOUS_RESULT_MESSAGE)
        # Spec 4.4/4.5: the protected-path text lives in _EVENTS, not _MESSAGES.
        self.assertEqual(bridge_module._EVENTS["template_protected"], TEMPLATE_PROTECTED)
        self.assertNotIn(TEMPLATE_PROTECTED, bridge_module._MESSAGES.values())
        self.assertEqual(style_violations(TEMPLATE_PROTECTED), [])
        for table in (bridge_module._MESSAGES, bridge_module._EVENTS):
            for name, message in table.items():
                with self.subTest(name=name):
                    self.assertTrue(message.strip())
                    self.assertEqual(style_violations(message), [])
                    for sentence in re.split(r"[.,]\s*", message):
                        self.assertFalse(sentence.rstrip().endswith(STYLE_ENDINGS), sentence)

    def test_every_hangul_literal_in_the_bridge_is_noun_style_and_events_use_tables(self):
        tree = ast.parse(Path(bridge_module.__file__).read_text(encoding="utf-8"))
        docstrings = {id(node.body[0].value) for node in ast.walk(tree)
                      if isinstance(node, (ast.Module, ast.ClassDef, ast.FunctionDef)) and node.body
                      and isinstance(node.body[0], ast.Expr) and isinstance(node.body[0].value, ast.Constant)}
        literals = [node.value for node in ast.walk(tree) if isinstance(node, ast.Constant)
                    and type(node.value) is str and id(node) not in docstrings and HANGUL.search(node.value)]
        self.assertTrue(literals)
        for literal in literals:
            with self.subTest(literal=literal):
                self.assertEqual(style_violations(literal), [])
        # Every _event(...) message is a table entry, never an inline literal.
        calls = [node for node in ast.walk(tree) if isinstance(node, ast.Call)
                 and isinstance(node.func, ast.Attribute) and node.func.attr == "_event"]
        self.assertTrue(calls)
        for call in calls:
            # Status comparisons (e.g. documents status == '완료') are not messages.
            compared = {id(node) for compare in ast.walk(call.args[2]) if isinstance(compare, ast.Compare)
                        for node in ast.walk(compare)}
            for argument in ast.walk(call.args[2]):
                if id(argument) in compared:
                    continue
                self.assertNotIsInstance(argument, ast.JoinedStr)
                self.assertFalse(isinstance(argument, ast.Constant) and type(argument.value) is str
                                 and HANGUL.search(argument.value), ast.unparse(call))

    def test_public_methods_tuple_is_exactly_the_twelve_names(self):
        self.assertEqual(len(bridge_module.PUBLIC_METHODS), 12)
        self.assertEqual(set(bridge_module.PUBLIC_METHODS), PUBLIC_METHODS)
        # Spec 4.5: the new name is appended last; the eleven keep their order.
        self.assertEqual(bridge_module.PUBLIC_METHODS, (
            "get_initial_state", "select_path", "save_key", "delete_key", "start_single", "start_batch",
            "get_job", "cancel_job", "start_export", "search_notice_clauses", "get_notice_clause",
            "save_input_template"))


class RealCatalogueSearchTests(unittest.TestCase):
    """The two lookups over the installed v16 catalogue (resources_v16/notices)."""

    @classmethod
    def setUpClass(cls):
        cls.cat = load_catalogue()
        cls.backend = FakeBackend()
        cls.bridge = DesktopBridge(key_manager=KeyManager(cls.backend, environ={}),
                                   catalogue_loader=lambda: cls.cat, client_factory=None, path_selector=None)
        cls.clauses = {clause["clause_id"]: clause for clause in cls.cat.clauses}
        cls.sources = {source["source_id"]: source for source in cls.cat.manifest["sources"]}

    def ok(self, response):
        self.assertEqual(set(response), {"ok", "data", "error"})
        self.assertIs(response["ok"], True, response)
        self.assertIsNone(response["error"])
        return response["data"]

    def search(self, query="", family=None, offset=0, limit=100):
        return self.ok(self.bridge.search_notice_clauses(dict(query=query, family=family, offset=offset, limit=limit)))

    def ids(self, query, family=None):
        data = self.search(query, family)
        found, offset = [item["clause_id"] for item in data["items"]], 100
        while offset < data["total"]:
            found += [item["clause_id"] for item in self.search(query, family, offset)["items"]]
            offset += 100
        return found

    def test_frontend_mock_matches_bridge_contract(self):
        """The UI tests' mock must have the same shapes as the real bridge (bridge-UI seam)."""
        from tests.test_v16_frontend import ROOT as UI_ROOT, catalogue_notices
        catalogue_data = self.ok(self.bridge.get_initial_state())["catalogue"]
        self.assertEqual(catalogue_notices(self.cat.manifest["sources"]), catalogue_data["notices"])
        self.assertEqual(dict(domain.NOTICE_TITLES), catalogue_data["notice_titles"])
        bridge_source = Path(bridge_module.__file__).read_text(encoding="utf-8")
        counts = re.search(r"counts=dict\((.*?)\),\s*paths=dict", bridge_source, re.S).group(1)
        bridge_keys = set(re.findall(r"(?:^|,)\s*([a-z_]+)=", counts))
        mock_source = (UI_ROOT / "tests" / "test_v16_frontend.py").read_text(encoding="utf-8")
        mock_keys = set(re.findall(r"([a-z_]+):", re.search(r"counts:\{([^}]*)\}", mock_source).group(1)))
        self.assertEqual(bridge_keys, {"total", "completed", "failed", "cancelled", "needs_review", "excluded"})
        self.assertEqual(mock_keys, bridge_keys)
        app = (UI_ROOT / "ui_v16" / "app.js").read_text(encoding="utf-8")
        labels = re.search(r"\{total:'전체'[^}]*\}", app).group(0)
        self.assertEqual(set(re.findall(r"([a-z_]+):'", labels)), bridge_keys)

    def test_installed_notices_titles_and_body_urls(self):
        catalogue_data = self.ok(self.bridge.get_initial_state())["catalogue"]
        self.assertEqual(catalogue_data["catalogue_id"], "kr-notices-v16-1")
        self.assertEqual(catalogue_data["notice_titles"], {number: title for number, (title, *_) in NOTICES.items()})
        expected = [
            ("산업통상자원부고시 제2025-71호", ["advanced"], ["advanced-annex", "advanced-body"], "advanced-body"),
            ("산업통상부고시 제2025-2호", ["core"], ["core-annex", "core-body"], "core-body"),
            ("산업통상부고시 제2026-101호", ["dual_use", "military"],
             ["strategic-body", "dual-use-annex", "military-annex", "dual-use-hwpx", "military-hwpx"], "strategic-body"),
        ]
        self.assertEqual(len(catalogue_data["notices"]), 3)
        for notice, (number, families, source_ids, body) in zip(catalogue_data["notices"], expected):
            with self.subTest(notice=number):
                self.assertEqual(notice["notice_number"], number)
                self.assertEqual((notice["title"], notice["announcement_date"], notice["effective_date"]), NOTICES[number])
                self.assertEqual(notice["families"], families)
                self.assertEqual(notice["source_ids"], source_ids)
                self.assertEqual(notice["official_url"], self.sources[body]["official_url"])
        for source in catalogue_data["sources"]:
            self.assertEqual(source["announcement_date"], NOTICE_DATES[source["notice_number"]][0])

    def test_whitespace_and_newline_insensitive_title_search(self):
        title = self.clauses["advanced:item:001"]["title"]
        self.assertIn("기\n\n 술", title)
        self.assertNotIn("형성기술", title)
        self.assertIn("advanced:item:001", self.ids("기술", "advanced"))
        for query in ("형성기술", "적층형성기술", "3차원 적층\n형성　기 술", "  형성\t기\r\n술  "):
            with self.subTest(query=query):
                self.assertIn("advanced:item:001", self.ids(query))
                self.assertIn("advanced:item:001", self.ids(query, "advanced"))

    def test_dual_use_body_text_search_casefold_and_nfkc(self):
        clause = self.clauses["dual_use:control:1A001"]
        self.assertIn("연료용 블래더(bladders)", clause["text"])
        self.assertNotIn("블래더", clause["title"] + clause["field"] + clause["code"])
        self.assertNotIn("SEALANTS", clause["text"])
        for query in ("블래더", "연료용블래더", "연료용 블래더", "SEALANTS", "Sealants", "１Ａ００１"):
            with self.subTest(query=query):
                self.assertIn("dual_use:control:1A001", self.ids(query))
                self.assertIn("dual_use:control:1A001", self.ids(query, "dual_use"))
                self.assertNotIn("dual_use:control:1A001", self.ids(query, "core"))
        wire = json.dumps(self.bridge.search_notice_clauses(dict(query="블래더", family=None, offset=0, limit=100)),
                          ensure_ascii=False)
        self.assertNotIn("[국문 원문]", wire)
        self.assertNotIn("연료용 블래더(bladders)", wire)

    def test_family_filter_paging_and_whitespace_only_query(self):
        order = [clause["clause_id"] for clause in self.cat.clauses]
        self.assertEqual(len(order), 571)
        self.assertEqual(self.ids(""), order)
        self.assertEqual(self.ids(" \n\t　 "), order)
        counts = Counter(clause["family"] for clause in self.cat.clauses)
        for family in domain.FAMILIES:
            with self.subTest(family=family):
                data = self.search("", family)
                self.assertEqual(data["total"], counts[family])
                self.assertTrue(all(item["family"] == family for item in data["items"]))
                self.assertEqual(self.ids("", family), [identity for identity in order
                                                        if self.clauses[identity]["family"] == family])
        page = self.search("", None, 120, 7)
        self.assertEqual((page["total"], page["offset"], page["limit"]), (571, 120, 7))
        self.assertEqual([item["clause_id"] for item in page["items"]], order[120:127])
        self.assertEqual(len(self.ok(self.bridge.search_notice_clauses({}))["items"]), 50)
        self.assertEqual(self.search("", None, 571)["items"], [])
        self.assertEqual(self.search("", None, 10 ** 9)["total"], 571)
        for item in self.search("", None, 0, 100)["items"]:
            self.assertEqual(set(item), SEARCH_ITEM_KEYS)
            clause = self.clauses[item["clause_id"]]
            self.assertEqual(item["notice_number"], self.sources[clause["source_id"]]["notice_number"])
            for key in ("family", "source_id", "code", "title", "field"):
                self.assertEqual(item[key], clause[key])

    def test_get_notice_clause_returns_verbatim_text_and_own_source_url(self):
        key_calls = list(self.backend.calls)
        for identity in ("dual_use:control:1A001", "advanced:item:001",
                         max(self.clauses, key=lambda key: len(self.clauses[key]["text"]))):
            with self.subTest(clause=identity):
                data = self.ok(self.bridge.get_notice_clause(identity))
                clause, source = self.clauses[identity], self.sources[self.clauses[identity]["source_id"]]
                self.assertEqual(set(data), CLAUSE_KEYS)
                self.assertEqual(data["text"], clause["text"])
                self.assertEqual(data["source_anchor"], clause["source_anchor"])
                self.assertEqual(data["official_url"], source["official_url"])
                self.assertEqual(data["notice_title"], domain.NOTICE_TITLES[source["notice_number"]])
                self.assertEqual((data["notice_title"], data["announcement_date"], data["effective_date"]),
                                 NOTICES[source["notice_number"]])
        advanced = self.ok(self.bridge.get_notice_clause("advanced:item:001"))
        self.assertEqual(advanced["source_id"], "advanced-annex")
        self.assertNotEqual(advanced["official_url"], self.sources["advanced-body"]["official_url"])
        self.assertGreater(len(self.ok(self.bridge.get_notice_clause(
            max(self.clauses, key=lambda key: len(self.clauses[key]["text"]))))["text"]), 70000)
        self.assertEqual(self.backend.calls, key_calls)


if __name__ == "__main__":
    unittest.main()
