# -*- coding: utf-8 -*-
"""통계 결과를 서식 있는 엑셀 리포트로 출력."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
NAVY = '1F3864'
BLUE = '2F5B95'
GREY = 'F2F2F2'
RED = 'C00000'
ORANGE = 'ED7D31'

thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _title(ws, row, col, text, span=4, color=NAVY, size=11):
    c = ws.cell(row, col, text)
    c.font = Font(name=FONT, bold=True, size=size, color='FFFFFF')
    c.fill = PatternFill('solid', start_color=color)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    for j in range(col, col + span):
        ws.cell(row, j).fill = PatternFill('solid', start_color=color)
    ws.row_dimensions[row].height = 20


def _hdr(ws, row, col, texts):
    for j, t in enumerate(texts):
        c = ws.cell(row, col + j, t)
        c.font = Font(name=FONT, bold=True, size=9, color='FFFFFF')
        c.fill = PatternFill('solid', start_color=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER


def _cell(ws, row, col, v, fmt=None, bold=False, center=False, wrap=False, color=None):
    c = ws.cell(row, col, v)
    c.font = Font(name=FONT, size=9, bold=bold, color=color or '000000')
    c.border = BORDER
    c.alignment = Alignment(horizontal='center' if center else 'left',
                            vertical='center', wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    return c


def _table(ws, start_row, start_col, headers, rows, fmts=None, total=None):
    """헤더 + 데이터 테이블. fmts: 열별 number_format 리스트."""
    _hdr(ws, start_row, start_col, headers)
    r = start_row + 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            fmt = fmts[j] if fmts else None
            _cell(ws, r, start_col + j, v, fmt=fmt,
                  center=isinstance(v, (int, float)) and not isinstance(v, bool))
        r += 1
    if total:
        for j, v in enumerate(total):
            fmt = fmts[j] if fmts else None
            c = _cell(ws, r, start_col + j, v, fmt=fmt, bold=True,
                      center=isinstance(v, (int, float)))
            c.fill = PatternFill('solid', start_color=GREY)
        r += 1
    return r  # 다음 빈 행


def build_report(S, out_path):
    wb = openpyxl.Workbook()
    total = S['총건수'] or 1

    # ═══════════ 시트 1: 통계분석 ═══════════
    ws = wb.active
    ws.title = '통계분석'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFGHIJ', [24, 22, 10, 10, 14, 24, 10, 10, 14, 10]):
        ws.column_dimensions[col].width = w

    _title(ws, 1, 1, '국외지원사항 자동 통계분석 대시보드', span=10, size=13)
    logs_hdr = S.get('전처리내역') or []
    pre_txt = ''
    if logs_hdr:
        rm = sum(lg['제거건수'] for lg in logs_hdr)
        pre_txt = f'   전처리로 오류값 {rm:,}건 제거됨 — 상세는 [전처리 내역] 시트'
    _cell(ws, 2, 1, f'작업일시: {datetime.datetime.now():%Y-%m-%d %H:%M}   '
                    f'(연구자 신고 {S["총건수"]:,}건 / 과제 매핑 {S["총과제매핑건수"]:,}건 기준)'
                    + pre_txt)
    ws.cell(2, 1).border = Border()

    # ── 기본 통계 ──
    r = 4
    _title(ws, r, 1, '기본 통계', span=10)
    r += 1
    basics = [
        ('총 국외지원사항 건수', f"{S['총건수']:,}건"),
        ('총 과제 매핑 국외지원사항 건수', f"{S['총과제매핑건수']:,}건"),
        ('고유 연구자 수', f"{S['고유연구자수']:,}명 (연책: {S['고유연책수']:,}명)"),
        ('지원·지급 국가 수', f"{S['국가수']:,}개국"),
        ('지원·지급 기관 수(별칭 통합)', f"{S['지원기관수']:,}개"),
        ('고유 과제 수(중복 제외)', f"{S['고유과제수']:,}건"),
        ('보고 연구자 정부연구비 총합', f"{S['정부연구비합']:,.0f}"),
        ('보고 연구자 정부연구비 평균', f"{S['정부연구비평균']:,.0f}"),
        ('과제매핑 미진행 연구자 수', f"{S['미매핑연구자수']:,}명"),
        ('과제매핑 미진행 신고 수', f"{S['미매핑신고수']:,}건"),
    ]
    for i, (k, v) in enumerate(basics):
        rr, cc = r + i // 2, 1 + (i % 2) * 5
        _cell(ws, rr, cc, k, bold=True).fill = PatternFill('solid', start_color=GREY)
        ws.merge_cells(start_row=rr, start_column=cc + 1, end_row=rr, end_column=cc + 3)
        _cell(ws, rr, cc + 1, v)
        for j in range(cc + 2, cc + 4):
            ws.cell(rr, j).border = BORDER
    r += len(basics) // 2 + 1

    # ── 상위 국가 ──
    _title(ws, r, 1, '지원·지급 상위 국가 (Top 20)', span=4)
    rows = [(i + 1, k, v, p) for i, (k, v, p) in enumerate(S['국가별'][:20])]
    tsum = sum(x[2] for x in rows)
    next1 = _table(ws, r + 1, 1, ['순위', '국가', '건수', '비율'], rows,
                   fmts=[None, None, '#,##0', '0.0%'],
                   total=['합계', '', tsum, tsum / total])

    # ── 연도별 (오른쪽 배치) ──
    _title(ws, r, 6, '연도별 통계 (NRI 등록일자 기준)', span=4)
    yrows = [(y, c, c / S['날짜인식건수'] if S['날짜인식건수'] else 0) for y, c in S['연도별']]
    next2 = _table(ws, r + 1, 6, ['연도', '건수', '비율', ''], [(a, b, c, '') for a, b, c in yrows],
                   fmts=[None, '#,##0', '0.0%', None],
                   total=['합계', S['날짜인식건수'], 1.0 if S['날짜인식건수'] else 0, ''])
    # ── 월별 ──
    _title(ws, next2 + 1, 6, '월별 통계', span=4)
    mrows = [(m, c, c / S['날짜인식건수'] if S['날짜인식건수'] else 0, '') for m, c in S['월별']]
    next2 = _table(ws, next2 + 2, 6, ['월', '건수', '전체비율', ''], mrows,
                   fmts=[None, '#,##0', '0.0%', None])

    r = max(next1, next2) + 1

    # ── 소속기관 Top 20 ──
    _title(ws, r, 1, '최다 수혜 연구자 소속기관 (Top 20, 산학협력단→본교 통합)', span=10)
    rows = [(i + 1, a, c, n, s) for i, (a, c, n, s) in enumerate(S['소속Top20'])]
    _hdr(ws, r + 1, 1, ['순위', '소속기관', '보고 건수', '연구자 수'])
    ws.merge_cells(start_row=r + 1, start_column=5, end_row=r + 1, end_column=10)
    _hdr(ws, r + 1, 5, ['지원기관(정규화) 목록'])
    for j in range(6, 11):
        ws.cell(r + 1, j).fill = PatternFill('solid', start_color=BLUE)
        ws.cell(r + 1, j).border = BORDER
    rr = r + 2
    for rank, aff, cnt, nres, sup in rows:
        _cell(ws, rr, 1, rank, center=True)
        _cell(ws, rr, 2, aff)
        _cell(ws, rr, 3, cnt, fmt='#,##0', center=True)
        _cell(ws, rr, 4, nres, fmt='#,##0', center=True)
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10)
        _cell(ws, rr, 5, sup, wrap=True)
        for j in range(6, 11):
            ws.cell(rr, j).border = BORDER
        rr += 1
    r = rr + 1

    # ── 연구자 Top 10 ──
    _title(ws, r, 1, '다수 국제지원 수령 연구자 (Top 10)', span=10)
    _hdr(ws, r + 1, 1, ['순위', '연구자명', '소속기관(정규화)', '건수'])
    ws.merge_cells(start_row=r + 1, start_column=5, end_row=r + 1, end_column=10)
    _hdr(ws, r + 1, 5, ['지원기관(정규화) 목록'])
    for j in range(6, 11):
        ws.cell(r + 1, j).fill = PatternFill('solid', start_color=BLUE)
        ws.cell(r + 1, j).border = BORDER
    rr = r + 2
    for i, (name, aff, sup, cnt) in enumerate(S['연구자Top10']):
        _cell(ws, rr, 1, i + 1, center=True)
        _cell(ws, rr, 2, name)
        _cell(ws, rr, 3, aff)
        _cell(ws, rr, 4, cnt, fmt='#,##0', center=True)
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10)
        _cell(ws, rr, 5, sup, wrap=True)
        for j in range(6, 11):
            ws.cell(rr, j).border = BORDER
        rr += 1
    r = rr + 1

    # ── 과제 관련성 / 유형별 / 산학연 ──
    _title(ws, r, 1, '과제와의 연관성', span=4)
    rel = S['관련성']
    rt = rel['합계'] or 1
    next1 = _table(ws, r + 1, 1, ['관련성 구분', '건수', '비율', ''],
                   [('관련 있음', rel['관련있음'], rel['관련있음'] / rt, ''),
                    ('관련 없음', rel['관련없음'], rel['관련없음'] / rt, '')],
                   fmts=[None, '#,##0', '0.0%', None],
                   total=['합계', rel['합계'], 1.0 if rel['합계'] else 0, ''])

    _title(ws, r, 6, '연구자 소속기관 유형별 (산·학·연·기타)', span=4)
    srows = [(t, c, c / total, '') for t, c in S['산학연']]
    next2 = _table(ws, r + 1, 6, ['유형', '건수', '비율', ''], srows,
                   fmts=[None, '#,##0', '0.0%', None],
                   total=['합계', total, 1.0, ''])
    r = max(next1, next2) + 1

    _title(ws, r, 1, '지원·지급 유형별 등록 수 (지원사유 키워드 기반)', span=4)
    trows = [(t, c, c / total, '') for t, c in S['유형별']]
    r = _table(ws, r + 1, 1, ['유형', '건수', '비율', ''], trows,
               fmts=[None, '#,##0', '0.0%', None],
               total=['합계', S['총건수'], sum(c for _, c in S['유형별']) / total, ''])
    c = _cell(ws, r, 1, '중복건수', bold=True)
    c.fill = PatternFill('solid', start_color=GREY)
    c = _cell(ws, r, 2, S['유형중복건수'], fmt='#,##0', bold=True, center=True)
    c.fill = PatternFill('solid', start_color=GREY)
    _cell(ws, r, 3, '', ); _cell(ws, r, 4, '')
    r += 1
    _cell(ws, r, 1, '※ 한 건의 사유에 여러 유형이 포함되면 각 유형에 모두 집계됨 (원본 대시보드와 동일)')
    ws.cell(r, 1).border = Border()
    r += 2

    # ── 연구보안 ──
    _title(ws, r, 1, '연구보안 우려 지원·지급 수 (DOE 지정 위험·민감국가)', span=10, color=RED)
    sec = S['보안']
    ssum = sec['위험국가'] + sec['민감국가']
    r = _table(ws, r + 1, 1, ['보안등급', '건수', '비율', ''],
               [('위험국가', sec['위험국가'], sec['위험국가'] / total, ''),
                ('민감국가', sec['민감국가'], sec['민감국가'] / total, '')],
               fmts=[None, '#,##0', '0.0%', None],
               total=['합계', ssum, ssum / total, ''])
    _cell(ws, r, 1, '※ 세부 리스트는 [연구보안 세부] 시트, 과제 미매핑 연구자는 [미매핑 연구자] 시트 참고')
    ws.cell(r, 1).border = Border()

    # ═══════════ 시트 2: 기관 Top 30 ═══════════
    ws2 = wb.create_sheet('기관 Top30')
    ws2.sheet_view.showGridLines = False
    for col, w in zip('ABCDE', [6, 38, 14, 60, 8]):
        ws2.column_dimensions[col].width = w
    _title(ws2, 1, 1, '지원·지급 횟수 상위 30개 기관 (별칭 통합 + 교차분석)', span=5, size=12)
    _hdr(ws2, 2, 1, ['순위', '기관명(별칭 통합)', '기관 소재 국가', '연구자 소속기관(정규화)', '건수'])
    rr = 3
    for i, (org, nat, affs, cnt) in enumerate(S['기관Top30']):
        _cell(ws2, rr, 1, i + 1, center=True)
        _cell(ws2, rr, 2, org)
        _cell(ws2, rr, 3, nat, center=True)
        _cell(ws2, rr, 4, affs, wrap=True)
        _cell(ws2, rr, 5, cnt, fmt='#,##0', center=True)
        rr += 1
    c = _cell(ws2, rr, 4, 'Top 30 합계', bold=True)
    c.fill = PatternFill('solid', start_color=GREY)
    c = _cell(ws2, rr, 5, S['기관Top30합계'], fmt='#,##0', bold=True, center=True)
    c.fill = PatternFill('solid', start_color=GREY)
    ws2.freeze_panes = 'A3'

    # ═══════════ 시트 3: 연구보안 세부 ═══════════
    ws3 = wb.create_sheet('연구보안 세부')
    ws3.sheet_view.showGridLines = False
    widths = [10, 12, 30, 24, 45, 45]
    for j, w in enumerate(widths, 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    _title(ws3, 1, 1, '연구보안 우려 세부 리스트 (위험·민감국가 수혜 내역)', span=6, size=12, color=RED)
    df = S['보안세부']
    _hdr(ws3, 2, 1, list(df.columns))
    for i, row in df.iterrows():
        for j, v in enumerate(row, 1):
            _cell(ws3, 3 + i, j, '' if v is None else str(v), wrap=j >= 5,
                  color=RED if j == 1 and v == '위험국가' else None,
                  bold=(j == 1 and v == '위험국가'))
    ws3.freeze_panes = 'A3'

    # ═══════════ 시트 4: 미매핑 연구자 ═══════════
    ws4 = wb.create_sheet('미매핑 연구자')
    ws4.sheet_view.showGridLines = False
    um = S['미매핑리스트']
    hdrs = ['번호', '소속기관명', '사업자등록번호', '성명', '국가연구자번호', '지원·지급 기간',
            '지원·지급 사유', '지원·지급 내용', '출처 국가', '출처 기관', 'NRI 등록일자']
    widths = [8, 28, 16, 12, 14, 22, 40, 40, 12, 30, 14]
    for j, w in enumerate(widths, 1):
        ws4.column_dimensions[get_column_letter(j)].width = w
    _title(ws4, 1, 1, f'미협약 또는 과제매핑 미진행 연구자 신고 내역 '
                      f'({S["미매핑연구자수"]:,}명 / {S["미매핑신고수"]:,}건)', span=11, size=12, color=ORANGE)
    _hdr(ws4, 2, 1, hdrs)
    for i, row in um.reset_index(drop=True).iterrows():
        for j, v in enumerate(row, 1):
            if hasattr(v, 'strftime'):
                v = v.strftime('%Y-%m-%d')
            _cell(ws4, 3 + i, j, '' if v is None else str(v), wrap=j in (7, 8))
    ws4.freeze_panes = 'A3'

    # ═══════════ 시트 5: 전처리 내역 ═══════════
    logs = S.get('전처리내역') or []
    if logs:
        ws5 = wb.create_sheet('전처리 내역')
        ws5.sheet_view.showGridLines = False
        for j, w in enumerate([20, 34, 60, 12], 1):
            ws5.column_dimensions[get_column_letter(j)].width = w
        _title(ws5, 1, 1, 'raw 파일 전처리 내역 (오류값 제거 및 정렬 기록)', span=4, size=12)
        r5 = 3
        for lg in logs:
            _title(ws5, r5, 1, f"{lg['구분']}  —  원본 {lg['원본건수']:,}건 → "
                               f"정제 {lg['정제건수']:,}건 (제거 {lg['제거건수']:,}건)", span=4)
            _cell(ws5, r5 + 1, 1, '시트', bold=True).fill = PatternFill('solid', start_color=GREY)
            _cell(ws5, r5 + 1, 2, lg.get('시트', ''))
            _cell(ws5, r5 + 1, 3, f"열 인식: {lg.get('열인식', '')}")
            _cell(ws5, r5 + 1, 4, '')
            _hdr(ws5, r5 + 2, 1, ['처리 단계', '해당 값(건수)', '', '제거 건수'])
            rr = r5 + 3
            for step, detail, cnt in lg['단계']:
                _cell(ws5, rr, 1, step)
                ws5.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
                _cell(ws5, rr, 2, detail, wrap=True)
                ws5.cell(rr, 3).border = BORDER
                _cell(ws5, rr, 4, cnt, fmt='#,##0', center=True)
                rr += 1
            r5 = rr + 1
            if lg['검토후보']:
                _title(ws5, r5, 1, f"검토 후보 — 제거 목록에 없지만 무효값으로 의심되는 표기 "
                                   f"({lg['구분']}, 제거하지 않음)", span=4, color=ORANGE)
                _hdr(ws5, r5 + 1, 1, ['적용 열', '값', '', '건수'])
                rr = r5 + 2
                for colname, val, cnt in lg['검토후보'][:40]:
                    _cell(ws5, rr, 1, colname)
                    ws5.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
                    _cell(ws5, rr, 2, val, wrap=True)
                    ws5.cell(rr, 3).border = BORDER
                    _cell(ws5, rr, 4, cnt, fmt='#,##0', center=True)
                    rr += 1
                _cell(ws5, rr, 1, '※ 확인 후 제거해야 할 값이면 매핑테이블.xlsx의 [전처리규칙] 시트에 추가하세요.')
                ws5.cell(rr, 1).border = Border()
                r5 = rr + 2

    wb.save(out_path)
