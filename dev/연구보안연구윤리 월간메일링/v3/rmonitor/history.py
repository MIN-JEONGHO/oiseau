# -*- coding: utf-8 -*-
"""실행 이력 관리.

매 실행(리포트 생성) 결과를 `history/history.jsonl` 에 한 줄씩(JSON Lines) 누적 기록한다.
발송이 없어진 v3 부터는 "누가 언제 무엇을 생성했는지" 를 대신 남기는 감사 기록이며,
`excel_log.py` 의 누적 데이터(수집된 이슈 원문)와는 별개다 — 혼동하지 말 것.

이 파일은 append-only 로만 다룬다(레코드 수정/삭제 기능 없음). 큰 단일 JSON 배열을
매번 통째로 다시 쓰는 방식과 달리, 한 줄 추가 실패가 기존 기록을 손상시키지 않는다.
"""
from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .profiles import ReportProfile, get_profile
from .utils import Period, now_kst

log = logging.getLogger("rmonitor")

MODE_LABELS = {"monthly": "월간", "weekly": "주간"}
STATUS_LABELS = {"success": "성공", "error": "실패"}

HEADERS = ["번호", "보고서유형", "기간모드", "수집시작일", "수집종료일", "발행일",
           "이슈건수", "생성일시", "상태", "JSON경로", "HTML경로", "비고"]
COL_WIDTHS = [6, 24, 8, 12, 12, 12, 8, 18, 8, 40, 40, 30]

THIN = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_FONT = Font(name="맑은 고딕", size=10)
HEAD_FONT = Font(name="맑은 고딕", size=10, bold=True)
HEAD_FILL_COLOR = "FFE8EEF4"


def _history_path(app_dir: Path, cfg: Dict[str, Any]) -> Path:
    rel = cfg.get("history", {}).get("path", "history/history.jsonl")
    path = app_dir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _append(path: Path, record: Dict[str, Any]) -> None:
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def record_run(result: Dict[str, Any], period: Period, profile: ReportProfile,
               cfg: Dict[str, Any], app_dir: Path, *, status: str = "success") -> Dict[str, Any]:
    """리포트 1건 생성 완료 후 이력을 기록한다."""
    record = {
        "report_key": profile.key,
        "report_name": profile.name,
        "period_mode": period.mode,
        "period_start": period.start_iso,
        "period_end": period.end_iso,
        "period_label": period.period_label,
        "report_date": period.report_date_iso,
        "issue_count": int(result.get("issue_count") or 0),
        "generated_at": now_kst().strftime("%Y-%m-%d %H:%M:%S"),
        "status": status,
        "json_path": result.get("json_path", ""),
        "html_path": result.get("html_path", ""),
        "note": "",
    }
    _append(_history_path(app_dir, cfg), record)
    log.info("[%s] 이력 기록 완료 (%s)", profile.name, record["generated_at"])
    return record


def record_error(report_key: str, period: Period, cfg: Dict[str, Any], app_dir: Path,
                 error: str) -> Dict[str, Any]:
    """실행 실패도 이력에 남긴다."""
    try:
        profile = get_profile(report_key)
        name = profile.name
    except KeyError:
        name = report_key
    record = {
        "report_key": report_key,
        "report_name": name,
        "period_mode": period.mode,
        "period_start": period.start_iso,
        "period_end": period.end_iso,
        "period_label": period.period_label,
        "report_date": period.report_date_iso,
        "issue_count": 0,
        "generated_at": now_kst().strftime("%Y-%m-%d %H:%M:%S"),
        "status": "error",
        "json_path": "",
        "html_path": "",
        "note": str(error)[:500],
    }
    _append(_history_path(app_dir, cfg), record)
    log.warning("[%s] 실패 이력 기록 (%s)", name, error)
    return record


def load_history(app_dir: Path, cfg: Dict[str, Any], *,
                 mode: Optional[str] = None, report_key: Optional[str] = None) -> List[Dict[str, Any]]:
    """이력 전체(또는 조건에 맞는 것만)를 읽어온다. 오래된 것부터 순서 유지."""
    path = _history_path(app_dir, cfg)
    if not path.exists():
        return []
    records: List[Dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if mode and rec.get("period_mode") != mode:
            continue
        if report_key and rec.get("report_key") != report_key:
            continue
        records.append(rec)
    return records


def export_history_excel(app_dir: Path, cfg: Dict[str, Any], out_path: str | Path, *,
                         mode: Optional[str] = None,
                         report_key: Optional[str] = None) -> Path:
    """이력을 엑셀 파일로 내보낸다."""
    records = load_history(app_dir, cfg, mode=mode, report_key=report_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "실행이력"
    ws.append(HEADERS)
    for i, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = HEAD_FONT
        cell.fill = PatternFill(start_color=HEAD_FILL_COLOR, end_color=HEAD_FILL_COLOR, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    for idx, rec in enumerate(records, start=1):
        values = [
            idx,
            rec.get("report_name", rec.get("report_key", "")),
            MODE_LABELS.get(rec.get("period_mode"), rec.get("period_mode", "")),
            rec.get("period_start", ""),
            rec.get("period_end", ""),
            rec.get("report_date", ""),
            rec.get("issue_count", 0),
            rec.get("generated_at", ""),
            STATUS_LABELS.get(rec.get("status"), rec.get("status", "")),
            rec.get("json_path", ""),
            rec.get("html_path", ""),
            rec.get("note", ""),
        ]
        row = idx + 1
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if col in (1, 2, 3, 4, 5, 6, 7, 8, 9) else "left")

    if records:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(records) + 1}"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("이력 엑셀 보고서 생성 완료: %d건 → %s", len(records), out_path)
    return out_path
