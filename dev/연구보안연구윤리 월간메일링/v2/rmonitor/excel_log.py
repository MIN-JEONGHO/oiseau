# -*- coding: utf-8 -*-
"""발송 이슈를 '_양식_동향_모니터링.xlsx' 양식의 누적데이터 시트에 계속 쌓는다.

열 구성(양식 그대로 유지):
  번호 | 발간일 | 국가 | 기관 | 제목 | 주요내용 | 자료 출처(URL) | n대분류 | 원문제목
"""
from __future__ import annotations

import logging
import shutil
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .profiles import COUNTRY_SHORT

log = logging.getLogger("rmonitor")

HEADERS = ["번호", "발간일", "국가", "기관", "제목", "주요내용",
           "자료 출처(URL)", "n대분류", "원문제목"]
COL_WIDTHS = [5.6, 10.6, 5.6, 14.25, 21.5, 107.25, 20.1, 13.5, 25.0]

THIN = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_FONT = Font(name="맑은 고딕", size=10)
LINK_FONT = Font(name="맑은 고딕", size=10, color="0563C1", underline="single")
HEAD_FONT = Font(name="맑은 고딕", size=10, bold=True)


def ensure_workbook(path: str | Path, sheet: str = "누적데이터",
                    template: str | Path | None = None) -> Path:
    """누적 파일이 없으면 양식을 복사하거나 새로 만든다."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        return path

    if template and Path(template).exists():
        shutil.copy(Path(template), path)
        log.info("누적 엑셀 파일을 양식에서 생성했습니다: %s", path)
        return path

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(HEADERS)
    for i, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    log.info("누적 엑셀 파일을 새로 생성했습니다: %s", path)
    return path


def _first_empty_row(ws) -> int:
    """A~I 열이 모두 비어 있는 첫 행 번호."""
    row = 2
    while row <= ws.max_row:
        if any(ws.cell(row=row, column=c).value not in (None, "") for c in range(1, 10)):
            row += 1
        else:
            return row
    return max(row, 2)


def _existing_urls(ws) -> set:
    urls = set()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=7).value
        if v:
            urls.add(str(v).strip())
    return urls


def _last_number(ws) -> int:
    last = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        try:
            last = max(last, int(v))
        except (TypeError, ValueError):
            continue
    return last


def append_issues(data: Dict[str, Any], excel_cfg: Dict[str, Any],
                  logger: logging.Logger | None = None) -> Dict[str, int]:
    """수집 데이터를 누적 시트에 추가. 추가/중복 건수를 반환."""
    logger = logger or log
    path = ensure_workbook(excel_cfg["path"], excel_cfg.get("sheet", "누적데이터"),
                           excel_cfg.get("template"))
    sheet_name = excel_cfg.get("sheet", "누적데이터")

    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    if ws.max_row == 1 and all(ws.cell(row=1, column=c).value in (None, "") for c in range(1, 10)):
        for i, h in enumerate(HEADERS, start=1):
            ws.cell(row=1, column=i, value=h).font = HEAD_FONT

    skip_dup = bool(excel_cfg.get("skip_duplicate_url", True))
    known = _existing_urls(ws) if skip_dup else set()
    row = _first_empty_row(ws)
    number = _last_number(ws)

    added = duplicated = 0
    report_date = data["period"]["report_date"]

    for issue in data["issues"]:
        url = issue.get("source_url", "").strip()
        if skip_dup and url in known:
            duplicated += 1
            continue
        number += 1
        values: List[Any] = [
            number,
            report_date,
            COUNTRY_SHORT.get(issue["country"], issue["country"]),
            issue.get("org", ""),
            issue.get("title_kr", ""),
            issue.get("content", ""),
            url,
            issue.get("category", ""),
            issue.get("title_original", ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if col in (1, 2, 3) else "left")
        url_cell = ws.cell(row=row, column=7)
        if url:
            url_cell.hyperlink = url
            url_cell.font = LINK_FONT
        known.add(url)
        row += 1
        added += 1

    wb.save(path)
    logger.info("엑셀 누적 저장: %d건 추가, %d건 중복 제외 → %s", added, duplicated, path)
    return {"added": added, "duplicated": duplicated, "path": str(path)}
