# -*- coding: utf-8 -*-
"""파일 판독 계층: 인코딩·구분자 추정, CSV/XLSX/ZIP 판독."""
from __future__ import annotations

import csv
import io
import os
import zipfile
from typing import Dict, List, Optional, Tuple

# 표준 라이브러리만으로 처리 가능한 인코딩 후보 (exe 용량 최소화를 위해 chardet 미사용)
ENCODING_CANDIDATES = ("utf-8", "cp949", "euc-kr", "utf-16")

csv.field_size_limit(10 ** 7)


def col_letter(idx: int) -> str:
    """0-based 열 번호를 엑셀 열문자로 변환. 0 -> A"""
    s = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def cell_ref(row: int, col: int) -> str:
    """0-based (행, 열) → 엑셀식 좌표. (0,0) -> A1"""
    return f"{col_letter(col)}{row + 1}"


def detect_encoding(path: str) -> Tuple[str, bytes]:
    """인코딩을 추정해 (인코딩명, 원본 바이트)를 돌려준다."""
    with open(path, "rb") as fh:
        raw = fh.read()
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", raw
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16", raw
    for enc in ENCODING_CANDIDATES:
        try:
            raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        return enc, raw
    return "unknown", raw


def decode(raw: bytes, encoding: str) -> str:
    try:
        return raw.decode(encoding, errors="replace")
    except LookupError:
        return raw.decode("utf-8", errors="replace")


def detect_delimiter(text: str) -> str:
    """1~5행을 표본으로 실제 사용된 구분자를 추정한다."""
    sample = "\n".join(text.splitlines()[:5])
    if not sample:
        return ","
    scores: Dict[str, int] = {}
    for cand in (",", ";", "\t", "|", "^"):
        scores[cand] = sample.count(cand)
    best = max(scores, key=lambda k: scores[k])
    return best if scores[best] > 0 else ","


def read_csv_text(text: str, delimiter: str = ",") -> List[List[str]]:
    """따옴표 규칙을 지켜 CSV 텍스트를 2차원 배열로 판독한다."""
    buf = io.StringIO(text, newline="")
    return [row for row in csv.reader(buf, delimiter=delimiter)]


def read_xlsx(path: str, max_rows: int = 20000) -> Tuple[List[List[str]], List[str]]:
    """XLSX 판독. (셀 배열, 시트명 목록)"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("XLSX 판독에는 openpyxl 이 필요합니다. pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows: List[List[str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        rows.append(["" if c is None else str(c) for c in row])
    names = list(wb.sheetnames)
    wb.close()
    return rows, names


def read_json(path: str) -> Tuple[List[List[str]], str, Dict]:
    """JSON 판독. 객체 배열을 표로 펼친다. (셀 배열, 인코딩, 원본구조 정보)"""
    import json
    enc, raw = detect_encoding(path)
    data = json.loads(decode(raw, enc if enc != "unknown" else "utf-8"))

    root_key = ""
    if isinstance(data, dict):
        for k, v in data.items():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                root_key, data = k, v
                break
    if not (isinstance(data, list) and data and isinstance(data[0], dict)):
        return [], enc, {"table": False}

    header: List[str] = []
    for row in data:
        for k in row:
            if k not in header:
                header.append(k)
    rows = [header]
    for row in data:
        rows.append(["" if row.get(k) is None else str(row.get(k, "")) for k in header])
    return rows, enc, {"table": True, "root_key": root_key, "count": len(data)}


def write_json(rows: List[List[str]], path: str, encoding: str = "utf-8",
               root_key: str = "") -> None:
    """표를 JSON(객체 배열)으로 되돌려 저장한다."""
    import json
    if not rows:
        return
    header, body = rows[0], rows[1:]
    records = [{h: (r[i] if i < len(r) else "") for i, h in enumerate(header)}
               for r in body]
    payload = {root_key: records} if root_key else records
    with open(path, "w", encoding=encoding or "utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)


def write_csv(rows: List[List[str]], path: str, encoding: str = "utf-8-sig",
              delimiter: str = ",") -> str:
    """표를 CSV로 저장한다. 지정 인코딩으로 쓸 수 없으면 UTF-8(BOM)로 대체한다."""
    try:
        with open(path, "w", encoding=encoding, newline="") as fh:
            csv.writer(fh, delimiter=delimiter).writerows(rows)
        return encoding
    except UnicodeEncodeError:
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh, delimiter=delimiter).writerows(rows)
        return "utf-8-sig"


def write_xlsx(rows: List[List[str]], path: str) -> None:
    """표를 XLSX로 저장한다."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def zip_entries(path: str) -> List[str]:
    """ZIP 내부 파일 목록(디렉터리 제외)."""
    with zipfile.ZipFile(path) as zf:
        return [n for n in zf.namelist() if not n.endswith("/")]


def ext_of(name: str) -> str:
    return os.path.splitext(name)[1].lower().lstrip(".")


def trim_trailing_blank_rows(rows: List[List[str]]) -> List[List[str]]:
    """파일 끝의 완전 공백 행을 잘라낸다."""
    end = len(rows)
    while end > 0 and all((c or "").strip() == "" for c in rows[end - 1]):
        end -= 1
    return rows[:end]
